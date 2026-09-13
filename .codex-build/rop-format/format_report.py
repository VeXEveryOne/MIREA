from pathlib import Path
from copy import deepcopy
from io import BytesIO
import sys, shutil, re
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'bd-format'))
from format_reports import node, font, para, section, section_break
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
from docx.enum.style import WD_STYLE_TYPE

ROOT=Path('D:/GitHub/MIREA'); WORK=ROOT/'.codex-build/rop-format'
DEST=ROOT/'4/РОП/РОП_Единый_отчет_АлбахтинИВ.docx'
BACKUP=WORK/'original.docx'
if not BACKUP.exists():shutil.copy2(DEST,BACKUP)
d=Document(BACKUP); ref=Document(ROOT/'4/ИСУРО/ИСУРО_1_АлбахтинИВ.docx')
first=next(p._p for p in d.paragraphs if p.text=='ПРАКТИЧЕСКОЕ ЗАНЯТИЕ № 1')
while d.element.body[0] is not first:d.element.body.remove(d.element.body[0])
for st in d.styles:
    if st.type==WD_STYLE_TYPE.PARAGRAPH:
        st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
for name,size,align,indent,line,bold,italic in [
    ('Основной текст отчёта',14,A.JUSTIFY,1.25,1.5,False,False),
    ('Подпись таблицы',12,A.LEFT,0,1,False,True),
    ('Подпись рисунка',14,A.CENTER,0,1,False,False),
    ('Текст таблицы',12,A.LEFT,0,1,False,False),
    ('Heading 1',14,A.CENTER,0,1.5,True,False),
    ('Heading 2',14,A.LEFT,0,1.5,True,False)]:
    st=d.styles[name] if name in d.styles else d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=bold;st.font.italic=italic
    f=st.paragraph_format;f.alignment=align;f.first_line_indent=Cm(indent);f.left_indent=Cm(0);f.right_indent=Cm(0);f.line_spacing=line;f.space_before=Pt(0);f.space_after=Pt(0)

# Uniform footer shared by all sections, with a blank title-page footer.
for sec in d.sections:
    for hdr in [sec.header,sec.first_page_header,sec.footer,sec.first_page_footer]:
        for p in hdr.paragraphs:p.clear()
foot=d.sections[0].footer
p=foot.paragraphs[0];para(p,A.CENTER,0,1)
fld=node('fldSimple',instr=' PAGE ');r=node('r');r.append(node('rPr'));r[0].append(node('rFonts',ascii='Times New Roman',hAnsi='Times New Roman'));r[0].append(node('sz',val=24));t=node('t');t.text='2';r.append(t);fld.append(r);p._p.append(fld)
base=deepcopy(d.sections[0]._sectPr)
footerref=deepcopy(base.find(qn('w:footerReference')))
for i,sec in enumerate(list(d.sections)):
    old=sec._sectPr;land=sec.page_width>sec.page_height
    new=section(base,land)
    for e in new.findall(qn('w:footerReference')):new.remove(e)
    new.insert(0,deepcopy(footerref))
    old.getparent().replace(old,new)

for p in list(d.paragraphs):
    if p._p.xpath('./w:pPr/w:sectPr'):
        para(p,A.LEFT,0,1);p.paragraph_format.line_spacing=Pt(1)
        for r in p.runs:font(r,1)
        continue
    if not p.text.strip() and not p._p.xpath('.//w:drawing|.//w:pict'):
        p._p.getparent().remove(p._p);continue
    old=p.style.name;para(p)
    for br in p._p.xpath('.//w:br[@w:type="page"]'):br.getparent().remove(br)
    if old.startswith('Heading'):
        p.style=old;para(p,A.CENTER if old=='Heading 1' else A.LEFT,0,1.5,12,6,True,True)
        for r in p.runs:font(r,14,True,False)
    elif re.match(r'^(Таблица|Рисунок)\s+\d+',p.text):
        kind,num,title=re.match(r'^(Таблица|Рисунок)\s+(\d+)\s*[-–—]\s*(.*)',p.text).groups()
        p.text=f'{kind} {num} – {title.rstrip(".")}';tab=kind=='Таблица'
        p.style='Подпись таблицы' if tab else 'Подпись рисунка'
        para(p,A.LEFT if tab else A.CENTER,0,1,6 if tab else 3,3 if tab else 6,tab,True)
        for r in p.runs:font(r,12 if tab else 14,False,tab)
    elif p._p.xpath('.//w:drawing|.//w:pict'):para(p,A.CENTER,0,1,6,0,True,True)
    elif old=='List Number':
        para(p,A.JUSTIFY,0,1.5);p.paragraph_format.left_indent=Cm(.75);p.paragraph_format.first_line_indent=Cm(-.75)
    else:p.style='Основной текст отчёта'

# Preserve existing landscapes and size figures against their own section.
for sh in d.inline_shapes:
    p=sh._inline.getparent()
    while p.tag!=qn('w:p'):p=p.getparent()
    e=p;sp=None
    while e is not None:
        found=e.xpath('./w:pPr/w:sectPr')
        if found:sp=found[0];break
        if e.tag==qn('w:sectPr'):sp=e;break
        e=e.getnext()
    landscape=sp.find(qn('w:pgSz')).get(qn('w:orient'))=='landscape'
    limit=Cm(25.2 if landscape else 16.5)
    if sh.width>limit:ratio=limit/sh.width;sh.width=limit;sh.height=int(sh.height*ratio)

widthsets=[[3.5,13],[3.5,6.5,6.5],[1.4,5.1,5,5],[1.5,8.1,4.2,2.7],[2.1,4.3,10.1],[8.2,3.4,3.4,3.4,3.4,3.4]]
for ti,table in enumerate(d.tables):
    widths=widthsets[ti];table.autofit=False;table.alignment=1
    pr=table._tbl.tblPr
    for tag in ['tblInd','tblW','tblBorders']:
        for e in pr.findall(qn('w:'+tag)):pr.remove(e)
    pr.append(node('tblW',w=Cm(sum(widths)).twips,type='dxa'))
    borders=node('tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:borders.append(node(edge,val='single',sz=6,color='000000'))
    pr.append(borders)
    for col,w in zip(table.columns,widths):col.width=Cm(w)
    for ri,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr()
        for tag in ['trHeight','cantSplit','tblHeader']:
            for e in trpr.findall(qn('w:'+tag)):trpr.remove(e)
        trpr.append(node('cantSplit'))
        if ri==0:trpr.append(node('tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]);cp=c._tc.get_or_add_tcPr()
            for tag in ['shd','tcMar','tcBorders']:
                for e in cp.findall(qn('w:'+tag)):cp.remove(e)
            mar=node('tcMar')
            for k,v in [('top',50),('bottom',50),('left',70),('right',70)]:mar.append(node(k,w=v,type='dxa'))
            cp.append(mar)
            for p in c.paragraphs:
                p.style='Текст таблицы';para(p,A.CENTER if (ti==5 and ci>0) or (ci==0 and ti in [2,3,4]) else A.LEFT,0,1,0,0,False,True)
                for r in p.runs:font(r,12,ri==0,False)

# Import the actual cover components, preserving their local typography.
for idx,orig in enumerate(list(ref.element.body)[:16]):
    el=deepcopy(orig)
    for pp in el.xpath('.//w:p') if el.tag!=qn('w:p') else [el]:
        p=Paragraph(pp,d._body)
        direct=deepcopy(p._p.pPr) if p._p.pPr is not None else node('pPr')
        merged=deepcopy(ref.styles['Normal'].element.pPr)
        for child in direct:
            for old in merged.findall(child.tag):merged.remove(old)
            merged.append(deepcopy(child))
        for old in list(merged):
            if old.tag in [qn('w:pStyle'),qn('w:sectPr')]:merged.remove(old)
        if p._p.pPr is not None:p._p.remove(p._p.pPr)
        p._p.insert(0,merged)
        for r in p.runs:
            if r.text:r.text=r.text.replace('ИНБО-01-17','').replace('Информационные системы управления ресурсами организации','Разработка обеспечивающих подсистем')
            font(r,r.font.size.pt if r.font.size else 14)
        if 'Шендяпин' in p.text:
            p.text='Кириллина Ю.В., Трушин С.М.'
            for r in p.runs:font(r,12,False,False)
        p.paragraph_format.keep_with_next=False;p.paragraph_format.page_break_before=False
    for e in el.iter():
        for attr,val in list(e.attrib.items()):
            if attr.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'):
                rel=ref.part.rels.get(val)
                if rel is None:continue
                if rel.is_external:new=d.part.relate_to(rel.target_ref,rel.reltype,True)
                elif rel.reltype.endswith('/image'):new,_=d.part.get_or_add_image(BytesIO(rel.target_part.blob))
                else:new=d.part.relate_to(rel.target_part,rel.reltype)
                e.set(attr,new)
    if idx==7:
        p=Paragraph(el,d._body);p.text='Практическое занятие № 1';p.alignment=A.CENTER
        for r in p.runs:font(r,14,True,False)
    if idx==10:Paragraph(el,d._body).paragraph_format.space_after=Pt(54)
    if idx==14:Paragraph(el,d._body).paragraph_format.space_after=Pt(72)
    if idx==15:Paragraph(el,d._body)._p.get_or_add_pPr().append(section(base,False,True))
    first.addprevious(el)

p=node('p');first.addprevious(p);p=Paragraph(p,d._body);p.text='Содержание';para(p,A.CENTER,0,1.5,0,12,True,True)
for r in p.runs:font(r,14,True,False)
p=node('p');first.addprevious(p);p.append(node('fldSimple',instr=' TOC \\o "1-2" \\h \\z \\u '))
first.addprevious(section_break(base,False))
for name in ['TOC 1','TOC 2']:
    st=d.styles[name] if name in d.styles else d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    st.font.name='Times New Roman';st.font.size=Pt(14);st.font.bold=False
    st.paragraph_format.line_spacing=1.0;st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.first_line_indent=Cm(0);st.paragraph_format.left_indent=Cm(0 if name=='TOC 1' else .7)
for i,el in enumerate(d.element.body.xpath('.//wp:docPr'),1):el.set('id',str(i))
d.settings.element.append(node('updateFields',val='true'))
d.save(WORK/DEST.name)
print(WORK/DEST.name)
