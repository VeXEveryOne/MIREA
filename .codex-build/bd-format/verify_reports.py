from pathlib import Path
from collections import Counter
from hashlib import sha256
from zipfile import ZipFile
import json, re
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from pypdf import PdfReader

BASE=Path(__file__).resolve().parent
ROOT=BASE.parents[1]
def digest(p): return sha256(p.read_bytes()).hexdigest()
def norm(t):
    t=re.sub(r'\s+',' ',t).strip()
    m=re.match(r'^(Таблица|Рисунок|Рис\.)\s*(\d+)\s*[.\-–—]?\s*(.*)',t)
    if m: return ('Рисунок' if m[1]=='Рис.' else m[1])+' '+m[2]+' '+m[3].rstrip('.')
    return t
def paragraphs(doc):
    return [Paragraph(p,doc._body) for p in doc.element.body.xpath('.//w:p')]
def effective_para(p,key,default=0):
    v=getattr(p.paragraph_format,key)
    st=p.style
    while v is None and st is not None:
        v=getattr(st.paragraph_format,key);st=st.base_style
    return default if v is None else v
def effective_font(p,r,key):
    v=getattr(r.font,key)
    for st in [r.style,p.style]:
        while v is None and st is not None:
            v=getattr(st.font,key);st=st.base_style
    return v
def media(path):
    with ZipFile(path) as z:
        return Counter(sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))

results=[]
for n,qa,count in [(1,'qa1-v5',41),(2,'qa2-v3',14)]:
    name=f'Практическая работа №{n}'
    source=BASE/'backups'/f'practice_{n}'/(name+'.docx')
    final=BASE/f'practice_{n}'/(name+'_fields.docx')
    original=ROOT/f'4/BD/practice_{n}'/(name+'.docx')
    assert digest(original)==digest(source),'Original changed while working'
    assert digest(original.with_suffix('.pdf'))==digest(source.with_suffix('.pdf'))
    before=Document(source);after=Document(final)
    bp=paragraphs(before);ap=paragraphs(after)
    missing=Counter(norm(p.text) for p in bp if p.text.strip())-Counter(norm(p.text) for p in ap if p.text.strip())
    assert not missing,dict(missing)
    original_code=Counter(p.text for p in bp if p.style.name in ['Код','Результат'] and p.text.strip())
    assert not original_code-Counter(p.text for p in ap),'Code text changed'
    assert not media(source)-media(final),'Media content lost'
    caps=[]
    for p in after.paragraphs:
        if not re.match(r'^(Таблица|Рисунок)\s+\d+\s+–',p.text): continue
        table=p.text.startswith('Таблица');f=p.paragraph_format
        assert int(effective_para(p,'alignment'))==(0 if table else 1),p.text
        assert all(effective_para(p,k)==0 for k in ['left_indent','right_indent','first_line_indent']),p.text
        for r in p.runs:
            if not r.text: continue
            assert effective_font(p,r,'name')=='Times New Roman' and effective_font(p,r,'size').pt==(12 if table else 14),p.text
            assert bool(effective_font(p,r,'italic'))==table and not effective_font(p,r,'bold'),p.text
        caps.append(p.text)
    for s in after.sections:
        assert [s.top_margin.twips,s.bottom_margin.twips,s.left_margin.twips,s.right_margin.twips]==[1134,1134,1701,850]
        assert sorted([s.page_width.twips,s.page_height.twips])==[11906,16838]
    pdf=PdfReader(BASE/qa/(name+'.pdf'))
    assert len(pdf.pages)==count
    texts=[p.extract_text() for p in pdf.pages]
    title=re.sub(r'\s+','',texts[0])
    assert 'ВоронцовЮ.А.' in title and 'ИНБО-12-23' in title
    assert 'Шендяпин' not in title and 'ИНБО-01-17' not in title
    for i,t in enumerate(texts[1:],2):
        assert re.search(r'(?m)^\s*'+str(i)+r'\s*$',t),(i,t[-80:])
    assert not re.search(r'Ошибка!|Error!', '\n'.join(texts))
    results.append(dict(practice=n,pages=count,original_paragraphs=len(bp),original_media=sum(media(source).values()),code_paragraphs=sum(original_code.values()),captions=len(caps),final_sha256=digest(final),text_preserved=True,media_preserved=True))
assert digest(ROOT/'4/ИСУРО/ИСУРО_1_АлбахтинИВ.docx')=='53a4b8ccc73cf44599e8b7c7fb75045afec85dd7f4b0a079b19a0d4ccb5249b1'
(BASE/'verification.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(results,ensure_ascii=False,indent=2))
