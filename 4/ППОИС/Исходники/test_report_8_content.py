"""Mutation tests inspect in-memory copies, never edit canonical documents."""
from hashlib import sha256
import unittest

from docx import Document
from docx.oxml.ns import qn

from build_report_8 import SUBJECT, inputs
from verify_report_8_content import inspect_content


class ContentAuditTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.db, cls.math, cls.models = inputs()

    def setUp(self):
        self.document = Document(SUBJECT / 'ППОИС_8_АлбахтинИВ.docx')

    def failed(self):
        return [c['name'] for c in inspect_content(self.document, self.db, self.math, self.models)
                if not c['passed']]

    def test_current_report_passes(self):
        self.assertEqual(self.failed(), [])

    def test_floor_instead_of_ceiling_detected(self):
        self.document.element.xpath('//m:dPr/m:begChr')[0].set(qn('m:val'), '⌊')
        self.assertIn('ceil not floor or ordinary brackets', self.failed())

    def test_commission_denominator_detected(self):
        self.document.element.xpath('//m:f/m:den//m:t')[0].text = '1 + c'
        self.assertIn('equation 6 numerator and denominator', self.failed())

    def test_sum_upper_limit_detected(self):
        self.document.element.xpath('//m:nary/m:sup//m:t')[0].text = 'n - 1'
        self.assertIn('equation 4 symbol limits and term', self.failed())

    def test_field_optional_mismatch_detected(self):
        self.document.tables[3].cell(1, 3).text = 'Нет'
        self.assertIn('type optionality form and mode component_group.id', self.failed())

    def test_control_quantity_mismatch_detected(self):
        self.document.tables[4].cell(1, 1).text = '0'
        self.assertIn('every executed quantity price and mass', self.failed())

    def test_definition_after_formula_detected(self):
        definition = next(p for p in self.document.paragraphs if p.text.startswith('Момент t'))
        eq_paragraph = self.document.element.xpath('//m:oMath')[0].getparent()
        eq_paragraph.addnext(definition._p)
        self.assertIn('definition precedes equation 1: Момент t', self.failed())

    def test_removed_embedded_figure_detected(self):
        # Leave its relationship/media in the package: orphan hashes are not
        # proof that the document actually displays a required illustration.
        form_hash = next(e['sha256'] for e in self.models['emitted'] if e['file'] == 'ПР8_F1.png')
        drawing = next(d for d in self.document.element.xpath('//w:drawing')
                       if any(sha256(self.document.part.rels[r].target_part.blob).hexdigest() == form_hash
                              for r in d.xpath('.//@r:embed')))
        drawing.getparent().remove(drawing)
        self.assertIn('all eleven current form and algorithm images embedded', self.failed())


if __name__ == '__main__':
    unittest.main()
