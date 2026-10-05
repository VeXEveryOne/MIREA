"""PPOIS7 report from the executed PostgreSQL catalog and native ER exports.

The builder refuses stale SQL, missing model coverage or a failed native export.
It does not turn a formatting audit into proof of academic completeness.
"""
from __future__ import annotations

from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path
import re

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from report_layout import Report, ROOT, SUBJECT, field, font

DATABASE = SUBJECT / 'База_данных'
DIAGRAMS = SUBJECT / 'Схемы'
ORDER = ['component_group', 'component', 'supplier', 'offer', 'configuration',
         'bom_version', 'bom_slot', 'card', 'card_revision', 'image',
         'revision_image', 'publication_attempt']
NAMES = {
    'component_group': 'Группа компонентов', 'component': 'Компонент',
    'supplier': 'Поставщик', 'offer': 'Предложение',
    'configuration': 'Конфигурация', 'bom_version': 'Версия состава',
    'bom_slot': 'Позиция состава', 'card': 'Карточка',
    'card_revision': 'Ревизия карточки', 'image': 'Изображение',
    'revision_image': 'Изображение ревизии', 'publication_attempt': 'Попытка публикации',
}
DESCRIPTIONS = {
    'component_group': 'Группа допустимых компонентов и её функциональная роль: CPU, BOARD, RAM, SSD, CASE, PSU или COOLER.',
    'component': 'Конкретная модель комплектующего: SKU, масса и применимые технические характеристики.',
    'supplier': 'Источник предложения с уникальным кодом. Сведения в учебной базе не являются реквизитами реального поставщика.',
    'offer': 'Наблюдение цены и остатка конкретного компонента у поставщика в зафиксированный момент времени.',
    'configuration': 'Стабильный объект компьютерной конфигурации; изменения её состава оформляются версиями.',
    'bom_version': 'Версия спецификации состава внутри конфигурации. Локальный номер версии дополняет ключ владельца.',
    'bom_slot': 'Позиция версии состава: группа либо точный компонент, количество и выбранное предложение.',
    'card': 'Стабильная товарная карточка с уникальным кодом продавца, не зависящим от номера ревизии.',
    'card_revision': 'Снимок содержания и экономических параметров карточки, связанный с одной версией состава.',
    'image': 'Метаданные изображения: URI, контрольная сумма и размеры. Файл хранится вне таблицы.',
    'revision_image': 'Связь ревизии и изображения с порядковым номером для отображения; разрешает отношение M:N.',
    'publication_attempt': 'Попытка публикации конкретной ревизии: состояние, внешний ID, ошибка и срок проверки результата.',
}
CHECKS = {
    'bom_slot_check': 'Заполнено ровно одно поле: group_id либо requested_component_id.',
    'bom_slot_line_no_check': 'line_no > 0.',
    'bom_slot_quantity_check': 'quantity > 0.',
    'bom_version_state_check': 'state ∈ {DRAFT, CHECKED, FROZEN}.',
    'bom_version_version_no_check': 'version_no > 0.',
    'card_revision_assembly_rub_check': 'assembly_rub ≥ 0.',
    'card_revision_commission_rate_check': '0 ≤ commission_rate < 1.',
    'card_revision_logistics_rub_check': 'logistics_rub ≥ 0.',
    'card_revision_markup_rate_check': '0 ≤ markup_rate ≤ 1.',
    'card_revision_packaging_mass_kg_check': 'packaging_mass_kg ≥ 0.',
    'card_revision_packaging_rub_check': 'packaging_rub ≥ 0.',
    'card_revision_revision_no_check': 'revision_no > 0.',
    'card_revision_round_step_rub_check': 'round_step_rub > 0.',
    'card_revision_state_check': 'state ∈ {DRAFT, READY, STALE}.',
    'component_mass_kg_check': 'mass_kg > 0.',
    'component_power_w_check': 'power_w > 0 при заполнении; NULL допустим.',
    'component_group_role_code_check': 'role_code ∈ {CPU, BOARD, RAM, SSD, CASE, PSU, COOLER}.',
    'image_height_px_check': 'height_px > 0.',
    'image_sha256_check': 'sha256 соответствует ^[0-9a-f]{64}$: 64 шестнадцатеричных символа.',
    'image_width_px_check': 'width_px > 0.',
    'offer_stock_check': 'stock ≥ 0.',
    'offer_unit_price_rub_check': 'unit_price_rub > 0.',
    'publication_attempt_attempt_no_check': 'attempt_no > 0.',
    'publication_attempt_check': 'ACCEPTED или PUBLISHED требуют external_task_id IS NOT NULL.',
    'publication_attempt_check1': 'REJECTED или TEMP_ERROR требуют error_code IS NOT NULL.',
    'publication_attempt_check2': 'confirmed_not_accepted = true допустимо только при TEMP_ERROR.',
    'publication_attempt_state_check': 'state ∈ {NEW, ACCEPTED, PUBLISHED, REJECTED, TEMP_ERROR, UNKNOWN}.',
    'revision_image_ordinal_check': 'ordinal > 0.',
}
TEST_NAMES = {
    'duplicate configuration code': 'Повтор кода конфигурации',
    'negative offer price': 'Отрицательная цена предложения',
    'missing supplier FK': 'Ссылка на отсутствующего поставщика',
    'zero quantity': 'Нулевое количество в позиции',
    'group mismatch': 'Несоответствие группы выбранному предложению',
    'missing selected offer FK': 'Ссылка на отсутствующее предложение',
    'selected offer belongs to another group': 'Предложение относится к другой группе',
    'frozen BOM cannot be edited': 'Изменение состава, используемого ревизией',
    'frozen BOM slot cannot move to another version': 'Перенос используемой позиции в другую версию',
    'economic snapshot cannot be overwritten': 'Перезапись экономических параметров ревизии',
    'used offer price is immutable': 'Изменение цены используемого предложения',
    'used component mass is immutable': 'Изменение массы используемого компонента',
    'accepted task requires external ID': 'ACCEPTED без внешнего идентификатора',
    'viewer cannot update catalog': 'Изменение каталога ролью viewer',
    'engineer cannot publish': 'Создание публикации ролью engineer',
    'manager cannot rewrite price policy': 'Изменение наценки ролью manager',
    'worker cannot delete attempts': 'Удаление попытки ролью worker',
    'viewer cannot create tables': 'Создание таблицы ролью viewer',
    'ppois_viewer allowed operation': 'Разрешённое чтение viewer',
    'ppois_manager allowed operation': 'Разрешённое обновление state ревизии manager',
    'ppois_worker allowed operation': 'Разрешённое обновление результата worker',
    'Q01 excludes future offer observations': 'Q01 исключает наблюдение из будущего',
}
QUERY_NAMES = [
    'Q01 Допустимые предложения для позиций состава',
    'Q02 Стоимость и масса выбранного состава',
    'Q03 Состояние ревизии и результат публикации',
    'Q04 Карточки, затронутые изменением компонента',
    'Q05 Очередь сверки внешнего результата',
    'Q06 Ревизии без изображений',
]


def digest(path):
    return sha256(Path(path).read_bytes()).hexdigest()


def load_inputs():
    evidence = json.loads((DATABASE / 'Результаты_проверки.json').read_text(encoding='utf-8'))
    audit = json.loads((DIAGRAMS / 'ПР7_Проверка_модели.json').read_text(encoding='utf-8'))
    manifest = json.loads((DIAGRAMS / 'ПР7_Экспорт_Чена.json').read_text(encoding='utf-8'))
    for name, expected in evidence['source_sha256'].items():
        if digest(DATABASE / name) != expected:
            raise ValueError('Stale database evidence: ' + name)
    if audit['database'] != evidence['database'] or audit['source_sha256'] != evidence['source_sha256']:
        raise ValueError('Native models were built from another database catalog')
    if audit['diagnostics'] or not all(test['passed'] for test in evidence['tests']):
        raise ValueError('Failed model or database checks')
    expected_columns = {c['table_name'] + '__' + c['column_name'] for c in evidence['columns']}
    if set(audit['logicalAttributeCoverage']) != expected_columns:
        raise ValueError('Not all logical attributes are shown')
    if {c['conname'] for c in evidence['constraints'] if c['contype'] == 'c'} != set(CHECKS):
        raise ValueError('CHECK description registry must be reconciled with the current catalog')
    if {t['name'] for t in evidence['tests']} != set(TEST_NAMES):
        raise ValueError('Test description registry must be reconciled')
    native = json.loads((Path(manifest['output']) / 'Отчёт_экспорта.json').read_text(encoding='utf-8'))
    if native['cancelled'] or len(native['items']) != len(audit['conceptualSheets']):
        raise ValueError('Incomplete Chen export')
    for item in native['items']:
        if item['status'] != 'completed' or len(item['files']) != 1:
            raise ValueError('Failed Chen sheet: ' + item['diagramId'])
        canonical = DIAGRAMS / (item['title'] + '.png')
        if digest(canonical) != digest(item['files'][0]):
            raise ValueError('Canonical Chen image differs from the native export')
    if len(evidence['columns']) != 71 or len(evidence['counts']) != 12:
        raise ValueError('Reconcile report scope with the database before building')
    for table, count in evidence['counts'].items():
        if len(evidence['fixtures'][table]) != count:
            raise ValueError('Fixture count mismatch: ' + table)
    return evidence, audit


def sql(r, text):
    """Editable SQL in TNR14, with no decorative boxes or rasterised listings."""
    lines = [line.rstrip() for line in text.strip().splitlines()]
    p = r.document.add_paragraph('\n'.join(lines), 'ReportReference')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_together = len(lines) <= 18
    for run in p.runs:
        font(run, 14)
    return p


def configure_toc_styles(document):
    # Word must recognise these as built-in TOC styles, not custom lookalikes.
    # Otherwise field refresh creates different native styles with default
    # formatting. An already materialised Word TOC uses lower-case names.
    for level, indent in [(1, 0), (2, .7)]:
        name = f'toc {level}'
        matching = [s for s in document.styles if s.name.casefold() == name]
        if len(matching) > 1:
            raise ValueError('Duplicate TOC styles: ' + name)
        style = matching[0] if matching else document.styles.add_style(name, 1, builtin=True)
        style.element.attrib.pop(qn('w:customStyle'), None)
        style.base_style = document.styles['ReportReference']
        style.font.name, style.font.size = 'Times New Roman', Pt(14)
        style.font.bold, style.font.italic = False, False
        style.font.color.rgb = document.styles['ReportReference'].font.color.rgb
        pf = style.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.left_indent, pf.right_indent, pf.first_line_indent = Cm(indent), Cm(0), Cm(0)
        pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(0), 1
        pf.tab_stops.clear_all()
        pf.tab_stops.add_tab_stop(Cm(16.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        for item in list(style.element.findall(qn('w:autoRedefine'))):
            style.element.remove(item)


def toc(r):
    style = r.document.styles.add_style('ReportTOCTitle', 1)
    style.base_style = r.document.styles['Heading 1']
    heading = r.document.add_paragraph('Содержание', 'ReportTOCTitle')
    outline = OxmlElement('w:outlineLvl')
    outline.set(qn('w:val'), '9')
    style.element.get_or_add_pPr().append(outline)
    configure_toc_styles(r.document)
    p = r.document.add_paragraph(style='ReportReference')
    field(p, 'TOC \\o "1-2" \\h \\z \\u')
    r.page(new_page=True)


def display(value):
    if value is None:
        return 'NULL'
    if isinstance(value, bool):
        return str(value).lower()
    if isinstance(value, str):
        uuid = re.fullmatch(r'00000000-0000-0000-0000-(\d{12})', value)
        if uuid:
            return 'U' + f'{int(uuid.group(1)):03d}'
        if re.match(r'^2026-\d\d-\d\d \d\d:', value):
            moment = datetime.fromisoformat(value)
            return moment.strftime('%d.%m.%Y\n%H:%M') + ' UTC'
    return str(value)


def physical_type(c):
    name = {'int4': 'integer', 'bool': 'boolean', 'bpchar': 'char'}.get(c['udt_name'], c['udt_name'])
    if name in ('varchar', 'char'):
        return f"{name}({c['character_maximum_length']})"
    if name == 'numeric':
        return f"numeric({c['numeric_precision']},{c['numeric_scale']})"
    return name


def build():
    evidence, audit = load_inputs()
    keys = evidence['key_catalog']
    checks = [c for c in evidence['constraints'] if c['contype'] == 'c']
    key_ids = {k['conname']: f'K{i:02d}' for i, k in enumerate(keys, 1)}
    check_ids = {c['conname']: f'C{i:02d}' for i, c in enumerate(checks, 1)}

    def columns(table):
        return sorted((c for c in evidence['columns'] if c['table_name'] == table), key=lambda c: c['ordinal_position'])

    def constraints(c, logical=False):
        table, name = c['table_name'], c['column_name']
        parts = []
        if logical:
            parts.append('Обязательный' if c['is_nullable'] == 'NO' else 'NULL допустим')
        for k in keys:
            if k['table_name'] == table and name in k['columns']:
                parts.append({'p': 'PK', 'u': 'UQ', 'f': 'FK'}[k['contype']] + ' ' + key_ids[k['conname']])
        for check in checks:
            if check['table_name'] == table and re.search(r'\b' + re.escape(name) + r'\b', check['definition']):
                parts.append(check_ids[check['conname']])
        if c['column_default'] is not None:
            parts.append('DEFAULT ' + c['column_default'])
        return '; '.join(parts) or 'Дополнительных нет'

    r = Report(7, 'Проектирование информационной модели ИС')
    toc(r)
    r.heading('Проектирование информационной модели ИС')
    r.heading('Цель и задание', 2)
    r.paragraph('Цель работы — разработать базу данных информационной системы формирования и актуализации карточек компьютерных систем ООО «Юкомс» по трёхэтапной методологии: концептуальное, логическое и физическое проектирование. Задание предусматривает ER-модель Чена со слабыми сущностями, логическую модель в третьей нормальной форме в нотации Crow’s Foot, описание полей физической модели, тестовые данные и выполненные SQL-запросы, включая не менее пяти DDL и пяти DCL [1].')
    r.paragraph('Разработаны 12 таблиц с 71 полем и 13 внешними ключами. Схема создана в отдельной учебной базе PostgreSQL, заполнена 48 тестовыми записями; выполнены шесть функциональных запросов и 22 проверки ограничений и прав. Центральная цепочка связывает конфигурацию, версию состава, выбранное предложение, ревизию карточки и попытку публикации.')
    r.paragraph('Модель продолжает предметную область практик 1–6. Версия состава содержит позиции с количеством и выбором группы либо точного компонента. Ревизия карточки фиксирует содержание и параметры расчёта; попытка публикации имеет отдельное состояние. Статус ACCEPTED означает принятие внешней задачи и не равен PUBLISHED. В этой работе обмен с внешней площадкой не выполняется.')
    r.heading('1 Концептуальная модель данных')
    r.heading('1.1 Сущности и границы модели', 2)
    r.table('Сущности информационной модели', ['Название сущности', 'Описание'],
            [[NAMES[t] + '\n' + t, DESCRIPTIONS[t]] for t in ORDER], [5, 11.5],
            lead='Назначение всех двенадцати сущностей приведено в таблице 1. Составные владельческие ключи позволяют сохранять историю без перезаписи прежних версий.')
    r.paragraph('Предложение рассматривается как отдельное наблюдение, а не как текущая цена, перезаписываемая в строке компонента. Новое наблюдение получает новый идентификатор. После использования состава в ревизии его позиции, выбранные наблюдения и характеристики компонентов защищены триггерами от обновления. Для изменения расчётных параметров создаётся новая ревизия.')
    r.heading('1.2 Нотация и зависимые сущности', 2)
    r.paragraph('Для концептуальной модели используются прямоугольники сущностей, овалы атрибутов и ромбы связей Чена [2]. Двойной прямоугольник обозначает зависимую сущность, двойной ромб — идентифицирующую связь. Сплошное подчёркивание обозначает первичный ключ сильной сущности, пунктирное — локальный дискриминатор слабой. Применено бинарное расширение min–max: 1, 0..1 и 0..N. Метка у конца связи показывает число объектов этого конца для одного объекта противоположного конца.')
    r.paragraph('Внешние ключи концептуально представлены связями и не дублируются как отдельные овалы. Все их поля и обязательность включены в логическую и физическую модели. Тринадцать листов Чена являются согласованными представлениями одной модели, а не отдельными базами. Атрибуты ревизии карточки разделены между двумя листами; повторённые владельцы и связи сохраняют тот же смысл.')
    r.table('Зависимые сущности и их идентификация', ['Сущность', 'Владелец и полный ключ', 'Локальная идентификация'], [
        ['bom_version', 'configuration; PK(configuration_id, version_no)', 'Номер версии внутри конфигурации.'],
        ['bom_slot', 'bom_version; PK(configuration_id, version_no, line_no)', 'Номер позиции внутри версии.'],
        ['card_revision', 'card; PK(card_id, revision_no)', 'Номер ревизии внутри карточки; BOM-ссылка не идентифицирует.'],
        ['publication_attempt', 'card_revision; PK(card_id, revision_no, attempt_no)', 'Номер попытки внутри ревизии.'],
        ['revision_image', 'card_revision и image; PK(card_id, revision_no, image_id)', 'Ключ — пара владельцев; ordinal образует альтернативный ключ.'],
    ], [4.7, 6.3, 5.5])
    r.page(True)
    r.heading('1.3 Диаграммы Чена', 2)
    conceptual_by_table = {t: [s for s in audit['conceptualSheets'] if s['entity'] == t] for t in ORDER}
    # Two compact strong-entity views share a page; crowded/weak views retain
    # a full page and are never shrunk to make room for another model.
    conceptual_pages = [
        ['component_group', 'supplier'], ['component'], ['offer'],
        ['configuration', 'card'], ['bom_version'], ['bom_slot'],
        ['card_revision'], ['image'], ['revision_image'], ['publication_attempt'],
    ]
    ordered_sheets = []
    page_index = 0
    for tables in conceptual_pages:
        sheets = [s for t in tables for s in conceptual_by_table[t]]
        sheet_pages = [[s] for s in sheets] if tables == ['card_revision'] else [sheets]
        for sheet_page in sheet_pages:
            if page_index:
                r.page(True, new_page=True)
            for sheet in sheet_page:
                ordered_sheets.append(sheet)
                title = 'Модель Чена для сущности «' + NAMES[sheet['entity']] + '»'
                if sheet['total'] > 1:
                    title += f" часть {sheet['part']} из {sheet['total']}"
                relation_note = ' и связи с владельцами' if sheet['relationships'] else ''
                r.figure(title, DIAGRAMS / sheet['file'], max_height=5.2 if len(sheet_page) == 2 else 13.2,
                         lead=f"На рисунке {r.figure_count+1} показаны атрибуты сущности «{NAMES[sheet['entity']]}»{relation_note}. "
                              + ('Идентифицирующие связи имеют двойной контур.' if sheet['entity'] in audit['weakEntities']
                                 else 'Сущность имеет самостоятельный первичный ключ.'))
            page_index += 1
    r.page(False)
    r.heading('2 Логическая и физическая модели')
    r.heading('2.1 Нормализация и преобразование связей', 2)
    r.paragraph('При переходе к отношениям каждой сущности соответствует таблица. Связь 1:N реализована внешним ключом на стороне N. Зависимые сущности включают ключ владельца в собственный составной PK. Отношение между ревизиями и изображениями разрешено таблицей revision_image. Уникальность её ordinal обеспечивается внутри одной ревизии, а не во всей базе.')
    r.paragraph('Обязательная ссылка даёт одного владельца, необязательная — 0..1; обратная мощность равна 0..N. Ссылки позиции на группу, точный компонент и предложение допускают NULL по отдельности. CHECK требует ровно один способ задания позиции: группу либо точный компонент.')
    r.paragraph('Первая нормальная форма обеспечивается атомарными полями: позиции состава и изображения вынесены в отдельные отношения, списки комплектующих не хранятся одной строкой. Во второй нормальной форме неключевые атрибуты зависимых сущностей определяются полным составным ключом: номер строки сам по себе не определяет позицию, а номер ревизии не определяет карточку. Имя конфигурации, поставщика и компонента хранится у владельца, а не повторяется в дочерних отношениях.')
    r.paragraph('Для третьей нормальной формы неключевые атрибуты не определяют другие неключевые атрибуты по принятым правилам предметной области. Например, offer.id определяет component_id, supplier_id, цену, остаток и время, но название поставщика вынесено в supplier. Уникальные code, sku и seller_code являются альтернативными ключами, поэтому зависимости от них не нарушают 3НФ. Итоговая цена и масса не хранятся рядом с исходными значениями, а вычисляются представлением. Экономические параметры ревизии относятся к её собственному снимку, а не к изменяемому справочнику политики.')
    r.table('Функциональные зависимости отношений', ['Отношение', 'Ключ и определяемые атрибуты'], [
        [t, '(' + ', '.join(next(k['columns'] for k in keys if k['table_name'] == t and k['contype'] == 'p'))
         + ') → ' + ', '.join(c['column_name'] for c in columns(t)
                              if c['column_name'] not in next(k['columns'] for k in keys if k['table_name'] == t and k['contype'] == 'p'))]
        for t in ORDER], [4.4, 12.1])
    r.paragraph('Перечень зависимостей основан на указанных бизнес-правилах; равенство значений в нескольких тестовых строках не считается новой функциональной зависимостью. Статусы ревизии и попытки публикации разделены: изменение внешнего результата не переписывает содержание карточки.')
    r.heading('2.2 Логические диаграммы Crow’s Foot', 2)
    r.paragraph('В Crow’s Foot кружок означает необязательность, черта — единичность, развилка — множество. PK обозначает первичный ключ, FK — внешний, UQ — уникальное ограничение; несколько меток у полей одного ключа не означают отдельную уникальность каждого поля. На повторённых справочных блоках показаны только ключи; полный набор атрибутов дан на основном листе и в таблицах.')
    for i, (group, title) in enumerate([('catalog', 'каталога и предложений'), ('bom', 'конфигурации и состава'),
                                      ('cards', 'карточки и расчёта'), ('publication', 'изображений и публикации')]):
        r.page(True, new_page=True)
        r.figure('Логическая модель ' + title, DIAGRAMS / f'ПР7_Логическая_{group}.png', max_height=13.6,
                 lead=f'На рисунке {r.figure_count+1} показана логическая модель {title}.')
    r.page(False)
    r.heading('2.3 Ключи и ограничения', 2)
    r.paragraph('Обозначения K01–K32 ниже являются локальными ссылками отчёта на ключевые ограничения каталога PostgreSQL. PK — первичный, UQ — уникальный ключ, FK — ссылка. Ограничение на комбинацию полей применяется совместно: метка возле одного поля означает участие в ключе, а не самостоятельную уникальность. NO ACTION используется для изменения и удаления ссылочных ключей [3].')
    key_rows = []
    for k in keys:
        text = {'p': 'PK', 'u': 'UQ', 'f': 'FK'}[k['contype']] + '(' + ', '.join(k['columns']) + ')'
        if k['contype'] == 'f':
            text += ' → ' + k['parent_table'] + '(' + ', '.join(k['parent_columns']) + ')'
        key_rows.append([key_ids[k['conname']], k['table_name'], text])
    r.table('Первичные уникальные и внешние ключи', ['Код', 'Сущность', 'Ограничение'], key_rows, [1.2, 4.4, 10.9])
    r.table('Ограничения допустимых значений', ['Код', 'Сущность', 'Правило CHECK'],
            [[check_ids[c['conname']], c['table_name'], CHECKS[c['conname']]] for c in checks], [1.2, 4.4, 10.9])
    r.heading('2.4 Атрибуты логической модели', 2)
    r.paragraph('В таблице ниже перечислены все 71 атрибут и ограничения. Коды K и C раскрыты в предыдущих таблицах; обязательность и допустимость NULL указаны явно. Составные ссылки сохранены без разбиения на независимые ограничения отдельных полей.')
    r.table('Атрибуты логической модели', ['Название сущности', 'Название атрибута', 'Ограничения'],
            [[c['table_name'], c['column_name'], constraints(c, True)] for t in ORDER for c in columns(t)], [4.4, 5.5, 6.6])
    r.heading('2.5 Типы полей физической модели', 2)
    r.paragraph('Физическая модель реализована в схеме ppois PostgreSQL 18.6. UUID идентифицирует самостоятельные объекты, integer хранит номера и количества, NUMERIC обеспечивает десятичные денежные расчёты, timestamptz — моменты времени. NULL технической характеристики означает неприменимость либо отсутствие значения, а не ноль. Типы и размеры полей подтверждены каталогом созданной базы.')
    r.page(True)
    r.table('Поля физической модели', ['Наименование сущности', 'Название поля', 'Тип данных', 'Обязательно', 'Ограничения'],
            [[c['table_name'], c['column_name'], physical_type(c), 'Да' if c['is_nullable'] == 'NO' else 'Нет', constraints(c)]
             for t in ORDER for c in columns(t)], [5, 6.2, 3.1, 3.1, 7.8], center=(3,))
    for i, (group, title) in enumerate([('catalog', 'каталога и предложений'), ('bom', 'конфигурации и состава'),
                                      ('cards', 'карточки и расчёта'), ('publication', 'изображений и публикации')]):
        r.page(True, new_page=True)
        if i == 0:
            r.heading('2.6 Физические диаграммы', 2)
        r.figure('Физическая модель ' + title, DIAGRAMS / f'ПР7_Физическая_{group}.png', max_height=13.6,
                 lead=f'На рисунке {r.figure_count+1} показана физическая модель {title}.')
    r.page(False)
    r.heading('2.7 Межтабличные инварианты и границы SQL', 2)
    r.paragraph('В CHECK проверяется только текущая строка. Соответствие выбранного предложения группе или точному компоненту проверяется триггером guard_bom_slot. Этот же триггер запрещает вставку, удаление и изменение позиций версии, на которую уже ссылается ревизия, включая перенос позиции в другую версию. guard_revision_snapshot разрешает менять только state ревизии; guard_used_catalog запрещает обновление предложения и компонента, используемых ревизией.')
    r.paragraph('Триггеры не заменяют проверку полной совместимости и готовности прикладным сервисом. DRAFT/READY/STALE в таблице являются хранимыми состояниями; представление card_status не доказывает полноту состава или наличие всех изображений. Роль manager может менять state, но не экономические поля. Перечень разрешённых переходов, конкурентная блокировка и повторная проверка перед внешней отправкой относятся к дальнейшей реализации сервисного слоя.')
    r.paragraph('Пять зависимых сущностей имеют составные ключи; все тринадцать FK защищают существование владельцев. Прикладным ролям не выданы DELETE, TRUNCATE и CREATE в схеме. Роли в учебном кластере созданы с NOLOGIN: вход выполняется выделенной лабораторной учётной записью, а проверки прав — с SET LOCAL ROLE внутри транзакции. Это проверка прав PostgreSQL, не готовая пользовательская аутентификация веб-приложения.')
    r.heading('3 Реализация и проверка базы данных')
    r.heading('3.1 Среда и тестовые данные', 2)
    r.paragraph(f"Скрипт verify_database.py создал отдельную базу {evidence['database']} в локальном учебном кластере на 127.0.0.1:55432. Сервер: {evidence['server_version']}. Выполнение зафиксировано {datetime.fromisoformat(evidence['executed_at']).strftime('%d.%m.%Y %H:%M UTC')}. Рабочие базы и существующие данные не сбрасывались; DROP в учебных SQL-скриптах отсутствует.")
    r.table('Заполнение таблиц тестовыми данными', ['Таблица', 'Число строк'],
            [[t, evidence['counts'][t]] for t in ORDER] + [['Всего', sum(evidence['counts'].values())]], [12.5, 4], center=(1,))
    r.paragraph('Данные созданы для проверки, а не получены из прайс-листов. Они включают два допустимых предложения CPU, несовместимый процессор, устаревшее наблюдение и предложение без остатка. Две карточки имеют одинаковый состав, но разные состояния READY и STALE; у второй отсутствует изображение. Внешний ID demo-task-001 является тестовой строкой, а не результатом обращения к Ozon.')
    r.paragraph('Для читаемости UUID в следующих таблицах сокращены до Uxxx. Например, U011 означает 00000000-0000-0000-0000-000000000011; числовой суффикс дополняется слева нулями до двенадцати цифр. В исполняемом SQL используются полные UUID. NULL обозначает отсутствие значения, время показано в UTC. Полный снимок строк, SQL и ответы сохранены в База_данных/Результаты_проверки.json.')
    r.page(True)
    r.heading('3.2 Содержимое тестовой базы', 2)
    fixture_order = ['component_group', 'supplier', 'component', 'configuration', 'card',
                     'offer', 'bom_version', 'bom_slot', 'card_revision', 'image',
                     'revision_image', 'publication_attempt']
    fixture_page_starts = {'component', 'offer', 'bom_version', 'card_revision', 'image', 'publication_attempt'}
    for t in fixture_order:
        records = evidence['fixtures'][t]
        fields = [c['column_name'] for c in columns(t)]
        if t in ('card_revision', 'publication_attempt', 'image'):
            headers = ['Поле'] + [f'Запись {i+1}' for i in range(len(records))]
            rows = [[name] + [display(record[name]) for record in records] for name in fields]
            widths = [6.5] + [18.7 / len(records)] * len(records)
        else:
            headers = fields
            rows = [[display(record[name]) for name in fields] for record in records]
            widths = [1.5 if name in ('id', 'group_id', 'component_id', 'supplier_id', 'offer_id', 'quantity', 'stock', 'version_no', 'line_no', 'revision_no', 'ordinal')
                      else 3.8 if name in ('name', 'observed_at', 'created_at', 'configuration_id', 'requested_component_id', 'seller_code') else 2.5
                      for name in fields]
        r.table('Тестовые записи ' + t, headers, rows, widths,
                keep_rows=True,
                lead=f"Все поля тестовых записей {t} приведены в таблице {r.table_count+1}.")
        if t in fixture_page_starts:
            # A break-only paragraph can overflow after a full-page table and
            # create a blank page. Attach the break to the actual lead instead.
            r.document.paragraphs[-2].paragraph_format.page_break_before = True
    r.page(False)
    r.heading('3.3 Выполненные запросы DDL', 2)
    r.paragraph('schema.sql содержит создание схемы, двенадцати таблиц, трёх операционных индексов, представления, трёх функций и четырёх триггеров. Ниже показаны девять реально выполненных операторов разных назначений. Они являются выдержками того же SQL, по которому получен каталог; повторно запускать их в уже созданной схеме нельзя. Полный файл включает остальные таблицы и проверки снимков.')
    source = (DATABASE / 'schema.sql').read_text(encoding='utf-8')
    patterns = [
        ('DDL1 Создание пространства имён', r'CREATE SCHEMA ppois;'),
        ('DDL2 Справочник групп', r'CREATE TABLE ppois\.component_group \([\s\S]*?\n\);'),
        ('DDL3 Наблюдение предложения', r'CREATE TABLE ppois\.offer \([\s\S]*?\n\);'),
        ('DDL4 Версия состава', r'CREATE TABLE ppois\.bom_version \([\s\S]*?\n\);'),
        ('DDL5 Позиция состава и составной внешний ключ', r'CREATE TABLE ppois\.bom_slot \([\s\S]*?\n\);'),
        ('DDL6 Индекс для поиска свежих предложений', r'CREATE INDEX offer_freshness_idx[^;]+;'),
        ('DDL7 Представление состояния карточки', r'CREATE VIEW ppois\.card_status AS[\s\S]*?\) p ON true;'),
        ('DDL8 Функция защиты снимка ревизии', r'CREATE FUNCTION ppois\.guard_revision_snapshot\(\)[\s\S]*?END \$\$;'),
        ('DDL9 Триггер защиты снимка', r'CREATE TRIGGER revision_snapshot_guard[\s\S]*?guard_revision_snapshot\(\);'),
    ]
    for title, pattern in patterns:
        match = re.search(pattern, source)
        if match is None:
            raise ValueError('DDL excerpt missing: ' + title)
        r.paragraph(title, keep=True)
        sql(r, match.group())
    r.paragraph('Создание объектов завершилось успешно; реальные имена полей, типы и ограничения получены из information_schema и pg_catalog. Первичные и уникальные ограничения создают собственные индексы; дополнительно заданы offer_freshness_idx, revision_bom_idx и частичный attempt_reconcile_idx. Последний ускоряет выбор состояний ACCEPTED, UNKNOWN и TEMP_ERROR, но сам по себе не разрешает повторную отправку.')
    r.heading('3.4 Выполненные запросы DCL', 2)
    r.paragraph('После создания четырёх NOLOGIN-ролей выполнены двенадцать операторов GRANT/REVOKE из privileges.sql. USAGE разрешает обращаться к именам в схеме, SELECT — читать данные, а INSERT/UPDATE выданы по рабочим обязанностям. Право UPDATE отдельных столбцов не даёт права изменять остальные поля [4].')
    dcl = (DATABASE / 'privileges.sql').read_text(encoding='utf-8')
    blocks = list(re.finditer(r'-- (DCL\d{2}): ([^\n]+)\n([\s\S]*?;)', dcl))
    if len(blocks) != 12:
        raise ValueError('Expected twelve executed DCL statements')
    dcl_titles = ['Закрыть неявный доступ PUBLIC', 'Разрешить доступ к именам схемы',
                  'Разрешить чтение viewer', 'Разрешить чтение engineer',
                  'Разрешить работу engineer с составом', 'Разрешить чтение manager',
                  'Разрешить создание карточек и запросов manager', 'Ограничить UPDATE manager состоянием ревизии',
                  'Разрешить чтение worker', 'Ограничить UPDATE worker полями результата',
                  'Не выдавать удаление и очистку таблиц', 'Не выдавать создание объектов схемы']
    for match, title in zip(blocks, dcl_titles):
        r.paragraph(match.group(1) + ' ' + title, keep=True)
        sql(r, match.group(3))
    r.paragraph('Проверки выполняются от прикладных ролей, а не от владельца схемы или суперпользователя. Права относятся к текущим объектам; новые объекты миграций требуют выдачи прав либо ALTER DEFAULT PRIVILEGES. GRANT PostgreSQL не управляет доступом к внешнему API.')
    r.page(True)
    r.heading('3.5 Функциональные SQL-запросы и результаты', 2)
    explanations = [
        'Запрос получает предложения для каждой позиции первой версии, проверяя соответствие группе или точному компоненту, остаток и интервал свежести 24 часа. Верхняя граница исключает наблюдения из будущего. Результат содержит семь строк: два предложения процессора и по одному для остальных позиций. Это множество кандидатов, а не завершённый глобальный подбор совместимого состава.',
        'Агрегируются именно выбранные offer_id. Обе версии содержат шесть позиций и дают 51 000 руб. стоимости комплектующих и 7,020 кг массы. INNER JOIN не включает неразрешённые позиции; такой результат нельзя использовать как доказательство полноты любого произвольного состава.',
        'Представление показывает независимые состояния ревизии и публикации. Для обеих карточек цена равна 76 400 руб., масса — 7,820 кг. READY первой карточки сочетается с ACCEPTED, а не с PUBLISHED; STALE второй не устраняется наличием расчёта.',
        'Обратный поиск по идентификатору CPU U011 возвращает обе карточки, использующие этот компонент. Запрос выявляет область влияния изменения; он не переписывает версии и не меняет их состояние.',
        'Для ACCEPTED требуется проверка результата, для UNKNOWN — сверка неопределённого исхода. TEMP_ERROR входит в выборку только при confirmed_not_accepted = true. Единственная строка тестовой базы относится к принятой задаче; повторная отправка для ACCEPTED не разрешается этим SELECT.',
        'LEFT JOIN и HAVING выявляют ревизию без связанного изображения: YK-PC-102, ревизия 1. Отсутствие изображения является причиной блокировать подготовку, но SELECT сам по себе не меняет состояние строки.',
    ]
    query_widths = [[2.4, 3.3, 1.9, 2.4, 3.6, 1.8, 5], [4.2, 2.4, 5, 5, 2.2],
                    [3.4, 2, 3.1, 3, 2.7, 4.5, 5], [4, 3, 5], [3.6, 2, 2, 3.4, 4, 4.6, 5], [8, 4]]
    for i, (query, title, explanation, widths) in enumerate(zip(evidence['queries'], QUERY_NAMES, explanations, query_widths)):
        if i:
            r.page(True, new_page=True)
        r.paragraph(title, keep=True)
        sql(r, '\n'.join(line for line in query['sql'].splitlines() if not line.startswith('--')))
        r.paragraph(explanation)
        if i == 0:
            r.page(True, new_page=True)
        fields = list(query['rows'][0])
        headers = [name.replace('confirmed_not_accepted', 'confirmed_\nnot_accepted') for name in fields]
        r.table('Результат ' + title.split()[0], headers,
                [[display(row[name]) for name in fields] for row in query['rows']], widths, keep_rows=True,
                lead=f"Фактический ответ PostgreSQL приведён в таблице {r.table_count+1}."
                     + (' Порядок Q02 без ORDER BY не гарантируется SQL.' if i == 1 else ''))
    r.page(False)
    r.heading('3.6 Проверки целостности и прав', 2)
    r.paragraph('Восемнадцать отрицательных сценариев должны завершаться ожидаемой ошибкой PostgreSQL: 23505 — уникальность, 23503 — внешний ключ, 23514 — CHECK или отказ триггера, 42501 — недостаточные права. Три положительных сценария проверяют разрешённые действия; дополнительный сценарий Q01 проверяет исключение будущей даты. Отрицательные сценарии и будущая запись выполняются с откатом. Разрешённые обновления повторяют существующие значения, поэтому итоговый набор из 48 записей сохраняется.')
    r.table('Фактические результаты проверок', ['№', 'Проверка', 'Ответ', 'Результат'],
            [[i, TEST_NAMES[t['name']], t.get('sqlstate', 'Выполнено без ошибки'), 'Ожидание подтверждено']
             for i, t in enumerate(evidence['tests'], 1)], [1, 8, 3.5, 4], center=(0,))
    r.paragraph('Все 22 проверки завершились ожидаемым результатом. Они подтверждают конкретные ограничения и операции на учебных данных; нагрузка, восстановление после сбоя, конкурентные транзакции и внешняя публикация этими сценариями не испытаны. Источник SQL привязан к журналу SHA-256, поэтому изменение схемы или запросов требует повторного выполнения проверки и пересборки моделей.')
    r.heading('Вывод', new_page=True)
    r.paragraph('Разработаны согласованные концептуальная модель Чена, логическая модель Crow’s Foot в 3НФ и физическая модель PostgreSQL. Выделены пять зависимых сущностей; описаны все 12 отношений, 71 атрибут, 32 ключевых ограничения и 28 CHECK. Создана тестовая база из 48 записей, выполнены девять приведённых DDL, двенадцать DCL и шесть функциональных SELECT. Все 22 проверки ограничений и прав получили ожидаемый результат.')
    r.paragraph('Версионирование сохраняет связь между составом, расчётом и содержанием ревизии. Разделение состояния карточки и попытки публикации исключает ошибочное отождествление принятия внешней задачи с успешной публикацией. Следующая практика использует эту схему для проектирования форм ввода и вывода и расчётных алгоритмов.')
    r.sources([
        'Практическая работа №7. Проектирование информационной модели ИС. Методические материалы ППОИС, предоставленные к заданию, 2 с.',
        'Song I. Y., Evans M., Park E. K. A Comparative Analysis of Entity-Relationship Diagrams. Journal of Computer and Software Engineering. 1995. Vol. 3, No. 4. P. 427–459. https://cci.drexel.edu/faculty/song/publications/p_Jcse-erd.PDF (дата обращения: 05.10.2026).',
        'PostgreSQL Global Development Group. PostgreSQL 18 Documentation. Constraints. https://www.postgresql.org/docs/18/ddl-constraints.html (дата обращения: 05.10.2026).',
        'PostgreSQL Global Development Group. PostgreSQL 18 Documentation. GRANT. https://www.postgresql.org/docs/18/sql-grant.html (дата обращения: 05.10.2026).',
        'Материалы выполнения практики: База_данных/schema.sql, privileges.sql, queries.sql, Результаты_проверки.json; Исходники/seed_database.py, verify_database.py, build_database_models.mts; Схемы/ПР7_Модель_данных.omni.',
    ], new_page=True)
    path = r.save()
    coverage = {
        'database': evidence['database'], 'source_sha256': evidence['source_sha256'],
        'model_sha256': digest(DIAGRAMS / 'ПР7_Модель_данных.omni'),
        'chen_sheets': len(ordered_sheets), 'logical_sheets': 4, 'physical_sheets': 4,
        'logical_attributes': len(evidence['columns']), 'physical_fields': len(evidence['columns']),
        'fixture_records': sum(evidence['counts'].values()), 'fixture_tables': ORDER,
        'ddl_excerpts': len(patterns), 'dcl_statements': len(blocks),
        'functional_queries': len(evidence['queries']), 'database_checks': len(evidence['tests']),
        'status': 'authored; final Word export and all-page visual review still required',
    }
    (ROOT / '.cache/ppois5-10/report7-content-coverage.json').write_text(
        json.dumps(coverage, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return path


if __name__ == '__main__':
    build()
