"""Regression checks that the report verifier rejects real formatting defects.

Creates disposable fixtures under .cache, never modifies canonical reports.
The caller's visual-review record is kept separate from automatic validation.
"""
from contextlib import redirect_stdout
from io import StringIO
import json
from pathlib import Path
import shutil
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from verify_report import ROOT, verify


REPORT = ROOT/'4/ППОИС/Ансамблевое_обучение/Ансамблевое_обучение_АлбахтинИВ.docx'
RENDER = ROOT/'.cache/ppois5-10/verified-ensemble'


class ReportVerifierTests(unittest.TestCase):
    def validate_fixture(self, mutate=None, pages=None):
        location = ROOT/'.cache/ppois5-10/verifier-tests'
        location.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(dir=location) as directory:
            fixture = Path(directory)/'fixture.docx'
            document = Document(REPORT)
            if mutate:
                mutate(document)
            document.save(fixture)
            shutil.copyfile(REPORT.with_suffix('.pdf'), fixture.with_suffix('.pdf'))
            output = Path(directory)/'result.json'
            try:
                with redirect_stdout(StringIO()):
                    verify(fixture, RENDER, pages if pages is not None else list(range(1,12)), output)
            except SystemExit as error:
                self.assertEqual(error.code, 1)
            return json.loads(output.read_text(encoding='utf-8'))

    def assert_rejected(self, mutation, reason):
        result = self.validate_fixture(mutation)
        self.assertEqual(result['status'], 'failed')
        self.assertTrue(any(reason in label for label in result['failed_checks']), result['failed_checks'])

    def test_valid_report(self):
        self.assertEqual(self.validate_fixture()['status'], 'passed')

    def test_review_cannot_be_inferred_from_formatting(self):
        result = self.validate_fixture(pages=[1,2])
        self.assertEqual(result['status'], 'failed')
        self.assertIn('Every final page explicitly reviewed', result['failed_checks'])

    def test_body_font_size(self):
        def change(document):
            next(p for p in document.paragraphs if p.style.name=='ReportBody').runs[0].font.size = Pt(10)
        self.assert_rejected(change, 'correct text font and size')

    def test_body_indent(self):
        def change(document):
            next(p for p in document.paragraphs if p.style.name=='ReportBody').paragraph_format.first_line_indent = Cm(0)
        self.assert_rejected(change, 'body first line 1.25cm')

    def test_cell_shading(self):
        def change(document):
            item = OxmlElement('w:shd')
            item.set(qn('w:val'), 'clear')
            item.set(qn('w:fill'), 'DDEEFF')
            document.tables[2].cell(0,0)._tc.get_or_add_tcPr().append(item)
        self.assert_rejected(change, 'no shading')

    def test_header_repeat(self):
        def change(document):
            properties = document.tables[2].rows[0]._tr.get_or_add_trPr()
            properties.remove(properties.find(qn('w:tblHeader')))
        self.assert_rejected(change, 'repeating header')

    def test_row_split(self):
        def change(document):
            properties = document.tables[2].rows[1]._tr.get_or_add_trPr()
            properties.remove(properties.find(qn('w:cantSplit')))
        self.assert_rejected(change, 'not split')

    def test_heading_style_is_builtin(self):
        def change(document):
            document.styles['Heading 2'].element.set(qn('w:customStyle'), '1')
        self.assert_rejected(change, 'built-in Word heading style')

    def test_footer_indent(self):
        def change(document):
            document.sections[1].footer.paragraphs[0].paragraph_format.left_indent = Cm(1)
        self.assert_rejected(change, 'centered unindented footer')

    def test_caption_number_and_pdf_match(self):
        def change(document):
            caption = next(p for p in document.paragraphs if p.style.name=='ReportFigureCaption')
            for run in caption.runs:
                if run.text == '1':
                    run.text = '99'
        self.assert_rejected(change, 'Figure sequential cached numbers')


if __name__=='__main__':
    unittest.main(verbosity=2)
