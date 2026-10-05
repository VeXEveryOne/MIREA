"""Shared PPOIS report layout, retaining only the ISURO cover page.

The cover source is never saved. All new content follows the repository's
report standard; no formatting from the ISURO main body is reused.
"""
from __future__ import annotations

from copy import deepcopy
from hashlib import sha256
from pathlib import Path
import json
import os

from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from revise_reports import prune_unused_images

ROOT = Path(__file__).resolve().parents[3]
SUBJECT = ROOT / '4/ППОИС'
REFERENCE = ROOT / '4/ИСУРО/ИСУРО_1_АлбахтинИВ.docx'
REFERENCE_HASH = '53a4b8ccc73cf44599e8b7c7fb75045afec85dd7f4b0a079b19a0d4ccb5249b1'
DISCIPLINE = 'Проектирование предметно-ориентированных информационных систем'


def font(run, size=14, bold=None, italic=None):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + key), 'Times New Roman')


def field(paragraph, instruction, cached=''):
    for kind in ('begin', 'instruction', 'separate', 'cached', 'end'):
        run = paragraph.add_run()
        font(run)
        if kind == 'instruction':
            element = OxmlElement('w:instrText')
            element.set(qn('xml:space'), 'preserve')
            element.text = ' ' + instruction + ' '
        elif kind == 'cached':
            run.text = str(cached)
            continue
        else:
            element = OxmlElement('w:fldChar')
            element.set(qn('w:fldCharType'), kind)
        run._r.append(element)


def set_body_section(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Cm(29.7 if landscape else 21)
    section.page_height = Cm(21 if landscape else 29.7)
    section.left_margin, section.right_margin = Cm(3), Cm(1.5)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    section.header_distance, section.footer_distance = Cm(.8), Cm(1)
    section.different_first_page_header_footer = False
    for numbering in list(section._sectPr.findall(qn('w:pgNumType'))):
        section._sectPr.remove(numbering)
    for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
        part.is_linked_to_previous = False
        for paragraph in part.paragraphs:
            paragraph.clear()
    paragraph = section.footer.paragraphs[0]
    paragraph.style = 'ReportFooter'
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.left_indent = 0
    paragraph.paragraph_format.right_indent = 0
    paragraph.paragraph_format.first_line_indent = 0
    field(paragraph, 'PAGE', '2')
    for run in paragraph.runs:
        font(run, 12)


def configure_styles(document):
    for name in ('ReportBody', 'ReportReference', 'ReportTable', 'ReportTableCaption',
                 'ReportFigure', 'ReportFigureCaption', 'ReportFooter', 'Heading 1', 'Heading 2'):
        style = document.styles[name] if name in document.styles else document.styles.add_style(
            name, 1, builtin=name.startswith('Heading'))
        style.base_style = document.styles['Normal']
        style.font.name, style.font.size = 'Times New Roman', Pt(14)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold, style.font.italic = False, False
        pf = style.paragraph_format
        pf.left_indent, pf.right_indent = Cm(0), Cm(0)
        pf.first_line_indent = Cm(1.25)
        pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(0), 1.5
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.widow_control, pf.keep_with_next, pf.keep_together = True, False, False
        pf.page_break_before = False
        if name.startswith('Heading'):
            style.font.bold = True
            pf.first_line_indent = Cm(0)
            pf.space_before, pf.space_after = Pt(12), Pt(6)
            pf.keep_with_next, pf.keep_together = True, True
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if name == 'Heading 1' else WD_ALIGN_PARAGRAPH.LEFT
            level = OxmlElement('w:outlineLvl')
            level.set(qn('w:val'), str(int(name[-1]) - 1))
            ppr = style.element.get_or_add_pPr()
            for prior in list(ppr.findall(qn('w:outlineLvl'))):
                ppr.remove(prior)
            ppr.append(level)
        elif name != 'ReportBody':
            pf.first_line_indent, pf.line_spacing = Cm(0), 1
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if name in ('ReportTable', 'ReportTableCaption', 'ReportFooter'):
                style.font.size = Pt(12)
            if name == 'ReportTableCaption':
                style.font.italic = True
                pf.space_before, pf.space_after = Pt(6), Pt(3)
                pf.keep_with_next, pf.keep_together = True, True
            if name in ('ReportFigure', 'ReportFigureCaption', 'ReportFooter'):
                pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if name == 'ReportFigure':
                pf.keep_with_next, pf.keep_together = True, True
            if name == 'ReportFigureCaption':
                pf.space_before, pf.space_after = Pt(3), Pt(6)
                pf.keep_together = True
            if name == 'ReportReference':
                pf.line_spacing = 1.5


class Report:
    def __init__(self, number, topic, filename=None):
        if sha256(REFERENCE.read_bytes()).hexdigest() != REFERENCE_HASH:
            raise ValueError('Cover reference changed; distillation must be repeated')
        self.number, self.topic = number, topic
        self.filename = filename or f'ППОИС_{number}_АлбахтинИВ.docx'
        self.document = Document(REFERENCE)
        body = self.document.element.body
        if len(body) < 17 or not body[15].xpath('./w:pPr/w:sectPr'):
            raise ValueError('Unexpected cover structure')
        # Keep the source page boundary and its original geometry. Its copy is
        # the final section, which will receive the new body's dimensions.
        final_section = deepcopy(body[15].xpath('./w:pPr/w:sectPr')[0])
        for child in list(body)[16:]:
            body.remove(child)
        body.append(final_section)
        title_text = f'ОТЧЕТ ПО ПРАКТИЧЕСКОЙ РАБОТЕ №{number}'
        for element in body[:16]:
            # Names in a DOCX may be split across runs, so replacing each w:t
            # independently is insufficient for the teacher slot.
            for paragraph in element.xpath('.//w:p'):
                texts = paragraph.xpath('.//w:t')
                value = ''.join(t.text or '' for t in texts)
                if 'Шендяпин А.В.' in value:
                    texts[0].text = value.replace('Шендяпин А.В.', 'Борзых Н.Ю.')
                    for text in texts[1:]:
                        text.text = ''
            for text in element.xpath('.//w:t'):
                if text.text:
                    text.text = text.text.replace('ИНБО-01-17', '')
                    text.text = text.text.replace('ОТЧЕТ ПО ПРАКТИЧЕСКИМ РАБОТАМ', title_text)
                    text.text = text.text.replace('Информационные системы управления ресурсами организации', DISCIPLINE)
                    text.text = text.text.replace('Шендяпин А.В.', 'Борзых Н.Ю.')
            for fonts in element.xpath('.//w:rFonts'):
                for key, value in list(fonts.attrib.items()):
                    if value == 'Liberation Serif':
                        fonts.set(key, 'Times New Roman')
        configure_styles(self.document)
        # The source title has no page number; the separate body footer is
        # unlinked, with continuous PAGE fields starting at physical page 2.
        cover = self.document.sections[0]
        cover.different_first_page_header_footer = True
        cover.first_page_footer.is_linked_to_previous = False
        for p in cover.first_page_footer.paragraphs:
            p.clear()
        set_body_section(self.document.sections[-1])
        self.landscape = False
        self.table_count = self.figure_count = 0
        self.inventory = []
        self.document.core_properties.author = 'Албахтин И.В.'
        self.document.core_properties.title = f'ППОИС Практическая работа {number} {topic}'
        self.document.core_properties.subject = DISCIPLINE
        settings = self.document.settings.element
        for prior in list(settings.findall(qn('w:updateFields'))):
            settings.remove(prior)
        update = OxmlElement('w:updateFields')
        update.set(qn('w:val'), 'true')
        settings.append(update)

    def page(self, landscape=False, new_page=False):
        if landscape != self.landscape:
            set_body_section(self.document.add_section(WD_SECTION_START.NEW_PAGE), landscape)
            self.landscape = landscape
        elif new_page:
            self.document.add_page_break()

    def heading(self, text, level=1, new_page=False):
        if level not in (1, 2):
            raise ValueError('Supported heading levels: 1 and 2')
        paragraph = self.document.add_paragraph(text, f'Heading {level}')
        if new_page:
            paragraph.paragraph_format.page_break_before = True
        return paragraph

    def paragraph(self, text, keep=False):
        p = self.document.add_paragraph(str(text), 'ReportBody')
        p.paragraph_format.keep_with_next = keep
        return p

    def table(self, title, headers, rows, widths=None, center=(), lead=None, keep_rows=False):
        self.table_count += 1
        number = self.table_count
        self.paragraph(lead or f'Сведения приведены в таблице {number}.', keep=True)
        caption = self.document.add_paragraph(style='ReportTableCaption')
        caption.add_run('Таблица ')
        field(caption, 'SEQ Table \\* ARABIC', number)
        caption.add_run(' – ' + title)
        for run in caption.runs:
            font(run, 12, False, True)
        table = self.document.add_table(rows=1, cols=len(headers))
        table.alignment, table.autofit = WD_TABLE_ALIGNMENT.CENTER, False
        table.style = 'Table Grid'
        total = 25.2 if self.landscape else 16.5
        ratios = widths or [1] * len(headers)
        if len(ratios) != len(headers) or any(w <= 0 for w in ratios):
            raise ValueError('Invalid table column widths')
        sizes = [total * w / sum(ratios) for w in ratios]
        for column, size in zip(table.columns, sizes):
            column.width = Cm(size)
        all_rows = [list(headers)] + [list(row) for row in rows]
        if any(len(row) != len(headers) for row in all_rows):
            raise ValueError('Table row length mismatch')
        for ri, source in enumerate(all_rows):
            row = table.rows[0] if ri == 0 else table.add_row()
            properties = row._tr.get_or_add_trPr()
            properties.append(OxmlElement('w:cantSplit'))
            if ri == 0:
                properties.append(OxmlElement('w:tblHeader'))
            for ci, (cell, value, size) in enumerate(zip(row.cells, source, sizes)):
                cell.width = Cm(size)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cell.text = str(value)
                props = cell._tc.get_or_add_tcPr()
                for prior in list(props.findall(qn('w:shd'))):
                    props.remove(prior)
                shading = OxmlElement('w:shd')
                shading.set(qn('w:val'), 'nil')
                shading.set(qn('w:fill'), 'auto')
                props.append(shading)
                cell_borders = OxmlElement('w:tcBorders')
                for side in ('top', 'left', 'bottom', 'right'):
                    item = OxmlElement('w:' + side)
                    for key, border_value in {'val': 'single', 'sz': '4', 'color': '808080'}.items():
                        item.set(qn('w:' + key), border_value)
                    cell_borders.append(item)
                props.append(cell_borders)
                margins = OxmlElement('w:tcMar')
                for side in ('top', 'bottom', 'left', 'right'):
                    item = OxmlElement('w:' + side)
                    item.set(qn('w:w'), '75' if side in ('top', 'bottom') else '85')
                    item.set(qn('w:type'), 'dxa')
                    margins.append(item)
                props.append(margins)
                for paragraph in cell.paragraphs:
                    paragraph.style = 'ReportTable'
                    paragraph.paragraph_format.keep_with_next = ri == 0 or (keep_rows and ri < len(all_rows)-1)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if ci in center else WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        font(run, 12, ri == 0, False)
        borders = OxmlElement('w:tblBorders')
        for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            item = OxmlElement('w:' + side)
            for key, value in {'val': 'single', 'sz': '4', 'color': '808080'}.items():
                item.set(qn('w:' + key), value)
            borders.append(item)
        table._tbl.tblPr.append(borders)
        self.inventory.append({'type': 'table', 'number': number, 'title': title, 'rows': len(rows)})
        return table

    def figure(self, title, path, max_height=18, width=None, lead=None):
        self.figure_count += 1
        number = self.figure_count
        self.paragraph(lead or f'На рисунке {number} представлен {title[0].lower() + title[1:]}.', keep=True)
        image = Image.open(path)
        iw, ih = image.size
        maximum = width or (25.2 if self.landscape else 16.5)
        physical_width = min(maximum, max_height * iw / ih)
        p = self.document.add_paragraph(style='ReportFigure')
        shape = p.add_run().add_picture(str(path), width=Cm(physical_width))
        shape._inline.docPr.set('descr', title)
        caption = self.document.add_paragraph(style='ReportFigureCaption')
        caption.add_run('Рисунок ')
        field(caption, 'SEQ Figure \\* ARABIC', number)
        caption.add_run(' – ' + title)
        for run in caption.runs:
            font(run, 14, False, False)
        self.inventory.append({'type': 'figure', 'number': number, 'title': title,
                               'source': str(Path(path).relative_to(SUBJECT)), 'width_cm': physical_width})

    def sources(self, items, new_page=False):
        self.heading('Список использованных источников', new_page=new_page)
        for number, source in enumerate(items, 1):
            p = self.document.add_paragraph(f'{number}. {source}', 'ReportReference')
            p.paragraph_format.keep_together = True

    def save(self, output=None):
        output = Path(output or os.environ.get('PPOIS_REPORT_OUTPUT', str(SUBJECT)))
        output.mkdir(parents=True, exist_ok=True)
        path = output / self.filename
        self.document.save(path)
        prune_unused_images(path)
        if sha256(REFERENCE.read_bytes()).hexdigest() != REFERENCE_HASH:
            raise AssertionError('Cover source was modified')
        cache = ROOT / '.cache/ppois5-10/report-evidence'
        cache.mkdir(parents=True, exist_ok=True)
        (cache / (path.stem + '.json')).write_text(json.dumps({
            'cover_source_sha256': REFERENCE_HASH, 'topic': self.topic,
            'tables': self.table_count, 'figures': self.figure_count,
            'objects': self.inventory, 'docx_sha256_before_word': sha256(path.read_bytes()).hexdigest(),
            'note': 'Build evidence only; final PDF and all-page visual inspection required'
        }, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        print(path)
        return path
