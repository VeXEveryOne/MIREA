"""Check the final Word/PDF pair without modifying either deliverable.

Automated checks are not visual QA. --reviewed-pages records only pages that
the caller has actually opened from the supplied, freshly rendered directory.
No report is deemed academically complete by this formatting verifier.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path
import re
import sys
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from PIL import Image
import pdfplumber

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / 'tools'))
from coursework_audit import inherited, paragraph_data
from report_layout import REFERENCE, REFERENCE_HASH


def digest(path):
    return sha256(Path(path).read_bytes()).hexdigest()


def effective(paragraph, key):
    value = getattr(paragraph.paragraph_format, key)
    return value if value is not None else inherited(paragraph.style, 'paragraph_format', key)


def zero(value):
    return value is None or abs(value.pt) < .1


def enabled(element):
    return element is not None and element.get(qn('w:val'), 'true') not in ('0', 'false', 'off')


def unshaded(element):
    return element is None or (
        element.get(qn('w:fill'), 'auto').lower() in ('auto', 'ffffff', 'none') and
        element.get(qn('w:val'), 'nil') in ('nil', 'clear'))


def parts(path):
    with ZipFile(path) as package:
        return {name: sha256(package.read(name)).hexdigest() for name in package.namelist()}


def cover_images(document):
    result = []
    for node in list(document.element.body)[:16]:
        for inline in node.xpath('.//wp:inline'):
            images = inline.xpath('.//a:blip')
            part = document.part.related_parts[images[0].get(qn('r:embed'))] if images else None
            extent = inline.xpath('./wp:extent')[0]
            result.append({'sha256': sha256(part.blob).hexdigest() if part is not None else None,
                           'cx': int(extent.get('cx')), 'cy': int(extent.get('cy'))})
    return result


def verify(docx, render_dir, reviewed_pages, output):
    docx, render_dir, output = [Path(p).resolve() for p in (docx, render_dir, output)]
    pdf = docx.with_suffix('.pdf')
    for path in (docx, pdf, REFERENCE):
        if not path.is_file():
            raise FileNotFoundError(path)
    document, reference = Document(docx), Document(REFERENCE)
    checks, errors = [], []

    def check(condition, label):
        checks.append(label)
        if not condition:
            errors.append(label)

    check(digest(REFERENCE) == REFERENCE_HASH, 'Original cover reference unchanged')
    check(cover_images(document) == cover_images(reference), 'Cover emblem bytes and dimensions preserved')
    # Compare the two retained tables, not the discarded reference body.
    for index in (0, 11):
        old, new = list(reference.element.body)[index], list(document.element.body)[index]
        check(old.tag == new.tag == qn('w:tbl'), f'Cover table {index} retains its position')
        check(old.xpath('./w:tblGrid/w:gridCol/@w:w') == new.xpath('./w:tblGrid/w:gridCol/@w:w'),
              f'Cover table {index} column widths unchanged')
        check(len(old.xpath('./w:tr')) == len(new.xpath('./w:tr')), f'Cover table {index} row count unchanged')
    for key in ('page_width', 'page_height', 'left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        check(abs(getattr(document.sections[0], key) - getattr(reference.sections[0], key)) < 1000,
              f'Cover geometry {key} preserved')
    for key in ('alignment', 'left_indent', 'first_line_indent', 'space_after', 'line_spacing'):
        check(getattr(document.styles['Normal'].paragraph_format, key) ==
              getattr(reference.styles['Normal'].paragraph_format, key), f'Cover Normal {key} preserved')
    cover_text = '\n'.join(''.join(n.xpath('.//w:t/text()')) for n in list(document.element.body)[:16])
    for value in ('МИНОБРНАУКИ РОССИИ', 'Институт информационных технологий',
                  'Албахтин И.В.', 'ИНБО-12-23', 'Борзых Н.Ю.', 'Москва 2026'):
        check(value in cover_text, 'Current cover text: ' + value)
    check('ИНБО-01-17' not in cover_text and 'Шендяпин' not in cover_text, 'Stale cover slots removed')
    check(document.sections[0].different_first_page_header_footer, 'First-page number hidden')
    check(not document.sections[0].first_page_footer.paragraphs[0].text.strip(), 'Cover footer empty')

    for index, section in enumerate(document.sections[1:], 1):
        dimensions = sorted([section.page_width.cm, section.page_height.cm])
        check(abs(dimensions[0] - 21) < .03 and abs(dimensions[1] - 29.7) < .03, f'Section {index} A4')
        for key, expected in [('left_margin', 3), ('right_margin', 1.5), ('top_margin', 2), ('bottom_margin', 2)]:
            check(abs(getattr(section, key).cm - expected) < .02, f'Section {index} {key}')
        check(not section._sectPr.xpath('./w:pgNumType/@w:start'), f'Section {index} continuous numbering')
        footer = section.footer.paragraphs[0]
        check('PAGE' in ''.join(footer._p.xpath('.//w:instrText/text()')), f'Section {index} real PAGE field')
        check(effective(footer, 'alignment') == 1 and all(zero(effective(footer, key)) for key in
              ('left_indent', 'right_indent', 'first_line_indent')), f'Section {index} centered unindented footer')
        check(all(r['font'] == 'Times New Roman' and r['size'] == 12 for r in paragraph_data(footer)['runs']),
              f'Section {index} footer TNR12')
    check(document.sections[-1].page_width < document.sections[-1].page_height, 'Final section restored to portrait')

    source_blocks = list(document.element.body)
    body = source_blocks[16:]
    body_text = '\n'.join(''.join(n.xpath('.//w:t/text()')) for n in body)
    check(not re.search(r'Ошибка!|Error!|TODO|FIXME|TBD|ИНБО-01-17', body_text), 'No placeholder or field-error text')
    check(not document.element.xpath('//w:ins | //w:del'), 'No tracked changes')
    check(not document.element.xpath('//wp:anchor'), 'All drawings inline')
    caps = {'Table': [], 'Figure': []}
    pdf_required_text = []
    table_total = 0
    for index, node in enumerate(body):
        label = f'Block {index + 16}'
        if node.tag == qn('w:p'):
            p = Paragraph(node, document._body)
            name = p.style.name
            if name.lower() in ('toc 1', 'toc 2'):
                # Word materialises the automatic TOC into native TOC styles.
                # Their intentional hierarchy is not foreign template prose.
                info = paragraph_data(p)
                check(all(r['font'] == 'Times New Roman' and r['size'] == 14 for r in info['runs']),
                      label + ' TOC TNR14')
                check(effective(p, 'alignment') in (None, 0) and
                      zero(effective(p, 'first_line_indent')), label + ' TOC left without first-line indent')
                expected_indent = 0 if name.lower() == 'toc 1' else .7
                indent = effective(p, 'left_indent')
                check(abs((indent.cm if indent is not None else 0) - expected_indent) < .03,
                      label + ' TOC hierarchy indent')
                continue
            if not name.startswith(('Report', 'Heading')):
                # Word creates empty section/page-break paragraphs. They are
                # not text content and must not silently carry old prose.
                check(not p.text.strip(), label + ' no foreign body style')
                continue
            info = paragraph_data(p)
            expected_size = 12 if name == 'ReportTableCaption' else 14
            check(all(r['font'] == 'Times New Roman' and r['size'] == expected_size for r in info['runs']),
                  label + ' correct text font and size')
            check(all(c.get(qn('w:val'), 'auto').lower() in ('000000', 'auto')
                      for c in node.xpath('.//w:color')), label + ' no colored text override')
            for key in ('left_indent', 'right_indent'):
                check(zero(effective(p, key)), label + ' ' + key)
            if name == 'ReportBody':
                check(effective(p, 'alignment') == 3 and effective(p, 'line_spacing') == 1.5,
                      label + ' body justified 1.5')
                indent = effective(p, 'first_line_indent')
                check(indent is not None and abs(indent.cm - 1.25) < .02, label + ' body first line 1.25cm')
                check(all(zero(effective(p, key)) for key in ('space_before', 'space_after')),
                      label + ' body spacing zero')
                check(effective(p, 'widow_control') is not False, label + ' widow control')
            if name.startswith('Heading'):
                pdf_required_text.append(p.text)
                check(p.style.builtin, label + ' built-in Word heading style')
                alignment = effective(p, 'alignment')
                check(alignment == 1 if name == 'Heading 1' else alignment in (None, 0), label + ' heading alignment')
                check(effective(p, 'keep_with_next') and effective(p, 'keep_together'), label + ' heading stays together')
                check(all(r['bold'] for r in info['runs']), label + ' heading bold')
            if name in ('ReportTableCaption', 'ReportFigureCaption', 'ReportFigure'):
                check(zero(effective(p, 'first_line_indent')) and effective(p, 'line_spacing') == 1,
                      label + ' object paragraph unindented single')
            if name == 'ReportTableCaption':
                pdf_required_text.append(p.text)
                match = re.match(r'Таблица (\d+) – ', p.text)
                check(match is not None, label + ' table caption numbered')
                if match:
                    number = int(match.group(1))
                    caps['Table'].append(number)
                    check(index > 0 and re.search(rf'таблице\s+{number}\b', ''.join(body[index-1].xpath('.//w:t/text()'))),
                          label + ' preceding table reference')
                check(index + 1 < len(body) and body[index+1].tag == qn('w:tbl'), label + ' caption above table')
                check(effective(p, 'alignment') in (None, 0) and effective(p, 'keep_with_next'), label + ' table caption left kept')
                check(all(r['italic'] and not r['bold'] for r in info['runs']), label + ' table caption italic not bold')
            if name == 'ReportFigureCaption':
                pdf_required_text.append(p.text)
                match = re.match(r'Рисунок (\d+) – ', p.text)
                check(match is not None, label + ' figure caption numbered')
                if match:
                    number = int(match.group(1))
                    caps['Figure'].append(number)
                    check(index > 1 and re.search(rf'рисунке\s+{number}\b', ''.join(body[index-2].xpath('.//w:t/text()'))),
                          label + ' preceding figure reference')
                check(index > 0 and body[index-1].xpath('.//wp:inline'), label + ' caption below inline image')
                check(effective(p, 'alignment') == 1 and effective(p, 'keep_together') and
                      not effective(p, 'keep_with_next'), label + ' figure caption centered not chained')
                check(all(not r['bold'] and not r['italic'] for r in info['runs']), label + ' figure caption regular')
            if name == 'ReportFigure':
                check(effective(p, 'alignment') == 1 and effective(p, 'keep_with_next'), label + ' image centered paired')
                # The section declaration following the paragraph governs its
                # page geometry, including the last body-level sectPr.
                declarations = [x for following in body[index:] for x in
                                ([following] if following.tag == qn('w:sectPr') else following.xpath('./w:pPr/w:sectPr'))]
                section = declarations[0]
                page_w = int(section.xpath('./w:pgSz/@w:w')[0])
                margins = section.xpath('./w:pgMar')[0]
                area_cm = (page_w - int(margins.get(qn('w:left'))) - int(margins.get(qn('w:right')))) / 567
                for inline in node.xpath('.//wp:inline'):
                    extent = inline.xpath('./wp:extent')[0]
                    cx, cy = int(extent.get('cx')), int(extent.get('cy'))
                    check(cx / 360000 <= area_cm + .03, label + ' image within page width')
                    image_part = document.part.related_parts[inline.xpath('.//a:blip')[0].get(qn('r:embed'))]
                    import io
                    with Image.open(io.BytesIO(image_part.blob)) as image:
                        check(abs((cx / cy) / (image.width / image.height) - 1) < .005, label + ' image aspect preserved')
        elif node.tag == qn('w:tbl'):
            table_total += 1
            table = Table(node, document._body)
            inherited_shading = []
            style = table.style
            seen_styles = set()
            while style is not None and style.style_id not in seen_styles:
                seen_styles.add(style.style_id)
                inherited_shading.extend(style.element.xpath('.//w:shd'))
                style = style.base_style
            check(all(unshaded(s) for s in inherited_shading + node.xpath('./w:tblPr//w:shd')),
                  label + ' table style does not add shading')
            check(enabled(table.rows[0]._tr.find('./' + qn('w:trPr') + '/' + qn('w:tblHeader'))), label + ' repeating header')
            for ri, row in enumerate(table.rows):
                check(enabled(row._tr.find('./' + qn('w:trPr') + '/' + qn('w:cantSplit'))), label + f' row {ri} not split')
                for ci, cell in enumerate(row.cells):
                    props = cell._tc.get_or_add_tcPr()
                    shading = props.find(qn('w:shd'))
                    check(unshaded(shading), label + f' cell {ri},{ci} no shading')
                    check(props.find(qn('w:tcBorders')) is not None, label + f' cell {ri},{ci} visible borders')
                    for p in cell.paragraphs:
                        info = paragraph_data(p)
                        check(all(r['font'] == 'Times New Roman' and r['size'] == 12 and
                                  bool(r['bold']) == (ri == 0) for r in info['runs']), label + f' cell {ri},{ci} TNR12 header bold')
                        check(zero(effective(p, 'first_line_indent')) and effective(p, 'line_spacing') == 1,
                              label + f' cell {ri},{ci} unindented single')
    for name, numbers in caps.items():
        check(numbers == list(range(1, len(numbers)+1)), f'{name} sequential cached numbers')
    check(table_total == len(caps['Table']), 'Every content table has a caption')

    page_details = []
    with pdfplumber.open(pdf) as published:
        count = len(published.pages)
        normalize = lambda value: re.sub(r'\s+', '', value)
        pdf_text = normalize('\n'.join(page.extract_text() or '' for page in published.pages))
        for value in pdf_required_text:
            check(normalize(value) in pdf_text, 'PDF contains current heading or caption: ' + value)
        check(reviewed_pages == list(range(1, count + 1)), 'Every final page explicitly reviewed')
        check(len(list(render_dir.glob('page-*.png'))) == count, 'Rendered image count equals PDF pages')
        check('Цель' not in (published.pages[0].extract_text() or ''), 'Body does not overflow cover')
        for n, page in enumerate(published.pages, 1):
            filename = render_dir / f'page-{n}.png'
            if not filename.is_file():
                errors.append(f'Missing render page {n}')
                continue
            text = page.extract_text() or ''
            check(len(text.strip()) >= 40, f'PDF page {n} not blank')
            width_cm, height_cm = page.width / 72 * 2.54, page.height / 72 * 2.54
            check(all(abs(a-b) < .04 for a,b in zip(sorted([width_cm,height_cm]), [21,29.7])), f'PDF page {n} A4')
            if n > 1:
                bottom = [c for c in page.chars if c['top'] > page.height - 55 and c['text'].strip()]
                check(''.join(c['text'] for c in bottom) == str(n), f'PDF page {n} actual footer number')
                if bottom:
                    center = (min(c['x0'] for c in bottom) + max(c['x1'] for c in bottom)) / 2
                    text_center = (3 / 2.54 * 72 + page.width - 1.5 / 2.54 * 72) / 2
                    check(abs(center - text_center) < 1.5, f'PDF page {n} footer actually centered in text area')
                    check(all(abs(c['size']-12) < .1 for c in bottom), f'PDF page {n} actual footer 12pt')
            with Image.open(filename) as image:
                check(image.width >= 900 and image.height >= 900, f'Page {n} review resolution')
            page_details.append({'number': n, 'render_sha256': digest(filename), 'manually_viewed': n in reviewed_pages})

    result = {'checked_at': datetime.now(timezone.utc).isoformat(), 'docx': docx.relative_to(ROOT).as_posix(),
              'docx_sha256': digest(docx), 'pdf_sha256': digest(pdf), 'cover_source_sha256': digest(REFERENCE),
              'automated_checks': len(checks), 'failed_checks': errors,
              'status': 'passed' if not errors else 'failed', 'tables': table_total, 'figures': len(caps['Figure']),
              'pages': page_details,
              'visual_reviewer': ('main agent opened every listed final PNG' if reviewed_pages == list(range(1, count+1))
                                  else 'Incomplete or not supplied'),
              'package_inventory': {'source': parts(REFERENCE), 'output': parts(docx)},
              'scope': 'Formatting and recorded visual review only; academic requirements checked separately'}
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({k: result[k] for k in ('status','automated_checks','failed_checks','tables','figures')}, ensure_ascii=False))
    if errors:
        raise SystemExit(1)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('docx', type=Path)
    parser.add_argument('render_dir', type=Path)
    parser.add_argument('--reviewed-pages', required=True, help='Comma-separated page numbers actually opened and inspected')
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    reviewed = sorted({int(x) for x in args.reviewed_pages.split(',')})
    verify(args.docx, args.render_dir, reviewed, args.output)
