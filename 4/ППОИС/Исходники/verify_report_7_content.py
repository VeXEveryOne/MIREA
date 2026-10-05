"""Check the saved PPOIS7 report against its executed catalog and data.

This is a content reconciliation, not a replacement for reading the assignment
or inspecting every rendered page. It never builds or edits the final report.
"""
from __future__ import annotations

from datetime import datetime, timezone
from hashlib import sha256
import json
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table

from build_report_7 import (CHECKS, DIAGRAMS, DATABASE, ORDER, ROOT,
                            SUBJECT, TEST_NAMES, digest, display, load_inputs,
                            physical_type)


def compact(text):
    return re.sub(r'\s+', ' ', text).strip()


def verify():
    evidence, models = load_inputs()
    path = SUBJECT / 'ППОИС_7_АлбахтинИВ.docx'
    document = Document(path)
    checks = []

    def check(name, condition):
        checks.append({'name': name, 'passed': bool(condition)})

    tables = {}
    for paragraph in document.paragraphs:
        match = re.match(r'Таблица (\d+)\s+[–—-]\s+(.+)', paragraph.text)
        if not match:
            continue
        sibling = paragraph._p.getnext()
        check('caption immediately precedes table ' + match[1],
              sibling is not None and sibling.tag == qn('w:tbl'))
        if sibling is not None and sibling.tag == qn('w:tbl'):
            table = Table(sibling, document._body)
            tables[match[2]] = [[cell.text for cell in row.cells] for row in table.rows]

    def rows(title):
        if title not in tables:
            raise ValueError('Missing content table: ' + title)
        return tables[title][1:]

    check('all twelve entities described',
          {row[0].split('\n')[-1] for row in rows('Сущности информационной модели')} == set(ORDER))
    check('five weak entities and identifying ownership described',
          {row[0] for row in rows('Зависимые сущности и их идентификация')} == set(models['weakEntities']))

    columns = {(c['table_name'], c['column_name']): c for c in evidence['columns']}
    logical = rows('Атрибуты логической модели')
    physical = rows('Поля физической модели')
    for label, records in [('logical', logical), ('physical', physical)]:
        check(label + ' covers every field exactly once',
              len(records) == len(columns) and {(r[0], r[1]) for r in records} == set(columns))
    for row in physical:
        catalog = columns[(row[0], row[1])]
        check('physical type and mandatory ' + row[0] + '.' + row[1],
              row[2] == physical_type(catalog)
              and row[3] == ('Да' if catalog['is_nullable'] == 'NO' else 'Нет'))
    for row in logical:
        catalog = columns[(row[0], row[1])]
        check('logical optionality ' + row[0] + '.' + row[1],
              ('Обязательный' if catalog['is_nullable'] == 'NO' else 'NULL допустим') in row[2])

    keys = rows('Первичные уникальные и внешние ключи')
    check('all actual key constraints described', len(keys) == len(evidence['key_catalog']))
    for row, key in zip(keys, evidence['key_catalog']):
        expected = {'p': 'PK', 'u': 'UQ', 'f': 'FK'}[key['contype']] + '(' + ', '.join(key['columns']) + ')'
        if key['contype'] == 'f':
            expected += ' → ' + key['parent_table'] + '(' + ', '.join(key['parent_columns']) + ')'
        check('key ' + key['conname'], row[1:] == [key['table_name'], expected])
    actual_checks = [c for c in evidence['constraints'] if c['contype'] == 'c']
    check_rows = rows('Ограничения допустимых значений')
    check('all actual CHECK constraints described', len(check_rows) == len(actual_checks))
    for row, constraint in zip(check_rows, actual_checks):
        check('CHECK ' + constraint['conname'],
              row[1:] == [constraint['table_name'], CHECKS[constraint['conname']]])

    for table in ORDER:
        fields = [c['column_name'] for c in sorted(
            (c for c in evidence['columns'] if c['table_name'] == table), key=lambda c: c['ordinal_position'])]
        records = evidence['fixtures'][table]
        title = 'Тестовые записи ' + table
        if table in ('card_revision', 'publication_attempt', 'image'):
            expected = [['Поле'] + [f'Запись {i+1}' for i in range(len(records))]]
            expected += [[name] + [display(record[name]) for record in records] for name in fields]
        else:
            expected = [fields] + [[display(record[name]) for name in fields] for record in records]
        check('every fixture field and value ' + table, tables[title] == expected)

    paragraphs = document.paragraphs
    ddl = []; dcl = []
    for i, p in enumerate(paragraphs[:-1]):
        if re.match(r'^DDL\d+ ', p.text):
            ddl.append(paragraphs[i+1].text)
        if re.match(r'^DCL\d+ ', p.text):
            dcl.append(paragraphs[i+1].text)
    schema = compact((DATABASE / 'schema.sql').read_text(encoding='utf-8'))
    privileges = compact((DATABASE / 'privileges.sql').read_text(encoding='utf-8'))
    check('nine DDL excerpts from executed source', len(ddl) == 9 and all(compact(s) in schema for s in ddl))
    check('twelve DCL excerpts from executed source', len(dcl) == 12 and all(compact(s) in privileges for s in dcl))
    paragraph_sql = {compact(p.text) for p in paragraphs}
    for i, query in enumerate(evidence['queries'], 1):
        sql = '\n'.join(line for line in query['sql'].splitlines() if not line.startswith('--'))
        check(f'Q{i:02d} exact executed SQL', compact(sql) in paragraph_sql)
        fields = list(query['rows'][0])
        expected = [[compact(f) for f in fields]] + [
            [display(row[name]) for name in fields] for row in query['rows']]
        actual = tables[f'Результат Q{i:02d}']
        actual = [[compact(h) for h in actual[0]]] + actual[1:]
        # Word wraps this long header at an underscore; it remains one field.
        actual[0] = [h.replace('confirmed_ not_accepted', 'confirmed_not_accepted') for h in actual[0]]
        check(f'Q{i:02d} all result columns and rows', actual == expected)
    expected_tests = [[str(i), TEST_NAMES[t['name']], t.get('sqlstate', 'Выполнено без ошибки'),
                       'Ожидание подтверждено'] for i, t in enumerate(evidence['tests'], 1)]
    check('all 22 database test results and descriptions',
          rows('Фактические результаты проверок') == expected_tests)

    figure_paths = [DIAGRAMS / s['file'] for s in models['conceptualSheets']]
    figure_paths += [DIAGRAMS / f'ПР7_{kind}_{group}.png'
                     for kind in ('Логическая', 'Физическая')
                     for group in ('catalog', 'bom', 'cards', 'publication')]
    package_image_hashes = {sha256(rel.target_part.blob).hexdigest()
                            for rel in document.part.rels.values()
                            if rel.reltype.endswith('/image')}
    check('all 21 required current diagram images embedded',
          len(figure_paths) == 21 and all(digest(p) in package_image_hashes for p in figure_paths))
    qa_path = ROOT / '.cache/ppois5-10/report-evidence/practice7-final-qa.json'
    qa = json.loads(qa_path.read_text(encoding='utf-8'))
    # Layout verification is reported separately, never inferred from content.
    check('latest separate formatting and visual audit passed',
          qa['status'] == 'passed' and qa['docx_sha256'] == digest(path)
          and qa['pdf_sha256'] == digest(path.with_suffix('.pdf'))
          and len(qa['pages']) == 60 and all(p['manually_viewed'] for p in qa['pages']))
    result = {'status': 'passed' if all(c['passed'] for c in checks) else 'failed',
              'verified_at': datetime.now(timezone.utc).isoformat(),
              'report_sha256': digest(path), 'pdf_sha256': digest(path.with_suffix('.pdf')),
              'model_sha256': digest(DIAGRAMS / 'ПР7_Модель_данных.omni'),
              'database': evidence['database'], 'source_sha256': evidence['source_sha256'],
              'assignment': 'practice 7 supplied PDF items 1–3; visually read both pages',
              'scope': 'artifact and executed-catalog reconciliation; no external API, load or recovery claims',
              'checks': checks}
    output = ROOT / '.cache/ppois5-10/report-evidence/practice7-content-qa.json'
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'status': result['status'], 'checks': len(checks),
                      'failed': [c['name'] for c in checks if not c['passed']]}, ensure_ascii=False))
    return result['status'] == 'passed'


if __name__ == '__main__':
    raise SystemExit(0 if verify() else 1)
