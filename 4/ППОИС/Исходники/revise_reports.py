"""Controlled PPOIS corrections; preserve the title and unrelated package parts."""
from __future__ import annotations
import argparse
from copy import deepcopy
import hashlib
import json
import posixpath
import re
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[3]
SUBJECT = ROOT / '4/ППОИС'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def prune_unused_images(path: Path):
    """Remove only unreferenced image relationships and their orphan media.

    Relationship references in all attributes (including VML fallbacks) are
    respected. Non-image relationships and all content XML remain untouched.
    """
    changed={}; dropped_relationships=[]
    with ZipFile(path) as original:
        names=set(original.namelist())
        for name in names:
            if not name.endswith('.rels') or '/_rels/' not in name: continue
            directory,leaf=name.rsplit('/_rels/',1)
            owner=directory+'/'+leaf.removesuffix('.rels')
            if owner not in names: continue
            content=etree.fromstring(original.read(owner))
            references={value for node in content.iter() for value in node.attrib.values()}
            rels=etree.fromstring(original.read(name)); removed=[]
            for relation in list(rels):
                if relation.get('Type','').endswith('/image') and relation.get('TargetMode')!='External' and relation.get('Id') not in references:
                    removed.append(relation.get('Id')); rels.remove(relation)
            if removed:
                changed[name]=etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
                dropped_relationships.extend(name+':'+rid for rid in removed)
        live=set()
        for name in names:
            if not name.endswith('.rels'): continue
            rels=etree.fromstring(changed.get(name,original.read(name)))
            base=name.rsplit('/_rels/',1)[0] if '/_rels/' in name else ''
            for relation in rels:
                target=relation.get('Target','')
                if relation.get('TargetMode')=='External': continue
                live.add(posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(base,target)))
        deleted={name for name in names if name.startswith('word/media/') and name not in live}
        if deleted:
            types=etree.fromstring(original.read('[Content_Types].xml'))
            for item in list(types):
                if item.get('PartName','').lstrip('/') in deleted: types.remove(item)
            changed['[Content_Types].xml']=etree.tostring(types,xml_declaration=True,encoding='UTF-8',standalone=True)
        temporary=path.with_suffix('.writing.docx')
        with ZipFile(temporary,'w') as result:
            for item in original.infolist():
                if item.filename not in deleted: result.writestr(item,changed.get(item.filename,original.read(item.filename)))
    temporary.replace(path)
    return {'removed_orphan_media':sorted(deleted),'removed_unused_relationships':dropped_relationships}


def merge_package(original: Path, edited: Path, destination: Path, permitted: set[str]):
    with ZipFile(original) as before, ZipFile(edited) as after:
        original_names = set(before.namelist())
        changes = {}
        for name in original_names & set(after.namelist()):
            if before.read(name) != after.read(name):
                changes[name] = name in permitted
        new_parts = set(after.namelist()) - original_names
        if any(not name.startswith('word/media/') for name in new_parts):
            raise ValueError(f'Unexpected new parts: {new_parts}')
        temporary = destination.with_suffix('.writing.docx')
        with ZipFile(temporary, 'w') as result:
            for item in before.infolist():
                result.writestr(item, after.read(item.filename) if item.filename in permitted else before.read(item.filename))
            for name in new_parts:
                result.writestr(after.getinfo(name), after.read(name))
    temporary.replace(destination)
    return {'changed_parts': sorted(name for name, allowed in changes.items() if allowed),
            'preserved_rewritten_parts': sorted(name for name, allowed in changes.items() if not allowed),
            'new_media': sorted(new_parts)}


def patch_text(source: Path, destination: Path, replacements: dict[str, str]):
    with ZipFile(source) as archive:
        xml = etree.fromstring(archive.read('word/document.xml'))
        found = set()
        accepted={value:key for key,value in replacements.items()}
        for paragraph in xml.xpath('//w:p', namespaces=NS):
            nodes = paragraph.xpath('.//w:t', namespaces=NS)
            text = ''.join(node.text or '' for node in nodes)
            if text in replacements:
                nodes[0].text = replacements[text]
                for node in nodes[1:]: node.text = ''
                found.add(text)
            elif text in accepted:
                found.add(accepted[text])
        if found != set(replacements):
            raise ValueError(f'Missing replacement paragraphs: {set(replacements) - found}')
        for fonts in xml.xpath('//w:rFonts', namespaces=NS):
            for key, value in list(fonts.attrib.items()):
                if value == 'Liberation Serif': fonts.set(key, 'Times New Roman')
        temporary = destination.with_suffix('.writing.docx')
        with ZipFile(temporary, 'w') as result:
            for item in archive.infolist():
                data = etree.tostring(xml, xml_declaration=True, encoding='UTF-8', standalone=True) if item.filename == 'word/document.xml' else archive.read(item.filename)
                result.writestr(item, data)
    temporary.replace(destination)


def replace_idef0_images(path: Path, figures: Path):
    with ZipFile(path) as original:
        xml=etree.fromstring(original.read('word/document.xml'))
        rels=etree.fromstring(original.read('word/_rels/document.xml.rels'))
        relationships={r.get('Id'):r.get('Target') for r in rels}
        drawing_ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
                    'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
        pictures=xml.xpath('//w:body//a:blip',namespaces={**NS,**drawing_ns})
        if len(pictures)!=4: raise ValueError('PPOIS2: expected title and three diagrams')
        replacements={}
        for blip,sheet in zip(pictures[1:],['A-0','A0','A2']):
            target='word/'+relationships[blip.get('{'+drawing_ns['r']+'}embed')]
            replacements[target]=(figures/f'ПР2_IDEF0_{sheet}.png').read_bytes()
        temporary=path.with_suffix('.writing.docx')
        with ZipFile(temporary,'w') as result:
            for item in original.infolist(): result.writestr(item,replacements.get(item.filename,original.read(item.filename)))
    temporary.replace(path)
    return sorted(replacements)


ANALYSIS_PARTS = [
    'ПР3_03_Последовательность_UC02_1.png',
    'ПР3_03_Последовательность_UC02_2.png',
    'ПР3_04_Последовательность_UC05_1.png',
    'ПР3_04_Последовательность_UC05_2.png',
    'ПР3_05_Кооперация_UC05.png',
]


def replace_analysis_images(path: Path, figures: Path):
    """Replace print views, not the title/use-case/class diagram or relationships."""
    with ZipFile(path) as original:
        xml=etree.fromstring(original.read('word/document.xml'))
        rels=etree.fromstring(original.read('word/_rels/document.xml.rels'))
        relationships={r.get('Id'):r.get('Target') for r in rels}
        drawing_ns={'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
                    'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
        pictures=xml.xpath('//w:body//a:blip',namespaces={**NS,**drawing_ns})
        if len(pictures)!=8: raise ValueError('PPOIS3: expected title and seven diagram images')
        replacements={}
        for blip,name in zip(pictures[3:],ANALYSIS_PARTS):
            target='word/'+relationships[blip.get('{'+drawing_ns['r']+'}embed')]
            replacements[target]=(figures/name).read_bytes()
        temporary=path.with_suffix('.writing.docx')
        with ZipFile(temporary,'w') as result:
            for item in original.infolist(): result.writestr(item,replacements.get(item.filename,original.read(item.filename)))
    temporary.replace(path)
    return sorted(replacements)


def font(run, size=14, bold=False, italic=False):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0,0,0)
    run.bold = bold
    run.italic = italic
    for kind in ('ascii','hAnsi','eastAsia','cs'):
        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:'+kind), 'Times New Roman')


def format_existing(path: Path, working: Path, figures: Path | None = None):
    """Normalize only the report body; keep first-page XML and unrelated parts."""
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    doc = Document(path)
    body = doc.element.body
    first = next(p._p for p in doc.paragraphs if p.text.startswith('ПРАКТИЧЕСКАЯ РАБОТА №'))
    prefix = list(body)[:list(body).index(first)]
    title = [etree.tostring(child) for child in prefix]
    for index,section in enumerate(doc.sections):
        if index==0: continue  # the template's title section is immutable
        section.different_first_page_header_footer=False
        for numbering in section._sectPr.findall(qn('w:pgNumType')):
            numbering.attrib.pop(qn('w:start'),None)
            numbering.set(qn('w:fmt'),'decimal')
    for child in list(body)[len(prefix):]:
        if child.tag == qn('w:p'):
            paragraph = Paragraph(child, doc._body)
            text = paragraph.text
            joined_text = ''.join(child.xpath('.//w:t/text()'))
            if joined_text == 'Формирование функциональных требований к ИСОбъектный анализ и модель анализа':
                paragraph.clear()
                paragraph.add_run('Объектный анализ и модель анализа')
                text = paragraph.text
            if 'Открытый сайт компании' in text:
                for run in paragraph.runs: run.text=run.text.replace('каналы продаж [1]','каналы продаж [2]')
                text=paragraph.text
            pf = paragraph.paragraph_format
            pf.left_indent = Cm(0); pf.right_indent = Cm(0)
            pf.first_line_indent = Cm(1.25); pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.5
            generated_reference=bool(re.fullmatch(r'(?:Результаты приведены в таблице \d+\.|На рисунке \d+ представлена модель\.)',text))
            pf.widow_control = True; pf.keep_with_next = generated_reference; pf.keep_together = False
            heading = text.casefold() == 'список использованных источников' or bool(re.match(r'^\d+(?:\.\d+)?\s', text)) or text.startswith('ПРАКТИЧЕСКАЯ РАБОТА №') or text in {
                'СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ',
                'Анализ результатов предварительного обследования. Формирование функциональных требований к ИС',
                'Функциональная модель информационной системы',
                'Формирование функциональных требований ИС (структурный анализ)',
                'Объектный анализ. Модель анализа',
                'Объектный анализ и модель анализа',
            }
            caption = text.startswith('Таблица ') or text.startswith('Рисунок ')
            picture = bool(child.xpath('.//w:drawing | .//w:pict'))
            if not text and not picture and not child.xpath('.//w:fldChar | .//w:instrText'):
                if child.xpath('./w:pPr/w:sectPr'):
                    pf.line_spacing=Pt(1); pf.first_line_indent=Cm(0)
                    for run in paragraph.runs: font(run,1)
                else:
                    body.remove(child)
                continue
            if heading:
                sub = bool(re.match(r'^\d+\.\d+\s',text))
                paragraph.style = 'Heading 2' if sub else 'Heading 1'
                pf.first_line_indent = Cm(0)
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT if sub else WD_ALIGN_PARAGRAPH.CENTER
                pf.space_before = Pt(12); pf.space_after = Pt(6)
                pf.keep_with_next = True; pf.keep_together = True
            elif caption:
                pf.first_line_indent = Cm(0); pf.line_spacing = 1
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT if text.startswith('Таблица ') else WD_ALIGN_PARAGRAPH.CENTER
                pf.space_before = Pt(6 if text.startswith('Таблица ') else 3)
                pf.space_after = Pt(3 if text.startswith('Таблица ') else 6)
                pf.keep_with_next = text.startswith('Таблица '); pf.keep_together = True
            elif picture:
                pf.alignment = WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent = Cm(0)
                pf.line_spacing = 1; pf.keep_with_next = True
            for run in paragraph.runs:
                font(run,12 if text.startswith('Таблица ') else 14,heading,text.startswith('Таблица '))
            if caption:
                number = re.match(r'^(Таблица|Рисунок) (\d+)', text)
                if number:
                    target = child
                    if number[1] == 'Рисунок':
                        previous = child.getprevious()
                        if previous is not None and previous.xpath('.//w:drawing | .//w:pict'): target = previous
                        # A split diagram has one first mention, not a new object
                        # on its continuation page. Do not orphan an extra intro
                        # before the explicit page break of the second image.
                        if '(окончание)' in text:
                            previous = target.getprevious()
                            if previous is not None and ''.join(previous.xpath('.//w:t/text()')) == f'На рисунке {number[2]} представлена модель.':
                                body.remove(previous)
                            continue
                    previous=target.getprevious()
                    context=[]
                    for _ in range(5):
                        if previous is None or previous.tag==qn('w:tbl'): break
                        previous_text=''.join(previous.xpath('.//w:t/text()'))
                        if previous_text.startswith(('Рисунок ','Таблица ')): break
                        context.append(previous_text)
                        previous=previous.getprevious()
                    word='таблиц[аеуы]' if number[1]=='Таблица' else 'рисун(?:ке|ок|ка)'
                    if re.search(rf'{word}\s+{number[2]}(?!\d)',' '.join(context),re.IGNORECASE):
                        continue
                    reference = Paragraph(target,doc._body).insert_paragraph_before(
                        f'Результаты приведены в таблице {number[2]}.' if number[1]=='Таблица' else f'На рисунке {number[2]} представлена модель.')
                    reference.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                    reference.paragraph_format.left_indent = Cm(0); reference.paragraph_format.right_indent = Cm(0)
                    reference.paragraph_format.first_line_indent = Cm(1.25)
                    reference.paragraph_format.space_before = Pt(0); reference.paragraph_format.space_after = Pt(0)
                    reference.paragraph_format.line_spacing = 1.5; reference.paragraph_format.keep_with_next = True
                    reference.paragraph_format.keep_together = False; reference.paragraph_format.widow_control = True
                    target_paragraph = Paragraph(target,doc._body)
                    if number[1] == 'Рисунок' and target_paragraph.paragraph_format.page_break_before:
                        reference.paragraph_format.page_break_before = True
                        target_paragraph.paragraph_format.page_break_before = False
                    for run in reference.runs: font(run)
        elif child.tag == qn('w:tbl'):
            table = Table(child,doc._body)
            priority = len(table.columns)==4 and table.cell(0,0).text=='Код' and 'Приоритет' in table.cell(0,3).text
            operations = len(table.columns)==8 and table.cell(0,0).text in ('Диаграмма, операция','Код')
            processes = len(table.columns)==3 and table.cell(0,0).text=='№' and 'бизнес-процесса' in table.cell(0,1).text
            widths = None
            if priority:
                widths = [1.6,8.0,4.1,2.8]
            elif operations:
                widths = [1.7,3.1,3.5,3.2,2.4,3.8,3.0,4.5]
                for cell,text in zip(table.rows[0].cells,['Код','Результат','Операция','Исполнитель','Период','Основание','Реестр','Примечание']): cell.text=text
                for row in table.rows[1:]:
                    if row.cells[2].text=='Зарегистрировать запрос': row.cells[2].text='Записать запрос'
                    if row.cells[4].text in ('По мере необходимости','По потребности'): row.cells[4].text='По запросу'
                    if row.cells[4].text=='При обновлении данных': row.cells[4].text='При загрузке цен'
            elif processes:
                widths = [1.2,7.5,7.8]
            if widths:
                table.autofit=False
                for column,width in zip(table.columns,widths): column.width=Cm(width)
                for row in table.rows:
                    for cell,width in zip(row.cells,widths): cell.width=Cm(width)
            for i,row in enumerate(table.rows):
                trpr=row._tr.get_or_add_trPr()
                if trpr.find(qn('w:cantSplit')) is None: trpr.append(OxmlElement('w:cantSplit'))
                if i==0 and trpr.find(qn('w:tblHeader')) is None: trpr.append(OxmlElement('w:tblHeader'))
                for j,cell in enumerate(row.cells):
                    if j==1 and row.cells[0].text=='Статусы карточек':
                        cell.text='Карточка: черновик, на проверке, ошибка проверки, готова к публикации, опубликована, требует актуализации. Обмен: отправляется, ожидает результата, временная ошибка, отказ, результат неизвестен.'
                    for shade in list(cell._tc.get_or_add_tcPr().findall(qn('w:shd'))): cell._tc.get_or_add_tcPr().remove(shade)
                    for paragraph in cell.paragraphs:
                        pf=paragraph.paragraph_format
                        pf.first_line_indent=Cm(0); pf.left_indent=Cm(0); pf.right_indent=Cm(0)
                        pf.space_before=Pt(0); pf.space_after=Pt(0); pf.line_spacing=1
                        pf.keep_with_next=i==0; pf.keep_together=True
                        pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if (priority and j in (0,3)) or (processes and j==0) or (operations and j==0) else WD_ALIGN_PARAGRAPH.LEFT
                        for run in paragraph.runs: font(run,12,i==0)
    if path.name=='ППОИС_2_АлбахтинИВ.docx' and figures:
        # The SVG layout can change its aspect ratio. Fit from the new media,
        # not the old document extents, so no diagram is stretched.
        for shape,sheet,height in zip(list(doc.inline_shapes)[1:],['A-0','A0','A2'],[12.63,12.8,12.4]):
            with Image.open(figures/f'ПР2_IDEF0_{sheet}.png') as image:
                width=min(25.2,height*image.width/image.height)
                shape.width=Cm(width)
                shape.height=int(Cm(width)*image.height/image.width)
    if path.name=='ППОИС_3_АлбахтинИВ.docx' and figures:
        for shape,name,height in zip(list(doc.inline_shapes)[3:],ANALYSIS_PARTS,[12.3,14.3,12.3,14.3,12.9]):
            with Image.open(figures/name) as image:
                width=min(25.2,height*image.width/image.height)
                shape.width=Cm(width)
                shape.height=int(Cm(width)*image.height/image.width)
    if [etree.tostring(child) for child in list(body)[:len(prefix)]] != title:
        raise ValueError('Formatting touched the title page')
    working.mkdir(parents=True,exist_ok=True)
    intermediate=working/(path.stem+'_formatted.docx'); doc.save(intermediate)
    return merge_package(path,intermediate,path,{'word/document.xml'})


def build_four(source: Path, destination: Path, figures: Path, working: Path):
    doc = Document(source)
    body = doc.element.body
    first = next(p._p for p in doc.paragraphs if p.text == 'ПРАКТИЧЕСКАЯ РАБОТА № 4')
    prefix = list(body)[:list(body).index(first)]
    title_xml = [etree.tostring(child) for child in prefix]
    final_section = deepcopy(body[-1])
    # The last section can inherit its footer. Capture the first body section's
    # explicit references before rebuilding, otherwise a second build can lose
    # continuous page numbering while leaving the title deceptively intact.
    references=[deepcopy(item) for item in doc.sections[1]._sectPr if item.tag in (qn('w:headerReference'),qn('w:footerReference'))]
    for item in list(final_section):
        if item.tag in (qn('w:headerReference'),qn('w:footerReference')): final_section.remove(item)
    for item in reversed(references): final_section.insert(0,item)
    for child in list(body)[len(prefix):]: body.remove(child)
    body.append(final_section)
    table_number = 0
    figure_number = 0

    def p(text='', *, heading=0, size=14, italic=False, align=None, indent=1.25, before=0, after=0, spacing=1.5, keep=False):
        paragraph = doc.add_paragraph(style=f'Heading {heading}' if heading else 'Normal')
        pf = paragraph.paragraph_format
        pf.alignment = align if align is not None else (WD_ALIGN_PARAGRAPH.CENTER if heading == 1 else WD_ALIGN_PARAGRAPH.LEFT if heading else WD_ALIGN_PARAGRAPH.JUSTIFY)
        pf.first_line_indent = Cm(0 if heading else indent)
        pf.left_indent = Cm(0); pf.right_indent = Cm(0)
        pf.space_before = Pt(12 if heading else before); pf.space_after = Pt(6 if heading else after)
        pf.line_spacing = spacing; pf.widow_control = True
        pf.keep_with_next = bool(heading or keep); pf.keep_together = bool(heading)
        pf.page_break_before = False
        font(paragraph.add_run(text), size, bool(heading), italic)
        return paragraph

    def section(landscape=False):
        current = doc.sections[-1]
        wanted = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        if current.orientation != wanted:
            current = doc.add_section(WD_SECTION_START.NEW_PAGE)
            separator = doc.paragraphs[-1]
            separator.paragraph_format.space_before = Pt(0); separator.paragraph_format.space_after = Pt(0)
            separator.paragraph_format.line_spacing = Pt(1)
            separator.paragraph_format.keep_with_next = False
            font(separator.add_run(),1)
        current.orientation = wanted
        current.page_width = Cm(29.7 if landscape else 21)
        current.page_height = Cm(21 if landscape else 29.7)
        current.left_margin = Cm(3); current.right_margin = Cm(1.5)
        current.top_margin = Cm(2); current.bottom_margin = Cm(2)
        current.header_distance = Cm(.8); current.footer_distance = Cm(1)
        current.different_first_page_header_footer = False
        for numbering in current._sectPr.findall(qn('w:pgNumType')): current._sectPr.remove(numbering)

    def figure(file, caption, *, landscape=False, new_page=True):
        nonlocal figure_number
        figure_number += 1
        changed = doc.sections[-1].orientation != (WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT)
        section(landscape)
        reference = p(f'На рисунке {figure_number} показана модель «{caption}».', keep=True)
        reference.paragraph_format.page_break_before = new_page and not changed
        image = Image.open(figures / file)
        width = min(25.2 if landscape else 16.5, (13 if landscape else 19) * image.width / image.height)
        paragraph = p(indent=0, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1, keep=True)
        shape=paragraph.add_run().add_picture(str(figures/file), width=Cm(width))
        shape._inline.graphic.graphicData.pic.nvPicPr.cNvPr.set('name',file)
        cap = p(f'Рисунок {figure_number} – {caption}', indent=0, align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1, before=3, after=6)
        cap.paragraph_format.keep_together = True

    def table(caption, headers, rows, widths):
        nonlocal table_number
        section()
        table_number += 1
        p(f'Результаты приведены в таблице {table_number}.', keep=True)
        p(f'Таблица {table_number} – {caption}', size=12, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, indent=0, spacing=1, before=6, after=3, keep=True)
        result = doc.add_table(rows=1, cols=len(headers))
        result.style = 'Table Grid'; result.autofit = False
        for column, width in zip(result.columns,widths): column.width = Cm(width)
        for index, row in enumerate([headers, *rows]):
            cells = result.rows[0].cells if index == 0 else result.add_row().cells
            for cell, width, text in zip(cells,widths,row):
                cell.width = Cm(width)
                paragraph = cell.paragraphs[0]
                pf = paragraph.paragraph_format
                pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                pf.first_line_indent = Cm(0); pf.left_indent = Cm(0); pf.right_indent = Cm(0)
                pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1
                pf.keep_with_next = index==0; pf.keep_together = True
                font(paragraph.add_run(str(text)),12,index==0)
                shading = cell._tc.get_or_add_tcPr().find(qn('w:shd'))
                if shading is not None: cell._tc.get_or_add_tcPr().remove(shading)
            trpr = result.rows[index]._tr.get_or_add_trPr()
            trpr.append(OxmlElement('w:cantSplit'))
            if index == 0: trpr.append(OxmlElement('w:tblHeader'))
        properties = result._tbl.tblPr
        margins = OxmlElement('w:tblCellMar')
        for name, value in [('top','65'),('bottom','65'),('left','85'),('right','85')]:
            element=OxmlElement('w:'+name); element.set(qn('w:w'),value); element.set(qn('w:type'),'dxa'); margins.append(element)
        properties.append(margins)
        return result

    section()
    p('ПРАКТИЧЕСКАЯ РАБОТА № 4',heading=1)
    p('Объектный анализ. Модель проектирования', heading=1)
    p('1 Цель и индивидуальное задание',heading=1)
    p('Цель работы — детализировать статическую структуру и поведение информационной системы формирования и актуализации карточек компьютерных систем ООО «Юкомс», продолжая модели практических работ № 1–3.')
    p('Индивидуальное задание: построить диаграммы классов проектирования, деятельности и состояний; определить типы атрибутов и операций, кратности связей, параллельные действия, условия проверки и безопасного повторения внешних операций.')
    p('2 Этапы выполнения работы',heading=1)
    for text in [
        '1. Классы анализа уточнены до классов-сущностей, границ и управления; для них определены типизированные данные и операции.',
        '2. Уточнены групповые и точные позиции BOM, условия выбора предложений и неизменяемые снимки ревизий.',
        '3. Независимая подготовка состава и обновление предложений разделены параллельно; зависимые проверки, расчёты и подготовка содержания выполняются после синхронизации.',
        '4. В жизненном цикле отделены приём задачи, подтверждённая публикация, окончательный отказ, временная ошибка и неизвестный исход.',
        '5. Классы, действия и переходы сопоставлены с функциями IDEF0 и сообщениями UC02/UC05.',
    ]: p(text,indent=0)
    section(True)
    p('3 Диаграммы классов проектирования',heading=1)
    for index, (file, caption) in enumerate([
        ('01_ПР4_01_Классы.png','Классы конфигураций, позиций BOM и предложений'),
        ('02_ПР4_02_Классы.png','Классы ревизий и попыток публикации'),
        ('03_ПР4_03_Классы.png','Границы и управление подготовкой состава'),
        ('04_ПР4_04_Классы.png','Границы и управление публикацией'),
    ]): figure(file,caption,landscape=True,new_page=index>0)
    section()
    p('Модель разделена на четыре представления одного состава классов. Повторное изображение класса на другом листе не создаёт новый класс. Классы-сущности имеют идентификатор id: UUID с ограничением {PK}. Инкапсулированные атрибуты закрыты, операции открыты. Формы не содержат бизнес-правил.')
    model = json.loads((figures/'ПР4_Модель_проектирования.omni').read_text(encoding='utf-8'))['model']
    elements = model['elements']; by_id = {e['id']:e for e in elements}
    descriptions = {
        'Configuration':'Код модельного ряда; поиск затронутых конфигураций.',
        'BomVersion':'Версия состава и результат проверки; после использования в ревизии её состав не меняется.',
        'BomSlot':'Роль и количество; режим GROUP или EXACT; выбор конкретного предложения.',
        'ComponentGroup':'Ограничения выбора, например чипсет B650; не является конкретным товаром.',
        'Component':'Нормализованные технические характеристики и масса конкретной модели.',
        'Offer':'Цена, наличие и время получения предложения конкретного поставщика.',
        'Card':'Постоянный код предложения продавца; история ревизий.',
        'CardRevision':'Неизменяемый снимок проверенного состава, цены, массы, атрибутов и изображений.',
        'Image':'Адрес файла и контрольная сумма; одна картинка может использоваться в разных ревизиях.',
        'PublicationAttempt':'Идентификатор попытки и внешней задачи, состояние и замечания; прошлые попытки сохраняются.',
        'BomForm':'Ввод черновика и показ замечаний; передача команды управлению BOM.',
        'CardForm':'Подтверждение отправки ревизии и показ состояния попытки.',
        'MarketplaceGateway':'Подготовка внешнего запроса, отправка и сверка результата. credentialRef — ссылка на секрет, не сам ключ.',
        'BomControl':'Версионирование, разрешение слотов и предварительный просмотр массовой замены.',
        'Compatibility':'Проверка конкретных моделей по сокету, габаритам, питанию и интерфейсам.',
        'Calculation':'Расчёт себестоимости, цены и массы по выбранным актуальным предложениям.',
        'Readiness':'Проверка совместимости, актуальности, обязательных полей и изображений.',
        'PublicationControl':'Регистрация попытки до отправки, классификация результата, сверка и ограниченные повторы.',
    }
    table('Описание классов проектирования',['Класс','Стереотип','Ответственность'],[[e['label'],e['properties']['role'],descriptions[e['id']]] for e in elements if e['kind']=='class'],[4.5,2.4,9.6])
    p('3.1 Типы данных и ограничения',heading=2)
    p('UUID — внутренний уникальный идентификатор; Integer — целое число; Decimal — точное десятичное число; DateTime — момент времени с часовым поясом; JSON — структурированная коллекция, а не строка произвольного текста. Boolean задаёт логическое условие. Void означает отсутствие возвращаемого значения. Необязательные значения, например externalTaskId до получения ответа, допускают отсутствие; отсутствие не заменяется пустым фиктивным идентификатором.')
    table('Структура значимых входных и выходных типов',['Тип','Состав и правила'],[
        ['CheckResult','isValid: Boolean; errors: коллекция {slotId?: UUID, code: String, message: String}. Ошибка содержит проверяемое условие и адресата исправления.'],
        ['CalculationResult','costRub: Decimal, priceRub: Decimal, massKg: Decimal, computedAt: DateTime. Рубли и килограммы; цена и масса положительны.'],
        ['PreparedRequest','revisionId: UUID; payload: JSON; checksum: String. Запрос строится только из одного сохранённого снимка.'],
        ['ExternalResult','kind: ACCEPTED / SUCCESS / REJECTED / TRANSIENT / UNKNOWN; taskId?: String; errors: коллекция замечаний. ACCEPTED не равен SUCCESS.'],
        ['JSON-параметр formVersion','configurationId: UUID; slots: коллекция {roleCode, selectionMode, groupId?, componentId?, quantity}. Для EXACT задана конкретная модель; для GROUP задана группа.'],
    ],[4.4,12.1])
    p('Количество позиции — целое число больше нуля; цена предложения неотрицательна; наличие — целое число не меньше нуля. В черновике допускаются незаполненные позиции. Для состояния CHECKED каждая обязательная роль присутствует, группа разрешена в конкретную комплектующую, выбрано актуальное предложение той же модели и проверена совместимость. Для EXACT group отсутствует; для GROUP выбранная модель удовлетворяет ограничениям группы. Внутренние статусы не являются названиями полей внешнего API.')
    relations=[]
    for relation in [e for e in elements if e['kind']=='association']:
        ends=sorted([e for e in elements if e['kind']=='end' and e['properties']['owner']==relation['id']],key=lambda e:int(e['properties']['order']))
        def multiplicity(e):
            props=e['properties']; return props['lower'] if props['lower']==props['upper'] else props['lower']+'..'+props['upper']
        relations.append([by_id[ends[0]['properties']['classifier']]['label'],multiplicity(ends[0])+' — '+multiplicity(ends[1]),'Композиция' if any(e['properties'].get('aggregation')=='composite' for e in ends) else 'Ассоциация',by_id[ends[1]['properties']['classifier']]['label']])
    table('Кратности связей классов-сущностей',['Класс','Кратность','Отношение','Связанный класс'],relations,[4.6,2.4,3.0,6.5])
    p('Ревизия всегда ссылается ровно на одну версию BOM. Цена, масса, атрибуты и контрольные суммы изображений входят в снимок. При изменении исходных данных создаётся новая ревизия; уже отправленный payload и прежние попытки не перезаписываются. У композиционных объектов единственный владелец; удаление истории публикаций физическим каскадом не допускается.')
    table('Соответствие сообщениям модели анализа',['Сценарий / сообщение','Операция проектирования','Результат'],[
        ['UC02: сохранить состав, получить предложения','BomControl:\nformVersion /\nresolveSlots','Черновик версии; конкретные модели и выбранные предложения'],
        ['UC02: проверить совместимость','Compatibility:\ncheckCompatibility','CheckResult; CHECKED только при отсутствии ошибок'],
        ['UC05: получить снимок, проверить готовность','CardRevision: getSnapshot;\nReadiness: checkReadiness','Снимок и проверяемое решение о возможности отправки'],
        ['UC05: создать запись отправки, передать снимок','PublicationControl:\nrequestPublication;\nMarketplaceGateway:\nprepareRequest / send','Попытка до обмена; ACCEPTED либо ошибка, но не фиктивный успех'],
        ['UC05: сохранить результат, обновить статус','PublicationAttempt:\nrecordResult;\nPublicationControl:\ncheckOutcome','Окончательный итог либо дальнейшая сверка UNKNOWN'],
    ],[5.1,5.8,5.6])
    heading=p('4 Диаграмма деятельности',heading=1)
    heading.paragraph_format.page_break_before = True
    figure('05_ПР4_05_Деятельность_BOM.png','Деятельность: подготовка и проверка состава',new_page=False)
    figure('06_ПР4_06_Деятельность_Карточка.png','Деятельность: подготовка содержания и передача на публикацию')
    p('Две первые ветви независимы: технический специалист задаёт требования к позициям, а система обновляет каталог предложений для выбранных поставщиков. Ветвь загрузки не получает объект BOM из параллельной ветви. После AND-синхронизации выбираются конкретные модели и предложения, проверяется совместимость и готовится содержание. Ошибки загрузки сохраняются как результат ветви и блокируют выбор непроверенных предложений; токен завершения ветви не теряется.')
    table('Действия подготовки ревизии',['Действие','Исполнитель','Результат'],[
        ['Задать позиции','Технический специалист','Роли, количество, группы либо точные модели'],
        ['Обновить предложения','Система / администратор интеграций','Данные с ценой, наличием, временем и диагностикой загрузки'],
        ['Разрешить слоты','Управление BOM','Конкретные модели и предложения; сохранён черновик'],
        ['Проверить совместимость','Проверка совместимости','Проверенная BOM либо адресные конфликты'],
        ['Рассчитать показатели','Расчёт показателей','Себестоимость, цена, масса'],
        ['Подготовить содержание','Дизайнер / менеджер / система','Изображения, название и обязательные атрибуты'],
        ['Проверить ревизию','Проверка готовности','Полный снимок либо замечания к черновику'],
        ['Передать на публикацию','Менеджер / управление публикацией','Записана попытка, отправлен снимок; публикация ещё не подтверждена'],
    ],[5.2,4.5,6.8])
    p('Положительное завершение деятельности означает передачу готовой ревизии модулю публикации, а не появление карточки на маркетплейсе. Продолжение внешнего обмена описано автоматом состояний. Отрицательная ветвь завершает текущий запуск и сохраняет черновик; исправление и повторная проверка — новый запуск.')
    section(True)
    p('5 Диаграммы состояний',heading=1)
    figure('ПР4_07_Состояния_Карточка.png','Состояния карточки с составным состоянием внешнего обмена',landscape=True,new_page=False)
    figure('ПР4_08_Состояния_Обмен.png','Состояния попытки внешнего обмена',landscape=True)
    section()
    p('Объект жизненного цикла — карточка с выбранной текущей ревизией. История прежних ревизий остаётся неизменной; состояние обмена уточняется в попытках. «Обмен с Ozon» — составное состояние, раскрытое отдельно. После приёма задачи карточка ожидает окончательного результата. Только подтверждённое создание или обновление даёт состояние «Опубликована».')
    transitions=json.loads((figures/'ПР4_Переходы.json').read_text(encoding='utf-8'))
    table('Переходы жизненного цикла',['ID','Исходное состояние','Событие и условие','Новое состояние'],transitions,[1.2,3.7,7.9,3.7])
    p('При UNKNOWN запрещена слепая повторная отправка: сначала выполняется сверка по известному внешнему идентификатору или идентификатору товара. Если исход надёжно установить нельзя, попытка остаётся на ручном разборе. Автоматически повторяются только подтверждённые временные отказы без принятия операции, с увеличением задержки и ограничением количества. Валидационные ошибки и отказы в доступе требуют исправления, а не бесконечной очереди.')
    p('Идемпотентность не предполагается свойством любого внешнего метода. Внутренний идентификатор попытки предотвращает дублирование заданий в своей очереди, но сам по себе не гарантирует отсутствие дубля у внешней системы. После изменения содержания либо правил готовность проверяется заново. Если лимит повторов исчерпан, операция остаётся в журнале для ответственного.')
    p('6 Согласованность проектной модели',heading=1)
    table('Связь с функциями IDEF0 и требованиями',['Функция / требование','Классы и правило'],[
        ['A1 — инициировать карточку','Формы создают запрос и черновик; границы приёма запроса из практики 2 сохранены.'],
        ['A2 — сформировать BOM','Управление BOM, конфигурация, версия, позиции, группы, комплектующие и предложения.'],
        ['A3 — проверить совместимость','Проверка совместимости; конфликты блокируют CHECKED и отправку.'],
        ['A4 — рассчитать и подготовить контент','Расчёт показателей, ревизия и изображения; содержание подготовлено до готовности.'],
        ['A5 — проверить и опубликовать','Проверка готовности, управление публикацией, шлюз и попытки.'],
        ['FR-19 / ФТ-04 — неуспешный обмен','TRANSIENT, REJECTED и UNKNOWN разделены; повтор допускается только после проверки условий.'],
        ['ФТ-02 — массовая замена','replacePreview показывает затронутые конфигурации; подтверждение создаёт новые версии, старые снимки сохраняются.'],
    ],[5.0,11.5])
    p('В модели анализа UC05 внешний обмен представлен обобщённой операцией. Проектная модель раскрывает эту операцию на отправку и уточнение результата, сохраняя инициатора, снимок, регистрацию попытки и ответственность адаптера. Синхронный ответ конкретного запроса не означает синхронное завершение всей внешней обработки. Детальные адреса и поля API в данной работе не специфицируются.')
    p('7 Вывод',heading=1)
    p('Разработана модель проектирования из 18 классов с типизированными атрибутами и операциями, ключами, кратностями и инвариантами группового либо точного выбора. Диаграмма деятельности различает независимые ветви, зависимые проверки, расчёты, подготовку содержания и передачу готовой ревизии.')
    p('Автомат карточки и попытки обмена разделяет подтверждённую публикацию, ожидание, окончательный отказ, временную ошибку и неизвестный исход. Ревизии и попытки сохраняют историю; внешние повторы ограничены и не подменяют исправление данных. Проектные решения прослеживаются к требованиям, IDEF0 и сценариям предыдущих практик.')
    p('Список использованных источников',heading=1)
    for text in [
        '1. Проектирование предметно-ориентированных информационных систем. Практическая работа № 4. Методические материалы.',
        '2. Албахтин И.В. ППОИС. Отчёт по практической работе № 1. Москва, 2026.',
        '3. Албахтин И.В. ППОИС. Отчёт по практической работе № 2. Москва, 2026.',
        '4. Албахтин И.В. ППОИС. Отчёт по практической работе № 3. Москва, 2026.',
    ]: p(text,indent=0,align=WD_ALIGN_PARAGRAPH.LEFT)
    if [etree.tostring(child) for child in list(body)[:len(prefix)]] != title_xml:
        raise ValueError('Title prefix changed')
    working.mkdir(parents=True,exist_ok=True)
    intermediate=working/'PPOIS4_generated.docx'; doc.save(intermediate)
    result=merge_package(source,intermediate,destination,{'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'})
    result.update({'figures':figure_number,'tables':table_number,'classes':len(descriptions),'transitions':len(transitions),'title_prefix_preserved':True})
    return result


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--baseline',type=Path,required=True)
    parser.add_argument('--figures',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--working',type=Path,required=True)
    args=parser.parse_args(); args.output.mkdir(parents=True,exist_ok=True)
    replacements={
        1:{
            'ОТЧЕТ ПО ПРАКТИЧЕСКИМ РАБОТАМ':'ОТЧЕТ ПО ПРАКТИЧЕСКОЙ РАБОТЕ № 1',
            'Общество с ограниченной ответственностью «Юкомс» зарегистрировано в Москве и осуществляет деятельность, связанную с компьютерной техникой, электронной торговлей и информационными технологиями. Для рассматриваемого проекта существенным направлением является формирование и продажа готовых компьютерных систем через электронные каналы. Компания использует сведения поставщиков комплектующих и публикует товарные предложения на маркетплейсе Ozon [1-3].':'Объект учебного проекта — формирование и продажа компьютерных систем ООО «Юкомс». Открытый сайт компании описывает компьютеры и электронные каналы продаж [2]. Внутренние роли, документооборот и связи с PC4Games и ITPartner ниже приняты как проектные предположения для автоматизации, а не результаты фактического интервью или доступа к внутренним регламентам.',
            'Инициирующее событие процесса - руководитель или менеджер принимает решение создать новое товарное предложение либо изменить существующее. Завершающее событие - Ozon принял новую или обновлённую карточку, а система сохранила результат публикации. При ошибке процесс завершается только после регистрации проблемы и назначения карточке статуса «ошибка публикации».':'Инициирующее событие — решение руководителя или менеджера создать либо изменить товарное предложение. Успешное завершение — подтверждено создание или обновление карточки и сохранён окончательный результат. Приём задачи без окончательного результата означает ожидание. Отказ, временная ошибка и неизвестный исход регистрируются отдельно; неизвестный исход требует сверки, а не слепой повторной отправки.',
            'Система должна помещать неуспешную внешнюю операцию в очередь повторной отправки.':'Система должна разделять окончательные, временные ошибки и неизвестный исход. Повторная отправка допускается только для подтверждённых временных отказов без принятия операции, с задержкой и лимитом попыток. Ошибки данных исправляются; неизвестный исход сначала сверяется.',
            'В ходе практической работы выполнено предварительное обследование ООО «Юкомс» в части формирования и актуализации карточек товаров компьютерных систем. Определены участники, границы и проблемы процесса. Основной причиной проектирования специализированной ИС является необходимость согласованно изменять BOM, точные модели поставщиков и опубликованные карточки при сохранении технической совместимости и актуальности цены.':'В ходе практической работы подготовлено учебное предпроектное описание процесса формирования и актуализации карточек компьютерных систем ООО «Юкомс». По открытым материалам и явно заданным проектным предположениям определены участники, границы и проблемы. Необходимость специализированной ИС связана с согласованным изменением BOM, моделей поставщиков и карточек при сохранении совместимости и актуальности цены. Фактические интервью и результаты внедрения не утверждаются.',
        },
        2:{
            'На диаграмме A0 основной процесс разложен на пять взаимосвязанных функций. Результат каждой функции служит входом следующей; публикация допускается только после успешного прохождения предыдущих этапов. Декомпозиция показана на рисунке 2.':'На рисунке 2 показана декомпозиция на пять функций. Публикация допускается только после проверки ревизии.',
            'После успешной проверки рассчитываются цена и вес, по шаблонам формируются изображения. Менеджер проверяет обязательные атрибуты и публикует карточку через API Ozon. При ошибке система сохраняет ответ в журнале и формирует повторную попытку. Процесс завершается принятой Ozon карточкой либо зарегистрированной ошибкой с назначенным ответственным.':'После проверки рассчитываются цена и масса, готовятся изображения и обязательные атрибуты. Менеджер подтверждает отправку сохранённой ревизии. Система регистрирует попытку до обмена, сохраняет ответ и уточняет окончательный результат. Приём внешней задачи не означает публикацию. Окончательный отказ требует исправления; подтверждённый временный отказ без принятия операции допускает ограниченный повтор с задержкой; неизвестный исход — только сверку. Неоконченный обмен сохраняется в состоянии ожидания с ответственным.',
            'Ошибка помещается в очередь':'Временный подтверждённый отказ — ограниченный повтор; ошибка данных — исправление; неизвестный исход — сверка',
            'Опубликованная карточка или зарегистрированная ошибка':'Подтверждённая публикация, ожидающая сверки попытка либо классифицированная ошибка',
        },
        3:{
            'В UC05 попытка регистрируется до отправки. Адаптер представляет обмен с Ozon как одну операцию анализа; при отложенной обработке итог уточняется отдельно. Подтверждённый приём даёт статус «опубликована», отказ — «ошибка публикации», отсутствие окончательного ответа — «ожидает подтверждения».':'В UC05 попытка регистрируется до отправки. Адаптер представляет внешний обмен как одну обобщённую операцию анализа. Приём задачи без окончательного результата даёт «ожидает подтверждения», а не «опубликована». Только подтверждённое создание или обновление карточки означает успех. Отказ требует классификации; неизвестный исход сначала сверяется. Проектная модель в практике 4 детализирует эту операцию, не меняя инициатора и ответственности участников.',
        },
    }
    log={}
    for number, mapping in replacements.items():
        name=f'ППОИС_{number}_АлбахтинИВ.docx'
        patch_text(args.baseline/name,args.output/name,mapping)
        package=format_existing(args.output/name,args.working,args.figures if number in (2,3) else None)
        log[name]={'replaced_paragraphs':len(mapping),'only_document_xml_changed':True,'formatting':package}
        if number==2:
            log[name]['changed_media']=replace_idef0_images(args.output/name,args.figures)
            log[name]['only_document_xml_changed']=False
        if number==3:
            log[name]['changed_media']=replace_analysis_images(args.output/name,args.figures)
            log[name]['only_document_xml_changed']=False
    name='ППОИС_4_АлбахтинИВ.docx'
    log[name]=build_four(args.baseline/name,args.output/name,args.figures,args.working)
    for name in log:
        log[name]['package_cleanup']=prune_unused_images(args.output/name)
    (args.working/'revision_log.json').write_text(json.dumps(log,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(log,ensure_ascii=False,indent=2))


if __name__=='__main__': main()
