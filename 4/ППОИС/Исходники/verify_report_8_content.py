"""Reconcile the saved report with executed data, figures and native equations.

The assignment and every rendered page must also be read. This audit does not
claim that mockups implement a client, nor that an external publication ran.
"""
from __future__ import annotations

from datetime import datetime, timezone
from hashlib import sha256
import json
import re
import unicodedata

from docx import Document
from docx.oxml.ns import nsmap, qn
from docx.table import Table
from lxml.etree import XPath

from build_report_7 import physical_type
from build_report_8 import DESCRIPTIONS, FORMS, MATH, ROOT, SCHEMES, SUBJECT, TEST_NAMES, digest, inputs


def normalize(value):
    return re.sub(r'\s+', '', unicodedata.normalize('NFKC', value).replace('−', '-'))


def xp(element, expression):
    return XPath(expression, namespaces=nsmap)(element)


def math_text(element):
    return normalize(''.join(xp(element, './/m:t/text()')))


def inspect_content(document, db, math, models):
    checks = []

    def check(name, condition):
        checks.append({'name': name, 'passed': bool(condition)})

    tables = {}
    for p in document.paragraphs:
        caption = re.match(r'Таблица (\d+)\s+[–—-]\s+(.+)', p.text)
        if caption:
            sibling = p._p.getnext()
            check('caption immediately precedes table ' + caption[1],
                  sibling is not None and sibling.tag == qn('w:tbl'))
            if sibling is not None and sibling.tag == qn('w:tbl'):
                tables[caption[2]] = [[c.text for c in row.cells]
                                      for row in Table(sibling, document._body).rows]

    def rows(title):
        return tables.get(title, [])[1:]

    check('exactly five content tables', len(tables) == 5)
    form_rows = rows('Назначение входных и результатных форм')
    expected_forms = [[f['id'], f['title'],
                       'Результат расчёта' if f['id'] == 'F7' else
                       'Результат внешней операции' if f['id'] == 'F8' else
                       'Ввод и просмотр сохранённых данных'] for f in models['forms']]
    check('all eight named forms and input/result kinds', form_rows == expected_forms)
    paragraphs = {normalize(p.text) for p in document.paragraphs}
    for form_id, (title, first, second) in DESCRIPTIONS.items():
        check('description and restrictions ' + form_id,
              all(normalize(s) in paragraphs for s in (form_id + ' ' + title, first, second)))

    mappings = {(m['table'] + '.' + m['column']): m for m in models['mapping']}
    columns = {(c['table_name'] + '.' + c['column_name']): c for c in db['columns']}
    matrix = rows('Связь всех полей БД с макетами')
    check('every actual field exactly once',
          len(matrix) == 71 and len(columns) == 71 and {r[0] for r in matrix} == set(columns))
    for row in matrix:
        if row[0] not in columns:
            check('unknown physical field ' + row[0], False)
            continue
        c = columns[row[0]]
        m = mappings[row[0]]
        mode = 'С' if m['mode'].startswith('Системное') else 'А' if 'администратором' in m['mode'] else 'В'
        check('type optionality form and mode ' + row[0],
              row[1:] == [m['form'], physical_type(c), 'Да' if c['is_nullable'] == 'NO' else 'Нет', mode])

    expected_lines = [[s['sku'], str(s['quantity']), s['price_rub'], s['line_cost_rub'],
                       s['mass_kg'], s['line_mass_kg']] for s in math['selected']]
    check('every executed quantity price and mass', rows('Позиции контрольного расчёта') == expected_lines)
    expected_totals = [[label, math['control'][key], sql, 'Совпадает']
                       for label, key, sql in zip(
                           ('Себестоимость, ₽', 'Цена, ₽', 'Масса, кг'),
                           ('cost_rub', 'price_rub', 'mass_kg'), math['sql_control'])]
    check('all Python and SQL control values', rows('Сопоставление Python и PostgreSQL') == expected_totals)
    expected_tests = [[str(i), TEST_NAMES[t['test'].split('.test_')[-1]], 'Пройдено']
                      for i, t in enumerate(math['tests'], 1)]
    check('all actual 26 test descriptions and results',
          math['passed'] and all(t['passed'] for t in math['tests'])
          and len(expected_tests) == 26 and rows('Результаты выполненных доменных тестов') == expected_tests)
    check('read-only SQL reconciliation was actually performed',
          math['runtime']['role'] == 'ppois_viewer' and math['runtime']['read_only'] == 'on')

    equations = document.element.xpath('//m:oMath')
    expected_text = [
        'Agej=t-tj', 'Ei={j∈Ji:0≤Agej≤τ}', 'ji=lexmin{(pj,uj,hj):j∈Ei}',
        'Cкомп=i=1nqipi', 'C=Cкомп+A+B', 'P0=C(1+m)+L1-c', 'P=P0dd',
        'W=i=1nqiwi+b', 'HБП≥(1+β)i∈KqiHi',
        'R=(s=TEMP_ERROR)∧f∧(1≤a<k)', '∧Z(tnext)∧Z(tnow)∧(tnext≤tnow)',
    ]
    check('eleven native editable equations', len(equations) == 11)
    for i, (eq, expected) in enumerate(zip(equations, expected_text), 1):
        check(f'equation {i} complete ordered text', math_text(eq) == normalize(expected))
        check(f'equation {i} no embedded image fallback', not xp(eq, './/w:drawing'))
    if len(equations) == 11:
        for eq_number, numerator, denominator in ((6, 'C(1+m)+L', '1-c'), (7, 'P0', 'd')):
            fractions = xp(equations[eq_number - 1], './/m:f')
            check(f'equation {eq_number} native fraction', len(fractions) == 1)
            if fractions:
                check(f'equation {eq_number} numerator and denominator',
                      math_text(xp(fractions[0], './m:num')[0]) == numerator
                      and math_text(xp(fractions[0], './m:den')[0]) == denominator)
        check('ceil not floor or ordinary brackets',
              xp(equations[6], './/m:dPr/m:begChr/@m:val') == ['⌈']
              and xp(equations[6], './/m:dPr/m:endChr/@m:val') == ['⌉'])
        for eq_number, lower, upper, term in ((4, 'i=1', 'n', 'qipi'), (8, 'i=1', 'n', 'qiwi'),
                                              (9, 'i∈K', '', 'qiHi')):
            sums = xp(equations[eq_number - 1], './/m:nary')
            check(f'equation {eq_number} one summation', len(sums) == 1)
            if sums:
                check(f'equation {eq_number} symbol limits and term',
                      xp(sums[0], './m:naryPr/m:chr/@m:val') == ['∑']
                      and math_text(xp(sums[0], './m:sub')[0]) == lower
                      and math_text(xp(sums[0], './m:sup')[0]) == upper
                      and math_text(xp(sums[0], './m:e')[0]) == term)
        expected_subscripts = [
            [('Age', 'j'), ('t', 'j')], [('E', 'i'), ('J', 'i'), ('Age', 'j')],
            [('j', 'i'), ('p', 'j'), ('u', 'j'), ('h', 'j'), ('E', 'i')],
            [('C', 'комп'), ('q', 'i'), ('p', 'i')], [('C', 'комп')], [('P', '0')], [('P', '0')],
            [('q', 'i'), ('w', 'i')], [('H', 'БП'), ('q', 'i'), ('H', 'i')], [],
            [('t', 'next'), ('t', 'now'), ('t', 'next'), ('t', 'now')],
        ]
        for i, (eq, expected) in enumerate(zip(equations, expected_subscripts), 1):
            actual = [(math_text(xp(s, './m:e')[0]), math_text(xp(s, './m:sub')[0]))
                      for s in xp(eq, './/m:sSub')]
            check(f'equation {i} all native subscripts', actual == expected)

    # Inspect definitions in actual document order, not just in the builder.
    prefix = ''; eq_number = 0
    definitions = {
        1: ('Пусть n', 'i — индекс позиции', 'qᵢ — требуемое', 'Индекс j', 'Jᵢ — множество',
            'pⱼ измеряется', 'uⱼ — код', 'hⱼ — его UUID', 'Момент t', 'tⱼ — время', 'τ — допустимый',
            'Eᵢ обозначает', 'Ageⱼ — их возраст'),
        3: ('jᵢ — индекс выбранного', 'Функция lexmin возвращает'),
        4: ('pᵢ — цена единицы', 'wᵢ — масса единицы', 'A — стоимость', 'B — стоимость',
            'L — логистические', 'm — доля', 'c — доля', 'd — положительный', 'b — неотрицательная',
            'Cкомп — стоимость', 'C — себестоимость', 'P₀ — цена', 'P — итоговая', 'W — масса'),
        9: ('K — множество индексов', 'Hᵢ — мощность', 'HБП — мощность', 'β — доля'),
        10: ('s — внутреннее', 'f — истинный', 'a — номер', 'k — предел', 't_next — назначенный',
             't_now — текущее', 'Z(t) означает', 'R — логический'),
    }
    for p in document.paragraphs:
        for _eq in p._p.xpath('.//m:oMath'):
            eq_number += 1
            for definition in definitions.get(eq_number, ()):
                check('definition precedes equation ' + str(eq_number) + ': ' + definition,
                      normalize(definition) in normalize(prefix))
        prefix += '\n' + p.text

    used_image_ids = set(document.element.xpath('//@r:embed'))
    media = {sha256(rel.target_part.blob).hexdigest() for key, rel in document.part.rels.items()
             if key in used_image_ids and rel.reltype.endswith('/image')}
    check('all eleven current form and algorithm images embedded',
          len(models['emitted']) == 11 and all(e['sha256'] in media for e in models['emitted']))
    check('all eight form headings and three algorithm captions',
          all(normalize(f"Рисунок {i} – {f['id']} {f['title']}") in paragraphs
              for i, f in enumerate(models['forms'] + models['algorithms'], 1)))
    return checks


def verify():
    db, math, models = inputs()
    path = SUBJECT / 'ППОИС_8_АлбахтинИВ.docx'
    checks = inspect_content(Document(path), db, math, models)
    qa_path = ROOT / '.cache/ppois5-10/report-evidence/practice8-final-qa.json'
    qa = json.loads(qa_path.read_text(encoding='utf-8'))
    checks.append({'name': 'current separate layout and complete visual inspection',
                   'passed': qa['status'] == 'passed' and qa['docx_sha256'] == digest(path)
                   and qa['pdf_sha256'] == digest(path.with_suffix('.pdf'))
                   and len(qa['pages']) == 28 and all(p['manually_viewed'] for p in qa['pages'])})
    result = {'status': 'passed' if all(c['passed'] for c in checks) else 'failed',
              'verified_at': datetime.now(timezone.utc).isoformat(),
              'docx_sha256': digest(path), 'pdf_sha256': digest(path.with_suffix('.pdf')),
              'database': db['database'], 'database_evidence_sha256': math['database_evidence_sha256'],
              'mathematics_sha256': digest(MATH), 'forms_sha256': digest(FORMS),
              'verifier_sha256': digest(__file__),
              'assignment': 'practice 8 supplied PDF items 1–3; its one page visually read',
              'scope': 'saved artifact reconciliation; mockups and bounded domain tests, not external API or client tests',
              'checks': checks}
    output = ROOT / '.cache/ppois5-10/report-evidence/practice8-content-qa.json'
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'status': result['status'], 'checks': len(checks),
                      'failed': [c['name'] for c in checks if not c['passed']]}, ensure_ascii=False))
    return result['status'] == 'passed'


if __name__ == '__main__':
    raise SystemExit(0 if verify() else 1)
