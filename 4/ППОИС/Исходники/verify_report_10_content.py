"""Reconcile the saved PPOIS10 report with its client contract and observations.

This read-only audit covers both tasks in the supplied three-page assignment.
The separate layout audit must record a complete review of the current PDF.
No new browser execution or external publication is inferred from this check.
"""
from __future__ import annotations

from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from pypdf import PdfReader

from build_report_10 import (CLIENT, FUNCTION_NOTES, MODELS, PROOF, ROOT,
                             ROUTE_NOTES, SCHEMES, SUBJECT, digest,
                             function_label, inputs)


def compact(value):
    return re.sub(r'\s+', ' ', value).strip()


def inspect_content(document, contract, proof, models, figures):
    checks = []

    def check(name, condition):
        checks.append({'name': name, 'passed': bool(condition)})

    tables = {}
    captions = []
    for paragraph in document.paragraphs:
        match = re.match(r'Таблица (\d+)\s+[–—-]\s+(.+)', paragraph.text)
        if match:
            sibling = paragraph._p.getnext()
            check('caption precedes table ' + match[1],
                  sibling is not None and sibling.tag == qn('w:tbl'))
            if sibling is not None and sibling.tag == qn('w:tbl'):
                table = Table(sibling, document._body)
                tables[match[2]] = [[compact(cell.text) for cell in row.cells]
                                    for row in table.rows]
        if paragraph.text.startswith('Рисунок '):
            captions.append(compact(paragraph.text))

    def rows(title):
        if title not in tables:
            check('required table ' + title, False)
            return []
        return tables[title][1:]

    for branch, title in [('main', 'Основные функции'), ('service', 'Служебные функции')]:
        expected = [[f['id'], compact(function_label(f)), compact(FUNCTION_NOTES[f['id']])]
                    for f in contract['functions'] if f['branch'] == branch]
        check(title + ' matches current functions and scope', rows(title) == expected)

    roles = list(contract['roles'])
    matrix = rows('Матрица доступных функций')
    check('role matrix covers all twelve functions exactly once',
          len(matrix) == 12 and [r[0] for r in matrix] == [f['id'] for f in contract['functions']])
    for function in contract['functions']:
        actual = [r for r in matrix if r[0] == function['id']]
        expected = [function['id']] + ['Да' if function['id'] in contract['roles'][role]['functions']
                                      else 'Нет' for role in roles]
        check('effective permissions for ' + function['id'], actual == [expected])

    expected_routes = [[route['method'] + ' ' + route['path'], route['function'],
                        compact(ROUTE_NOTES[route['path']])] for route in contract['routes']]
    check('all twelve implemented HTTP routes and permissions',
          rows('HTTP-маршруты клиента') == expected_routes)
    check('all eighteen observed dialog checkpoints',
          rows('Контрольные точки пользовательского диалога') ==
          [[point['id'], compact(point['description'])] for point in proof['browser_checks']])
    check('all twelve database counts before and after the observed scenario',
          rows('Состав данных до и после браузерного сценария') ==
          [[name, str(count), str(proof['final_counts'][name])]
           for name, count in proof['initial_counts'].items()])
    check('canonical tree and four role views verified against the contract',
          models['canonical_tree']['unique_parent'] and models['canonical_tree']['reachable'] == 24
          and len(models['sheets']) == 11
          and models['roles'] == {role: spec['functions'] for role, spec in contract['roles'].items()})

    used = set(document.element.xpath('//@r:embed'))
    image_hashes = {sha256(rel.target_part.blob).hexdigest()
                    for key, rel in document.part.rels.items()
                    if key in used and rel.reltype.endswith('/image')}
    check('exactly twelve required figure captions', len(captions) == len(figures) == 12)
    for figure in figures:
        check('current image embedded for figure ' + str(figure['number']),
              digest(SUBJECT / figure['source']) in image_hashes)
        expected = compact(f"Рисунок {figure['number']} – {figure['title']}")
        check('caption of figure ' + str(figure['number']), expected in captions)
    full_screens = {'ПР10_Экран_входа.jpg', 'ПР10_Экран_инженера.jpg', 'ПР10_Результат_расчёта.jpg',
                    'ПР10_История_ревизий.jpg', 'ПР10_Экран_администратора.jpg'}
    check('all five screenshots use full source frames',
          {Path(item['source']).name for item in figures if item['source'].endswith('.jpg')} == full_screens)
    check('no screenshot uses nonzero Word crop',
          all(all(int(value) == 0 for value in rect.attrib.values())
              for rect in document.element.xpath('//a:srcRect')))

    text = compact('\n'.join(p.text for p in document.paragraphs))
    for heading in ['Задание 1 Дерево функций и доступ пользователей',
                    'Задание 1 Модели сценариев диалога', 'Задание 2 Структурная схема ИС']:
        check('required task section ' + heading, heading in text)
    for value in ['53 500,00', '76 400,00', '7,820', '48 + 1 + 6 + 1 = 56',
                  '05.10.2026 10:00 +03:00', '28 автоматических тестов', '111 проверок']:
        check('observed result stated in report ' + value, value in text)
    check('the scope retains DRAFT and bounded external exchange',
          'добавляет ревизию DRAFT' in text and 'утверждение и публикация в этой итерации не реализованы' in text)
    check('no unrelated ensemble or Colab requirements',
          not re.search(r'ансамблев|бэггинг|бустинг|Colab', text, re.I))
    return checks


def verify():
    contract, proof, models = inputs()
    path = SUBJECT / 'ППОИС_10_АлбахтинИВ.docx'
    inventory = json.loads((ROOT / '.cache/ppois5-10/report-evidence' /
                            (path.stem + '.json')).read_text(encoding='utf-8'))
    figures = [item for item in inventory['objects'] if item['type'] == 'figure']
    checks = inspect_content(Document(path), contract, proof, models, figures)
    qa = json.loads((ROOT / '.cache/ppois5-10/report-evidence/practice10-final-qa.json')
                    .read_text(encoding='utf-8'))
    checks.append({'name': 'current separate formatting audit and complete visual review',
                   'passed': qa['status'] == 'passed' and qa['docx_sha256'] == digest(path)
                   and qa['pdf_sha256'] == digest(path.with_suffix('.pdf'))
                   and len(qa['pages']) == len(PdfReader(path.with_suffix('.pdf')).pages)
                   and all(p['manually_viewed'] for p in qa['pages'])})
    result = {'status': 'passed' if all(c['passed'] for c in checks) else 'failed',
              'verified_at': datetime.now(timezone.utc).isoformat(),
              'docx_sha256': digest(path), 'pdf_sha256': digest(path.with_suffix('.pdf')),
              'contract_sha256': digest(CLIENT / 'client_contract.json'),
              'models_sha256': digest(MODELS), 'client_evidence_sha256': digest(PROOF),
              'verifier_sha256': digest(Path(__file__)),
              'assignment': 'PPOIS practice 10, supplied three-page PDF, tasks 1 and 2',
              'scope': 'saved report, current contract and recorded observations; no new browser-run claim',
              'checks': checks}
    output = ROOT / '.cache/ppois5-10/report-evidence/practice10-content-qa.json'
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({'status': result['status'], 'checks': len(checks),
                      'failed': [c['name'] for c in checks if not c['passed']]}, ensure_ascii=False))
    return result['status'] == 'passed'


if __name__ == '__main__':
    raise SystemExit(0 if verify() else 1)
