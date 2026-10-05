"""Structural regression checks for the narrow OMML path before Word QA."""
import unittest
from docx import Document
from docx.oxml.ns import qn, nsmap
from lxml.etree import XPath
from report_math import delimiter, element, fraction, sub, summation, text


def select(node, expression):
    # Generic m:* nodes are lxml elements, unlike specialised w:* classes:
    # explicitly bind namespaces rather than assuming python-docx's override.
    return XPath(expression, namespaces=nsmap)(node)


class MathTests(unittest.TestCase):
    def test_fraction_preserves_both_expressions(self):
        value = fraction([sub('P', '0')], [text('d')])
        self.assertEqual(value.tag, qn('m:f'))
        self.assertEqual(select(value, './m:num/m:sSub/m:sub/m:r/m:t/text()'), ['0'])
        self.assertEqual(select(value, './m:den/m:r/m:t/text()'), ['d'])

    def test_sum_has_defined_limits_and_body(self):
        value = summation([sub('q', 'i'), sub('p', 'i')])
        self.assertEqual(select(value, './m:sub/m:r/m:t/text()'), ['i = 1'])
        self.assertEqual(select(value, './m:sup/m:r/m:t/text()'), ['n'])
        self.assertEqual(len(select(value, './m:e/m:sSub')), 2)

    def test_ceiling_is_structured_not_literal_latex(self):
        value = delimiter([fraction([sub('P', '0')], [text('d')])], '⌈', '⌉')
        self.assertEqual(select(value, './m:dPr/m:begChr/@m:val'), ['⌈'])
        self.assertEqual(select(value, './m:dPr/m:endChr/@m:val'), ['⌉'])
        self.assertEqual(len(select(value, './m:e/m:f')), 1)

    def test_children_are_not_moved_or_shared(self):
        token = text('C')
        a, b = element('oMath', token), element('oMath', token)
        select(a[0], './m:t')[0].text = 'W'
        self.assertEqual(select(b, './m:r/m:t/text()'), ['C'])
        self.assertIsNone(token.getparent())

    def test_package_round_trip(self):
        import io
        document = Document()
        paragraph = document.add_paragraph()
        paragraph._p.append(element('oMath', text('P = '),
                                    delimiter([fraction([sub('P', '0')], [text('d')])], '⌈', '⌉'), text('d')))
        stream = io.BytesIO()
        document.save(stream)
        stream.seek(0)
        restored = Document(stream)
        self.assertEqual(len(restored.element.xpath('//m:oMath/m:d/m:e/m:f')), 1)
        self.assertEqual(restored.element.xpath('//m:oMath/m:r/m:t/text()'), ['P = ', 'd'])


if __name__ == '__main__':
    unittest.main()
