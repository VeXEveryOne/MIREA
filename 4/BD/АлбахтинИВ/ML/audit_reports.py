"""Structural audit and file registry after native Word export and page review."""
from pathlib import Path
from datetime import datetime,timezone
import hashlib,json,re,zipfile
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader
ROOT=Path(__file__).resolve().parent;BD=ROOT.parents[1]
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def effective(style,group,prop):
    while style is not None:
        value=getattr(getattr(style,group),prop)
        if value is not None:return value
        style=style.base_style
def plain(t):return re.sub(r'\s+','',t)
results={}
for number,expected_pages,figures,tables in [(5,20,8,10),(6,9,5,5),(7,10,3,8)]:
    docx=BD/f'ТИАБД_ПР{number}_АлбахтинИВ.docx';pdf=docx.with_suffix('.pdf');doc=Document(docx)
    reader=PdfReader(pdf);texts=[p.extract_text() for p in reader.pages]
    assert len(texts)==expected_pages
    assert 'Приходько Н.А.' in texts[0] and 'ИНБО-12-23' in texts[0] and 'Албахтин И.В.' in texts[0] and 'Москва 2026' in texts[0]
    assert 'Содержание' in texts[1] and '1 Цель' in texts[2]
    for old in ['ИНБО-01-17','Шендяпин','Воронцов','Содержание обновляется при открытии']:
        assert old not in '\n'.join(texts)
    for sec in doc.sections:
        assert abs(sec.page_width.cm-21)<.02 and abs(sec.page_height.cm-29.7)<.02
        for prop,want in [('left_margin',3),('right_margin',1.5),('top_margin',2),('bottom_margin',2)]:assert abs(getattr(sec,prop).cm-want)<.02
    for name,size,align,first,line in [('Report Body',14,3,1.25,1.5),('Report Table',12,0,0,1),('Report Table Caption',12,0,0,1),('Report Figure Caption',14,1,0,1)]:
        st=doc.styles[name]
        assert effective(st,'font','name')=='Times New Roman'
        assert effective(st,'font','size').pt==size
        assert effective(st,'paragraph_format','alignment')==align
        assert abs(effective(st,'paragraph_format','first_line_indent').cm-first)<.02
        assert effective(st,'paragraph_format','line_spacing')==line
    assert effective(doc.styles['Report Table Caption'],'font','italic')
    xml=doc.element.xml
    assert not doc.element.xpath('.//a:srcRect')
    for f in doc.element.xpath('.//w:rFonts'):
        for key in ['ascii','hAnsi','eastAsia','cs']:
            assert f.get(qn('w:'+key)) in [None,'Times New Roman']
    for shd in doc.element.xpath('.//w:tcPr/w:shd'):
        assert shd.get(qn('w:fill')) in [None,'auto','FFFFFF'] and shd.get(qn('w:val')) in [None,'nil','clear']
    captions=[p.text for p in doc.paragraphs if p.style.name=='Report Figure Caption']
    tablecaps=[p.text for p in doc.paragraphs if p.style.name=='Report Table Caption']
    assert len(captions)==figures and len(tablecaps)==tables
    assert [int(re.match(r'Рисунок (\d+)',t)[1]) for t in captions]==list(range(1,figures+1))
    assert [int(re.match(r'Таблица (\d+)',t)[1]) for t in tablecaps]==list(range(1,tables+1))
    text='\n'.join(texts)
    for i,page in enumerate(texts):
        assert len(page.strip())>10
        assert not '\ufffd' in page
    headings=[p.text for p in doc.paragraphs if p.style.name in ['Heading 1','Heading 2']]
    toc='\n'.join(plain(line) for line in texts[1].splitlines());toc_checked={}
    for heading in headings:
        target=next(i+1 for i,t in enumerate(texts[2:],start=2) if plain(heading) in plain(t))
        match=re.search(re.escape(plain(heading))+r'\.+(\d+)\s*$',toc,re.M)
        assert match and int(match[1])==target,(number,heading,target,match[1] if match else None)
        toc_checked[heading]=target
    run=json.loads((ROOT/f'results/pr{number}/execution.json').read_text())
    verification=json.loads((ROOT/f'results/pr{number}/verification.json').read_text())
    assert run['status']==verification['status']=='passed'
    assert digest(ROOT/f'Practice_{number}.ipynb')==run['notebook_sha256']==verification['notebook_sha256']
    results[str(number)]={'pages':len(texts),'figures':figures,'tables':tables,'toc':toc_checked,'teacher':'Приходько Н.А.','notebook_code_cells':run['code_cells'],'errors':run['errors'],'docx_sha256':digest(docx),'pdf_sha256':digest(pdf),'notebook_sha256':run['notebook_sha256'],'structure_status':'passed'}
    if number==7:
        cloud=json.loads((ROOT/'results/pr7/colab/verification.json').read_text(encoding='utf-8'))
        assert cloud['status']=='passed' and cloud['calculation_cells_match_vm']==5
        assert '3.13.16' in text and '1.6.1' in text and '7 Полный запуск в Google Colab' in text
        results[str(number)]['google_colab']='passed'
if 5 in [int(n) for n in results]:
    q=plain(''.join(PdfReader(BD/'ТИАБД_ПР5_АлбахтинИВ.pdf').pages[i].extract_text() for i in range(13,20)))
    answers=(BD/'ТИАБД_ПР5_Ответы.md').read_text(encoding='utf-8').split('## II')[0]
    questions=re.findall(r'(?m)^(\d+)\. \*\*(.*?)\*\*',answers)
    assert len(questions)==24
    for number,text in questions:assert plain(f'{number}. {text}') in q
ref=BD.parents[1]/'4/ИСУРО/ИСУРО_1_АлбахтинИВ.docx'
assert digest(ref)=='53a4b8ccc73cf44599e8b7c7fb75045afec85dd7f4b0a079b19a0d4ccb5249b1'
registry={}
for directory in [ROOT,BD/'ml_labs']:
    for path in sorted(directory.rglob('*')):
        if path.is_file() and not any(p in ['data','models','work','__pycache__'] for p in path.relative_to(directory).parts) and path.name!='delivery.json':registry[path.relative_to(BD).as_posix()]={'bytes':path.stat().st_size,'sha256':digest(path)}
for pattern in ['ТИАБД_ПР[567]_АлбахтинИВ.*','ТИАБД_ПР[567]_Ответы.md','ТИАБД_Лекция_8_Ответы.md']:
    for path in BD.glob(pattern):registry[path.relative_to(BD).as_posix()]={'bytes':path.stat().st_size,'sha256':digest(path)}
for path in [BD/'АлбахтинИВ/Исходники/build_ml_reports.py',BD/'АлбахтинИВ/README.md',BD/'AGENTS.md',BD/'Материалы/ПР5-8_Проверка_соответствия.md',*[BD/'Методички'/f'Технологии и инструментарий анализа больших данных_Лекции №{n}.pdf' for n in [5,6,7,8]]]:registry[path.relative_to(BD).as_posix()]={'bytes':path.stat().st_size,'sha256':digest(path)}
record={'verified_utc':datetime.now(timezone.utc).isoformat(),'reports':results,'files':registry}
(ROOT/'delivery.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
print(json.dumps(results,ensure_ascii=False,indent=2))
