"""Build personal PR3 from a completed fresh-kernel execution and current evidence."""
from pathlib import Path
from datetime import datetime,timedelta,timezone
import hashlib,json,re,sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm,Pt

PACK=Path(__file__).resolve().parents[1]
BD=PACK.parent
ROOT=BD.parents[1]
LAB=PACK/'Spark'
sys.path.insert(0,str(BD/'spark_lab'))
from report_layout import Report,font,spacing,A
from build_pr2_storage_report import cover

def num(value,digits=2):
    return 'нет данных' if value is None else f'{value:,.{digits}f}'.replace(',',' ').replace('.',',')

def questions(number):
    paragraphs=[p.text for p in Document(BD/'spark_lab/assignment'/f'Практика №{number}.docx').paragraphs]
    start=next(i for i,t in enumerate(paragraphs) if 'Контрольные вопросы' in t)
    return [t for t in paragraphs[start+1:] if re.match(r'^\d+\. ',t)][:16]

def schema_rows(before,after):
    return [[b['name'],f"{b['type']} / {a['type']}",f"{b['nullable']} / {a['nullable']}"] for b,a in zip(before['fields'],after['fields'])]

def environment_table(d,r):
    e=r['environment']
    d.table('Окружение выполнения', ['Параметр','Значение'],[
        ['VM / ОС','TIABD-Albakhtin-IV / Ubuntu Server 22.04.5 LTS'],
        ['Ресурсы','2 vCPU; 4 ГБ RAM; диск 40 ГБ'],['Python',e['python']],
        ['Java',e['java']],['Spark',e['spark']],['Master / parallelism',f"{e['master']} / {e['parallelism']}"],
        ['Driver memory',e['driver_memory']],['Shuffle partitions / AQE',f"{e['shuffle_partitions']} / {e['aqe']}"]],[5.5,11])
    d.p('Local[*] использует два виртуальных ядра одного узла. Методичка рекомендует Python 3.11; фактически применены совместимые Python 3.10, Spark 3.5.7 и JDK 17. Локальный режим сохраняет модель partitions, transformations и actions, но не проверяет сеть физического кластера.')

class PR3Report(Report):
    def __init__(self,number,topic):
        super().__init__(number,topic)
        cover(self.doc)
        for section in self.doc.sections:
            section.page_width=Cm(21);section.page_height=Cm(29.7)
            section.left_margin=Cm(3);section.right_margin=Cm(1.5)
            section.top_margin=Cm(2);section.bottom_margin=Cm(2)
        for level in [1,2]:
            style=self.doc.styles[f'Heading {level}']
            style.element.set(qn('w:customStyle'),'0')
            style.element.find(qn('w:name')).set(qn('w:val'),f'heading {level}')
            ppr=style.element.get_or_add_pPr()
            for old in ppr.findall(qn('w:outlineLvl')):ppr.remove(old)
            outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),str(level-1));ppr.append(outline)
        self.doc.styles['Report Code'].paragraph_format.keep_together=True
        for p in self.doc.paragraphs:
            if p.text.startswith('Москва 2026'): p.paragraph_format.space_before=Pt(0)
        self.doc.core_properties.title='Практическая работа 3 Распределенная модель обработки данных'
    def save(self,path):
        add_toc(self.doc)
        self.doc.save(BD/'ТИАБД_ПР3_АлбахтинИВ.docx')
    def table(self,*args,**kwargs):
        table=super().table(*args,**kwargs)
        for row in table.rows[1:]:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.alignment=A.CENTER if re.fullmatch(r'[\d ,.%-]+',p.text) else A.LEFT
        return table

def add_toc(doc):
    """Insert a real Word field after the cover and start the body on a page."""
    body = doc.element.body
    city = next(p for p in doc.paragraphs if p.text.startswith('Москва 2026'))
    cover_end = list(body).index(city._p)
    first = next(node for node in list(body)[cover_end+1:] if node.tag == qn('w:p')
                 and ''.join(node.xpath('.//w:t/text()')).strip())
    for sec in doc.sections:
        for p in [*sec.header.paragraphs, *sec.first_page_header.paragraphs]:
            p.text = ''
    heading = doc.add_paragraph('Содержание')
    spacing(heading, A.CENTER, 0, 1.5, 0, 6, True)
    heading.paragraph_format.page_break_before = True
    for r in heading.runs: font(r, bold=True)
    field_p = doc.add_paragraph()
    spacing(field_p, A.LEFT, 0, 1, 0, 0)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), ' TOC \\o "1-2" \\h \\z \\u ')
    run = OxmlElement('w:r'); txt = OxmlElement('w:t')
    txt.text = 'Содержание обновляется при открытии документа в Word'
    run.append(txt); field.append(run); field_p._p.append(field)
    first.addprevious(heading._p); first.addprevious(field_p._p)
    first.get_or_add_pPr().append(OxmlElement('w:pageBreakBefore'))
    for name in ['TOC 1', 'TOC 2']:
        if name not in doc.styles:
            from docx.enum.style import WD_STYLE_TYPE
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st = doc.styles[name]; font(st, 14)
        f = st.paragraph_format
        f.first_line_indent = Cm(0); f.left_indent = Cm(0 if name == 'TOC 1' else .5)
        f.right_indent = Cm(0); f.line_spacing = 1; f.space_after = Pt(3)
        f.space_before = Pt(0); f.keep_with_next = False
    settings = doc.settings.element
    for old in settings.findall(qn('w:updateFields')): settings.remove(old)
    update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true'); settings.append(update)

ANSWERS3=[
'local[*] использует доступные логические ядра одного узла. В физическом кластере executors размещены на разных узлах; DataFrame, partitions, задачи и ленивое планирование сохраняются в обоих режимах.',
'Driver создаёт SparkContext, строит планы и планирует задачи. Executors выполняют задачи и хранят промежуточные данные; в local[*] вычислительная часть работает внутри локальной JVM.',
'Partition — логическая часть данных, над которой выполняется задача стадии. Малое число разделов ограничивает параллелизм, большое увеличивает расходы на планирование коротких задач.',
'Transformation строит план: select, filter, withColumn. Action требует результат: count, collect, запись через write. Одно действие может вызвать несколько внутренних Jobs.',
'Lazy evaluation откладывает вычисления до action. Создание withColumn формирует выражение, поэтому его время без action не измеряет чтение, вычисление и обмен данными.',
'CSV хранит текст и разделители, но не строгую схему. StructType полезен при стабильном контракте входа и позволяет избежать отдельного прохода inferSchema; некорректные значения всё равно требуют контроля.',
'Timestamp задаёт временную семантику для datediff, month и сравнения дат. Строковое сравнение зависит от одинакового формата и не проверяет корректность даты.',
'Null может обозначать незавершённую операцию. В Olist большинство отсутствующих дат доставки относится к shipped, canceled и промежуточным статусам; только восемь пропусков относятся к delivered.',
'Select выбирает столбцы и выражения, filter выбирает строки. WithColumn добавляет либо заменяет именованный столбец, а выбор существующего столбца не создаёт нового признака.',
'Рост количества строк может указывать на неуникальный справочник или связь многие-ко-многим. Уменьшение при INNER JOIN может означать потерю фактов без подходящего ключа справа.',
'INNER JOIN сохраняет только совпавшие ключи. LEFT JOIN сохраняет все левые факты и делает отсутствие справочной записи измеримым через null; выбор определяется целью витрины.',
'COUNT(*) считает позиции заказа, COUNT(DISTINCT order_id) — уникальные заказы. Несколько позиций одного заказа дают разный результат этих двух агрегатов.',
'Оба интерфейса используют Spark SQL и Catalyst. SQL задаёт декларативный текст запроса, DataFrame API строит эквивалентные выражения; это не два независимых движка.',
'Каждая выходная задача записывает свой part-файл. Каталог соответствует распределённой записи и включает служебные файлы, например _SUCCESS.',
'Для 74 строк потеря параллелизма несущественна и один CSV удобен. Для большого факта coalesce(1) создаёт единственную задачу записи и узкое место.',
'Parquet хранит типы, поддерживает колоночное чтение, сжатие и проталкивание фильтров. CSV требует восстановления типов; при чтении Parquet nullable может стать true даже при сохранении имён и типов полей.'
]

def build():
    r=json.loads((LAB/'results'/'pr3.json').read_text(encoding='utf-8')); k=r['komus']
    d=PR3Report(3,'Распределенная модель обработки данных')
    d.h('1 Цель и постановка задачи')
    d.p('Цель работы — выполнить распределённую табличную обработку в PySpark: подготовить источник SQLite, задать схемы, исследовать пропуски, построить агрегаты и JOIN, сопоставить Spark SQL с DataFrame API и проверить экспорт. Учебная часть выполнена на Olist, самостоятельная — на продажах КОМУС.')
    d.p('Итогом стали аналитическая витрина Olist из 74 категорий и витрина КОМУС по товарным группам. Результаты получены полным последовательным запуском notebook в новом ядре; исходные таблицы не изменялись.')
    d.h('2 Окружение и источники')
    environment_table(d,r)
    d.p('Ecommerce.sqlite восстановлена из CSV Olist, использованных во второй практике, поскольку исходный файл SQLite не был передан. Восстановление и staging выполняются отдельным скриптом. Фактическая версия содержит 100 000 отзывов, тогда как в примере методички указано 97 621; эта разница не скрывается и не исправляется под ожидаемые числа. Справочник категорий восстановлен до исходных 71 записей.')
    d.table('Выгруженные учебные таблицы',['Таблица','Число строк'],[[name,num(n,0)] for name,n in r['source_counts'].items()],[11,5.5])
    d.p('Spark не читает SQLite как штатный файловый формат. Модуль sqlite3 и Pandas выгружают таблицы частями по 50 000 строк; последующая обработка производится в Spark. Данные КОМУС получены из предоставленного архива «Транзакции КОМУС.zip». Клиентские реквизиты, адреса и номера машин исключены ещё при подготовке входа. Архивы transactions и northwind_small не использовались: выбранный набор КОМУС удовлетворяет требованиям самостоятельной части.')
    d.p('Вход учебной части: ecommerce.sqlite → staging CSV → связанные DataFrame orders, order_items, products, customers, payments и translations. Самостоятельная часть: продажи КОМУС → fact; товары КОМУС → dim; соединение fact.product_key = dim.product_key. Связи orders–customers и items–products контролируются по идентификаторам, а не по названиям.')
    d.h('3 Схемы и преобразования')
    d.p('Orders при чтении CSV имеет строковые столбцы. Для order_items задан StructType: order_item_id — integer, price и freight_value — double, идентификаторы — string. Автоматический вывод схемы дополнительно определяет shipping_limit_date как timestamp, поэтому схемы не обязаны полностью совпадать. Пять временных полей orders преобразованы функцией to_timestamp.')
    d.code('items_schema = StructType([\n    StructField("order_id", StringType(), True),\n    StructField("order_item_id", IntegerType(), True),\n    StructField("product_id", StringType(), True),\n    StructField("seller_id", StringType(), True),\n    StructField("shipping_limit_date", StringType(), True),\n    StructField("price", DoubleType(), True),\n    StructField("freight_value", DoubleType(), True)\n])')
    d.p(f'Применены select, filter, withColumn, drop, dropDuplicates и orderBy. Для позиций price ≥ 500 вычислена цена с доставкой gross_item_value; условию удовлетворяют {num(r["expensive_count"],0)} строк. Это стоимость позиции, а не всего заказа.')
    d.table('Самые дорогие позиции с доставкой',['Позиция','Цена','Доставка','Итого'],[[x['order_item_id'],num(x['price']),num(x['freight_value']),num(x['gross_item_value'])] for x in r['expensive_top']],[3,4.5,4.5,4.5])
    d.p('Вывод: строковые временные значения приведены к timestamp; отбор дорогих позиций сохраняет уровень наблюдения order_id + order_item_id. Результаты printSchema() и сравнение с inferSchema сохранены в выполненном notebook.')
    d.h('4 Пропуски и предметная интерпретация')
    d.table('Пропуски временных полей orders',['Поле','Null'],[[name,n] for name,n in r['nulls'].items()],[12,4.5])
    d.p('Фактическая дата доставки отсутствует в 2 965 заказах. По статусам это 1 107 shipped, 619 canceled, 609 unavailable, 314 invoiced, 301 processing, 8 delivered, 5 created и 2 approved. Дата не заполнена вымышленным значением и не служит причиной удаления незавершённых заказов.')
    delivery=r['delivery']
    d.p(f'Из {num(delivery["delivered"],0)} заказов delivered дата известна у {num(delivery["known"],0)}. Просрочено {num(delivery["late"],0)} заказов, или {num(delivery["late"]/delivery["known"]*100)}% известной выборки. В products у 610 товаров отсутствует категория; она обозначена unknown, чтобы сохранить позиции при соединении.')
    d.figure(LAB/'screenshots'/'pr3'/'02_quality.png','Проверка пропусков и сохранения строк в консоли Ubuntu',14)
    d.h('5 Агрегации и соединения')
    d.table('Статусы заказов',['Статус','Заказы'],[[x['order_status'],num(x['count'],0)] for x in r['status']],[9,7.5])
    d.table('Агрегаты платёжных строк',['Способ оплаты','Строки','Сумма','Среднее'],[[x['payment_type'],num(x['payments'],0),num(x['total_value']),num(x['avg_value'])] for x in r['payments']],[5,3,5,3.5])
    d.p('Платежи credit_card преобладают по числу строк и сумме. COUNT(*) в order_payments нельзя интерпретировать как число заказов: один заказ может иметь несколько платежей.')
    d.p(f'LEFT JOIN order_items → products → translations сохранил {num(r["join"]["after"],0)} из {num(r["join"]["before"],0)} позиций; {num(r["join"]["unknown"],0)} позиций имеют unknown. Категория выбирается через coalesce: английский перевод, исходное название, затем unknown.')
    d.table('Пять штатов с высокой долей просрочки',['Штат','Доставлено','Просрочено','Доля %'],[[x['customer_state'],x['delivered'],x['late'],num(x['late_pct'])] for x in r['late_state'][:5]],[3,4.5,4.5,4.5])
    d.p('AL имеет 23,93% просрочки при 397 известных доставках. Доля сравнивается вместе с размером группы; сама таблица не доказывает причинное влияние штата на доставку.')
    d.h('6 Spark SQL и сохранение витрины')
    d.code('sales.createOrReplaceTempView("sales")\nSELECT category, COUNT(*) AS items,\n       COUNT(DISTINCT order_id) AS orders,\n       ROUND(SUM(price), 2) AS revenue,\n       ROUND(AVG(price), 2) AS avg_price\nFROM sales\nGROUP BY category\nORDER BY revenue DESC')
    d.table('Топ категорий Olist по сумме price',['Категория','Позиции','Сумма price','Средняя цена'],[[x['category'],num(x['items'],0),num(x['revenue']),num(x['avg_price'])] for x in r['category'][:5]],[6.5,2.5,4.5,3])
    d.p('Двусторонний exceptAll между полными результатами SQL и DataFrame API дал 0 различий. Витрина содержит 74 группы и записана в Parquet и CSV; для маленького CSV использован coalesce(1). После чтения Parquet совпали число строк, имена и типы полей. Признак nullable стал true для ранее гарантированно заполненных агрегатов, поэтому полное равенство StructType отдельно отмечено как false.')
    d.table('Схема итоговой витрины Olist', ['Поле', 'Тип до / после Parquet', 'Nullable до / после'], schema_rows(r['export']['schema_before'], r['export']['schema_after']), [5, 6, 5.5])
    d.figure(LAB/'screenshots'/'pr3'/'03_olist.png','Вычисленные категории Olist и проверки экспорта',14)
    d.h('7 Самостоятельный анализ продаж КОМУС')
    d.p(f'Основная таблица содержит {num(k["rows"],0)} строк и восемь полей: себестоимость, дата, чек, товарный ключ, филиал, количество, сумма покупки и скидка. Справочник содержит {num(k["dim_unique"],0)} уникальных ключей товара и признаки бренда и группы. Названия и условия предоставлены в пользовательском архиве; публичная лицензия на исходные строки не заявлена, поэтому в Git сохраняются код, агрегаты и обезличенные результаты, а не исходные транзакции.')
    d.h('7.1 Чтение, схема и преобразования', 2)
    d.code('''raw = (spark.read.option("header", True)
    .option("escape", '"').option("mode", "FAILFAST")
    .schema(own_schema).csv(str(own_path / "sales.csv")))
dim_raw = (spark.read.option("header", True)
    .option("escape", '"').option("mode", "FAILFAST")
    .csv(str(own_path / "products.csv")))''')
    d.p(f'CSV записан Pandas с двойной кавычкой как экранирующим символом; этот же символ явно задан читателю Spark. Все четыре поля {num(k["dim_raw"],0)} товаров сверены со стандартным csv.DictReader: расхождений {k["source_product_mismatches"]}. В полном справочнике {k["source_dimension_groups"]} разных групп; продажи относятся к {k["output_rows"]} из них. Проверка числа строк JOIN дополняется проверкой содержимого справочника.')
    d.table('Типы полей продаж КОМУС', ['Поле', 'Сырой вход', 'После преобразования'], [[f['name'], 'string', next(t['type'] for t in k['schema']['fields'] if t['name']==f['name'])] for f in k['raw_schema']['fields']], [5.5, 5, 6])
    d.p('Применены select, filter, withColumn, dropDuplicates и orderBy. Содержательные признаки: month для динамики, margin = amount − cost для строк с известной себестоимостью. Категориальные признаки филиала, бренда и группы используются отдельно от денежных величин.')
    d.p('Явный StructType фиксирует строковую схему сырого входа. Дата преобразована из dd.MM.yyyy, десятичная запятая заменена точкой перед приведением четырёх числовых полей. Филиал очищен trim, вычислены месяц покупки и margin = amount − cost. Категориальные пропуски справочника маркируются unknown только в соответствующих полях.')
    d.h('7.2 Качество данных и JOIN', 2)
    d.table('Пропуски КОМУС после приведения типов',['Поле','Null'],[[name,num(n,0)] for name,n in k['nulls'].items()],[11,5.5])
    d.p('У 3 934 строк отсутствует номер чека, но продажи сохранены: фиктивный чек не создаётся. У 86 431 строки отсутствует себестоимость; такие строки участвуют в выручке, но их маржа остаётся null. SUM(margin) описывает только строки с известной себестоимостью, а cost_known показывает охват. Фильтр валидной даты и суммы не удалил строки, поскольку оба поля полностью заполнены.')
    j=k['joins']
    d.p(f'До JOIN: {num(j["before"],0)} строк; LEFT JOIN: {num(j["left"],0)}; INNER JOIN: {num(j["inner"],0)}. Несопоставленных ключей — {j["unmatched"]}, конфликтующих ключей справочника — {k["conflicts"]}. Полный охват объясняет одинаковые размеры обоих соединений.')
    d.h('7.3 Агрегации, Spark SQL и экспорт', 2)
    d.table('Товарные группы КОМУС с высокой выручкой',['Группа','Число строк','С cost','Выручка','Известная маржа','Средняя сумма'],[[x['product_group'],num(x['lines'],0),num(x['cost_known'],0),num(x['revenue']),num(x['margin']),num(x['avg_line'])] for x in k['groups'][:10]],[5,1.7,1.7,2.9,2.9,2.3])
    d.table('Выручка по филиалам',['Филиал','Строки','Выручка'],[[x['branch'],num(x['lines'],0),num(x['revenue'])] for x in k['branches']],[8,3,5.5])
    d.p('Выполнены группировки по товарной группе и филиалу. Первая включает count, sum, avg и счётчик известной себестоимости. Два SQL-запроса формируют рейтинг групп и динамику по месяцам; полный рейтинг SQL совпал с DataFrame API по двустороннему exceptAll.')
    d.code('SELECT product_group, COUNT(*) AS lines,\n       COUNT(cost) AS cost_known,\n       ROUND(SUM(amount), 2) AS revenue,\n       ROUND(SUM(margin), 2) AS margin,\n       ROUND(AVG(amount), 2) AS avg_line\nFROM komus_sales\nGROUP BY product_group\nORDER BY revenue DESC')
    d.code('SELECT month, COUNT(*) AS lines,\n       ROUND(SUM(amount), 2) AS revenue,\n       ROUND(SUM(discount), 2) AS discount\nFROM komus_sales\nGROUP BY month\nORDER BY month')
    d.p(f'Итоговая витрина содержит {num(k["output_rows"],0)} товарных групп. Parquet и CSV сохранены в exports. Повторное чтение подтвердило количество строк и типы, изменение nullable зафиксировано отдельно. Расчётная маржа не является чистой прибылью: дополнительные расходы и налоги в источнике отсутствуют.')
    d.table('Первые и последние месяцы продаж КОМУС', ['Месяц', 'Строки', 'Выручка', 'Скидка'], [[x['month'],num(x['lines'],0),num(x['revenue']),num(x['discount'])] for x in k['monthly'][:3]+k['monthly'][-3:]], [3, 3, 5.5, 5])
    d.table('Схема итоговой витрины КОМУС', ['Поле', 'Тип до / после Parquet', 'Nullable до / после'], schema_rows(k['output_schema_before'],k['output_schema_after']), [5, 6, 5.5])
    d.code('''group_stats.write.mode("overwrite").parquet(
    str(OUT / "pr3_komus_parquet"))
(group_stats.coalesce(1).write.mode("overwrite")
    .option("header", True).csv(str(OUT / "pr3_komus_csv")))
own_back = spark.read.parquet(str(OUT / "pr3_komus_parquet"))''')
    d.p(f'Вывод: число строк после фильтра и обоих JOIN равно {num(k["rows"],0)}; исходный справочник прочитан без искажений. Суммарная выручка {num(k["total_revenue"])} совпадает с независимым пересчётом исходного архива. Получено {k["output_rows"]} групп; лидер — {k["groups"][0]["product_group"]}. SQL/API: {k["sql_diff"]} различий. После чтения Parquet совпали число строк, имена и типы полей; nullable проверен отдельно.')
    d.figure(LAB/'screenshots'/'pr3'/'04_komus.png','Реальные проверки данных КОМУС и обоих JOIN в Ubuntu',14)
    d.figure(LAB/'screenshots'/'pr3'/'01_spark_jobs.png','Spark UI фактически выполненного приложения третьей практики',14)
    d.h('8 Воспроизводимость и выводы')
    run=json.loads((LAB/'results/pr3_execution.json').read_text(encoding='utf-8'))
    completed=datetime.fromisoformat(run['completed_utc']).astimezone(timezone(timedelta(hours=3))).strftime('%d.%m.%Y %H:%M:%S МСК')
    d.p(f'Полный Run All выполнен {completed}: новое ядро, {run["code_cells"]} кодовых ячеек в порядке {run["execution_counts"]}, ошибок {run["errors"]}, длительность {num(run["seconds"],2)} с. В последней ячейке выполнен spark.stop(). Сохраняются notebook, JSON результатов и выполнения, каталоги Parquet/CSV и свежий снимок Spark UI.')
    d.code('cd /home/albakhtin/tiabd/АлбахтинИВ/Spark\n/home/albakhtin/tiabd/.venv/bin/python run_pr3.py')
    d.p('Проверка исходного CSV независимым парсером необходима наряду с SQL/API и round-trip: одинаковые агрегаты двух интерфейсов не обнаруживают ошибку, если оба читают уже искажённые поля.')
    for text in ['1. Явная схема и преобразование timestamp сделали расчёты воспроизводимыми; автоматическая схема CSV не равна контракту данных.', '2. Пропуски даты доставки связаны со статусом, поэтому массовое удаление незавершённых заказов исказило бы анализ.', '3. LEFT JOIN Olist сохранил 112 650 позиций, включая 1 603 позиции неизвестной категории.', '4. Полные результаты Spark SQL и DataFrame API совпали; обе реализации используют один вычислительный движок.', f'5. Все {num(k["rows"],0)} продаж КОМУС сохранены, в том числе продажи без себестоимости; оценка маржи требует отдельного контроля охвата.', '6. Справочник КОМУС покрывает все товарные ключи факта, поэтому INNER JOIN и LEFT JOIN дали одинаковое количество строк.', '7. Parquet сохранил количество групп и типы полей, но nullable после чтения отличается от исходной схемы агрегатов.']:
        d.p(text)
    d.p('Ограничения: локальная VM не проверяет работу физического кластера; восстановленная версия Olist отличается от SQLite примера по числу отзывов. Для КОМУС неизвестны полные условия формирования выборки, а неполная себестоимость ограничивает интерпретацию маржи. Выводы о сезонности и причинных связях по этим агрегатам не делаются.')
    d.questions(questions(3),ANSWERS3)
    d.h('Источники')
    for text in ['1. Практическая работа №3. Распределенная модель обработки данных. Предоставленная методичка.', '2. Транзакции КОМУС.zip. Предоставленный учебный архив; SHA-256 зафиксирован в source_manifest.json.', '3. Apache Spark 3.5.7. Installation. https://spark.apache.org/docs/3.5.7/api/python/getting_started/install.html', '4. Таблицы Olist из второй практики; версия и количество строк зафиксированы в source_manifest.json.']:
        d.p(text)
    d.save(ROOT/'4'/'BD'/'ТИАБД_ПР3_АлбахтинИВ.docx')

if __name__=='__main__':
    run=json.loads((LAB/'results/pr3_execution.json').read_text(encoding='utf-8'))
    assert run['status']=='passed' and run['errors']==0 and run['fresh_kernel']
    assert run['notebook_sha256']==hashlib.sha256((LAB/'Practice_3.ipynb').read_bytes()).hexdigest()
    assert hashlib.sha256((ROOT/'4/ИСУРО/ИСУРО_1_АлбахтинИВ.docx').read_bytes()).hexdigest()=='53a4b8ccc73cf44599e8b7c7fb75045afec85dd7f4b0a079b19a0d4ccb5249b1'
    build()
    print('PR3 DOCX built from verified execution')
