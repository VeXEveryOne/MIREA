"""Build the three non-Spark reports from the new personal VM results."""
from __future__ import annotations

import hashlib
import json
import re
import sys
from pathlib import Path

import pandas as pd
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH as A

PACK = Path(__file__).resolve().parents[1]
BD = PACK.parent
REPO = BD.parents[1]
sys.path.insert(0, str(BD / 'spark_lab'))
from report_layout import Report, font, spacing

STUDENT = 'Албахтин И.В.'
GROUP = 'ИНБО-12-23'
SUBJECT = 'Технологии и инструментарий анализа больших данных'
REFERENCE = REPO / '4/ИСУРО/ИСУРО_1_АлбахтинИВ.docx'
SOURCE_HASH = '53a4b8ccc73cf44599e8b7c7fb75045afec85dd7f4b0a079b19a0d4ccb5249b1'


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def n(value, digits=0):
    return f'{value:,.{digits}f}'.replace(',', ' ').replace('.', ',')


def plain_md(text):
    return re.sub(r'\[([^]]+)\]\(([^)]+)\)', r'\1 (\2)', text).replace('**', '').replace('`', '')


def markdown_section(report, path, level=1, new_page=False):
    in_code = False
    code = []
    buffered = []

    def flush():
        if buffered:
            report.p(plain_md(' '.join(buffered)))
            buffered.clear()

    for line in Path(path).read_text(encoding='utf-8-sig').splitlines():
        if line.strip().startswith('```'):
            flush()
            if in_code:
                report.code('\n'.join(code)); code.clear()
            in_code = not in_code
        elif in_code:
            code.append(line)
        elif line.startswith('# '):
            flush()
            p = report.h(plain_md(line[2:]), level)
            if new_page:
                p.paragraph_format.page_break_before = True
        elif re.match(r'^\d+\. \*\*', line):
            flush()
            p = report.doc.add_paragraph(plain_md(line), 'Report Body')
            spacing(p, A.LEFT, 0, 1.5, 6, 2, True)
            for r in p.runs: font(r, bold=True)
        elif not line.strip():
            flush()
        else:
            buffered.append(line.strip())
    flush()


def personalize(doc):
    for part in [doc.part, *[rel.target_part for rel in doc.part.rels.values()
                             if not rel.is_external and rel.reltype.endswith(('/header', '/footer'))]]:
        for node in part.element.xpath('.//w:t'):
            if node.text:
                node.text = (node.text.replace('Албахтин И.В.', STUDENT)
                             .replace('Албахтин И. В.', STUDENT)
                             .replace('АлбахтинИВ', 'АлбахтинИВ')
                             .replace('Анализ больших данных', SUBJECT)
                             .replace('ИНБО-01-17', ''))
    doc.core_properties.author = STUDENT
    doc.core_properties.last_modified_by = STUDENT
    doc.core_properties.subject = SUBJECT
    doc.core_properties.keywords = 'ТИАБД, Албахтин И.В., ИНБО-12-23'
    doc.core_properties.comments = ''
    # Signature fields must retain cover spacing instead of inheriting body indents.
    city = next(p for p in doc.paragraphs if p.text.startswith('Москва 2026'))
    cover_nodes = list(doc.element.body)[:list(doc.element.body).index(city._p)]
    signature = next(t for t in doc.tables if t._tbl in cover_nodes
                     and any(r.cells[0].text.strip().startswith('Студент группы') for r in t.rows))
    for row in signature.rows:
        label = row.cells[0].text.strip()
        if label.startswith('Студент группы'):
            row.cells[0].text = 'Студент группы'
            row.cells[1].text = f'{GROUP}, {STUDENT}'
        elif label.startswith('Преподаватель'):
            row.cells[1].text = 'Воронцов Ю.А.'
    for row in signature.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                spacing(p, A.LEFT, 0, 1)
                for run in p.runs: font(run, 12)
    # On Word the template's inherited empty-paragraph rhythm differs from its
    # saved reference render. Keep Moscow/year near the bottom of the cover.
    city.paragraph_format.space_before = Pt(100)


def add_toc(doc):
    """Insert a real Word field after the cover and start the body on a page."""
    body = doc.element.body
    city = next(p for p in doc.paragraphs if p.text.startswith('Москва 2026'))
    cover_end = list(body).index(city._p)
    first = next(node for node in list(body)[cover_end+1:] if node.tag == qn('w:p')
                 and ''.join(node.xpath('.//w:t/text()')).strip())
    for sec in doc.sections:
        for p in [*sec.header.paragraphs, *sec.first_page_header.paragraphs]:
            p.text = ''
    heading = doc.add_paragraph('Содержание')
    spacing(heading, A.CENTER, 0, 1.5, 0, 6, True)
    heading.paragraph_format.page_break_before = True
    for r in heading.runs: font(r, bold=True)
    field_p = doc.add_paragraph()
    spacing(field_p, A.LEFT, 0, 1, 0, 0)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), ' TOC \\o "1-2" \\h \\z \\u ')
    run = OxmlElement('w:r'); txt = OxmlElement('w:t')
    txt.text = 'Содержание обновляется при открытии документа в Word'
    run.append(txt); field.append(run); field_p._p.append(field)
    first.addprevious(heading._p); first.addprevious(field_p._p)
    first.get_or_add_pPr().append(OxmlElement('w:pageBreakBefore'))
    for name in ['TOC 1', 'TOC 2']:
        if name not in doc.styles:
            from docx.enum.style import WD_STYLE_TYPE
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        st = doc.styles[name]; font(st, 14)
        f = st.paragraph_format
        f.first_line_indent = Cm(0); f.left_indent = Cm(0 if name == 'TOC 1' else .5)
        f.right_indent = Cm(0); f.line_spacing = 1; f.space_after = Pt(3)
        f.space_before = Pt(0); f.keep_with_next = False
    settings = doc.settings.element
    for old in settings.findall(qn('w:updateFields')): settings.remove(old)
    update = OxmlElement('w:updateFields'); update.set(qn('w:val'), 'true'); settings.append(update)


def save(doc, name):
    personalize(doc)
    # Native heading roles survive Word's localization without becoming a
    # custom "Heading 2 1" style, and TOC includes the intended subsections.
    for style in doc.styles:
        match = re.match(r'^(?:Heading|Заголовок)\s*([12])(?:1)?$', style.name, re.I)
        if match:
            level = int(match.group(1))
            style.element.set(qn('w:customStyle'), '0')
            style.element.find(qn('w:name')).set(qn('w:val'), f'heading {level}')
            ppr = style.element.get_or_add_pPr()
            for old in ppr.findall(qn('w:outlineLvl')): ppr.remove(old)
            outline = OxmlElement('w:outlineLvl'); outline.set(qn('w:val'), str(level-1)); ppr.append(outline)
    add_toc(doc)
    doc.core_properties.title = name.replace('_', ' ')
    targets = {
        'ТИАБД_ПР1_АлбахтинИВ': BD / 'practice_1/Практическая работа №1.docx',
        'ТИАБД_ПР2_Визуализация_АлбахтинИВ': BD / 'practice_2/Практическая работа №2.docx',
        'ТИАБД_Лекции_1–2_АлбахтинИВ': BD / 'ТИАБД_Лекции_1–2_АлбахтинИВ.docx',
        'ТИАБД_ПР3_АлбахтинИВ': BD / 'Практическая работа №3.docx',
        'ТИАБД_ПР4_АлбахтинИВ': BD / 'Практическая работа №4.docx',
        'ТИАБД_ПР2_Хранение_АлбахтинИВ': BD / 'Практическая работа №2.docx',
    }
    doc.save(targets[name])
    print(name)


class AlbakhtinReport(Report):
    def figure(self, path, caption, width=16):
        self.figures += 1
        self.p(f'На рисунке {self.figures} представлен результат анализа.').paragraph_format.keep_with_next = True
        p = self.doc.add_paragraph(); spacing(p, A.CENTER, 0, 1, next=True)
        from PIL import Image
        with Image.open(path) as im: ratio = im.height / im.width
        p.add_run().add_picture(str(path), width=Cm(min(width, 16.5, 17 / ratio)))
        cap = self.doc.add_paragraph(f'Рисунок {self.figures} – {caption}', 'Report Figure Caption')
        spacing(cap, A.CENTER, 0, 1, 3, 6); cap.paragraph_format.keep_together = True


def pandas_report():
    root = BD / 'practice_1'
    output = PACK / 'Результаты/ПР1'
    fs = read_json(output / 'fires/summary.json')
    bs = read_json(output / 'bike_sharing/summary.json')
    raw_f = pd.read_csv(root / 'FiresRu.csv')
    raw_b = pd.read_csv(root / 'data/bike_sharing/day.csv')
    assert raw_f.shape == (26681, 9) and raw_b.shape == (731, 16)
    assert raw_f.isna().sum().sum() == raw_b.isna().sum().sum() == 0
    assert raw_f.duplicated().sum() == raw_b.duplicated().sum() == 0
    assert raw_b['cnt'].eq(raw_b.casual + raw_b.registered).all()
    r = AlbakhtinReport(1, 'Введение в инструментарий анализа данных')
    r.h('1 Цель и результаты работы')
    r.p('Цель работы — выполнить полный цикл анализа табличных данных средствами NumPy и Pandas: диагностику, обоснованную предобработку, фильтрацию, агрегации, разведочный анализ и проверку экспортированных файлов. Учебная часть выполнена для FiresRu; самостоятельная — для дневного набора Bike Sharing. Обе части представлены в едином выполненном Jupyter Notebook.')
    r.p('FiresRu содержит 26 681 наблюдение и девять исходных полей; Bike Sharing — 731 наблюдение и 16 полей. Пропусков и полных дубликатов нет. Для каждого набора получены четыре графика и три экспорта; количество строк и значения после повторного чтения проверены.')
    r.h('2 Окружение и воспроизводимость')
    r.p('Notebook выполнен сверху вниз в чистом ядре отдельной виртуальной машины TIABD-Albakhtin-IV (Ubuntu 22.04.5, 2 vCPU, 4 ГБ RAM). Источники сохранены отдельно от очищенных таблиц и графиков. Версии библиотек выводятся в первой ячейке каждой части; в ноутбуке сохранены результаты всех 31 кодовых ячеек без ошибок. Случайные выборки для графиков используют random_state=42.')
    r.code('import platform, numpy as np, pandas as pd\nprint(platform.python_version(), np.__version__, pd.__version__)\ndf = pd.read_csv("FiresRu.csv", encoding="utf-8-sig")\ndf.head()\ndf.info(memory_usage="deep")\ndf.describe(include="all")')
    r.h('3 Учебный набор FiresRu')
    r.h('3.1 Паспорт и качество данных', 2)
    r.p('Источник — предоставленный с заданием файл FiresRu.csv. Одна строка описывает категорию события, координаты и погодные характеристики. Даты и времени наблюдения в таблице нет; временную динамику по номеру строки получить нельзя. Для влажности и солнечной радиации единицы не подтверждены отдельными метаданными и не дописываются по предположению.')
    fields = [('type_id', 'Код категории события'), ('type_name', 'Название категории'),
              ('lon, lat', 'Географические координаты'), ('temperature_c', 'Температура, °C'),
              ('precipitation_mm', 'Осадки, мм'), ('relative_humidity', 'Влажность; единица не подтверждена'),
              ('wind_speed_ms', 'Скорость ветра, м/с'), ('solar_radiation', 'Солнечная радиация; единица не подтверждена')]
    r.table('Поля исходного набора FiresRu', ['Поле', 'Содержание'], fields, [5, 11.5])
    r.p('Проверки isna(), duplicated(), nunique() и перекрёстная таблица type_id/type_name подтверждают отсутствие пропусков, полных повторов и противоречий между кодом и названием. Числовой type_id трактуется как категория. Строки без основания не удалялись.')
    r.h('3.2 Типы и вычисляемые признаки', 2)
    r.code('clean_df = df.copy()\nclean_df["type_name"] = clean_df.type_name.astype("category")\nclean_df["type_id"] = pd.to_numeric(clean_df.type_id, downcast="integer")\nclean_df["wind_speed_kmh"] = (clean_df.wind_speed_ms * 3.6).round(2)\nclean_df["is_forest_fire"] = clean_df.type_name.eq("Forest fire")')
    r.p(f'Память исходной таблицы — {n(fs["memory_before_bytes"])} байт; после оптимизации типов — {n(fs["memory_after_bytes"])} байт. Снижение составляет {n(fs["memory_reduction_pct"], 2)}%. Сравнение выполнено на одинаковых девяти исходных столбцах до добавления вычисляемых полей. Затем созданы скорость ветра в км/ч и признак лесного пожара.')
    r.h('3.3 Фильтрация и групповые статистики', 2)
    r.code('subset = df.loc[(df.relative_humidity < 40) & (df.wind_speed_ms > 2)]\nassert subset.index.equals(df.query("relative_humidity < 40 and wind_speed_ms > 2").index)\ntype_stats = df.type_name.value_counts()\ngroup_stats = clean_df.groupby("type_name", observed=True).agg(\n    n=("type_name", "size"), temp_mean=("temperature_c", "mean"),\n    humidity_mean=("relative_humidity", "mean"), wind_mean=("wind_speed_ms", "mean"))')
    r.p(f'Составной фильтр отобрал {n(fs["filter_rows"])} строк ({n(fs["filter_rows"]/fs["source_rows"]*100, 2)}%). Пороги используются для учебной демонстрации условия и не объявляются нормативом пожарной опасности. Выполнены частотная группировка и агрегация пяти погодных характеристик по категории события.')
    group = pd.read_csv(output / 'fires/group_stats.csv')
    r.table('Сравнение категорий событий', ['Категория', 'Строк', 'Доля %', 'Средняя T °C', 'Ветер м/с'],
            [[x.type_name, n(x.n), n(x.n/len(raw_f)*100, 2), n(x.temp_mean, 2), n(x.wind_mean, 2)] for x in group.itertuples()], [5, 2, 2.5, 3.5, 3.5])
    r.h('3.4 Корреляции и потенциальные выбросы', 2)
    corr = raw_f[['temperature_c', 'precipitation_mm', 'relative_humidity', 'wind_speed_ms', 'solar_radiation']].corr()
    q1, q3 = raw_f.wind_speed_ms.quantile([.25, .75]); iqr = q3-q1
    r.p(f'Корреляция температуры и влажности по всем строкам составляет {n(corr.loc["temperature_c", "relative_humidity"], 3)}; температуры и ветра — {n(corr.loc["temperature_c", "wind_speed_ms"], 3)}. Слабая линейная связь не исключает нелинейной зависимости и не показывает причинности.')
    r.p(f'Для ветра Q1={n(q1, 3)}, Q3={n(q3, 3)}, IQR={n(iqr, 3)}. Правило [Q1−1,5·IQR; Q3+1,5·IQR] выделило {n(fs["wind_iqr_candidates"])} кандидата на выбросы. Они сохранены: редкое или асимметричное значение без предметной проверки не считается ошибкой.')
    r.h('3.5 Визуальный анализ', 2)
    figures = [('01_types.png', 'Число наблюдений по категориям событий', 'Forest fire занимает 69,02% наблюдений; Peat fire представлен 27 строками. Средние малых групп требуют осторожной интерпретации.'),
               ('02_temperature.png', 'Распределение температуры', 'Основная масса температур положительна; левый хвост содержит отрицательные значения. Без сведений о времени и источнике их нельзя автоматически исключать.'),
               ('03_temperature_humidity.png', 'Температура и влажность в выборке из 5 000 строк', 'Широкое облако точек согласуется со слабой линейной связью. Для коэффициента используются все строки, для рисунка — воспроизводимая выборка.'),
               ('04_coordinates.png', 'Пространственное расположение событий', 'Координаты показывают неоднородное расположение наблюдений. График не позволяет сравнить риски регионов без площади, населения и полноты регистрации.')]
    for file, caption, conclusion in figures:
        r.figure(output / 'fires/figures' / file, caption, 15)
        r.p(conclusion)
    r.h('3.6 Экспорт и проверка повторного чтения', 2)
    f_formats = pd.read_csv(output / 'fires/format_comparison.csv')
    r.table('Экспорт очищенного FiresRu', ['Формат', 'Байт', 'Чтение мс', 'Размер', 'Значения'],
            [[x['Формат'], n(x['Байт']), n(x['Чтение, мс'], 3), f'{x["Строк"]} × {x["Столбцов"]}', 'Совпали'] for _, x in f_formats.iterrows()], [2.5, 3.5, 3, 3.5, 4])
    r.p('CSV и JSON требуют восстановления типов при чтении. Parquet хранит схему и обеспечивает выборочное чтение столбцов. В данном сохранённом запуске он занял меньше места и прочитан быстрее; это наблюдение об этих данных, а не универсальная оценка форматов.')
    r.h('4 Самостоятельная работа с Bike Sharing')
    r.h('4.1 Источник и паспорт данных', 2)
    r.p('Bike Sharing — дневные наблюдения Capital Bikeshare в Вашингтоне за 2011–2012 годы. Автор набора — Hadi Fanaee-T; источник UCI Machine Learning Repository, DOI 10.24432/C5W894, лицензия CC BY 4.0. Файл day.csv содержит 731 строку и 16 столбцов, включая дату, погодные показатели, категориальные коды и число прокатов. Требования о 500 строках, пяти информативных полях, двух числовых признаках и категориальном признаке выполнены.')
    r.table('Группы признаков Bike Sharing', ['Поля', 'Содержание'], [
        ['dteday, yr, mnth, weekday', 'Дата и календарные признаки'],
        ['season, holiday, workingday, weathersit', 'Категориальные коды сезона, праздника, рабочего дня и погоды'],
        ['temp, atemp, hum, windspeed', 'Нормированные погодные показатели'],
        ['casual, registered, cnt', 'Прокаты по типу пользователя и общий итог'],
        ['instant', 'Технический индекс; исключён из содержательного анализа']
    ], [6.5, 10])
    r.h('4.2 Диагностика и преобразования', 2)
    r.p('Проверены head(), info(), describe(include="all"), пропуски, уникальные значения, полные повторы, повторные даты и идентификаторы. Даты уникальны; год, месяц и день недели согласуются с dteday; отрицательных счётчиков нет; cnt=casual+registered для всех строк. Нормированные погодные показатели находятся в диапазоне [0; 1]. Нулевая влажность обнаружена в одной строке и сохранена с ограничением: это числовое значение, а не NULL.')
    r.p('Дата преобразована в datetime, календарные и погодные коды — в category, бинарные признаки — в bool, счётчики — в подходящие меньшие целочисленные типы. Для season в разных описаниях встречается различная словесная расшифровка; анализ использует исходные коды. Температура остаётся нормированной, поскольку обратный перевод в °C по противоречивым описаниям создавал бы неопределённость.')
    r.code('clean_df["dteday"] = pd.to_datetime(df.dteday, format="%Y-%m-%d")\nclean_df["humidity_pct"] = clean_df.hum * 100\nclean_df["casual_share_pct"] = clean_df.casual / clean_df.cnt * 100\nclean_df["humidity_zero_flag"] = clean_df.hum.eq(0)')
    r.p(f'Память до оптимизации — {n(bs["memory_before_bytes"])} байт, после оптимизации — {n(bs["memory_after_bytes"])} байт; уменьшение {n(bs["memory_reduction_pct"], 2)}%. Сравнение выполнено на одинаковых 16 исходных столбцах. Затем добавлены влажность в процентах, доля незарегистрированных пользователей и флаг нулевой влажности. Доля casual сравнивает структуру спроса в дни с различным общим количеством поездок.')
    r.h('4.3 Фильтрация и агрегации', 2)
    r.p(f'Составной фильтр workingday==True, temp>0,5 и humidity_pct<65 отобрал {n(bs["filter_rows"])} дней. Условия отражают рабочие дни с повышенной нормированной температурой и влажностью ниже 65%. Выполнены группировки по погоде, типу дня, году и календарному месяцу. Группировка по погоде включает size, mean, median, std и sum.')
    weather = pd.read_csv(output / 'bike_sharing/weather_stats.csv')
    r.table('Число прокатов в разных погодных категориях', ['Код погоды', 'Дней', 'Среднее', 'Медиана', 'Итого'],
            [[int(x.weathersit), int(x.days), n(x.mean_rentals, 2), n(x.median_rentals), n(x.total_rentals)] for x in weather.itertuples()], [3, 2, 3.5, 3.5, 4.5])
    r.p('В ясной погодной категории средний спрос выше, чем в категориях облачности и небольших осадков. Группы имеют разный размер; это наблюдаемое сопоставление, а не доказательство отдельного причинного эффекта погоды.')
    years = pd.read_csv(output / 'bike_sharing/yearly_stats.csv')
    r.table('Годовая динамика Bike Sharing', ['Год', 'Дней', 'Всего прокатов', 'В среднем в день'],
            [[int(x.dteday), int(x.days), n(x.total_rentals), n(x.mean_daily_rentals, 2)] for x in years.itertuples()], [2.5, 2.5, 5.5, 6])
    r.h('4.4 Числовой анализ и выбросы', 2)
    r.p(f'Корреляция нормированной температуры с количеством прокатов составляет {n(bs["temperature_demand_correlation"], 3)}. Между temp и atemp наблюдается близкая к единице связь, поэтому эти признаки несут во многом сходную информацию. Полное число прокатов за период — {n(raw_b.cnt.sum())}; медиана — {n(raw_b.cnt.median())}, среднее — {n(raw_b.cnt.mean(), 2)} в день.')
    iqr_data = pd.read_csv(output / 'bike_sharing/iqr_stats.csv')
    r.table('Проверка числовых показателей по IQR', ['Признак', 'Q1', 'Q3', 'Кандидатов'],
            [[x['Признак'], n(x['Q1'], 3), n(x['Q3'], 3), int(x['Кандидатов'])] for _, x in iqr_data.iterrows()], [6, 3.5, 3.5, 3.5])
    r.p('Для cnt и temp кандидатов по IQR нет; для влажности найдено два, для ветра — 13. Значения сохранены и не интерполируются без проверки первичного источника. Отсутствие NULL не исключает смысловую аномалию, как показывает нулевая влажность.')
    r.h('4.5 Четыре визуализации самостоятельной части', 2)
    b_figures = [('01_weather.png', 'Среднее число прокатов по погодным категориям', 'Сравнение категорий дополняет таблицу групповых средних; категория осадков содержит только 21 день.'),
                 ('02_demand.png', 'Распределение дневного спроса', 'Дневной спрос существенно меняется. Одного среднего недостаточно; нужны медиана, распределение и календарные срезы.'),
                 ('03_temperature_demand.png', 'Температура и число прокатов', 'Рост спроса связан с увеличением нормированной температуры, но разброс отражает влияние календаря, времени и других факторов.'),
                 ('04_monthly.png', 'Средний дневной спрос по календарным месяцам', 'Среднесуточный максимум — сентябрь 2012 года: 7 285,77 проката в день, 218 573 за весь месяц. Сравнение одноимённых месяцев отделяет сезонность от общего роста системы.')]
    for file, caption, conclusion in b_figures:
        r.figure(output / 'bike_sharing/figures' / file, caption, 15)
        r.p(conclusion)
    r.h('4.6 Экспорт самостоятельной части', 2)
    b_formats = pd.read_csv(output / 'bike_sharing/format_comparison.csv')
    r.table('Экспорт очищенного Bike Sharing', ['Формат', 'Байт', 'Чтение мс', 'Размер', 'Значения'],
            [[x['Формат'], n(x['Байт']), n(x['Чтение, мс'], 3), f'{x["Строк"]} × {x["Столбцов"]}', 'Совпали'] for _, x in b_formats.iterrows()], [2.5, 3.5, 3, 3.5, 4])
    r.p('Все три формата после чтения сохранили 731 строку, 19 столбцов и значения. Типы не совпали полностью без приведения: в частности, категории потребовали восстановления по сохранённой схеме BikeSharing_schema.json. Проверка dtypes выполняется отдельно от проверки значений; факт чтения файла сам по себе не подтверждает равенство типов.')
    r.h('5 Выводы и ограничения')
    conclusions = [
        '1. Учебный набор FiresRu содержит пять резко несбалансированных категорий; Forest fire доминирует, а выводы о Peat fire ограничены 27 наблюдениями.',
        f'2. Обоснованные категориальные и целочисленные типы сократили память исходных столбцов FiresRu на {n(fs["memory_reduction_pct"], 2)}%, Bike Sharing — на {n(bs["memory_reduction_pct"], 2)}% при сохранении всех значений. Вычисляемые признаки добавлены после сравнения памяти.',
        '3. Bike Sharing покрывает 731 последовательный день. Пропуски, полные дубликаты и нарушения равенства cnt=casual+registered не найдены.',
        '4. В погодной категории 1 среднее равно 4 876,79 проката; в категории 3 — 1 803,29. Погода связана со спросом, но группы существенно различаются по численности.',
        '5. Средний дневной спрос вырос с 3 405,76 в 2011 году до 5 599,93 в 2012 году — примерно на 64,43%. Сравнивается среднее за день с учётом високосного года.',
        '6. Месячный максимум — сентябрь 2012 года, 218 573 проката. Для анализа по месяцам использована реальная дата dteday.',
        '7. Корреляция temp и cnt составляет 0,627. Это линейная связь в наблюдаемом наборе, а не причинная модель или прогноз.',
        '8. Проверка IQR выявила аномальные кандидаты влажности и ветра; их автоматическое удаление изменило бы состав наблюдений без доказательства ошибки.',
        '9. CSV, JSON и Parquet сохранили значения при повторном чтении. Parquet оказался компактнее в сохранённом запуске; категории проверены и восстановлены отдельно.'
    ]
    for c in conclusions: r.p(c)
    r.p('Ограничения FiresRu: отсутствует время, не подтверждены некоторые единицы и полнота регистрации событий. По координатам нельзя оценивать риск региона без внешних данных о территории и охвате наблюдений.')
    r.p('Ограничения Bike Sharing: только один город и два года; рост инфраструктуры и календарь могут смешиваться с погодой. Нулевая влажность и неоднозначная расшифровка season ограничивают предметную интерпретацию. IQR и корреляция не заменяют проверку первоисточника.')
    markdown_section(r, root / 'Контрольные_вопросы.md')
    # Match shared answer examples to the executed personal notebook and runner.
    for node in r.doc.element.body.xpath('.//w:t'):
        if node.text:
            node.text = node.text.replace('53,5%', n(fs['memory_reduction_pct'], 2) + '%')
            node.text = node.text.replace('season_name', 'season')
            node.text = node.text.replace(
                'В этой работе скрипт scripts/run_all.py собирает единый notebook и выполняет его в чистом ядре bd-vscode. Файл сохраняется только после успешного выполнения всех ячеек.',
                'В этой работе Исходники/execute_all.py выполняет notebook в новом ядре python3 личной VM. После каждой ячейки сохраняется текущий результат; успешное завершение всех ячеек отдельно подтверждено в Результаты/execution.json.'
            )
    r.h('Источники и приложенные материалы')
    for text in ['1. Методические материалы практической работы №1 и предоставленный FiresRu.csv.',
                 '2. H. Fanaee-T. Bike Sharing. UCI Machine Learning Repository. https://doi.org/10.24432/C5W894. CC BY 4.0.',
                 '3. Ноутбуки/ТИАБД_ПР1_АлбахтинИВ.ipynb: 31 выполненная кодовая ячейка, восемь графиков и проверки повторного чтения.',
                 '4. Результаты/ПР1/fires и Результаты/ПР1/bike_sharing: очищенные таблицы, групповые статистики, графики и JSON-сводки.']:
        r.p(text)
    save(r.doc, 'ТИАБД_ПР1_АлбахтинИВ')


def lecture_report():
    r = AlbakhtinReport('1–2', 'Ответы на контрольные вопросы лекций')
    r.doc.paragraphs[6].text = 'Лекции №1–2\nОтветы на контрольные вопросы'
    for run in r.doc.paragraphs[6].runs: font(run, 14, bold=True)
    r.doc.paragraphs[4].text = 'ОТВЕТЫ НА КОНТРОЛЬНЫЕ ВОПРОСЫ'
    for run in r.doc.paragraphs[4].runs: font(run, 16, bold=True)
    r.h('Назначение работы')
    r.p('Документ содержит ответы на все 16 контрольных вопросов двух предоставленных лекций: архитектура больших данных, масштабирование и аналитические слои; SQL, оконные функции, планы выполнения и модели NoSQL. Нумерация вопросов соответствует страницам 70 первой лекции и 107 второй.')
    markdown_section(r, BD / 'practice_1/Лекция_1_Контрольные_вопросы.md', new_page=False)
    markdown_section(r, BD / 'practice_2/Лекция_2_Контрольные_вопросы.md', new_page=True)
    r.h('Источники')
    r.p('Технологии и инструментарий анализа больших данных. Лекция №1. Предоставленный PDF, контрольные вопросы на странице 70.')
    r.p('Технологии и инструментарий анализа больших данных. Лекция №2. Предоставленный PDF, контрольные вопросы на странице 107.')
    save(r.doc, 'ТИАБД_Лекции_1–2_АлбахтинИВ')


def visualization_report():
    output = PACK / 'Результаты/Визуализация'
    summary = read_json(output / 'summary.json')
    data = summary['bike_sharing']
    timing = pd.read_csv(output / 'dimension_reduction_timings.csv')
    r = AlbakhtinReport(2, 'Визуализация данных в Python')
    r.h('1 Цель и выполненные задания')
    r.p('Цель работы — выбрать визуальные представления для табличных и многомерных данных, настроить оформление Plotly и Matplotlib, исследовать t-SNE и UMAP при разных параметрах и измерить время их выполнения. Для временных и категориальных графиков использован Bike Sharing; для проекций — набор рукописных цифр load_digits. Все 11 кодовых ячеек выполнены в чистом ядре отдельной VM TIABD-Albakhtin-IV без ошибок.')
    r.h('2 Источники и диагностика данных')
    r.p('Bike Sharing содержит 731 дневное наблюдение и 16 исходных столбцов за 2011–2012 годы. Источник — UCI Machine Learning Repository, H. Fanaee-T, DOI 10.24432/C5W894, CC BY 4.0. Признаки описывают дату, календарь, погодные условия и число прокатов случайных и зарегистрированных пользователей.')
    r.code('df = pd.read_csv(DATA_PATH, parse_dates=["dteday"])\ndf.head()\ndf.info()\ndf.isna().sum()\ndf.duplicated().sum()\nassert (df.casual + df.registered == df.cnt).all()')
    r.table('Первичная диагностика Bike Sharing', ['Показатель', 'Результат'], [
        ['Размер', '731 × 16'], ['Период', f'{data["date_min"]} — {data["date_max"]}'],
        ['Пропуски и полные дубликаты', '0 и 0'], ['Всего прокатов', n(data['total_rentals'])],
        ['Среднее за день', n(data['mean_daily_rentals'], 2)], ['Медиана', n(data['median_daily_rentals'])],
        ['Максимум за день', n(data['max_daily_rentals'])]
    ], [7, 9.5])
    r.p('Удаление или интерполяция строк не требуются. Дата имеет тип datetime; коды сезона и погоды дополнены категориальными подписями. Температура и ветер сохраняются в нормированной шкале. Влажность переводится в проценты по метаданным. Для season применены исходные коды, поскольку архивное и обновлённое описания дают разные словесные расшифровки. Обработанная таблица сохранена в CSV и Parquet.')
    r.h('3 Столбчатая диаграмма Plotly')
    r.p('Данные сгруппированы по календарному месяцу; по X отложен месяц, по Y — сумма cnt. Цвет столбца зависит от этой же суммы. Использован модуль plotly.graph_objects с цветовой шкалой, чёрными границами толщиной 2, заголовком по центру размером 20, подписями осей размером 16 и метками размером 14. Угол −45° эквивалентен 315°. Высота 700 px, ширина адаптивная, сетка ivory толщиной 2 и минимальные внешние поля.')
    r.code('bar = go.Figure(go.Bar(\n    x=monthly.index, y=monthly["cnt"],\n    marker=dict(color=monthly["cnt"], coloraxis="coloraxis",\n                line=dict(color="black", width=2))))\nbar.update_layout(autosize=True, height=700,\n    title=dict(x=0.5, font=dict(size=20)))\nbar.update_xaxes(tickangle=-45, tickfont=dict(size=14),\n    gridcolor="ivory", gridwidth=2)')
    r.figure(output / 'figures/01_monthly_bar.png', 'Число прокатов по календарным месяцам', 16)
    r.p(f'Максимум — {summary["monthly_peak"]["month"]}: {n(summary["monthly_peak"]["rentals"])} проката. Месяцы 2012 года в целом имеют больший спрос, чем одноимённые месяцы 2011 года. Цвет повторяет количественный показатель и облегчает поиск максимумов.')
    r.h('4 Круговая диаграмма Plotly')
    r.p('Для круговой диаграммы cnt просуммирован по четырём кодам сезона. go.Pie использует читаемые подписи и проценты, чёрные границы секторов толщиной 2, согласованную цветовую схему и заголовок по центру. Категорий всего четыре, поэтому дополнительное объединение мелких объектов не требуется.')
    r.figure(output / 'figures/02_season_pie.png', 'Доли прокатов по исходным кодам сезона', 16)
    r.table('Количество прокатов по кодам сезона', ['Категория', 'Прокатов', 'Доля %'],
            [[key, n(value), n(value/data['total_rentals']*100, 2)] for key, value in summary['season_totals'].items()], [7, 5, 4.5])
    r.p('Максимум приходится на код 3. Словесные названия сезонов не используются для вывода, чтобы неоднозначность метаданных не превратилась в неверное предметное утверждение. Круговая диаграмма показывает структуру общей суммы, но не отражает хронологический порядок.')
    r.h('5 Линейные графики Matplotlib')
    r.p('По месяцу построены три связанных показателя: casual, registered и cnt. Каждый ряд показан отдельными осями одной Figure с линией crimson, белыми маркерами, чёрными границами маркеров толщиной 2 и сеткой mistyrose толщиной 2. Объектная модель Figure/Axes позволяет единообразно настроить подписи и шкалы.')
    r.figure(output / 'figures/03_monthly_lines.png', 'Месячная динамика трёх показателей спроса', 16)
    c = summary['monthly_correlations']
    r.p(f'Корреляция месячного cnt с registered равна {n(c["cnt"]["registered"], 3)}, с casual — {n(c["cnt"]["casual"], 3)}. Общий спрос теснее меняется вместе с зарегистрированной аудиторией. Это частично обусловлено арифметическим равенством cnt=casual+registered; показатели нельзя считать независимыми.')
    r.h('6 Визуализация t-SNE')
    r.p('Использован встроенный набор sklearn.datasets.load_digits: 1 797 изображений десяти классов размером 8 × 8 пикселей, то есть 64 числовых признака на объект. Это допустимый альтернативный готовый набор по условию задания. Признаки стандартизованы; метки применяются только для цвета точек и не участвуют в обучении проекции. Все настройки используют random_state=42.')
    r.p('t-SNE выполняется для perplexity=5, 30 и 50. Перплексия управляет эффективным размером локального соседства. Малое значение выделяет более мелкие группы; увеличение делает структуру менее раздробленной. Расстояния между отдельными островками и их абсолютные размеры не имеют прямого смысла исходных расстояний.')
    r.figure(output / 'figures/04_tsne_perplexities.png', 'Проекции t-SNE для трёх значений перплексии', 16)
    r.h('7 Визуализация UMAP')
    r.p('UMAP применяется к тем же стандартизованным объектам с тремя конфигурациями: n_neighbors=5 и min_dist=0,0; 15 и 0,1; 50 и 0,5. Первая настройка подчёркивает очень локальные компактные группы; последняя оставляет больше пространства между точками и использует более широкое окружение. Сравнение проводится по одинаковому набору и цветовым меткам.')
    r.figure(output / 'figures/05_umap_parameters.png', 'Проекции UMAP при разных n_neighbors и min_dist', 16)
    r.h('8 Время и качество проекций')
    r.p('Время каждого fit_transform измерено time.perf_counter(). Дополнительно рассчитана trustworthiness при k=10: она оценивает сохранение локальных соседей в двумерной проекции; выше означает лучшее сохранение этого свойства. Метрика не доказывает истинность кластеров. Первый UMAP-запуск включает JIT-компиляцию Numba, поэтому холодные и последующие измерения не полностью сопоставимы.')
    r.table('Результаты измерения алгоритмов', ['Алгоритм', 'Параметры', 'Время с', 'Trustworthiness'],
            [[x['algorithm'], x['parameters'], n(x['seconds'], 4), n(x['trustworthiness_k10'], 4)] for _, x in timing.iterrows()], [2.8, 7.2, 3, 3.5])
    r.figure(output / 'figures/06_runtime_comparison.png', 'Сопоставление времени шести конфигураций', 16)
    dim = summary['dimension_reduction']
    r.p(f'Среднее время трёх конфигураций t-SNE — {n(dim["tsne_mean_seconds"], 4)} с, UMAP — {n(dim["umap_mean_seconds"], 4)} с. Отношение средних {n(dim["umap_to_tsne_ratio"], 2)} включает холодный UMAP и разные настройки. Его нельзя считать универсальным коэффициентом скорости алгоритмов.')
    r.table('Версии окружения измерений', ['Компонент', 'Версия'],
            [[key, dim[key]] for key in ['python', 'pandas', 'numpy', 'matplotlib', 'plotly', 'scikit_learn', 'umap_learn']], [7, 9.5])
    r.h('9 Выводы и ограничения')
    for text in [
        '1. Для количественного сравнения месяцев подходит столбчатая диаграмма; для структуры четырёх категорий — круговая; для динамики трёх показателей — линейные графики.',
        '2. Сумма за период составляет 3 292 679 прокатов; пик — сентябрь 2012 года. Динамика позволяет одновременно увидеть колебания внутри года и общий рост.',
        '3. Plotly добавляет интерактивные подсказки и масштабирование; Matplotlib обеспечивает точную статическую компоновку. Для каждой библиотеки выполнены требования оформления из задания.',
        '4. t-SNE и UMAP показывают разделение известных классов цифр, но форма и плотность островков меняются при смене параметров.',
        '5. Trustworthiness дополняет визуальный анализ измеримой оценкой локальных соседств. Она не оценивает все аспекты сохранения геометрии и не заменяет предметную интерпретацию.',
        '6. Измерения зависят от прогрева Numba, размера набора и нагрузки. Результат одной серии на 1 797 объектах не переносится без проверки на более крупные данные.',
        'Ограничения: Bike Sharing относится к одному городу и двум годам; исходная расшифровка сезонов неоднозначна. Двумерные нелинейные проекции теряют часть информации и не являются доказательством существования истинных кластеров.'
    ]: r.p(text)
    markdown_section(r, BD / 'practice_2/Контрольные вопросы.md')
    r.h('Источники и приложения')
    r.p('Методическое задание «Практика №2» — визуализация данных в Python, восемь пунктов задания.')
    r.p('Bike Sharing. UCI Machine Learning Repository. https://doi.org/10.24432/C5W894. CC BY 4.0.')
    r.p('sklearn.datasets.load_digits — встроенный набор рукописных цифр. Выполненный Notebook, CSV с измерениями, шесть PNG и два интерактивных HTML приложены к работе.')
    save(r.doc, 'ТИАБД_ПР2_Визуализация_АлбахтинИВ')


def main():
    assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest() == SOURCE_HASH
    PACK.mkdir(parents=True, exist_ok=True)
    pandas_report(); visualization_report(); lecture_report()


if __name__ == '__main__': main()
