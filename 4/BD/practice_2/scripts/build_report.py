from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT.parent / "practice_1" / "Практическая работа №1.docx"
FINAL = ROOT / "Практическая работа №2.docx"
SUMMARY = json.loads((ROOT / "output" / "summary.json").read_text(encoding="utf-8"))
TIMINGS = pd.read_csv(ROOT / "output" / "dimension_reduction_timings.csv")
DATA = pd.read_csv(ROOT / "data" / "day.csv")
FIGURES = ROOT / "output" / "figures"
EXPECTED_REFERENCE_HASH = "6b00be9fc25d556825599bd05e3a3db45fa36f66bcfc2300d15eb26685d63595"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, name="Times New Roman", size=14, bold=None, italic=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc: Document, text: str, style="Normal", bold_lead: str | None = None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        lead.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_code(doc: Document, code: str):
    p = doc.add_paragraph(style="Код")
    p.paragraph_format.keep_together = True
    p.add_run(code.strip("\n"))
    return p


def add_caption(doc: Document, text: str):
    p = doc.add_paragraph(text, style="Caption")
    p.paragraph_format.keep_together = True
    return p


def add_figure(doc: Document, filename: str, caption: str, width=6.55):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIGURES / filename), width=Inches(width))
    add_caption(doc, caption)


def add_table(doc: Document, headers, rows, widths=None, font_size=10.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        cell.text = str(value)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if widths:
            cell.width = Inches(widths[i])
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                set_font(run, size=font_size, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            if widths:
                cells[i].width = Inches(widths[i])
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                if i > 0 and len(str(value)) < 24:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_font(run, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def configure_styles(doc: Document) -> None:
    for style_name, size, bold in [
        ("Normal", 14, False),
        ("Title", 18, True),
        ("Heading 1", 16, True),
        ("Heading 2", 14, True),
        ("Caption", 12, True),
        ("List Bullet", 14, False),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"].paragraph_format
    normal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.line_spacing = 1.2
    normal.space_after = Pt(6)


def clear_body(doc: Document) -> None:
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def parse_questions(path: Path):
    questions = []
    current = None
    for raw in path.read_text(encoding="utf-8").splitlines():
        match = re.match(r"(\d+)\. \*\*(.+?)\*\*", raw)
        if match:
            current = [int(match.group(1)), match.group(2), ""]
            questions.append(current)
        elif current and raw.strip():
            clean = raw.strip().replace("`", "")
            current[2] += (" " if current[2] else "") + clean
    return questions


assert sha256(REFERENCE) == EXPECTED_REFERENCE_HASH, "Reference document changed; distill again."

doc = Document(REFERENCE)
clear_body(doc)
configure_styles(doc)

for section in doc.sections:
    for p in section.header.paragraphs:
        if "Практическая работа №1" in p.text:
            p.text = p.text.replace("Практическая работа №1", "Практическая работа №2")
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            set_font(run, size=8)

doc.core_properties.title = "Практическая работа №2. Визуализация данных в Python"
doc.core_properties.subject = "Matplotlib, Plotly, t-SNE и UMAP"

title = doc.add_paragraph(style="Title")
title.add_run("Практическая работа №2. Визуализация данных в Python")

doc.add_heading("1. Цель и задачи работы", level=1)
add_text(
    doc,
    "Цель работы — освоить статическую и интерактивную визуализацию данных в Python, "
    "исследовать влияние параметров t-SNE и UMAP на двумерное представление "
    "многомерных объектов и сравнить измеренное время работы алгоритмов.",
)
add_text(doc, "В работе выполнены следующие задачи:")
for item in [
    "загрузка и диагностика многомерного набора Bike Sharing средствами Pandas;",
    "построение столбчатой и круговой диаграмм с помощью Plotly graph_objects;",
    "построение линейных графиков с помощью Matplotlib;",
    "визуализация набора рукописных цифр при разных значениях перплексии t-SNE;",
    "визуализация тех же объектов UMAP при разных n_neighbors и min_dist;",
    "измерение времени и сравнение возможностей библиотек.",
]:
    add_text(doc, item, style="List Bullet")

doc.add_heading("2. Исходные данные и первичная диагностика", level=1)
add_text(
    doc,
    "Основной набор — дневная таблица Capital Bikeshare за 2011–2012 годы. Она "
    "содержит 731 наблюдение и 16 исходных столбцов: календарные признаки, сезон, "
    "погодные условия, нормированные температуру, влажность и скорость ветра, а "
    "также число поездок случайных и зарегистрированных пользователей. Источник: "
    "UCI Machine Learning Repository, Bike Sharing Dataset, автор Hadi Fanaee-T.",
)
add_caption(doc, "Таблица 1. Основные группы признаков Bike Sharing")
add_table(
    doc,
    ["Группа", "Поля", "Смысл"],
    [
        ["Идентификатор и дата", "instant, dteday", "Номер записи и календарная дата"],
        ["Календарь", "season, yr, mnth, holiday, weekday, workingday", "Сезон, год, месяц и тип дня"],
        ["Погода", "weathersit, temp, atemp, hum, windspeed", "Погодный класс и нормированные измерения"],
        ["Спрос", "casual, registered, cnt", "Поездки по типам пользователей и итог"],
    ],
    widths=[1.45, 2.65, 2.5],
)

add_code(
    doc,
    '''df = pd.read_csv("data/day.csv", parse_dates=["dteday"])
df.head()
df.info()
df.isna().sum()
df.duplicated().sum()''',
)
add_caption(doc, "Таблица 2. Первые пять строк исходного набора")
head_rows = []
for row in DATA.head()[["dteday", "season", "weathersit", "temp", "hum", "casual", "registered", "cnt"]].itertuples(index=False):
    head_rows.append([row[0], row[1], row[2], f"{row[3]:.3f}", f"{row[4]:.3f}", row[5], row[6], row[7]])
add_table(
    doc,
    ["Дата", "Сезон", "Погода", "temp", "hum", "casual", "registered", "cnt"],
    head_rows,
    widths=[0.95, 0.55, 0.62, 0.65, 0.6, 0.7, 0.9, 0.7],
    font_size=9.5,
)

summary = SUMMARY["bike_sharing"]
add_caption(doc, "Таблица 3. Результат первичной диагностики")
add_table(
    doc,
    ["Показатель", "Значение"],
    [
        ["Размер", f'{summary["rows"]} строк × {summary["columns_source"]} столбцов'],
        ["Период", f'{summary["date_min"]} — {summary["date_max"]}'],
        ["Пропущенные значения", summary["missing_values"]],
        ["Полные дубликаты", summary["duplicate_rows"]],
        ["Всего поездок", f'{summary["total_rentals"]:,}'.replace(",", " ")],
        ["Среднее в день", f'{summary["mean_daily_rentals"]:.2f}'],
        ["Медиана в день", f'{summary["median_daily_rentals"]:.0f}'],
        ["Максимум в день", summary["max_daily_rentals"]],
    ],
    widths=[3.4, 2.2],
)
add_text(
    doc,
    "Таблица не содержит пропусков и полных дубликатов, поэтому удаление или "
    "интерполяция строк не требуются. Проверка casual + registered = cnt выполнена "
    "для всех 731 наблюдения и не выявила расхождений.",
)

doc.add_heading("3. Предобработка данных", level=1)
add_text(
    doc,
    "Дата преобразована в тип datetime. Числовые коды сезонов и погоды дополнены "
    "читаемыми категориальными названиями. Нормированные погодные показатели "
    "переведены в исходные единицы согласно описанию набора.",
)
add_code(
    doc,
    '''clean["temp_c"] = clean["temp"] * 41
clean["atemp_c"] = clean["atemp"] * 50
clean["humidity_pct"] = clean["hum"] * 100
clean["windspeed_kmh"] = clean["windspeed"] * 67
assert (clean["casual"] + clean["registered"] == clean["cnt"]).all()''',
)
add_text(
    doc,
    "Очищенная таблица сохранена в CSV и Parquet. Исходные признаки не удаляются: "
    "производные поля добавлены отдельно, чтобы сохранить проверяемую связь с источником.",
)

doc.add_heading("4. Столбчатая диаграмма Plotly", level=1)
add_text(
    doc,
    "Дневные значения агрегированы по месяцам, чтобы даты оставались читаемыми. "
    "Высота и цвет столбца кодируют суммарное количество поездок за месяц. Граница "
    "каждого столбца чёрная, толщиной 2; высота интерактивного рисунка равна 700 px.",
)
add_code(
    doc,
    '''go.Bar(
    x=monthly["month_label"], y=monthly["cnt"],
    marker=dict(color=monthly["cnt"], coloraxis="coloraxis",
                line=dict(color="black", width=2))
)
fig.update_layout(
    title=dict(x=0.5, font=dict(size=20)), height=700,
    margin=dict(l=10, r=10, t=65, b=10)
)
fig.update_xaxes(title_font_size=16, tickfont_size=14,
                 tickangle=315, gridwidth=2, gridcolor="ivory")
fig.update_yaxes(title_font_size=16, tickfont_size=14,
                 gridwidth=2, gridcolor="ivory")''',
)
add_figure(doc, "01_monthly_bar.png", "Рис. 1. Количество поездок по месяцам")
peak = SUMMARY["monthly_peak"]
add_text(
    doc,
    f'Пиковый месяц — {peak["month"]}: {peak["rentals"]:,} поездок. '.replace(",", " ")
    + "В 2012 году столбцы в основном выше соответствующих месяцев 2011 года; "
      "внутри каждого года виден подъём от зимы к тёплому сезону и спад к концу года.",
)

doc.add_heading("5. Круговая диаграмма Plotly", level=1)
add_text(
    doc,
    "Для круговой диаграммы поездки сгруппированы по четырём сезонам. Небольшое "
    "число категорий позволяет разместить названия и доли без наложений.",
)
add_code(
    doc,
    '''go.Pie(
    labels=season_totals.index,
    values=season_totals.values,
    textinfo="label+percent",
    marker=dict(colors=pie_colors,
                line=dict(color="black", width=2))
)''',
)
add_figure(doc, "02_season_pie.png", "Рис. 2. Распределение поездок по сезонам", width=5.9)
season_total = SUMMARY["season_totals"]
add_text(
    doc,
    f'Осень формирует наибольший объём — {season_total["Осень"]:,} поездок '
    f'({season_total["Осень"] / summary["total_rentals"] * 100:.1f}%), а весна — '
    f'наименьший: {season_total["Весна"]:,} '
    f'({season_total["Весна"] / summary["total_rentals"] * 100:.1f}%).'.replace(",", " "),
)

doc.add_heading("6. Линейные графики Matplotlib", level=1)
add_text(
    doc,
    "По общей временной оси сопоставлены три месячных показателя: случайные, "
    "зарегистрированные и все поездки. Каждый ряд вынесен в отдельную область "
    "построения и оформлен линией crimson, белыми маркерами с чёрной границей "
    "толщиной 2 и сеткой mistyrose толщиной 2.",
)
add_code(
    doc,
    '''ax.plot(x, y, color="crimson", linewidth=2, marker="o",
        markerfacecolor="white", markeredgecolor="black",
        markeredgewidth=2)
ax.grid(linewidth=2, color="mistyrose")''',
)
add_figure(doc, "03_monthly_lines.png", "Рис. 3. Динамика спроса по типам пользователей")
corrs = SUMMARY["monthly_correlations"]
add_text(
    doc,
    "Все три ряда имеют сходную сезонную форму и заметный рост во втором году. "
    f'Корреляция месячного общего спроса с зарегистрированными пользователями равна '
    f'{corrs["cnt"]["registered"]:.3f}, со случайными — {corrs["cnt"]["casual"]:.3f}. '
    "Следовательно, основную динамику общего ряда задаёт зарегистрированная аудитория.",
)

doc.add_heading("7. Визуализация многомерных данных t-SNE", level=1)
add_text(
    doc,
    "Для снижения размерности использован встроенный набор load_digits: 1 797 "
    "изображений цифр 0–9, каждое описано 64 интенсивностями пикселей. Пропуски "
    "отсутствуют; признаки стандартизированы. Во всех запусках зафиксирован "
    "random_state=42, изменялась только перплексия.",
)
add_code(
    doc,
    '''TSNE(n_components=2, perplexity=perplexity,
     init="pca", learning_rate="auto", max_iter=1000,
     random_state=42).fit_transform(X_scaled)''',
)
tsne_rows = []
for row in TIMINGS[TIMINGS["algorithm"] == "t-SNE"].itertuples(index=False):
    tsne_rows.append([row.parameters, f"{row.seconds:.4f}", f"{row.trustworthiness_k10:.4f}"])
add_caption(doc, "Таблица 4. Результаты t-SNE")
add_table(doc, ["Параметры", "Время, с", "Trustworthiness@10"], tsne_rows, widths=[3.0, 1.25, 1.7])
add_figure(doc, "04_tsne_perplexities.png", "Рис. 4. t-SNE при разных значениях перплексии")
add_text(
    doc,
    "При перплексии 5 локальные группы сильнее дробятся. Значения 30 и 50 дают "
    "более цельные кластеры; лучший измеренный показатель сохранения локальных "
    "соседств получен при perplexity=30. Частичные пересечения объяснимы сходным "
    "написанием некоторых цифр. Положение далёких кластеров нельзя трактовать как "
    "точное расстояние в исходном 64-мерном пространстве.",
)

doc.add_heading("8. Визуализация многомерных данных UMAP", level=1)
add_text(
    doc,
    "UMAP применён к тем же стандартизированным объектам. Уменьшение n_neighbors "
    "усиливает локальные детали, а уменьшение min_dist делает группы компактнее. "
    "Чтобы исключить многопоточную невоспроизводимость, использованы random_state=42 "
    "и n_jobs=1.",
)
add_code(
    doc,
    '''umap.UMAP(n_components=2, n_neighbors=n_neighbors,
          min_dist=min_dist, metric="euclidean",
          random_state=42, n_jobs=1).fit_transform(X_scaled)''',
)
umap_rows = []
for row in TIMINGS[TIMINGS["algorithm"] == "UMAP"].itertuples(index=False):
    umap_rows.append([row.parameters, f"{row.seconds:.4f}", f"{row.trustworthiness_k10:.4f}"])
add_caption(doc, "Таблица 5. Результаты UMAP")
add_table(doc, ["Параметры", "Время, с", "Trustworthiness@10"], umap_rows, widths=[3.3, 1.2, 1.7])
add_figure(doc, "05_umap_parameters.png", "Рис. 5. UMAP при разных n_neighbors и min_dist")
add_text(
    doc,
    "Конфигурация n_neighbors=5, min_dist=0 формирует самые компактные группы. "
    "При n_neighbors=50 и min_dist=0.5 облака становятся шире, а локальная метрика "
    "trustworthiness снижается. Поэтому параметры следует выбирать под задачу, а "
    "не по визуальной привлекательности одной проекции.",
)

doc.add_heading("9. Сравнение времени работы", level=1)
add_caption(doc, "Таблица 6. Измеренное время всех запусков")
all_rows = [
    [row.algorithm, row.parameters, f"{row.seconds:.4f}", f"{row.trustworthiness_k10:.4f}"]
    for row in TIMINGS.itertuples(index=False)
]
add_table(
    doc,
    ["Алгоритм", "Параметры", "Время, с", "Trustworthiness@10"],
    all_rows,
    widths=[0.9, 3.15, 1.05, 1.55],
    font_size=9.7,
)
add_figure(doc, "06_runtime_comparison.png", "Рис. 6. Сравнение времени снижения размерности")
runtime = SUMMARY["dimension_reduction"]
warm_umap = TIMINGS[(TIMINGS["algorithm"] == "UMAP") & (TIMINGS["seconds"] < TIMINGS.loc[TIMINGS["algorithm"] == "UMAP", "seconds"].max())]["seconds"].mean()
add_text(
    doc,
    f'Среднее время t-SNE составило {runtime["tsne_mean_seconds"]:.4f} с. Среднее '
    f'время всех запусков UMAP — {runtime["umap_mean_seconds"]:.4f} с, а двух запусков '
    f'после первого — {warm_umap:.4f} с. В данном небольшом наборе t-SNE оказался '
    "быстрее. Первый запуск UMAP включает JIT-компиляцию Numba, поэтому результат "
    "нельзя обобщать на большие наборы и другие аппаратные среды.",
)

doc.add_heading("10. Сравнение библиотек", level=1)
add_caption(doc, "Таблица 7. Сильные стороны и ограничения инструментов")
add_table(
    doc,
    ["Инструмент", "Сильная сторона", "Ограничение"],
    [
        ["Matplotlib", "Точный контроль статического рисунка, подграфиков и публикационного экспорта", "Интерактивность требует дополнительных средств; много ручных настроек"],
        ["Plotly", "Интерактивные подсказки, масштабирование и автономный HTML", "graph_objects многословнее; статический экспорт зависит от браузерного движка"],
        ["t-SNE", "Хорошо подчёркивает локальные соседства и разделение кластеров", "Чувствителен к перплексии; глобальные расстояния ненадёжны; нет штатного transform для новых объектов"],
        ["UMAP", "Управляет балансом локальной и глобальной структуры и допускает transform", "Результат зависит от параметров; первый запуск несёт расходы JIT-компиляции"],
    ],
    widths=[1.0, 2.9, 2.75],
    font_size=10.2,
)

doc.add_heading("11. Итоговые выводы", level=1)
for item in [
    f'Набор Bike Sharing содержит {summary["rows"]} полных дневных наблюдений за два года; пропуски и полные дубликаты отсутствуют.',
    f'За период зарегистрировано {summary["total_rentals"]:,} поездок; максимум месячного спроса достигнут в {peak["month"]}.'.replace(",", " "),
    "Спрос сезонен и во втором году заметно выше; общий месячный ряд теснее связан с зарегистрированной аудиторией.",
    "Plotly удобен для интерактивного исследования и точных подсказок, Matplotlib — для контролируемой статической композиции.",
    "t-SNE при perplexity=30 лучше всего сохранил локальные соседства среди проверенных настроек; малое значение сильнее дробит группы.",
    "UMAP при малых n_neighbors и min_dist создаёт компактные локальные кластеры; рост параметров делает проекцию более разреженной.",
    "На этом небольшом наборе t-SNE был быстрее UMAP; сравнение времени зависит от JIT-прогрева, объёма данных и окружения.",
    "Двумерные проекции являются инструментом исследования, а не доказательством истинных кластеров или точных межкластерных расстояний.",
]:
    add_text(doc, item, style="List Bullet")

doc.add_heading("12. Ответы на контрольные вопросы", level=1)
add_text(
    doc,
    "Ниже приведены краткие ответы по объектной модели Matplotlib, возможностям "
    "Plotly и интерпретации методов снижения размерности.",
)
for number, question, answer in parse_questions(ROOT / "Контрольные вопросы.md"):
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(5)
    heading.paragraph_format.space_after = Pt(2)
    heading.paragraph_format.keep_with_next = True
    set_font(heading.add_run(f"{number}. {question}"), size=12.5, bold=True)
    answer_paragraph = add_text(doc, answer)
    answer_paragraph.paragraph_format.line_spacing = 1.05
    answer_paragraph.paragraph_format.space_after = Pt(4)
    for run in answer_paragraph.runs:
        set_font(run, size=12)

doc.save(FINAL)
assert sha256(REFERENCE) == EXPECTED_REFERENCE_HASH, "Reference document was modified."
print(FINAL)
