"""Remove trailing empty body paragraphs that can create a blank final page."""

from pathlib import Path

from docx import Document


path = Path(__file__).resolve().parents[1] / "Практическая работа №1.docx"
document = Document(path)
removed = 0

while document.paragraphs and not document.paragraphs[-1].text.strip():
    paragraph = document.paragraphs[-1]._element
    paragraph.getparent().remove(paragraph)
    removed += 1

document.save(path)
print(f"{path}: removed {removed} trailing empty paragraph(s)")
