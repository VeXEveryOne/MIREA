"""Replace the final control-question list with detailed beginner answers."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "Практическая работа №1.docx"
QUESTIONS = ROOT / "Контрольные_вопросы.md"


def set_font(run, *, size: float = 12, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def parse_questions(path: Path):
    questions = []
    current = None
    paragraph_lines = []

    def flush_paragraph() -> None:
        nonlocal paragraph_lines
        if current and paragraph_lines:
            current[2].append(" ".join(paragraph_lines))
            paragraph_lines = []

    for raw in path.read_text(encoding="utf-8").splitlines():
        match = re.match(r"(\d+)\. \*\*(.+?)\*\*", raw)
        if match:
            flush_paragraph()
            current = [int(match.group(1)), match.group(2), []]
            questions.append(current)
        elif current and raw.strip():
            paragraph_lines.append(raw.strip())
        elif current:
            flush_paragraph()
    flush_paragraph()
    return questions


document = Document(REPORT)
heading = next(
    paragraph
    for paragraph in document.paragraphs
    if paragraph.text.strip() == "13. Контрольные вопросы"
)
body = document._element.body
heading_index = list(body).index(heading._p)
for child in list(body)[heading_index:]:
    if child.tag != qn("w:sectPr"):
        body.remove(child)

document.add_heading("13. Ответы на контрольные вопросы", level=1)
intro = document.add_paragraph(
    "Ответы даны с пояснением терминов, механики работы и примеров из выполненного анализа."
)
intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for run in intro.runs:
    set_font(run)

for number, question, answer_paragraphs in parse_questions(QUESTIONS):
    question_paragraph = document.add_paragraph()
    question_paragraph.paragraph_format.space_before = Pt(6)
    question_paragraph.paragraph_format.space_after = Pt(3)
    question_paragraph.paragraph_format.keep_with_next = True
    set_font(question_paragraph.add_run(f"{number}. {question}"), size=12.5, bold=True)

    for answer in answer_paragraphs:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        paragraph.paragraph_format.line_spacing = 1.1
        paragraph.paragraph_format.space_after = Pt(5)
        label = re.match(r"\*\*(.+?)\*\*\s*(.*)", answer)
        if label:
            set_font(paragraph.add_run(label.group(1)), bold=True)
            if label.group(2):
                set_font(paragraph.add_run(" " + label.group(2).replace("`", "")))
        else:
            set_font(paragraph.add_run(answer.replace("`", "")))

document.save(REPORT)
print(REPORT)
