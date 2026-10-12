"""Prepare the two source sections for each combined practical report №2.

The section documents are build inputs in the student's Исходники directory.
Temporary, edited copies are written only to the repository .cache directory.
The final join and PDF export are performed by build_combined_pr2.ps1 in Word.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document


BD = Path(__file__).resolve().parents[1]
REPO = BD.parents[1]
CACHE = REPO / ".cache" / "bd-combined-pr2"

REF = re.compile(r"(?i)\b(рисун(?:ок|ка|ке|ки|ков|кам|ками|ках)|таблиц(?:а|у|е|ы|ой|ах|ами))\s+(\d+)\b")
CAPTION = re.compile(r"^(Рисунок|Таблица)\s+(\d+)\b")
NUMBERED_HEADING = re.compile(r"^(\d+)(\.\d+)?(\.?\s+)(.*)$")


def replace_runs(paragraph, transform):
    before = paragraph.text
    expected = transform(before)
    if before == expected:
        return
    for run in paragraph.runs:
        run.text = transform(run.text)
    if paragraph.text != expected:
        # Cross-run references are rare. Preserve the paragraph style if one occurs.
        paragraph.text = expected
    assert paragraph.text == expected


def set_heading(paragraph, title):
    if paragraph.runs:
        paragraph.runs[0].text = title
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(title)


def prepare(student):
    source = BD / student / "Исходники" / "Компоненты_ПР2"
    visual = Document(source / f"ТИАБД_ПР2_Визуализация_{student}.docx")
    storage = Document(source / f"ТИАБД_ПР2_Хранение_{student}.docx")
    target_dir = CACHE / student
    target_dir.mkdir(parents=True, exist_ok=True)

    cover = next(p for p in visual.paragraphs if p.text.startswith("Практическая работа №2\n"))
    assert len(cover.runs) >= 2
    cover.runs[1].text = "\nВизуализация и хранение данных"
    visual.core_properties.title = "Практическая работа №2. Визуализация и хранение данных"

    for p in visual.paragraphs:
        if p.style.name.startswith(("Heading", "Заголовок")):
            if p.text.startswith("Практическая работа №2. Подробные ответы"):
                set_heading(p, "Ответы на контрольные вопросы по визуализации")
            elif p.text == "Источники и приложения":
                set_heading(p, "Источники и приложения по визуализации")

    offsets = {"Рисунок": 0, "Таблица": 0}
    for p in visual.paragraphs:
        match = CAPTION.match(p.text)
        if match:
            offsets[match[1]] = max(offsets[match[1]], int(match[2]))
    assert offsets == {"Рисунок": 6, "Таблица": 4}, offsets

    first = next(p for p in storage.paragraphs if p.style.name.startswith(("Heading 1", "Заголовок 1")))
    first_node = first._p
    for node in list(storage.element.body):
        if node is first_node:
            break
        storage.element.body.remove(node)
    first.paragraph_format.page_break_before = True

    for p in storage.paragraphs:
        if p.style.name.startswith(("Heading", "Заголовок")):
            match = NUMBERED_HEADING.match(p.text)
            if match:
                number = int(match[1]) + 9
                title = match[4]
                if p._p is first_node:
                    title = "Цель и задачи системы хранения данных"
                elif "Ответы на контрольные вопросы" in title:
                    title = "Ответы на контрольные вопросы по хранению данных"
                elif "Использованные источники" in title:
                    title = "Источники и приложения по хранению данных"
                set_heading(p, f"{number}{match[2] or ''}{match[3]}{title}")
            elif p.text == "Ответы на контрольные вопросы":
                set_heading(p, "Ответы на контрольные вопросы по хранению данных")
            elif p.text == "Источники и приложения":
                set_heading(p, "Источники и приложения по хранению данных")

        def transform(text):
            def repl(match):
                word, old = match.groups()
                kind = "Таблица" if word.casefold().startswith("таблиц") else "Рисунок"
                return f"{word} {int(old) + offsets[kind]}"
            return REF.sub(repl, text)

        replace_runs(p, transform)

    visual_path = target_dir / "visual.docx"
    storage_path = target_dir / "storage.docx"
    visual.save(visual_path)
    storage.save(storage_path)
    print(student, offsets, visual_path, storage_path)


if __name__ == "__main__":
    for student_name in ("АлбахтинИВ", "ГубаревСА"):
        prepare(student_name)
