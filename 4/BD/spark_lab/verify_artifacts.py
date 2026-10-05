"""Read-only structural checks. Visual QA of every PDF page is still required."""
from pathlib import Path
import ast
import json
import re
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH as A

LAB = Path(__file__).resolve().parent
BD = LAB.parent

for script in LAB.glob('*.py'):
    ast.parse(script.read_text(encoding='utf-8'), filename=str(script))

for number, expected_questions, expected_figures in [(3, 16, 4), (4, 30, 8)]:
    notebook = json.loads((LAB / f'Practice_{number}.ipynb').read_text(encoding='utf-8'))
    code = [c for c in notebook['cells'] if c['cell_type'] == 'code']
    assert [c['execution_count'] for c in code] == list(range(1, len(code) + 1))
    assert not any(o['output_type'] == 'error' for c in code for o in c.get('outputs', []))
    assert 'spark.stop()' in ''.join(code[-1]['source'])

    d = Document(BD / f'Практическая работа №{number}.docx')
    # The source title page retains its section boundary; the body inherits
    # the continuous footer, but not the first-page footer exception.
    assert len(d.sections) == 2
    for s in d.sections:
        for value, cm in [(s.page_width, 21), (s.page_height, 29.7), (s.left_margin, 3),
                          (s.right_margin, 1.5), (s.top_margin, 2), (s.bottom_margin, 2)]:
            assert abs(value - Cm(cm)) < Cm(0.01)
    s = d.sections[0]
    assert s.different_first_page_header_footer
    assert s.footer._element.xpath('.//w:fldSimple/@w:instr') == [' PAGE ']
    text = '\n'.join(d.element.body.xpath('.//w:t/text()'))
    assert 'ИНБО-01-17' not in text and 'Шендяпин' not in text
    assert 'ИНБО-12-23, Албахтин И.В.' in text and 'Воронцов Ю.А.' in text
    assert not d.element.body.xpath('.//w:rPr/w:vanish[not(@w:val="0")]')
    for r in d.element.body.xpath('.//w:r[w:t]'):
        fonts = r.find('w:rPr/w:rFonts', r.nsmap)
        color = r.find('w:rPr/w:color', r.nsmap)
        if fonts is not None:
            assert fonts.get(qn('w:ascii')) == 'Times New Roman'
            assert color is not None and color.get(qn('w:val')) == '000000'
        else:
            # Captions intentionally inherit their complete run formatting
            # from paragraph styles rather than duplicate it on every run.
            from docx.text.paragraph import Paragraph
            p = Paragraph(r.getparent(), d)
            assert p.style.name in ['Report Figure Caption', 'Report Table Caption']
            assert p.style.font.name == 'Times New Roman'
            assert str(p.style.font.color.rgb) == '000000'
    normal = d.styles['Report Body']
    assert normal.font.name == 'Times New Roman' and normal.font.size.pt == 14
    assert normal.paragraph_format.alignment == A.JUSTIFY
    assert normal.paragraph_format.line_spacing == 1.5
    assert abs(normal.paragraph_format.first_line_indent - Cm(1.25)) < Cm(0.01)
    assert normal.paragraph_format.space_before.pt == normal.paragraph_format.space_after.pt == 0
    figures = [p for p in d.paragraphs if p.style.name == 'Report Figure Caption']
    assert len(figures) == expected_figures
    for p in figures:
        f = p.paragraph_format
        assert f.alignment == A.CENTER and f.first_line_indent == 0
        assert f.left_indent == 0 and f.right_indent == 0 and f.line_spacing == 1
        assert p.style.font.size.pt == 14 and not p.style.font.bold and not p.style.font.italic
    for p in d.paragraphs:
        if p.style.name == 'Report Table Caption':
            assert p.paragraph_format.alignment == A.LEFT
            assert p.paragraph_format.first_line_indent == 0
            assert p.style.font.size.pt == 12 and p.style.font.italic and not p.style.font.bold
    for table in d.tables[2:]:
        assert table.rows[0]._tr.xpath('./w:trPr/w:tblHeader')
        assert sum(col.width for col in table.columns) <= Cm(16.51)
        for row in table.rows:
            for cell in row.cells:
                assert all(x.get(qn('w:val')) == 'nil' for x in cell._tc.xpath('./w:tcPr/w:shd'))
                for p in cell.paragraphs:
                    assert p.paragraph_format.first_line_indent == 0
                    assert p.paragraph_format.line_spacing == 1
                    assert all(r.font.name == 'Times New Roman' and r.font.size.pt == 12 for r in p.runs)
    index = next(i for i,p in enumerate(d.paragraphs) if p.text == 'Ответы на контрольные вопросы')
    end = next(i for i,p in enumerate(d.paragraphs[index+1:], index+1) if p.text == 'Источники')
    questions = [p.text for p in d.paragraphs[index+1:end] if re.match(r'^\d+\. ', p.text)]
    assert len(questions) == expected_questions
    assert [int(q.split('.')[0]) for q in questions] == list(range(1, expected_questions+1))
    print(f'Practice {number}: {len(code)} executed cells, {expected_questions} answers, layout checks PASS')

r3 = json.loads((LAB / 'results/pr3.json').read_text(encoding='utf-8'))
r4 = json.loads((LAB / 'results/pr4.json').read_text(encoding='utf-8'))
assert r3['join']['before'] == r3['join']['after'] == 112650
assert r3['export']['rows'] == 74 and r3['export']['sql_diff'] == 0
assert r3['komus']['rows'] == 865222 and r3['komus']['output_rows'] == 2864
assert r3['komus']['joins']['before'] == r3['komus']['joins']['left'] == r3['komus']['joins']['inner']
assert r4['lazy']['before'] == r4['lazy']['after_transform'] == []
assert r4['joins']['rows'] == 112650
assert r4['variant']['join_rows'] == 112650
assert 'SortMergeJoin' in (LAB / 'results/plans/smj.txt').read_text()
assert 'BroadcastHashJoin' in (LAB / 'results/plans/bhj.txt').read_text()
print('Computed result consistency PASS')
