"""Shared report layout, derived only from the ISURO title page."""
from copy import deepcopy
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image

ROOT=Path(__file__).resolve().parents[3]
REFERENCE=ROOT/'4'/'ИСУРО'/'ИСУРО_1_АлбахтинИВ.docx'

def font(item,size=14,bold=False,italic=False):
    item.font.name='Times New Roman'; item.font.size=Pt(size)
    item.font.bold=bold; item.font.italic=italic; item.font.color.rgb=RGBColor(0,0,0)
    node=item._element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(node.attrib):
        if 'Theme' in key: del node.attrib[key]
    for key in ['ascii','hAnsi','eastAsia','cs']: node.set(qn('w:'+key),'Times New Roman')

def spacing(p,align=A.JUSTIFY,first=1.25,line=1.5,before=0,after=0,next=False):
    f=p.paragraph_format; f.alignment=align; f.first_line_indent=Cm(first)
    f.left_indent=Cm(0); f.right_indent=Cm(0); f.line_spacing=line
    f.space_before=Pt(before); f.space_after=Pt(after); f.keep_with_next=next
    f.widow_control=True

class Report:
    def __init__(self,number,topic):
        self.doc=Document(REFERENCE)
        old_section=deepcopy(self.doc.sections[0]._sectPr)
        body=self.doc.element.body
        for node in list(body)[16:]: body.remove(node)
        body.append(old_section)
        d=self.doc
        title=d.paragraphs
        for i,text in [(1,'Кафедра прикладной математики'),(4,'ОТЧЕТ ПО ПРАКТИЧЕСКОЙ РАБОТЕ'),(5,'по дисциплине «Технологии и инструментарий анализа больших данных»'),(6,f'Практическая работа №{number}\n{topic}')]:
            title[i].text=text
            for r in title[i].runs: font(r,16 if i==4 else 14,bold=i in [1,4,6])
            spacing(title[i],A.CENTER,0,1,after=3)
        # Clear the old white group instead of retaining it in the document.
        signature=d.tables[1]
        # The template has a non-rectangular grid. Address row cells directly:
        # Table.cell(1, 1) otherwise selects the signature instead of the name.
        signature.rows[0].cells[0].text='Студент группы'
        signature.rows[0].cells[1].text='ИНБО-12-23, Албахтин И.В.'
        signature.rows[1].cells[1].text='Воронцов Ю.А.'
        for table in d.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            # Retain existing title-block sizes and emphasis.
                            size=r.font.size.pt if r.font.size else 12
                            font(r,size,r.bold or False,r.italic or False)
                    for shd in cell._tc.get_or_add_tcPr().findall(qn('w:shd')): cell._tc.get_or_add_tcPr().remove(shd)
        for p in title:
            for r in p.runs:
                font(r,r.font.size.pt if r.font.size else 14,r.bold or False,r.italic or False)
        sec=d.sections[0]
        sec.page_width=Cm(21); sec.page_height=Cm(29.7)
        sec.left_margin=Cm(3); sec.right_margin=Cm(1.5); sec.top_margin=Cm(2); sec.bottom_margin=Cm(2)
        sec.different_first_page_header_footer=True
        for header in [sec.header,sec.first_page_header]:
            for p in header.paragraphs: p.text=''
        footer=sec.footer.paragraphs[0]; footer.text=''; spacing(footer,A.CENTER,0,1)
        field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),' PAGE ')
        r=OxmlElement('w:r'); pr=OxmlElement('w:rPr'); fonts=OxmlElement('w:rFonts')
        for k in ['ascii','hAnsi']: fonts.set(qn('w:'+k),'Times New Roman')
        sz=OxmlElement('w:sz'); sz.set(qn('w:val'),'24'); pr.extend([fonts,sz]); r.append(pr)
        txt=OxmlElement('w:t'); txt.text='2'; r.append(txt); field.append(r); footer._p.append(field)
        sec.first_page_footer.paragraphs[0].text=''
        normal=d.styles['Normal']; font(normal)
        f=normal.paragraph_format; f.alignment=A.JUSTIFY; f.first_line_indent=Cm(1.25); f.left_indent=Cm(0); f.right_indent=Cm(0); f.line_spacing=1.5; f.space_before=Pt(0); f.space_after=Pt(0); f.widow_control=True
        for name,align in [('Heading 1',A.CENTER),('Heading 2',A.LEFT)]:
            if name not in d.styles: d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
            st=d.styles[name]; font(st,bold=True); f=st.paragraph_format; f.alignment=align; f.first_line_indent=Cm(0); f.left_indent=Cm(0); f.right_indent=Cm(0); f.line_spacing=1.5; f.space_before=Pt(12); f.space_after=Pt(6); f.keep_with_next=True; f.keep_together=True
        for name,size,align,italic,line in [('Report Body',14,A.JUSTIFY,False,1.5),('Report Table',12,A.LEFT,False,1),('Report Table Caption',12,A.LEFT,True,1),('Report Figure Caption',14,A.CENTER,False,1),('Report Code',12,A.LEFT,False,1)]:
            st=d.styles[name] if name in d.styles else d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
            st.base_style=normal; font(st,size,italic=italic)
            f=st.paragraph_format; f.alignment=align; f.first_line_indent=Cm(1.25 if name=='Report Body' else 0); f.left_indent=Cm(0); f.right_indent=Cm(0); f.line_spacing=line; f.space_before=Pt(0); f.space_after=Pt(0)
        # Drop image relationships not used by the retained title page.
        used=set(body.xpath('.//@r:embed'))
        for rel in list(d.part.rels.values()):
            if rel.reltype.endswith('/image') and rel.rId not in used: d.part.drop_rel(rel.rId)
        # A break-only paragraph can overflow the tightly laid-out title page
        # and create an empty second page. Put the break on the first heading.
        self.body_started=False
        self.figures=0; self.tables=0

    def h(self,text,level=1):
        p=self.doc.add_heading(text,level)
        if not self.body_started:
            p.paragraph_format.page_break_before=True
            self.body_started=True
        for r in p.runs: font(r,bold=True)
        return p

    def p(self,text):
        p=self.doc.add_paragraph(text,'Report Body')
        for r in p.runs: font(r)
        return p

    def code(self,text):
        p=self.doc.add_paragraph(text.strip(),'Report Code'); spacing(p,A.LEFT,0,1,before=4,after=6)
        for r in p.runs: font(r,12)

    def table(self,title,headers,rows,widths=None):
        self.tables+=1
        self.p(f'Результаты приведены в таблице {self.tables}.')
        cap=self.doc.add_paragraph(f'Таблица {self.tables} – {title}','Report Table Caption')
        spacing(cap,A.LEFT,0,1,6,3,True); cap.paragraph_format.keep_together=True
        t=self.doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER
        if widths:
            for column,width in zip(t.columns,widths): column.width=Cm(width)
        values_by_row=[headers]+list(rows)
        for row_index,values in enumerate(values_by_row):
            row=t.rows[0] if row_index==0 else t.add_row()
            trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
            if row_index==0: trpr.append(OxmlElement('w:tblHeader'))
            for i,(cell,value) in enumerate(zip(row.cells,values)):
                if widths: cell.width=Cm(widths[i])
                cell.text=str(value); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr()
                for shd in tcpr.findall(qn('w:shd')): tcpr.remove(shd)
                shd=OxmlElement('w:shd'); shd.set(qn('w:val'),'nil'); tcpr.append(shd)
                for p in cell.paragraphs:
                    p.style=self.doc.styles['Report Table']; spacing(p,A.CENTER if row_index==0 or i>0 else A.LEFT,0,1)
                    p.paragraph_format.keep_with_next=(row_index==0 or (len(values_by_row)<=6 and row_index<len(values_by_row)-1))
                    for r in p.runs: font(r,12,bold=row_index==0)
        return t

    def figure(self,path,caption,width=16):
        self.figures+=1
        self.p(f'На рисунке {self.figures} показан результат выполнения на стенде.').paragraph_format.keep_with_next=True
        p=self.doc.add_paragraph(); spacing(p,A.CENTER,0,1,next=True)
        with Image.open(path) as im: ratio=im.height/im.width
        width=min(width,17/ratio)
        p.add_run().add_picture(str(path),width=Cm(width))
        cap=self.doc.add_paragraph(f'Рисунок {self.figures} – {caption}','Report Figure Caption')
        spacing(cap,A.CENTER,0,1,3,6); cap.paragraph_format.keep_together=True

    def questions(self,questions,answers):
        self.h('Ответы на контрольные вопросы')
        assert len(questions)==len(answers)
        for question,answer in zip(questions,answers):
            p=self.doc.add_paragraph(question,'Report Body'); spacing(p,A.LEFT,0,1.5,6,2,True)
            for r in p.runs: font(r,bold=True)
            self.p(answer)

    def save(self,path):
        self.doc.core_properties.author='Албахтин И.В.'
        self.doc.core_properties.title=path.stem
        settings=self.doc.settings.element
        f=OxmlElement('w:updateFields'); f.set(qn('w:val'),'true'); settings.append(f)
        self.doc.save(path)
