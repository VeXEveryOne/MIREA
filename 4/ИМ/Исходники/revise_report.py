"""Controlled revision of the unified IM report, preserving its original title.

Use the original unified report as --baseline (not an already revised output).
Models/print figures are built first. Work and trial outputs belong in .cache.
"""
from __future__ import annotations
import argparse
from copy import deepcopy
from decimal import Decimal, ROUND_FLOOR, ROUND_HALF_UP
import json
from pathlib import Path
import re
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from PIL import Image

def font(run, size=14, bold=False, italic=False):
    run.font.name='Times New Roman';run.font.size=Pt(size)
    run.font.color.rgb=RGBColor(0,0,0);run.bold=bold;run.italic=italic
    for name in ('ascii','hAnsi','eastAsia','cs'):
        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:'+name),'Times New Roman')

def body_format(p):
    f=p.paragraph_format;f.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    f.left_indent=Cm(0);f.right_indent=Cm(0);f.first_line_indent=Cm(1.25)
    f.line_spacing=1.5;f.space_before=Pt(0);f.space_after=Pt(0)
    f.widow_control=True;f.keep_together=False;f.keep_with_next=False
    for r in p.runs:font(r)

def table_format(t,widths):
    t.autofit=False
    t.alignment=WD_TABLE_ALIGNMENT.LEFT
    pr=t._tbl.tblPr
    for old in pr.findall(qn('w:tblInd')):pr.remove(old)
    indent=OxmlElement('w:tblInd');indent.set(qn('w:w'),'0');indent.set(qn('w:type'),'dxa');pr.append(indent)
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for row_i,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr()
        if row_i==0 and not pr.find(qn('w:tblHeader')) is not None:pr.append(OxmlElement('w:tblHeader'))
        if pr.find(qn('w:cantSplit')) is None:pr.append(OxmlElement('w:cantSplit'))
        for cell,w in zip(row.cells,widths):
            cell.width=Cm(w)
            for fill in cell._tc.get_or_add_tcPr().findall(qn('w:shd')):cell._tc.get_or_add_tcPr().remove(fill)
            for p in cell.paragraphs:
                f=p.paragraph_format;f.alignment=WD_ALIGN_PARAGRAPH.LEFT
                f.left_indent=Cm(0);f.right_indent=Cm(0);f.first_line_indent=Cm(0)
                f.line_spacing=1.05;f.space_before=Pt(0);f.space_after=Pt(0)
                f.keep_with_next=row_i==0 or (len(t.rows)<=4 and row_i<len(t.rows)-1);f.keep_together=True;f.widow_control=True
                for r in p.runs:font(r,12,row_i==0)

def replace_picture(doc,p,file,max_width,max_height):
    # Replace only this drawing's media and its dimensions; title parts untouched.
    blips=p._p.xpath('.//*[local-name()="blip"]')
    if len(blips)!=1:raise ValueError('Expected one body picture')
    rid=blips[0].get(qn('r:embed'));doc.part.related_parts[rid]._blob=file.read_bytes()
    with Image.open(file) as im:ratio=im.width/im.height
    width=min(max_width,max_height*ratio);height=width/ratio
    for e in p._p.xpath('.//*[local-name()="extent" or local-name()="ext"]'):
        if e.get('cx') and e.get('cy'):e.set('cx',str(int(Cm(width))));e.set('cy',str(int(Cm(height))))
    f=p.paragraph_format;f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.first_line_indent=Cm(0)
    f.line_spacing=1;f.space_before=Pt(0);f.space_after=Pt(0);f.keep_with_next=True;f.keep_together=True

def new_paragraph(doc,anchor,text='',style=None):
    p=doc.add_paragraph(text,style);anchor.addprevious(p._p);body_format(p);return p

def section_break(doc,anchor,orientation):
    p=new_paragraph(doc,anchor)
    sect=deepcopy(doc.sections[-1]._sectPr)
    for old in list(sect):
        if old.tag in [qn('w:headerReference'),qn('w:footerReference')]:sect.remove(old)
    size=sect.find(qn('w:pgSz'));margin=sect.find(qn('w:pgMar'))
    if orientation=='portrait':
        size.set(qn('w:w'),str(Cm(21).twips));size.set(qn('w:h'),str(Cm(29.7).twips));size.attrib.pop(qn('w:orient'),None)
        values={'left':3,'right':2.5,'top':2.72,'bottom':0.49}
    else:
        size.set(qn('w:w'),str(Cm(29.7).twips));size.set(qn('w:h'),str(Cm(21).twips));size.set(qn('w:orient'),'landscape')
        values={'left':0.49,'right':2.72,'top':3,'bottom':2.5}
    for k,v in values.items():margin.set(qn('w:'+k),str(Cm(v).twips))
    st=sect.find(qn('w:type'))
    if st is None:st=OxmlElement('w:type');sect.insert(0,st)
    st.set(qn('w:val'),'nextPage')
    p._p.get_or_add_pPr().append(sect)
    p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=False
    for r in p.runs:font(r,1)
    return p

def money(value):return f'{Decimal(value):,.2f}'.replace(',',' ').replace('.',',')

def revise(args):
    doc=Document(args.baseline);ps=list(doc.paragraphs);ts=list(doc.tables)
    assert ps[55].text=='Практическая работа № 1' and ps[519].text=='Заключение'
    title_nodes=list(doc.element.body)[:list(doc.element.body).index(ps[39]._p)]
    title_before=[etree.tostring(n) for n in title_nodes]
    replacements={
      60:'1. Систематизированы открытые сведения об организации и проектные предположения о процессе автоматизации.',
      71:'Процесс начинается с решения создать или изменить карточку. Результат — подтверждённая публикация, классифицированная ошибка с ответственным либо незавершённая попытка, ожидающая проверки. Приём внешней задачи не означает завершённую публикацию. Границы процесса представлены в таблице 1.2.',
      91:'На рисунке 1.1 представлена организационная структура ООО «Юкомс» из исходной модели ППОИС [2]. Внутренние роли и подчинение приняты как предположения учебного проекта; штатное расписание организации не проверялось.',
      175:'На рисунке 3.1 представлена ИСР: четыре стадии и десять пакетов. Сроки и трудозатраты раскрыты в таблицах 3.4 и 3.5.',
      191:'План разработчика по неделям Н1–Н12: 0, 0, 28, 28, 16, 36, 32, 12, 16, 4, 4, 16 ч. Итого 192 ч; максимум 36 ч/нед. при доступности 40 ч/нед. Остальные 240 ч приходятся на другие роли. Рабочая версия 3.1 передаётся для 3.2 в начале Н6.',
      210:'Начало процесса — задание на новую карточку или изменение состава, цены либо наличия. Окончание — подтверждённая публикация или документированное решение об отказе либо приостановке. Сборка, продажи, доставка и бухгалтерия находятся за границей модели.',
      213:'На рисунке 4.1 представлен контекст процесса AS-IS.',
      255:'На рисунке 5.1 представлен контекст TO-BE: утверждённая версия, подтверждённый результат или адресная задача, протокол и показатели.',
      259:'В таблице 5.2 раскрыты стрелки ICOM. Документы контроля N01–N05 относятся к выходу O3 — протоколу и истории операций, а не к O4 с управленческими показателями.',
      271:'В таблице 5.4 раскрыты документы выхода O3. Они дополняют результаты практики 1 и обеспечивают контроль условий выполнения операций.',
      288:'При 100 операциях в месяц активное время уменьшается со 100 до 50 ч: 100 × (60 − 30) / 60 = 50 ч/мес. Высвобождается 50% времени подготовки. При оценке часа 700 руб. его эквивалент — 35 000 руб./мес.; это доступная ёмкость, а не гарантированная экономия расходов или штата.',
      289:'Предварительное отношение 592 000 / 35 000 = 16,91 месяца не учитывает эксплуатацию и не является уточнённой окупаемостью. Расчёт с эксплуатационными расходами и проверкой монетизации времени приведён в практике 12.',
      320:'Проведён обзор официальных описаний и тарифов МойСклад, «1С:Управление нашей фирмой» и «1С-Битрикс: Управление сайтом». Источники повторно проверены 05.10.2026. Сравниваются типовые функции и потребность в адаптации; фактические испытания и внедрение этих продуктов не выполнялись.',
      345:'Модельный ряд имеет несколько версий; версия содержит строки состава, а строка — компонент. Компонент входит в разные версии. Карточка ссылается на выбранную версию, сохраняя прежние связи в истории. Попытки относятся к зафиксированным пакетам. Одно изображение допускается в нескольких карточках после подтверждения соответствия.',
      346:'Контроллеры импорта, состава, расчёта, карточки и публикации координируют сценарии; сущности хранят состояние, экран служит границей взаимодействия. Объект modelRangeVersion:ВерсияМодельногоРяда фиксирует состав, publication:ПопыткаПубликации — отправленный пакет. Предметные правила отделены от Ozon.',
      429:'На рисунке 9.1 показаны три зоны ответственности. Неготовая карточка возвращается на доработку. Если окончательный результат не получен, сохраняются ожидание и срок нового опроса той же операции; повторная публикация не запускается.',
      435:'На рисунке 9.2 показаны регистрация и передача пакета (1–8), затем отдельный запрос и сохранение результата (9–12). Пунктиром обозначены ответы. Интервал между 8 и 9 определяется регламентом, а не масштабом рисунка.',
      440:'Проверка готовности ведёт к доработке с завершением текущего запуска либо к отправке. Проверка окончательного результата ведёт к сохранению итога или ожиданию и циклу опроса через узел слияния. Параллельных потоков и безусловной повторной отправки в модели нет.',
      453:'На рисунке 10.1 показаны компоненты и зависимости: веб-интерфейс и импорт используют прикладные сервисы; сервисы обращаются к контролю доступа, репозиторию и адаптеру Ozon. Подписи раскрывают назначение каждого обращения.',
      461:'На рисунке 10.2 показаны узлы, среды исполнения и вложенные артефакты. Сплошные линии обозначают каналы. Соответствие артефактов логическим компонентам раскрыто в таблицах 10.2 и 10.3; отношения manifestation сохранены в исходной UML-модели.',
      475:'На рисунке 11.1 показаны 60 рабочих дней, зависимости FS/SS и перекрытие пакетов 3.1–3.3. Тёмными полосами отмечены критические работы. FS означает начало после завершения предшественника; SS с лагом — начало после его старта и указанной задержки.',
      487:'План согласован с ИСР: 432 чел.-ч, 54 чел.-дн и 60 рабочих дней. Недельная загрузка разработчика не превышает 40 ч; отклонения устраняются пересмотром сроков.',
      493:'Полные затраты = часы × полная ставка; оплата = полные затраты / 1,302; отчисления = полные затраты − оплата. Итог оплаты округляется до копеек. На уровне назначений остаток распределён методом наибольших дробных частей; суммы по ролям и пакетам поэтому совпадают без копеечных расхождений.',
      518:'Смета 592 000 руб. включает резерв; эксплуатация — 10 000 руб./мес. Окупаемость 23,68 месяца условна: на пилоте нужно подтвердить объём операций и полезное использование высвобождённого времени.',
      520:'В практиках 1–12 разработан учебный проект ИС подготовки и актуализации карточек ООО «Юкомс». По открытым сведениям и проектным предположениям описаны участники, девять ожидаемых результатов, 14 текущих потоков и шесть проблем. AS-IS и TO-BE показывают переход от файлов к связанным версиям, расчётам и управляемым публикациям.',
      521:'FURPS+ включает 27 проверяемых требований. Обзор трёх продуктов обосновывает специализированное приложение для учебного проекта; окончательный выбор требует оценки адаптации и полной стоимости альтернатив. Подготовлены десять прецедентов с альтернативами, UML-модели поведения и архитектуры, ИСР и Гант. Таблицы согласованы с диаграммами.',
      522:'План составляет 12 недель и 432 чел.-ч; лимит с резервом — 592 000 руб., эксплуатация — 120 000 руб./год. При 100 операциях в месяц и экономии 30 минут расчётная окупаемость составляет 23,68 месяца после приёмки. Отчёт не подтверждает фактические интервью, внедрение или достижение метрик.',
      523:'Далее требуется согласовать требования и данные с участниками, проверить внешний обмен и исходные показатели, реализовать ограниченную поставку и провести пилот. Решение принимается по качеству, трудозатратам и полезному использованию высвобождённого времени.',
    }
    for i,value in replacements.items():ps[i].text=value
    for p in ps[55:]:
        if '32 ч/нед' in p.text:p.text=p.text.replace('32 ч/нед','40 ч/нед')
        if p.text.startswith('В ходе практической работы выполнено экспресс-обследование'):
            p.text=p.text.replace('выполнено экспресс-обследование','подготовлено учебное предпроектное описание')
    for row in ts[14].rows:
        for c in row.cells:
            if '32 ч' in c.text:c.text=c.text.replace('32 ч','40 ч')
    for i in range(530,537):ps[i].text=ps[i].text.replace('27.09.2026','05.10.2026')
    # Move explanatory figure references before the drawing, not after it.
    for lead,picture in [(429,427),(435,432),(453,451),(461,459),(475,473)]:
        ps[picture]._p.addprevious(ps[lead]._p)
    # Isolate wide architecture figures into A4 landscape sections.
    for lead,picture,caption in [(453,451,452),(461,459,460)]:
        section_break(doc,ps[lead]._p,'portrait')
        anchor=ps[caption]._p.getnext();section_break(doc,anchor,'landscape')
    files={93:('ПР1_Оргструктура',26.49,12),214:('ПР4_AS_IS_context',26.49,11.5),221:('ПР4_AS_IS_decomposition',26.49,11.8),256:('ПР5_TO_BE_context',26.49,11.5),263:('ПР5_TO_BE_decomposition',26.49,11.8),350:('ПР8_Прецеденты_данные',26.49,11.8),354:('ПР8_Прецеденты_публикация',26.49,11.8),427:('ПР9_Деятельность',15.5,21.7),432:('ПР9_Последовательность',26.49,12.8),451:('ПР10_Компоненты',26.49,12.8),459:('ПР10_Развёртывание',26.49,12.8),473:('ПР11_Гант',26.49,12.8)}
    # Identify use-case pictures by captions to avoid relying on an unchecked index.
    for i in list(files):
        if i in [350,354] and not ps[i]._p.xpath('.//*[local-name()="blip"]'):del files[i]
    for i,p in enumerate(ps):
        if p.text.startswith('Рисунок 8.1'):files[i-1]=('ПР8_Прецеденты_данные',26.49,11.8)
        if p.text.startswith('Рисунок 8.2'):files[i-1]=('ПР8_Прецеденты_публикация',26.49,11.8)
    # Three introductory lines and the heading consume space above A0.
    # Leave explicit room for its two-line caption in the landscape section.
    for i in [221,263]:
        name,w,_=files[i];files[i]=(name,w,10.6)
    for i,h in [(427,19.0),(432,11.4),(451,11.4),(459,11.4),(473,11.4)]:
        name,w,_=files[i];files[i]=(name,w,h)
    for i,(name,w,h) in files.items():replace_picture(doc,ps[i],args.figures/(name+'.png'),w,h)
    # Bound the existing WBS so its lead, image and caption share the same page.
    drawing=ps[176]._p.xpath('.//*[local-name()="blip"]')[0]
    media=doc.part.related_parts[drawing.get(qn('r:embed'))]
    tmp=args.working/'wbs.png';tmp.write_bytes(media.blob);replace_picture(doc,ps[176],tmp,26.49,10.5)
    # Add the previously omitted FURPS+ classification as two landscape parts.
    anchor=ps[315]._p;section_break(doc,anchor,'portrait')
    for part,name,caption in [('а','ПР6_FURPS_Функции','функциональность'),('б','ПР6_FURPS_Качества','качества и ограничения')]:
        lead=new_paragraph(doc,anchor,f'На рисунке 6.1, {part}, представлена классификация требований FURPS+.')
        if part=='б':lead.paragraph_format.page_break_before=True
        pic=new_paragraph(doc,anchor)
        with Image.open(args.figures/(name+'.png')) as im:ratio=im.width/im.height
        width=min(26.49,12.8*ratio);pic.add_run().add_picture(str(args.figures/(name+'.png')),width=Cm(width))
        cap=new_paragraph(doc,anchor,f'Рисунок 6.1, {part} – FURPS+: {caption}')
    section_break(doc,anchor,'landscape')
    # Split actors from include relations, preserving all associations.
    includes=[[c.text for c in r.cells] for r in ts[34].rows[-2:]]
    for row in list(ts[34].rows)[-2:]:ts[34]._tbl.remove(row._tr)
    include_cap=new_paragraph(doc,ps[348]._p,'Таблица 8.3 – Обязательные включаемые прецеденты')
    new=new_paragraph(doc,include_cap._p,'Обязательные включения раскрыты в таблице 8.3 и показаны на рисунке 8.1.')
    inc=doc.add_table(rows=1,cols=3);inc.style='Table Grid';include_cap._p.addnext(inc._tbl)
    for c,v in zip(inc.rows[0].cells,['Связь','Тип','Смысл']):c.text=v
    for values in includes:
        for c,v in zip(inc.add_row().cells,values):c.text=v
    table_format(inc,[4,3,8.5])
    # Match the guide's numbered dictionaries and five-column activity mapping.
    for n in [35,36,37]:
        old=ts[n];data=[[c.text for c in r.cells] for r in old.rows]
        if n==35:data[3][0]='svc : Сервис публикации'
        if n==36:
            for row in data:
                if row[0]=='Сторожевое условие':row[1]='[да] / [else] для готовности; [итог получен] / [else] для результата.'
                if row[0]=='Решение':row[1]='Готовность карточки или наличие окончательного подтверждённого результата.'
            data.append(['Узел слияния','Объединяет первоначальный и повторный опрос одной попытки без синхронизации параллельных ветвей.'])
        if n==37:
            data=[data[0],
              ['Выбрать и подтвердить карточку','Проверить поля и версии','Менеджер','Карточка'],
              ['Проверить поля и версии','Готовность: [да] — сохранить; [else] — исправить','ИС','Карточка, версия'],
              ['Исправить замечания','Завершить текущий запуск; после подготовки — новый запуск','Менеджер','Черновик'],
              ['Сохранить попытку до вызова','Передать пакет','ИС','Попытка, пакет'],
              ['Передать пакет','Принять пакет в обработку','ИС','Зафиксированный пакет'],
              ['Принять пакет в обработку','Узел слияния → получить статус','Ozon','Внешняя операция'],
              ['Получить статус / тайм-аут','Итог: [итог получен] — сохранить; [else] — ожидание','ИС','Ответ, попытка'],
              ['Сохранить ожидание той же операции','Дождаться срока следующего опроса','ИС','Попытка, время опроса'],
              ['Дождаться срока следующего опроса','Через слияние снова получить статус той же операции','ИС','Та же попытка'],
              ['Сохранить итог и замечания','Завершить текущий сценарий','ИС','Итог, замечания']]
        t=doc.add_table(rows=1,cols=len(data[0])+1);t.style='Table Grid';old._tbl.addprevious(t._tbl)
        for c,v in zip(t.rows[0].cells,['№']+data[0]):c.text=v
        for i,values in enumerate(data[1:],1):
            for c,v in zip(t.add_row().cells,[str(i)]+values):c.text=v
        old._tbl.getparent().remove(old._tbl)
        table_format(t,[1,5,9.5] if n in [35,36] else [1,3.5,4.4,2.6,4])
    # Financial cents are allocated once per native assignment, then aggregated.
    model=json.loads((args.models/'ИМ_ПР11_Гант.omni').read_text(encoding='utf-8'))['model']['elements']
    rates=dict(zip(['analyst','lead','developer','qa','experts','admin'],[1000,1500,1200,1000,1000,1600]))
    assignments=[e for e in model if e['kind']=='assignment'];denominator=Decimal('1.302')
    exact=[Decimal(a['properties']['effortHours'])*rates[a['properties']['resource']]/denominator*100 for a in assignments]
    cents=[int(v.to_integral_value(rounding=ROUND_FLOOR)) for v in exact]
    target=int(sum(exact).to_integral_value(rounding=ROUND_HALF_UP))
    for i in sorted(range(len(exact)),key=lambda i: (-(exact[i]-cents[i]),assignments[i]['id']))[:target-sum(cents)]:cents[i]+=1
    role={};package={}
    for a,c in zip(assignments,cents):
        p=a['properties'];full=Decimal(p['effortHours'])*rates[p['resource']];payment=Decimal(c)/100
        for group,key in [(role,p['resource']),(package,p['task'])]:
            v=group.setdefault(key,[Decimal(0),Decimal(0)]);v[0]+=payment;v[1]+=full-payment
    assert sum(v[0] for v in role.values())==Decimal('393241.17') and sum(v[1] for v in role.values())==Decimal('118758.83')
    for row,r in zip(ts[45].rows[1:-1],rates):
        row.cells[3].text=money(role[r][0]);row.cells[4].text=money(role[r][1])
    for row in ts[48].rows[1:-1]:
        key='wp'+row.cells[0].text.replace('.','_');row.cells[1].text=money(package[key][0]);row.cells[2].text=money(package[key][1])
    for row in ts[44].rows[1:]:row.cells[2].text=row.cells[2].text.replace('.',',')
    # Reapply the IM-only style, leaving the immutable title and TOC content alone.
    for p in doc.paragraphs:
        if p._p in title_nodes or list(doc.element.body).index(p._p)<list(doc.element.body).index(ps[55]._p):continue
        text=p.text;is_picture=bool(p._p.xpath('.//*[local-name()="blip"]'))
        if p._p.xpath('./w:pPr/w:sectPr'):continue
        old_heading=p.style.name=='IM Head' or 'heading' in p.style.name.lower() or 'заголовок' in p.style.name.lower()
        body_format(p)
        if text.startswith('Практическая работа №') or text in ['Заключение','Список информационных источников']:
            p.style='Heading 1';p.paragraph_format.page_break_before=True;p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            for r in p.runs:font(r,16,True)
        elif old_heading:
            p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.left_indent=Cm(1.27);p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.keep_with_next=True
            for r in p.runs:font(r,14,True)
        elif text.startswith('Таблица '):
            p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(6);p.paragraph_format.keep_with_next=True;p.paragraph_format.keep_together=True
            for r in p.runs:font(r,12,False,True)
        elif text.startswith('Рисунок '):
            p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(6);p.paragraph_format.keep_together=True
            for r in p.runs:font(r,14)
        elif is_picture:
            p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1;p.paragraph_format.keep_with_next=True;p.paragraph_format.keep_together=True
        elif re.search(r'(?:На рисунке|В таблице|приведены в таблице)',text):p.paragraph_format.keep_with_next=True
        # Empty body paragraphs are not accidental spacer lines.
        elif not text:p.paragraph_format.line_spacing=1;[font(r,1) for r in p.runs]
    widths={8:[3.6,6,5.9],30:[1.8,4.57,4.57,4.56],34:[4,3,8.5],42:[1.6,4.5,2.8,2.2,3.4,4,7.99],50:[3.2,3.2,3,3,3.1]}
    for i,t in enumerate(ts):
        if i==0 or t._tbl.getparent() is None:continue
        total=sum(c.width.cm for c in t.columns if c.width)
        ws=widths.get(i,[c.width.cm for c in t.columns])
        if not ws or any(w is None for w in ws) or len(ws)!=len(t.columns):raise ValueError('Missing table widths')
        table_format(t,ws)
        for row in t.rows:
            for cell in row.cells:
                if cell.text=='Идемпотентная отправка, опрос результата и повторы':
                    cell.paragraphs[0].text='Уникальная локальная операция, сверка результата и безопасные повторы'
                    body_format(cell.paragraphs[0]);cell.paragraphs[0].paragraph_format.first_line_indent=Cm(0);cell.paragraphs[0].paragraph_format.line_spacing=1.05
                    for run in cell.paragraphs[0].runs:font(run,12)
    # Titles and source metadata must not be altered by this controlled revision.
    assert title_before==[etree.tostring(n) for n in title_nodes]
    args.output.parent.mkdir(parents=True,exist_ok=True);doc.save(args.output)
    (args.working/'revision_log.json').write_text(json.dumps({'title_preserved':True,'replaced_paragraphs':sorted(replacements),'financial_roles':role,'financial_packages':package,'total_payment':str(sum(v[0] for v in role.values())),'total_contributions':str(sum(v[1] for v in role.values()))},ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    print(args.output)

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__)
    for name in ['baseline','figures','models','output','working']:p.add_argument('--'+name,type=Path,required=True)
    args=p.parse_args();args.working.mkdir(parents=True,exist_ok=True);revise(args)
