from pathlib import Path
from copy import deepcopy
import json,hashlib,zipfile,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START,WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image
from content import blocks,ROOT
from config import BUILD_DIR as TMP, TITLE_TEMPLATE as template, REPORT_NAME
TMP.mkdir(parents=True, exist_ok=True)
doc=Document(template)
body=doc.element.body
sect=deepcopy(body[15].xpath('.//w:sectPr')[0])
for child in list(body)[16:]:body.remove(child)
body.append(sect)
teacher=doc.tables[1].rows[1].cells[1]
for pp in teacher.paragraphs:
 if pp.runs:
  pp.runs[0].text='Кириллина Ю.В.,\nТрушин С.М.'
  for rr in pp.runs[1:]:rr.text=''
for e in body.xpath('.//w:t'):
 if e.text:
  e.text=e.text.replace('ИНБО-01-17','').replace('Информационные системы управления ресурсами организации','Разработка обеспечивающих подсистем').replace('Шендяпин А.В.','Кириллина Ю.В., Трушин С.М.')
for rfont in body.xpath('.//w:rFonts'):
 for key in ('ascii','hAnsi','eastAsia','cs'):rfont.set(qn('w:'+key),'Times New Roman')
# Only the first page is retained from ISURO. All new body styles follow the repository standard.
for name in ['Normal','Body Text','Caption','Heading 1','Heading 2','Heading 3','Title']:
 if name not in doc.styles:doc.styles.add_style(name,1)
 st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(14);st.font.color.rgb=RGBColor(0,0,0)
 pf=st.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=1.5;pf.left_indent=Cm(0);pf.right_indent=Cm(0);pf.first_line_indent=Cm(1.25);pf.widow_control=True
 if name.startswith('Heading'):
  lvl=OxmlElement('w:outlineLvl');lvl.set(qn('w:val'),str(int(name[-1])-1));st.element.get_or_add_pPr().append(lvl)
  st.font.bold=True;pf.first_line_indent=0;pf.keep_with_next=True;pf.keep_together=True;pf.space_before=Pt(12);pf.space_after=Pt(6);pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if name=='Heading 1' else WD_ALIGN_PARAGRAPH.LEFT
  pf.page_break_before=name=='Heading 1'
 else:pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ['TableCaptionROP','FigureCaptionROP','TableROP','ReferencesROP']:
 if name not in doc.styles:doc.styles.add_style(name,1)
 st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(14 if name=='FigureCaptionROP' else 12);st.font.color.rgb=RGBColor(0,0,0)
 st.font.italic=name=='TableCaptionROP';st.font.bold=False
 pf=st.paragraph_format;pf.left_indent=0;pf.right_indent=0;pf.first_line_indent=0;pf.line_spacing=1;pf.space_before=Pt(3 if name=='FigureCaptionROP' else 0);pf.space_after=Pt(6 if name=='FigureCaptionROP' else 0);pf.alignment=WD_ALIGN_PARAGRAPH.CENTER if name=='FigureCaptionROP' else WD_ALIGN_PARAGRAPH.LEFT
 if name=='TableCaptionROP':pf.keep_with_next=True;pf.space_before=Pt(6);pf.space_after=Pt(3)
 if name=='FigureCaptionROP':pf.keep_together=True
 if name=='ReferencesROP':
  st.font.size=Pt(14);pf.alignment=WD_ALIGN_PARAGRAPH.LEFT;pf.line_spacing=1.5;pf.first_line_indent=0;pf.left_indent=0;pf.space_after=Pt(0)
def field(p,code,cached):
 r=p.add_run();begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');r._r.append(begin)
 instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' '+code+' ';r._r.append(instr)
 sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r._r.append(sep)
 p.add_run(str(cached));end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');p.add_run()._r.append(end)
def setup_section(s,land=False):
 s.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT;s.page_width=Cm(29.7 if land else 21);s.page_height=Cm(21 if land else 29.7)
 s.left_margin=Cm(3);s.right_margin=Cm(1.5);s.top_margin=Cm(2);s.bottom_margin=Cm(2);s.header_distance=Cm(.8);s.footer_distance=Cm(1)
 s.different_first_page_header_footer=False
 for pg in s._sectPr.findall(qn('w:pgNumType')):s._sectPr.remove(pg)
 for hdr in [s.header,s.first_page_header,s.footer,s.first_page_footer]:
  hdr.is_linked_to_previous=False
  for p in hdr.paragraphs:p.clear()
 fp=s.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;fp.paragraph_format.first_line_indent=0;field(fp,'PAGE','2')
 for r in fp.runs:r.font.name='Times New Roman';r.font.size=Pt(12)
for s in doc.sections:setup_section(s)
doc.sections[0].different_first_page_header_footer=True
doc.core_properties.author='Албахтин И.В.';doc.core_properties.title='РОП Практические занятия 1–8 ООО Юкомс'
setting=OxmlElement('w:updateFields');setting.set(qn('w:val'),'true');doc.settings.element.append(setting)
current=False
def page(land=False,force=False):
 global current
 if land!=current:
  setup_section(doc.add_section(WD_SECTION_START.NEW_PAGE),land);current=land
 elif force:doc.add_page_break()
def para(text):return doc.add_paragraph(text,'Body Text')
def in_sentence(title):
 # Preserve leading abbreviations such as UC, RBAC and AS IS.
 return title if len(title)>1 and title[:2].isupper() else title[0].lower()+title[1:]
toc=doc.add_paragraph('Содержание','Title');toc.paragraph_format.first_line_indent=0;toc.alignment=WD_ALIGN_PARAGRAPH.CENTER
field(doc.add_paragraph(),'TOC \\o "1-2" \\h \\z \\u','Содержание')
figure=table=0;md=['# РОП Практические занятия 1–8','\nАлбахтин И.В. ИНБО-12-23. ООО «Юкомс».\n']
pending=[]
for bi,b in enumerate(blocks):
 if b['type']=='h':
  if bi+1<len(blocks) and blocks[bi+1]['type'] in ('fig','table'):
   pending.append(b);md+=['\n'+'#'*(b['level']+1)+' '+b['text']+'\n'];continue
  page(False)
  pp=doc.add_paragraph(b['text'],'Heading '+str(b['level']));md+=['\n'+'#'*(b['level']+1)+' '+b['text']+'\n']
 elif b['type']=='p':
  page(False);para(b['text']);md+=[b['text']+'\n']
 elif b['type']=='source':
  page(False);doc.add_paragraph(b['text'],'ReferencesROP');md+=[b['text']+'\n']
 elif b['type']=='table':
  table+=1;heads=b['heads'];rows=b['rows'];wide=len(heads)>=5 or b['title'] in ['Сущности логической модели','Матрица RBAC']
  # Each table starts on a clean page. Long tables are split before insertion,
  # because Word and LibreOffice paginate tall mixed-width rows differently.
  page(wide,True)
  for heading in pending:
   doc.add_paragraph(heading['text'],'Heading '+str(heading['level']))
  pending=[]
  para(f'В таблице {table} представлены '+in_sentence(b['title'])+'.')
  caption=doc.add_paragraph(style='TableCaptionROP');caption.add_run('Таблица ');field(caption,'SEQ Table \\* ARABIC',table);caption.add_run(' – '+b['title'])
  width=25.2 if wide else 16.5;ww=b.get('widths') or [1]*len(heads);ww=[width*v/sum(ww) for v in ww]
  capacity=21 if wide else 32
  chunks=[];chunk=[];used=0
  for source_row in rows:
   lines=max(max(1,(len(str(tx))+max(8,int(ww[i]*4.2))-1)//max(8,int(ww[i]*4.2))) for i,tx in enumerate(source_row))+1
   if chunk and used+lines>capacity:
    chunks.append(chunk);chunk=[];used=0
   chunk.append((source_row,lines));used+=lines
  if chunk:chunks.append(chunk)
  for chunk_index,chunk_rows in enumerate(chunks):
   if chunk_index:
    page(wide,True)
    continuation=doc.add_paragraph(style='TableCaptionROP')
    continuation.add_run(f'Продолжение таблицы {table} – '+b['title'])
   tb=doc.add_table(rows=1, cols=len(heads));tb.autofit=False;tb.style='Table Grid'
   for c,w in zip(tb.columns,ww):c.width=Cm(w)
   for i,tx in enumerate(heads):tb.rows[0].cells[i].text=str(tx)
   for source_row,row_lines in chunk_rows:
    cs=tb.add_row().cells
    for i,tx in enumerate(source_row):cs[i].text=str(tx)
   for ri,row in enumerate(tb.rows):
    trPr=row._tr.get_or_add_trPr()
    row_lines=1 if ri==0 else chunk_rows[ri-1][1]
    if ri==0 or row_lines<=capacity:
     nr=OxmlElement('w:cantSplit');trPr.append(nr)
    if ri==0:
     rp=OxmlElement('w:tblHeader');trPr.append(rp)
    for ci,cell in enumerate(row.cells):
     cell.width=Cm(ww[ci]);pr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
     for side in ['top','left','bottom','right']:
      z=OxmlElement('w:'+side);z.set(qn('w:w'),'65' if side in ['top','bottom'] else '80');z.set(qn('w:type'),'dxa');marg.append(z)
     pr.append(marg)
     for pp in cell.paragraphs:
      pp.style='TableROP';pp.paragraph_format.keep_with_next=False
      for r in pp.runs:r.font.name='Times New Roman';r.font.size=Pt(12);r.bold=ri==0
  md+=['\n*Таблица '+str(table)+' – '+b['title']+'*\n','| '+' | '.join(heads)+' |','| '+' | '.join('---' for x in heads)+' |']
  md+=['| '+' | '.join(str(x).replace('\n','<br>').replace('|','/') for x in row)+' |' for row in rows];md+=['']
 elif b['type']=='fig':
  figure+=1;land=b.get('landscape',False);same_page_orientation=land==current
  if not same_page_orientation:page(land)
  had_heading=bool(pending)
  for hi,heading in enumerate(pending):
   hp=doc.add_paragraph(heading['text'],'Heading '+str(heading['level']))
   if hi==0 and same_page_orientation:hp.paragraph_format.page_break_before=True
  pending=[]
  pp=para(f'На рисунке {figure} показана модель «{b["title"]}».');pp.paragraph_format.keep_with_next=False;pp.paragraph_format.keep_together=True
  if same_page_orientation and not had_heading:pp.paragraph_format.page_break_before=True
  paths=b.get('parts') or [b['path']]
  for part_index,image_path in enumerate(paths):
   image=Image.open(image_path);wmax=25.2 if land else 16.5
   if land and len(paths)>1:hmax=8.6 if had_heading and part_index==0 else 9.2
   else:hmax=(11.0 if had_heading and part_index==0 else 11.6) if land else (18 if had_heading and part_index==0 else 19)
   iw,ih=image.size;w=min(wmax,hmax*iw/ih)
   pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf=pp.paragraph_format;pf.first_line_indent=0;pf.space_after=0;pf.space_before=0;pf.line_spacing=1;pf.keep_with_next=True
   if part_index:pf.page_break_before=True
   pp.add_run().add_picture(image_path,width=Cm(w))
   cp=doc.add_paragraph(style='FigureCaptionROP')
   if part_index==0:
    cp.add_run('Рисунок ');field(cp,'SEQ Figure \\* ARABIC',figure);cp.add_run(' – '+b['title']+(f' (часть 1 из {len(paths)})' if len(paths)>1 else ''))
   else:
    cp.add_run(f'Продолжение рисунка {figure} – {b["title"]} (часть {part_index+1} из {len(paths)})')
  md+=['\n![Рисунок '+str(figure)+' — '+b['title']+']('+str(Path(b['path']).relative_to(ROOT) if Path(b['path']).is_relative_to(ROOT) else Path(b['path']))+')\n']
output=ROOT/(REPORT_NAME+'.docx');doc.save(output)
(ROOT/(REPORT_NAME+'.md')).write_text('\n'.join(md),encoding='utf-8')
(TMP/'artifact.md').write_text(f'''# Контракт титульного листа
Источник: {template}
SHA256: {hashlib.sha256(template.read_bytes()).hexdigest()}
Сохранена только первая страница, body[0:16], включая эмблему, шапку, линии и подписи.
Реквизиты: Албахтин И.В., ИНБО-12-23; Кириллина Ю.В., Трушин С.М.; Москва 2026.
Удалён белый старый текст ИНБО-01-17. Изменена дисциплина.
Источник оформления основной части: Оформление_отчётов.md, прочитан полностью.
A4, поля 3/1.5/2/2 см; TNR 14, 1.5, абзац 1.25; таблицы 12 одинарный.
Широкие модели и таблицы: A4 альбомный, затем возврат в книжный раздел.
Титульная страница считается, номер скрыт. Остальные PAGE по центру 12 пт.
TOC по Heading 1–2, подписи SEQ Figure и SEQ Table, номера в ссылках рассчитаны при сборке.
Исходник не изменяется. Для рендера используется Word + Poppler через адаптер render_docx.py.
Рисунков: {figure}; таблиц: {table}.
''',encoding='utf-8')
print(output,figure,'figures',table,'tables',len(doc.sections),'sections')
