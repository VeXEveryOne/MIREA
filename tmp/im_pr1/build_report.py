from __future__ import annotations

from pathlib import Path
from xml.etree import ElementTree as ET

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


BASE = Path(r"D:\GitHub\MIREA")
OUT_DIR = BASE / "4" / "ИМ"
TMP_DIR = BASE / "tmp" / "im_pr1"
DOCX_PATH = OUT_DIR / "ИМ_Практическая_1_АлбахтинИВ.docx"
DRAWIO_PATH = OUT_DIR / "ИМ_Практическая_1_Диаграммы.drawio"
ORG_PNG = TMP_DIR / "org_structure.png"
FLOW_PNG = TMP_DIR / "functional_contour.png"

FONT = "Times New Roman"
ACCENT = "1F4E79"
ACCENT_LIGHT = "EAF2F8"
GRID = "D9D9D9"


def set_run_font(run, size=14, bold=False, italic=False, color="000000", name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_paragraph_format(p, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=True,
                         before=0, after=6, line=1.5, keep_with_next=False,
                         keep_together=False):
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE if line == 1.5 else WD_LINE_SPACING.SINGLE
    if line != 1.5:
        pf.line_spacing = line
    pf.first_line_indent = Cm(1.25) if first_line else None
    pf.keep_with_next = keep_with_next
    pf.keep_together = keep_together
    pf.widow_control = True
    return p


def add_text(doc, text, *, bold_prefix=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             first_line=True, before=0, after=6, line=1.5, keep_together=False):
    p = doc.add_paragraph()
    set_paragraph_format(
        p, align=align, first_line=first_line, before=before, after=after,
        line=line, keep_together=keep_together,
    )
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), bold=True)
        set_run_font(p.add_run(text[len(bold_prefix):]))
    else:
        set_run_font(p.add_run(text))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.first_line_indent = None
    pf.space_before = Pt(12 if level == 1 else 8)
    pf.space_after = Pt(5)
    pf.keep_with_next = True
    pf.widow_control = True
    for r in p.runs:
        set_run_font(r, size=14, bold=True)
    if not p.runs:
        set_run_font(p.add_run(text), size=14, bold=True)
    else:
        p.runs[0].text = text
    return p


def set_cell_border(cell, color=GRID, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_text(cell, text, *, bold=False, color="000000", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.first_line_indent = None
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    set_run_font(p.add_run(text), size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    set_cell_border(cell)


def add_table(doc, headers, rows, widths=None, caption=None, font_size=10.5):
    if caption:
        p = doc.add_paragraph()
        set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, before=5, after=5, line=1.0, keep_with_next=True)
        set_run_font(p.add_run(caption), size=12)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_header(table.rows[0])
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER, size=font_size)
        shade_cell(table.rows[0].cells[i], ACCENT)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        table.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i, value in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], str(value), align=align, size=font_size)
            if ridx % 2:
                shade_cell(cells[i], ACCENT_LIGHT)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_numbered_list(doc, items):
    for i, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        set_paragraph_format(p, first_line=False, after=3, line=1.5)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.hanging_indent = Cm(0.75)
        set_run_font(p.add_run(f"{i}. "))
        set_run_font(p.add_run(item))


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph()
        set_paragraph_format(p, first_line=False, after=3, line=1.5)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.hanging_indent = Cm(0.5)
        set_run_font(p.add_run("• "))
        set_run_font(p.add_run(item))


def add_caption(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, before=4, after=9, line=1.0)
    set_run_font(p.add_run(text), size=12)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_run_font(run, size=12)


def load_font(size, bold=False):
    font_name = "arialbd.ttf" if bold else "arial.ttf"
    return ImageFont.truetype(str(Path(r"C:\Windows\Fonts") / font_name), size=size)


def pbox(draw, xy, text, *, fill="#EAF2F8", outline="#1F4E79", font_size=36,
         bold=False, text_fill="#111111", radius=22, width=4):
    draw.rounded_rectangle(xy, radius=radius, fill=fill, outline=outline, width=width)
    font = load_font(font_size, bold)
    x1, y1, x2, y2 = xy
    bbox = draw.multiline_textbbox((0, 0), text, font=font, align="center", spacing=6)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.multiline_text(((x1 + x2 - tw) / 2, (y1 + y2 - th) / 2 - 3), text,
                        font=font, fill=text_fill, align="center", spacing=6)


def parrow(draw, start, end, *, fill="#6B7280", width=5):
    import math
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill=fill, width=width)
    angle = math.atan2(y2 - y1, x2 - x1)
    length = 22
    wing = 0.48
    points = [
        (x2, y2),
        (x2 - length * math.cos(angle - wing), y2 - length * math.sin(angle - wing)),
        (x2 - length * math.cos(angle + wing), y2 - length * math.sin(angle + wing)),
    ]
    draw.polygon(points, fill=fill)


def dashed_round_rect(draw, xy, fill="#6B7280", width=4, dash=18, gap=12, radius=24):
    # The boundary is deliberately simple and remains legible after Word scaling.
    x1, y1, x2, y2 = xy
    for x in range(x1 + radius, x2 - radius, dash + gap):
        draw.line((x, y1, min(x + dash, x2 - radius), y1), fill=fill, width=width)
        draw.line((x, y2, min(x + dash, x2 - radius), y2), fill=fill, width=width)
    for y in range(y1 + radius, y2 - radius, dash + gap):
        draw.line((x1, y, x1, min(y + dash, y2 - radius)), fill=fill, width=width)
        draw.line((x2, y, x2, min(y + dash, y2 - radius)), fill=fill, width=width)


def make_org_chart():
    img = Image.new("RGB", (2400, 1500), "white")
    draw = ImageDraw.Draw(img)
    pbox(draw, (900, 70, 1500, 210), "Генеральный директор", fill="#1F4E79",
         outline="#1F4E79", font_size=46, bold=True, text_fill="white")
    second = [
        ((100, 430, 500, 600), "Коммерческий\nдиректор"),
        ((670, 430, 1070, 600), "Технический\nдиректор"),
        ((1260, 430, 1660, 600), "Склад и\nлогистика"),
        ((1850, 430, 2250, 600), "Финансы и\nбухгалтерия"),
    ]
    centers = []
    for xy, label in second:
        pbox(draw, xy, label, font_size=39, bold=True)
        cx = (xy[0] + xy[2]) // 2
        centers.append(cx)
        parrow(draw, (1200, 210), (cx, 430))
    commercial = [
        ((40, 880, 350, 1080), "Отдел продаж\nна маркетплейсах"),
        ((385, 880, 695, 1080), "Отдел\nзакупок"),
        ((730, 880, 1040, 1080), "Контент и\nдизайн"),
    ]
    for xy, label in commercial:
        pbox(draw, xy, label, font_size=31)
        parrow(draw, (300, 600), ((xy[0] + xy[2]) // 2, 880))
    technical = [
        ((1100, 880, 1510, 1080), "Группа конфигурации\nи совместимости"),
        ((1550, 880, 1840, 1080), "ИТ служба"),
    ]
    for xy, label in technical:
        pbox(draw, xy, label, font_size=30)
        parrow(draw, (870, 600), ((xy[0] + xy[2]) // 2, 880))
    dashed_round_rect(draw, (20, 350, 1870, 1240))
    draw.text((55, 1270), "Пунктиром выделены подразделения, непосредственно участвующие в обследуемом процессе",
              font=load_font(31), fill="#4B5563")
    img.save(ORG_PNG, dpi=(220, 220))


def make_functional_contour():
    img = Image.new("RGB", (2400, 1300), "white")
    draw = ImageDraw.Draw(img)
    pbox(draw, (40, 220, 470, 390), "Поставщики\nцены, наличие, модели", font_size=34)
    pbox(draw, (40, 800, 470, 970), "Складской учёт\nостатки и резервы", font_size=34)
    pbox(draw, (650, 90, 1750, 1080), "", fill="#F8FAFC", outline="#1F4E79")
    title_font = load_font(43, True)
    title = "Проектируемая информационная система"
    tb = draw.textbbox((0, 0), title, font=title_font)
    draw.text(((2400 - (tb[2] - tb[0])) / 2, 145), title, font=title_font, fill="#1F4E79")
    modules = [
        ((760, 350, 1160, 520), "BOM и каталог\nкомпонентов"),
        ((1240, 350, 1640, 520), "Правила\nсовместимости"),
        ((760, 680, 1160, 850), "Цена, вес и\nполнота данных"),
        ((1240, 680, 1640, 850), "Карточка, публикация\nи аудит"),
    ]
    for xy, label in modules:
        pbox(draw, xy, label, font_size=31)
    pbox(draw, (1930, 220, 2360, 390), "Менеджер\nмаркетплейсов", font_size=34)
    pbox(draw, (1930, 800, 2360, 970), "Ozon Seller API\nкарточка и статус", font_size=34)
    for s, e in [
        ((470, 305), (760, 415)), ((470, 885), (760, 765)),
        ((1930, 305), (1640, 415)), ((1640, 475), (1930, 335)),
        ((1640, 765), (1930, 885)), ((1930, 915), (1640, 815)),
        ((1160, 435), (1240, 435)), ((960, 520), (960, 680)),
        ((1440, 520), (1440, 680)), ((1160, 765), (1240, 765)),
    ]:
        parrow(draw, s, e)
    note = "Единый объект изменения: BOM → проверка → товарная карточка → публикация → контроль статуса"
    nf = load_font(31)
    nb = draw.textbbox((0, 0), note, font=nf)
    draw.text(((2400 - (nb[2] - nb[0])) / 2, 1165), note, font=nf, fill="#374151")
    img.save(FLOW_PNG, dpi=(220, 220))


def mx_cell(parent, cell_id, value="", style="", vertex=False, edge=False, parent_id="1",
            x=None, y=None, w=None, h=None, source=None, target=None):
    attrs = {"id": cell_id, "value": value, "style": style, "parent": parent_id}
    if vertex:
        attrs["vertex"] = "1"
    if edge:
        attrs["edge"] = "1"
    if source:
        attrs["source"] = source
    if target:
        attrs["target"] = target
    cell = ET.SubElement(parent, "mxCell", attrs)
    if vertex:
        ET.SubElement(cell, "mxGeometry", {"x": str(x), "y": str(y), "width": str(w), "height": str(h), "as": "geometry"})
    elif edge:
        ET.SubElement(cell, "mxGeometry", {"relative": "1", "as": "geometry"})
    return cell


def create_page(mxfile, name, width=1169, height=827):
    diagram = ET.SubElement(mxfile, "diagram", {"name": name, "id": name.replace(" ", "_")})
    model = ET.SubElement(diagram, "mxGraphModel", {
        "dx": "1200", "dy": "800", "grid": "1", "gridSize": "10", "guides": "1",
        "tooltips": "1", "connect": "1", "arrows": "1", "fold": "1", "page": "1",
        "pageScale": "1", "pageWidth": str(width), "pageHeight": str(height), "math": "0", "shadow": "0",
    })
    root = ET.SubElement(model, "root")
    ET.SubElement(root, "mxCell", {"id": "0"})
    ET.SubElement(root, "mxCell", {"id": "1", "parent": "0"})
    return root


def make_drawio():
    mxfile = ET.Element("mxfile", {"host": "app.diagrams.net", "modified": "2026-09-13T00:00:00.000Z", "agent": "Codex", "version": "24.7.17", "compressed": "false"})
    style_main = "rounded=1;whiteSpace=wrap;html=1;fillColor=#1F4E79;fontColor=#FFFFFF;strokeColor=#1F4E79;fontStyle=1;fontSize=14;"
    style_box = "rounded=1;whiteSpace=wrap;html=1;fillColor=#EAF2F8;strokeColor=#1F4E79;fontColor=#000000;fontSize=12;"
    style_edge = "edgeStyle=orthogonalEdgeStyle;rounded=0;orthogonalLoop=1;jettySize=auto;html=1;endArrow=block;endFill=1;strokeColor=#6B7280;"

    root = create_page(mxfile, "Организационная структура")
    mx_cell(root, "g", "Генеральный директор", style_main, True, x=430, y=40, w=300, h=60)
    positions = {
        "c": (80, 180, 210, 70, "Коммерческий директор"),
        "t": (350, 180, 210, 70, "Технический директор"),
        "w": (620, 180, 190, 70, "Склад и логистика"),
        "f": (850, 180, 190, 70, "Финансы и бухгалтерия"),
        "m": (20, 360, 180, 70, "Отдел продаж на маркетплейсах"),
        "p": (220, 360, 150, 70, "Отдел закупок"),
        "d": (390, 360, 150, 70, "Контент и дизайн"),
        "b": (570, 360, 210, 70, "Группа конфигурации и совместимости"),
        "i": (810, 360, 150, 70, "ИТ служба"),
    }
    for cid, (x, y, w, h, label) in positions.items():
        mx_cell(root, cid, label, style_box, True, x=x, y=y, w=w, h=h)
    for n, (src, dst) in enumerate((("g", "c"), ("g", "t"), ("g", "w"), ("g", "f"), ("c", "m"), ("c", "p"), ("c", "d"), ("t", "b"), ("t", "i")), start=1):
        mx_cell(root, f"e{n}", style=style_edge, edge=True, source=src, target=dst)

    root = create_page(mxfile, "Функциональный контур")
    positions2 = {
        "sup": (20, 170, 210, 80, "Поставщики<br>цены, наличие, модели"),
        "stock": (20, 430, 210, 80, "Складской учёт<br>остатки и резервы"),
        "sys": (330, 90, 500, 500, "Проектируемая информационная система"),
        "bom": (370, 210, 180, 80, "BOM и каталог компонентов"),
        "rules": (600, 210, 180, 80, "Правила совместимости"),
        "calc": (370, 390, 180, 80, "Цена, вес и полнота"),
        "pub": (600, 390, 180, 80, "Карточка, публикация и аудит"),
        "mgr": (900, 170, 210, 80, "Менеджер маркетплейсов"),
        "ozon": (900, 430, 210, 80, "Ozon Seller API<br>карточка и статус"),
    }
    for cid, (x, y, w, h, label) in positions2.items():
        st = "rounded=1;whiteSpace=wrap;html=1;fillColor=#F8FAFC;strokeColor=#1F4E79;fontColor=#1F4E79;fontStyle=1;fontSize=15;verticalAlign=top;spacingTop=10;" if cid == "sys" else style_box
        mx_cell(root, cid, label, st, True, x=x, y=y, w=w, h=h)
    for n, (src, dst) in enumerate((("sup", "bom"), ("stock", "calc"), ("mgr", "rules"), ("pub", "mgr"), ("pub", "ozon"), ("ozon", "pub"), ("bom", "rules"), ("bom", "calc"), ("rules", "pub"), ("calc", "pub")), start=20):
        mx_cell(root, f"e{n}", style=style_edge, edge=True, source=src, target=dst)

    tree = ET.ElementTree(mxfile)
    ET.indent(tree, space="  ")
    tree.write(DRAWIO_PATH, encoding="utf-8", xml_declaration=True)


def add_picture_with_alt(doc, path, width_cm, alt_text):
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, before=3, after=2, line=1.0, keep_together=True)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(path), width=Cm(width_cm))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", alt_text)
    return p


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(14)
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.color.rgb = RGBColor(0, 0, 0)
    doc.styles["Title"].font.size = Pt(16)
    doc.styles["Title"].font.bold = True
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)


def make_document():
    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    section.different_first_page_header_footer = True
    add_page_number(section.footer.paragraphs[0])

    # Title page
    for text, size, bold, after in [
        ("МИНОБРНАУКИ РОССИИ", 12, True, 2),
        ("Федеральное государственное бюджетное образовательное учреждение", 12, False, 0),
        ("высшего образования", 12, False, 0),
        ("«МИРЭА - Российский технологический университет»", 12, False, 0),
        ("РТУ МИРЭА", 12, True, 10),
        ("Институт информационных технологий", 12, False, 0),
        ("Кафедра квантовых информационных технологий,", 12, False, 0),
        ("практической и прикладной информатики", 12, False, 20),
    ]:
        p = doc.add_paragraph()
        set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=after, line=1.0)
        set_run_font(p.add_run(text), size=size, bold=bold)

    p = doc.add_paragraph(style="Title")
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, before=8, after=10, line=1.0)
    set_run_font(p.add_run("ОТЧЁТ ПО ПРАКТИЧЕСКОМУ ЗАНЯТИЮ"), size=16, bold=True)
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=4, line=1.0)
    set_run_font(p.add_run("Практическое занятие № 1"), size=14, bold=True)
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=12, line=1.0)
    set_run_font(p.add_run("по дисциплине «Информационный менеджмент»"), size=14)
    p = doc.add_paragraph()
    set_paragraph_format(p, align=WD_ALIGN_PARAGRAPH.CENTER, first_line=False, after=20, line=1.0)
    set_run_font(p.add_run("Тема: «Обследование процесса формирования и актуализации карточек товаров компьютерных систем в ООО «Юкомс»»"), size=14, bold=True)

    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    labels = (("Студент группы ИНБО-12-23", "Албахтин Илья Владиславович"), ("Преподаватель", "Пяткин Василий Владимирович"))
    for ridx, row in enumerate(labels):
        for cidx, value in enumerate(row):
            cell = table.rows[ridx].cells[cidx]
            set_cell_text(cell, value, size=12)
            set_cell_border(cell, color="FFFFFF", size="0")
            cell.width = Cm(7.2 if cidx == 0 else 8.0)
    table.rows[0].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    table.rows[0].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    table.rows[1].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    table.rows[1].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(58)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("Москва 2026"), size=12)
    doc.add_page_break()

    add_heading(doc, "ПРАКТИЧЕСКОЕ ЗАНЯТИЕ № 1", 1)
    add_heading(doc, "1 Цель работы и индивидуальное задание", 1)
    add_text(doc, "Цель работы - выполнить экспресс-обследование предприятия, описать организацию и объект автоматизации, построить организационную структуру и обосновать необходимость информатизации отдельных подразделений и процессов.")
    add_text(doc, "Индивидуальное задание: выполнить обследование процесса формирования и актуализации карточек товаров компьютерных систем в ООО «Юкомс». Описать организацию, участников и информационные потоки, построить организационную структуру и обосновать автоматизацию процесса публикации товарных предложений на маркетплейсе Ozon.")
    add_text(doc, "Результат работы - описание организации и текущего процесса, структурно-функциональные схемы, перечень проблем, границы автоматизации и ожидаемые управленческие результаты.")

    add_heading(doc, "2 Этапы выполнения работы", 1)
    add_numbered_list(doc, [
        "Собрана предварительная информация о деятельности ООО «Юкомс» и выбранном процессе сбыта.",
        "Определены границы объекта автоматизации, участники, документы и информационные потоки.",
        "Построена модельная организационная структура и распределены функции подразделений.",
        "Выявлены проблемы текущей организации работы и сформирован целевой функциональный контур системы.",
        "Определены ожидаемые результаты автоматизации и целевые показатели пилотного внедрения.",
    ])

    add_heading(doc, "3 Предварительная информация об организации", 1)
    add_heading(doc, "3.1 Краткая характеристика", 2)
    add_text(doc, "ООО «Юкомс» рассматривается как организация, работающая в сфере компьютерной техники, электронной торговли и информационных технологий. Для учебного обследования выделено направление подготовки и продажи готовых компьютерных систем через маркетплейс Ozon. Компания получает сведения о комплектующих от поставщиков, формирует конфигурации компьютеров, рассчитывает цену и массу, готовит изображения и характеристики, после чего публикует товарные карточки.")
    add_text(doc, "Ключевым информационным объектом процесса является спецификация изделия BOM. Она связывает корпус с набором комплектующих и правилами их совместимости. Один и тот же корпус может использоваться для нескольких конфигураций, а элемент BOM может задаваться точной моделью или группой характеристик. Поэтому изменение доступности, цены или модели компонента способно затронуть сразу несколько товарных предложений.")
    add_table(doc,
        ["Характеристика", "Содержание"],
        [
            ("Организационно-правовая форма", "Общество с ограниченной ответственностью"),
            ("Основное направление", "Компьютерная техника, готовые компьютерные системы и электронная торговля"),
            ("Канал сбыта в границах обследования", "Маркетплейс Ozon"),
            ("Основной объект управления", "BOM компьютерной системы и связанная товарная карточка"),
            ("Источники данных", "Поставщики комплектующих, складской учёт, внутренние справочники, требования Ozon"),
            ("Потребитель результата", "Покупатель маркетплейса и подразделения, сопровождающие продажу"),
        ], widths=[5.2, 10.3], caption="Таблица 1 - Основные характеристики организации и процесса", font_size=10.5)

    add_heading(doc, "3.2 Границы обследования", 2)
    add_text(doc, "Процесс начинается с решения создать новую карточку компьютерной системы либо изменить существующую. Завершающим результатом является принятая Ozon карточка с актуальными характеристиками, ценой, массой и изображениями либо зарегистрированная ошибка публикации, назначенная ответственному исполнителю.")
    add_table(doc,
        ["В границах обследования", "За границами обследования"],
        [
            ("Создание и изменение BOM", "Заключение договоров с поставщиками"),
            ("Сопоставление внутренних компонентов и предложений поставщиков", "Бухгалтерский и налоговый учёт"),
            ("Проверка технической совместимости", "Физическая сборка и диагностика компьютеров"),
            ("Расчёт цены, массы и обязательных характеристик", "Приём платежей и финансовое закрытие заказа"),
            ("Подготовка изображений и текста карточки", "Доставка покупателю и гарантийное обслуживание"),
            ("Публикация и контроль статуса через Ozon Seller API", "Управление рекламными кампаниями"),
        ], widths=[7.75, 7.75], caption="Таблица 2 - Границы обследуемого процесса", font_size=10.2)

    add_heading(doc, "4 Описание объекта автоматизации", 1)
    add_heading(doc, "4.1 Текущая организация процесса", 2)
    add_text(doc, "В модели текущего состояния сведения о составе компьютера, предложениях поставщиков, правилах совместимости, расчётах и карточке Ozon хранятся в нескольких инструментах. Менеджер маркетплейсов координирует участников и переносит итоговые данные в кабинет продавца. Технический специалист проверяет конфигурацию, закупки и склад уточняют доступность и стоимость, а контент-группа готовит изображения и описание.")
    add_numbered_list(doc, [
        "Менеджер регистрирует потребность в новой карточке или изменении существующей.",
        "Технический специалист формирует BOM и выбирает компоненты.",
        "Сведения поставщиков и склада сопоставляются с позициями BOM.",
        "Совместимость, цена, масса и полнота характеристик проверяются вручную или в разрозненных файлах.",
        "Контент-группа готовит изображения, название и описание.",
        "Менеджер переносит данные в Ozon и отправляет карточку на проверку.",
        "Результат публикации контролируется, а найденные ошибки передаются на исправление.",
    ])

    add_heading(doc, "4.2 Документы и информационные потоки", 2)
    add_table(doc,
        ["Информационный объект", "Источник", "Использование"],
        [
            ("Запрос на создание или изменение карточки", "Менеджер маркетплейсов", "Определяет инициатора, основание, срок и приоритет"),
            ("BOM компьютерной системы", "Технический специалист", "Фиксирует корпус, слоты и допустимые компоненты"),
            ("Предложение поставщика", "Поставщик / отдел закупок", "Содержит модель, артикул, цену, наличие и дату актуальности"),
            ("Остатки и резервы", "Складской учёт", "Подтверждают доступность внутренних компонентов"),
            ("Правила совместимости", "Технический специалист", "Ограничивают сочетания по сокету, форм-фактору, мощности и габаритам"),
            ("Карточка товара", "Менеджер и контент-группа", "Объединяет название, описание, характеристики, цену, массу и изображения"),
            ("Протокол публикации", "Ozon Seller API", "Хранит идентификатор операции, ответ и итоговый статус"),
        ], widths=[4.2, 3.5, 7.8], caption="Таблица 3 - Документы и информационные потоки", font_size=9.8)

    doc.add_page_break()
    add_heading(doc, "4.3 Выявленные проблемы", 2)
    add_table(doc,
        ["Код", "Проблема", "Причина", "Последствие"],
        [
            ("P-01", "Повторный ввод данных", "BOM, расчёты и карточка ведутся раздельно", "Ошибки и увеличение срока подготовки"),
            ("P-02", "Ручная проверка совместимости", "Правила не формализованы в едином справочнике", "Риск некорректной конфигурации"),
            ("P-03", "Несогласованная замена компонента", "Нет анализа всех связанных BOM", "Часть карточек сохраняет устаревший состав"),
            ("P-04", "Неактуальная цена или наличие", "Источники обновляются в разное время", "Недостоверное товарное предложение"),
            ("P-05", "Поздняя проверка обязательных атрибутов", "Полнота контролируется перед публикацией", "Отклонение Ozon и возврат на исправление"),
            ("P-06", "Нет сквозной истории изменений", "Действия распределены между инструментами", "Трудно найти причину расхождения"),
        ], widths=[1.2, 3.5, 5.0, 5.8], caption="Таблица 4 - Проблемы текущего процесса", font_size=9.2)

    add_heading(doc, "5 Организационная структура", 1)
    add_text(doc, "Для целей учебного обследования построена модельная функциональная структура в границах рассматриваемого процесса. Она показывает не штатное расписание организации, а распределение ответственности между ролями, необходимое для управления сбытовой деятельностью и будущей системой.")
    add_picture_with_alt(doc, ORG_PNG, 16.0, "Модельная организационная структура ООО Юкомс")
    add_caption(doc, "Рисунок 1 - Организационная структура и участники обследуемого процесса")
    add_text(doc, "Владельцем процесса выступает коммерческий блок. Менеджер маркетплейсов отвечает за запуск и завершение работы с карточкой; технический блок - за корректность BOM и совместимость; закупки и склад - за доступность и цену; контент-группа - за визуальные и текстовые материалы; ИТ-служба - за интеграции, права доступа и работоспособность системы.")
    add_table(doc,
        ["Подразделение / роль", "Основная функция", "Ответственность в системе"],
        [
            ("Генеральный директор", "Стратегические решения и ресурсы", "Утверждение проекта и контроль итоговых показателей"),
            ("Коммерческий директор", "Управление сбытовой деятельностью", "Владелец процесса и правил приоритизации"),
            ("Менеджер маркетплейсов", "Подготовка и публикация предложений", "Запрос, проверка карточки, публикация, статусы"),
            ("Отдел закупок", "Работа с поставщиками", "Цены, наличие, сроки и сопоставление предложений"),
            ("Склад и логистика", "Учёт доступности компонентов", "Остатки, резервы и подтверждение наличия"),
            ("Технический специалист", "Конфигурация компьютерных систем", "BOM, совместимость и массовые замены"),
            ("Контент и дизайн", "Подготовка содержания карточки", "Изображения, описание и шаблоны"),
            ("ИТ-служба", "Эксплуатация информационных средств", "Пользователи, интеграции, журнал ошибок и доступность"),
            ("Финансы и бухгалтерия", "Финансовый контроль", "Согласование правил цены и получение итоговых данных"),
        ], widths=[4.0, 5.0, 6.5], caption="Таблица 5 - Распределение функций между участниками", font_size=9.4)

    add_heading(doc, "6 Обоснование необходимости автоматизации", 1)
    add_text(doc, "Автоматизация необходима прежде всего в точках, где одно изменение влияет на несколько объектов. Замена комплектующего должна автоматически находить связанные BOM, повторно проверять совместимость, пересчитывать цену и массу, отмечать затронутые карточки и сохранять историю. Без единого механизма зависимостей трудозатраты растут вместе с ассортиментом, а качество результата зависит от ручного контроля.")
    add_text(doc, "Целевая система должна стать единым источником данных о товарной конфигурации. В ней карточка Ozon связывается с конкретной версией BOM, предложения поставщиков - с внутренними компонентами, а правила совместимости, цены и веса - со справочниками. Неуспешная внешняя операция не должна приводить к потере подготовленных данных.")
    add_picture_with_alt(doc, FLOW_PNG, 15.0, "Функциональный контур автоматизации подготовки товарных карточек")
    add_caption(doc, "Рисунок 2 - Функциональный контур проектируемой системы")
    add_table(doc,
        ["Функция", "Результат", "Основной пользователь"],
        [
            ("Создание и версионирование BOM", "Единая актуальная спецификация конфигурации", "Технический специалист"),
            ("Проверка совместимости", "Перечень нарушенных правил до публикации", "Технический специалист"),
            ("Загрузка цен и наличия", "Актуальные предложения с датой получения", "Закупки / система"),
            ("Массовая замена", "Предварительный список затронутых BOM и конфликтов", "Технический специалист"),
            ("Расчёт цены и массы", "Воспроизводимые расчётные значения", "Менеджер / система"),
            ("Формирование карточки", "Название, описание, характеристики и изображения", "Контент / система"),
            ("Публикация через API", "Карточка и сохранённый ответ Ozon", "Менеджер маркетплейсов"),
            ("Аудит и уведомления", "История изменений и назначенные события", "Администратор / ответственные"),
        ], widths=[4.2, 7.0, 4.3], caption="Таблица 6 - Основные функции автоматизируемой системы", font_size=9.5)

    add_heading(doc, "7 Ожидаемые результаты реализации системы", 1)
    add_text(doc, "Ожидаемый управленческий эффект состоит в переходе от координации по разрозненным файлам к управлению единым объектом изменения. Руководитель получает прозрачные статусы и показатели, а исполнители - общий набор данных, правил и контрольных процедур.")
    add_table(doc,
        ["Результат", "Целевой показатель пилота", "Способ контроля"],
        [
            ("Сокращение цикла подготовки карточки", "Не менее 30% относительно зафиксированного базового времени", "Сравнение медианного времени до и после внедрения"),
            ("Предварительный контроль качества", "100% карточек проверены по совместимости и обязательным атрибутам", "Журнал автоматических проверок"),
            ("Согласованность массовых изменений", "Все затронутые BOM найдены и повторно проверены", "Протокол массовой операции"),
            ("Актуальность цен и наличия", "Каждое предложение содержит источник и дату обновления", "Отчёт по просроченным сведениям"),
            ("Прослеживаемость", "100% критичных изменений имеют автора, время, старое и новое значение", "Журнал аудита"),
            ("Устойчивость публикации", "Ошибка внешнего API сохраняется без потери карточки и допускает повтор", "Журнал интеграционных операций"),
        ], widths=[4.4, 5.8, 5.3], caption="Таблица 7 - Ожидаемые результаты и показатели", font_size=9.3)
    add_text(doc, "К качественным результатам относятся единая терминология, закреплённые зоны ответственности, снижение зависимости от памяти отдельных специалистов и возможность дальнейшего анализа продаж, причин отклонений и частоты замен компонентов.")

    add_heading(doc, "8 Выводы", 1)
    add_text(doc, "В практической работе выполнено обследование процесса формирования и актуализации карточек товаров компьютерных систем в ООО «Юкомс». Построена модельная организационная структура, определены участники и информационные потоки, описаны границы процесса и выявлены шесть системных проблем текущей организации работы.")
    add_text(doc, "Обследование подтвердило целесообразность автоматизации: ассортимент и взаимозависимость BOM, поставщицких предложений и карточек делают ручное согласование ненадёжным и плохо масштабируемым. Предложенный функциональный контур объединяет данные о конфигурации, совместимости, цене, массе, контенте и публикации, а ожидаемые результаты заданы через измеримые показатели пилотного внедрения.")

    add_heading(doc, "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", 1)
    sources = [
        "Информационный менеджмент. Материалы для практических занятий № 1. РТУ МИРЭА, 2026.",
        "Албахтин И. В. Проектирование информационной системы поддержки бизнес-процесса формирования и актуализации карточек товаров компьютерных систем на примере ООО «Юкомс». Практическая работа № 1 по дисциплине «Проектирование предметно-ориентированных информационных систем». Москва, 2026.",
        "Албахтин И. В. Анализ процесса формирования и актуализации карточек товаров компьютерных систем. Практическое занятие № 1 по дисциплине «Разработка обеспечивающих подсистем». Москва, 2026.",
        "Официальный сайт ООО «Юкомс». URL: https://ucoms.ru/ (дата обращения: 09.09.2026).",
        "Ozon Seller API. Документация для продавцов. URL: https://docs.ozon.ru/api/seller/ (дата обращения: 09.09.2026).",
    ]
    add_numbered_list(doc, sources)

    # Keep headings black and normalize all runs.
    for p in doc.paragraphs:
        for run in p.runs:
            if not run.font.name:
                set_run_font(run)
    doc.core_properties.title = "Практическое занятие № 1 по информационному менеджменту"
    doc.core_properties.subject = "Обследование процесса формирования и актуализации карточек товаров компьютерных систем в ООО Юкомс"
    doc.core_properties.author = "Албахтин Илья Владиславович"
    doc.core_properties.keywords = "информационный менеджмент, автоматизация, организационная структура, Ozon, BOM"
    doc.save(DOCX_PATH)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    make_org_chart()
    make_functional_contour()
    make_drawio()
    make_document()
    print(DOCX_PATH)
    print(DRAWIO_PATH)
    print(ORG_PNG)
    print(FLOW_PNG)


if __name__ == "__main__":
    main()
