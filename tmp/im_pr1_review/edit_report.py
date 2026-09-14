from pathlib import Path
from copy import deepcopy
import re, shutil, hashlib
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table

B=Path('D:/GitHub/MIREA'); T=B/'tmp/im_pr1_review'
dst=B/'4/ИМ/ИМ_Практическая_1_АлбахтинИВ.docx'
backup=T/'original.docx'
if not backup.exists():shutil.copy2(dst,backup)
src=Document(backup); d=Document(B/'4/ИСУРО/ИСУРО_1_АлбахтинИВ.docx')
refhash=hashlib.sha256((B/'4/ИСУРО/ИСУРО_1_АлбахтинИВ.docx').read_bytes()).hexdigest()
with (T/'artifact.md').open('a',encoding='utf-8') as f:f.write('\nSHA256 титульного образца: '+refhash+'\n')
body=d._element.body
for e in list(body)[16:]:body.remove(e)
title=list(body)
endsect=deepcopy(d.sections[0]._sectPr)
for e in list(endsect):
 if e.tag in [qn('w:headerReference'),qn('w:footerReference')]:endsect.remove(e)
body.append(endsect)

def font(r,size=None,bold=None,italic=None):
 r.font.name='Times New Roman'
 for a in ['ascii','hAnsi','eastAsia','cs']:r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:'+a),'Times New Roman')
 if size:r.font.size=Pt(size)
 if bold is not None:r.bold=bold
 if italic is not None:r.italic=italic
 r.font.color.rgb=RGBColor(0,0,0)
def replace(p,text):
 r=p.runs[0] if p.runs else p.add_run()
 r.text=text
 for x in list(p.runs)[1:]:p._p.remove(x._r)

for e in title:
 for pe in e.xpath('.//w:p') if e.tag!=qn('w:p') else [e]:
  p=Paragraph(pe,d)
  for r in list(p.runs):
   if 'ИНБО-01-17' in r.text:p._p.remove(r._r)
   else:font(r)
  if p.text=='ОТЧЕТ ПО ПРАКТИЧЕСКИМ РАБОТАМ':replace(p,'ОТЧЕТ ПО ПРАКТИЧЕСКОЙ РАБОТЕ № 1');font(p.runs[0],16,True)
  elif p.text.startswith('по дисциплине'):replace(p,'по дисциплине «Информационный менеджмент»');font(p.runs[0],14,False)
  elif 'Шендяпин' in p.text:replace(p,'Пяткин В.В.');font(p.runs[0],12,False)
 # normalize all title font properties including empty paragraph marks
 for rf in e.xpath('.//w:rFonts'):
  for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Times New Roman')

def style(name,size,align,first,line,bold=False,italic=False,before=0,after=0,keep=False):
 s=d.styles[name] if name in d.styles else d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=bold;s.font.italic=italic;s.font.color.rgb=RGBColor(0,0,0)
 f=s.paragraph_format;f.alignment=align;f.left_indent=Cm(0);f.right_indent=Cm(0);f.first_line_indent=Cm(first);f.line_spacing=line;f.space_before=Pt(before);f.space_after=Pt(after);f.widow_control=True;f.keep_with_next=keep;f.keep_together=keep
 return s
style('Основной текст отчёта',14,WD_ALIGN_PARAGRAPH.JUSTIFY,1.25,1.5)
style('Подпись таблицы',12,WD_ALIGN_PARAGRAPH.LEFT,0,1,italic=True,before=6,after=3,keep=True)
style('Подпись рисунка',14,WD_ALIGN_PARAGRAPH.CENTER,0,1,before=3,after=6)
style('Текст таблицы',12,WD_ALIGN_PARAGRAPH.LEFT,0,1)
style('Heading 1',14,WD_ALIGN_PARAGRAPH.CENTER,0,1.5,bold=True,before=12,after=6,keep=True)
style('Heading 2',14,WD_ALIGN_PARAGRAPH.LEFT,0,1.5,bold=True,before=12,after=6,keep=True)

changes={
18:None,
22:'3. Построена организационная структура и определены функции подразделений.',
23:'4. Выявлены проблемы текущей организации работы и определены функции, требующие автоматизации.',
27:'ООО «Юкомс» работает в сфере компьютерной техники, электронной торговли и информационных технологий. Обследование охватывает подготовку и продажу готовых компьютерных систем через маркетплейс Ozon. Компания получает сведения о комплектующих от поставщиков, формирует конфигурации компьютеров, рассчитывает цену и массу, готовит изображения и характеристики, после чего публикует товарные карточки. Основные характеристики организации и процесса приведены в таблице 1.',
32:src.paragraphs[32].text+' Границы процесса представлены в таблице 2.',
37:'В рассматриваемой модели текущего состояния сведения о составе компьютера, предложениях поставщиков, правилах совместимости, расчётах и карточке Ozon хранятся в нескольких инструментах. Менеджер маркетплейсов координирует участников и переносит итоговые данные в кабинет продавца. Производственный участок проверяет конфигурацию, складской отдел уточняет доступность компонентов, а дизайнер готовит изображения. Подготовка карточки включает следующие действия.',
39:'2. Сотрудник производственного участка формирует BOM и выбирает компоненты.',
42:'5. Дизайнер готовит изображения, а менеджер маркетплейсов — название и описание.',
53:'Организационная структура ООО «Юкомс» представлена на рисунке 1 [2]. Она включает отдел электронной коммерции, производственный участок со складским отделом, отдел клиентской поддержки и сервисного обслуживания, отдел ИТ-поддержки и технического сопровождения, а также отдел логистики.',
55:'Рисунок 1 – Организационная структура ООО «Юкомс»',
56:'За подготовку и публикацию карточек отвечает отдел электронной коммерции. Менеджер маркетплейсов организует работу с карточкой, дизайнер готовит изображения. Производственный участок проверяет состав и совместимость комплектующих, складской отдел предоставляет сведения об остатках. Отдел ИТ-поддержки и технического сопровождения обеспечивает интеграции, права доступа и работоспособность системы. Распределение функций приведено в таблице 5.',
61:src.paragraphs[61].text+' Основные функции системы представлены в таблице 6.',
62:None,63:None,
67:src.paragraphs[67].text+' Ожидаемые результаты и целевые показатели пилота приведены в таблице 7.',
72:src.paragraphs[72].text.replace('модельная организационная','организационная'),
73:src.paragraphs[73].text.replace('Предложенный функциональный контур объединяет','Предлагаемая система объединяет'),
}
tables=[deepcopy(t._tbl) for t in src.tables[1:]]
role_rows=[
['Подразделение / роль','Основная функция','Ответственность в системе'],
['Генеральный директор','Управление компанией','Утверждение проекта и контроль результатов'],
['Отдел электронной коммерции / менеджер маркетплейсов','Подготовка и публикация предложений','Создание карточек, проверка данных, публикация и контроль статусов'],
['Дизайнер','Подготовка визуальных материалов','Изображения товаров и шаблоны оформления'],
['Производственный участок / руководитель и сборщики','Подготовка конфигураций компьютерных систем','BOM, совместимость компонентов и согласование замен'],
['Складской отдел','Учёт компонентов','Остатки, резервы и подтверждение наличия'],
['Отдел клиентской поддержки и сервисного обслуживания','Обработка обращений клиентов','Передача замечаний о характеристиках и комплектации'],
['Отдел ИТ-поддержки и технического сопровождения / разработчики','Разработка и сопровождение системы','Интеграции, права доступа, журнал ошибок и доступность'],
['Отдел логистики / водитель','Доставка компьютерных систем','Использование сведений о массе и составе заказа'],
]
def para(text,sty='Основной текст отчёта'):
 p=d.add_paragraph(text,sty);return p
ti=0
for i,p in enumerate(src.paragraphs):
 if i<14 or i in [54] or (i in changes and changes[i] is None):continue
 text=changes.get(i,p.text)
 if not text.strip():continue
 if i==55:
  sec=d.add_section();sec.page_width=Cm(29.7);sec.page_height=Cm(21)
 if i==56:
  sec=d.add_section();sec.page_width=Cm(21);sec.page_height=Cm(29.7)
 if i==55:
  ip=para('');ip.alignment=WD_ALIGN_PARAGRAPH.CENTER;ip.paragraph_format.first_line_indent=Cm(0);ip.paragraph_format.line_spacing=1;ip.paragraph_format.keep_with_next=True
  ip.add_run().add_picture(str(T/'ppo_org.png'),width=Cm(25.2))
 if text.startswith('Таблица '):sty='Подпись таблицы';text=text.replace(' - ',' – ')
 elif text.startswith('Рисунок '):sty='Подпись рисунка'
 elif p.style.name.startswith('Heading'):sty=p.style.name
 else:sty='Основной текст отчёта'
 np=para(text,sty)
 if i==49:np.paragraph_format.page_break_before=True
 if sty=='Подпись рисунка':np.paragraph_format.keep_together=True
 if i==14:np.paragraph_format.space_before=Pt(0)
 if i in [45,49]:
   para('Документы и информационные потоки приведены в таблице 3.' if i==45 else 'Проблемы процесса, их причины и последствия приведены в таблице 4.')
 if text.startswith('Таблица '):
  tbl=Table(tables[ti],d);body.insert(len(body)-1,tbl._tbl)
  if ti==4:
   for row in list(tbl.rows)[1:]:tbl._tbl.remove(row._tr)
   for vals in role_rows[1:]:
    cells=tbl.add_row().cells
    for c,v in zip(cells,vals):c.text=v
  for row in tbl.rows:
   for c in row.cells:
    c.text=c.text.replace('Технический специалист','Производственный участок').replace('Менеджер и контент-группа','Менеджер и дизайнер').replace('Поставщик / отдел закупок','Поставщик / менеджер').replace('Закупки / система','Менеджер / система').replace('Контент / система','Дизайнер / система')
  ti+=1

# Format only body tables, preserve title tables exactly.
widthsets=[[5.4,11.1],[8.25,8.25],[5.2,4.3,7.0],[1.5,4.5,5.3,5.2],[5.6,4.4,6.5],[5.5,6.2,4.8],[5.1,6.6,4.8]]
for tbl,widths in zip(d.tables[2:],widthsets):
 tbl.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl.autofit=False
 for col,w in zip(tbl.columns,widths):col.width=Cm(w)
 for ri,row in enumerate(tbl.rows):
  pr=row._tr.get_or_add_trPr()
  for h in list(pr.findall(qn('w:trHeight'))):pr.remove(h)
  if pr.find(qn('w:cantSplit')) is None:pr.append(OxmlElement('w:cantSplit'))
  if ri==0 and pr.find(qn('w:tblHeader')) is None:pr.append(OxmlElement('w:tblHeader'))
  for c,w in zip(row.cells,widths):
   c.width=Cm(w);c.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
   cp=c._tc.get_or_add_tcPr()
   borders=cp.find(qn('w:tcBorders'))
   if borders is None:borders=OxmlElement('w:tcBorders');cp.append(borders)
   for edge in ['top','left','bottom','right']:
    be=borders.find(qn('w:'+edge))
    if be is None:be=OxmlElement('w:'+edge);borders.append(be)
    be.set(qn('w:val'),'single');be.set(qn('w:sz'),'8');be.set(qn('w:color'),'D9D9D9')
   margins=cp.find(qn('w:tcMar'))
   if margins is None:margins=OxmlElement('w:tcMar');cp.append(margins)
   for edge,value in [('top',90),('bottom',90),('start',120),('end',120)]:
    me=margins.find(qn('w:'+edge))
    if me is None:me=OxmlElement('w:'+edge);margins.append(me)
    me.set(qn('w:w'),str(value));me.set(qn('w:type'),'dxa')
   shd=c._tc.get_or_add_tcPr().find(qn('w:shd'))
   if shd is not None:shd.set(qn('w:fill'),'E7E6E6' if ri==0 else 'FFFFFF')
   for p in c.paragraphs:
    p.style='Текст таблицы'
    for child in list(p._p.get_or_add_pPr()):
     if child.tag!=qn('w:pStyle'):p._p.pPr.remove(child)
    for r in p.runs:font(r,12,ri==0,False)
    if re.fullmatch(r'P-\d+',p.text):p.alignment=WD_ALIGN_PARAGRAPH.CENTER

for idx,s in enumerate(d.sections):
 if s.page_width < s.page_height:s.page_width=Cm(21);s.page_height=Cm(29.7)
 else:
  s.page_width=Cm(29.7);s.page_height=Cm(21)
  va=OxmlElement('w:vAlign');va.set(qn('w:val'),'center');s._sectPr.append(va)
 s.left_margin=Cm(3);s.right_margin=Cm(1.5);s.top_margin=Cm(2);s.bottom_margin=Cm(2)
 s.header_distance=Cm(1.27);s.footer_distance=Cm(1.27)
 s.different_first_page_header_footer=False
 s.footer.is_linked_to_previous=False
 for p in s.footer.paragraphs:p.clear()
 if idx:
  p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
  r=p.add_run();font(r,12,False,False)
  fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
 # continuous page numbering
 pn=s._sectPr.find(qn('w:pgNumType'))
 if pn is not None:s._sectPr.remove(pn)

# Remove reviewer anchors and comment parts, hidden template fragments.
for e in d._element.xpath('.//w:commentRangeStart | .//w:commentRangeEnd | .//w:commentReference'):
 e.getparent().remove(e)
for rid,rel in list(d.part.rels.items()):
 if 'comment' in rel.reltype.lower():d.part.drop_rel(rid)
used={v for e in d._element.iter() for k,v in e.attrib.items() if k.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}')}
for rid,rel in list(d.part.rels.items()):
 if 'image' in rel.reltype and rid not in used:d.part.drop_rel(rid)
d.core_properties.title='Практическая работа 1 по информационному менеджменту'
d.core_properties.subject='Обследование ООО «Юкомс»'
d.core_properties.comments=''
d.save(dst)
print(dst)
