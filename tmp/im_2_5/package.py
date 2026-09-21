from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
import shutil, re, json, hashlib
from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader

BASE=Path(__file__).resolve().parent;ROOT=BASE.parent.parent;OUT=ROOT/'4/ИМ'
files=[];checks=[]
for n,count,tables,figs in [(2,8,4,1),(3,11,7,1),(4,8,3,2),(5,12,5,2)]:
    docx=OUT/f'ИМ_Практическая_{n}_АлбахтинИВ.docx'
    folder=BASE/(f'final_render_{n}' if n==2 else f'delivery_render_{n}')
    pdf=folder/f'ИМ_Практическая_{n}_АлбахтинИВ.pdf'
    d=Document(docx);r=PdfReader(pdf)
    assert len(r.pages)==count
    body='\n'.join(p.text for p in d.paragraphs)
    full=''.join(d._element.body.xpath('.//w:t/text()'))
    assert 'Шендяпин' not in full and 'ИНБО-01-17' not in full
    assert 'Пяткин В.В.' in full and 'Албахтин' in full and 'ИНБО-12-23' in full
    assert len(d.tables)==tables+2
    captions=[p.text for p in d.paragraphs if p.style.name=='Подпись таблицы']
    assert len(captions)==tables
    assert [int(re.search(r'Таблица (\d+)',t)[1]) for t in captions]==list(range(1,tables+1))
    figures=[p.text for p in d.paragraphs if p.style.name=='Подпись рисунка']
    assert len(figures)==figs
    for s in d.sections:
        assert abs(s.left_margin.cm-3)<.01 and abs(s.right_margin.cm-1.5)<.01
        assert abs(s.top_margin.cm-2)<.01 and abs(s.bottom_margin.cm-2)<.01
    for p in d.paragraphs:
        if p.style.name=='Основной текст отчёта':
            st=p.style
            assert st.font.name=='Times New Roman' and st.font.size.pt==14
            assert abs(st.paragraph_format.first_line_indent.cm-1.25)<.01
            assert st.paragraph_format.line_spacing==1.5
    for t in d.tables[2:]:
        assert t.rows[0]._tr.get_or_add_trPr().find(qn('w:tblHeader')) is not None
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for run in p.runs:assert run.font.size.pt==12
    assert all(len(p.extract_text().strip())>100 for p in r.pages)
    assert 'Пяткин' in r.pages[0].extract_text()
    assert 'Источники' in '\n'.join(p.extract_text() for p in r.pages)
    dst=OUT/pdf.name;shutil.copy2(pdf,dst);files.extend([docx,dst])
    checks.append({'practice':n,'pages':count,'tables':tables,'figures':figs,'docx_sha256':hashlib.sha256(docx.read_bytes()).hexdigest(),'pdf_sha256':hashlib.sha256(dst.read_bytes()).hexdigest()})

doc3=Document(OUT/'ИМ_Практическая_3_АлбахтинИВ.docx')
matrix=doc3.tables[-1]
matrix_rows=[[int(c.text) for c in row.cells[1:]] for row in list(matrix.rows)[1:-1]]
assert all(sum(row[:-1])==row[-1] for row in matrix_rows)
totals=[sum(row[i] for row in matrix_rows) for i in range(7)]
assert totals==[64,64,192,56,40,16,432]
assert sum([0,0,24,32,24,24,24,24,16,8,0,16])==192
assert sum([64*1000,64*1500,192*1200,56*1000,40*1000,16*1600])==512000
assert 512000+20000+60000==592000
drawio=OUT/'ИМ_Практические_2-5_Диаграммы.drawio'
assert len(ET.parse(drawio).getroot().findall('diagram'))==6
files.append(drawio)
ref=ROOT/'4/ИСУРО/ИСУРО_1_АлбахтинИВ.docx'
assert hashlib.sha256(ref.read_bytes()).hexdigest() in (BASE/'artifact.md').read_text(encoding='utf-8')

archive=OUT/'ИМ_Практические_2-5_АлбахтинИВ_комплект.zip'
with ZipFile(archive,'w',ZIP_DEFLATED) as z:
    for p in files:z.write(p,p.name)
with ZipFile(archive) as z:
    assert z.testzip() is None and len(z.namelist())==9
(BASE/'checks.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'reports':checks,'archive':str(archive),'archive_bytes':archive.stat().st_size,'files':len(files)},ensure_ascii=False,indent=2))
