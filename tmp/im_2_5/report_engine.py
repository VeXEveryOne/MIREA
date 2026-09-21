from pathlib import Path
from copy import deepcopy
import hashlib, json
from zipfile import ZipFile
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(__file__).resolve().parent
ROOT=BASE.parent.parent
OUT=ROOT/'4/ИМ'
REFERENCE=ROOT/'4/ИСУРО/ИСУРО_1_АлбахтинИВ.docx'

def font(r,size=14,bold=False,italic=False):
    r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=bold;r.italic=italic;r.font.color.rgb=RGBColor(0,0,0)
    rf=r._element.get_or_add_rPr().rFonts
    for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Times New Roman')

def contract():
    sha=hashlib.sha256(REFERENCE.read_bytes()).hexdigest()
    with ZipFile(REFERENCE) as z:
        inv={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}
    (BASE/'reference_parts.json').write_text(json.dumps(inv,indent=2),encoding='utf-8')
    (BASE/'artifact.md').write_text(f'''# Контракт оформления практик ИМ 2–5
Источник: {REFERENCE.as_posix()}
SHA-256: {sha}
Область использования — исключительно титульная страница по указанию AGENTS.md.
Визуальное свидетельство: reference_render/page-1.png. DOCX имеет два раздела; основной текст источника не переносится.
Сохранить элементы body 0–15 (две таблицы, эмблема, двойная линия, институт, кафедра, подписи и город).
Заменить body/p[5] (индекс абзаца 4): название отчёта на одну практику с номером, TNR 16 полужирный.
Заменить body/p[6] (индекс абзаца 5): дисциплина, TNR 14 обычный; сохранить позицию блока.
Из таблицы подписей удалить текст ИНБО-01-17, фамилию преподавателя заменить на Пяткин В.В. по первой практике ИМ.
Сохранить актуальные ИНБО-12-23, Албахтин И.В., Москва 2026 г. Университетская шапка и рисунок сохраняют исходные XML и отношения.
Нормализовать Liberation Serif в TNR, удалить белый старый текст. Основная часть источника и её разделы удаляются.
Основная часть создаётся по Оформление_отчётов.md: A4, 3/1.5/2/2 см; TNR 14, 1.5, по ширине, первая строка 1.25 см; интервалы 0.
Заголовки Heading 1/2 TNR 14 полужирный, первый центр, второй слева; 12/6 пт, keep-with-next.
Таблицы 12 пт, одинарный, нулевые отступы; повтор заголовка и запрет разрыва строк; подпись сверху 12 курсив.
Рисунки в тексте, подпись снизу 14 пт, центр. Нумерация таблиц и рисунков отдельная с 1 в каждом отчёте.
Альбомные разделы допустимы для широких таблиц/схем; после них восстановить книжный раздел. Нижний колонтитул PAGE 12, титульный скрыт.
Проверить все страницы итогового Word-экспорта; титульная страница одна. Сохранить исходный DOCX неизменным.
''',encoding='utf-8')

class Report:
    def __init__(self,n,theme):
        self.n=n;self.doc=Document(REFERENCE);self.land=False;self.t=0;self.f=0
        body=self.doc._element.body
        for e in list(body)[16:]:body.remove(e)
        sec=deepcopy(Document(REFERENCE).sections[0]._sectPr)
        for x in list(sec.findall(qn('w:sectPrChange'))):sec.remove(x)
        body.append(sec)
        for p in self.doc.paragraphs:
            for old in list(p._p.xpath('.//w:sectPr')):old.getparent().remove(old)
        ps=self.doc.paragraphs
        ps[4].clear();font(ps[4].add_run(f'ОТЧЕТ ПО ПРАКТИЧЕСКОЙ РАБОТЕ № {n}'),16,True)
        ps[5].clear();font(ps[5].add_run('по дисциплине «Информационный менеджмент»'))
        # Keep title-page run sizes and emphasis from the source.
        for p in self.doc.paragraphs:
            for r in p.runs:
                if r.font.name=='Liberation Serif':r.font.name='Times New Roman'
        for tx in body.xpath('.//w:t'):
            if tx.text and 'ИНБО-01-17' in tx.text:tx.text=tx.text.replace('ИНБО-01-17','')
            if tx.text and 'Шендяпин' in tx.text:tx.text=tx.text.replace('Шендяпин А.В.','Пяткин В.В.')
        for pp in self.doc.tables[1]._tbl.xpath('.//w:p'):
            nodes=pp.xpath('.//w:t')
            if 'Шендяпин' in ''.join(x.text or '' for x in nodes):
                nodes[0].text='Пяткин В.В.'
                for x in nodes[1:]:x.text=''
        for rf in body.xpath('.//w:rFonts'):
            for a in list(rf.attrib):
                if rf.attrib[a]=='Liberation Serif':rf.set(a,'Times New Roman')
        for st in self.doc.styles:
            if st.type==1 and st.name in ['Normal','Heading 1','Heading 2','Title']:
                st.font.name='Times New Roman';st.font.size=Pt(14);st.font.color.rgb=RGBColor(0,0,0)
        from docx.enum.style import WD_STYLE_TYPE
        for name in ['Основной текст отчёта','Текст таблицы','Подпись таблицы','Подпись рисунка','Heading 1','Heading 2']:
            if name not in self.doc.styles:self.doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        for name,size,align,indent,line,bold,italic in [
            ('Основной текст отчёта',14,WD_ALIGN_PARAGRAPH.JUSTIFY,1.25,1.5,False,False),
            ('Текст таблицы',12,WD_ALIGN_PARAGRAPH.LEFT,0,1,False,False),
            ('Подпись таблицы',12,WD_ALIGN_PARAGRAPH.LEFT,0,1,False,True),
            ('Подпись рисунка',14,WD_ALIGN_PARAGRAPH.CENTER,0,1,False,False),
            ('Heading 1',14,WD_ALIGN_PARAGRAPH.CENTER,0,1.5,True,False),
            ('Heading 2',14,WD_ALIGN_PARAGRAPH.LEFT,0,1.5,True,False)]:
            st=self.doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.bold=bold;st.font.italic=italic;st.font.color.rgb=RGBColor(0,0,0)
            pf=st.paragraph_format;pf.alignment=align;pf.first_line_indent=Cm(indent);pf.left_indent=Cm(0);pf.right_indent=Cm(0);pf.line_spacing=line;pf.space_before=Pt(0);pf.space_after=Pt(0);pf.widow_control=True
            if name.startswith('Heading'):pf.space_before=Pt(12);pf.space_after=Pt(6);pf.keep_with_next=True;pf.keep_together=True
            if name.startswith('Heading'):
                pr=st.element.get_or_add_pPr();ol=pr.find(qn('w:outlineLvl'))
                if ol is None:ol=OxmlElement('w:outlineLvl');pr.append(ol)
                ol.set(qn('w:val'),'0' if name=='Heading 1' else '1')
            if name=='Подпись таблицы':pf.space_before=Pt(6);pf.space_after=Pt(3);pf.keep_with_next=True;pf.keep_together=True
            if name=='Подпись рисунка':pf.space_before=Pt(3);pf.space_after=Pt(6);pf.keep_together=True
        s=self.doc.sections[0];self.setup(s,False);s.different_first_page_header_footer=True
        for p in s.footer.paragraphs:p.clear()
        p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.left_indent=Cm(0);p.paragraph_format.right_indent=Cm(0);p.paragraph_format.line_spacing=1
        r=p.add_run();font(r,12);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),' PAGE ');r._r.addnext(field)
        for p in s.first_page_footer.paragraphs:p.clear()
        self.page('Практическая работа '+str(n)+'\n'+theme)
        self.doc.core_properties.author='Албахтин И.В.';self.doc.core_properties.title=f'Информационный менеджмент. Практическая работа {n}. {theme}'
    def setup(self,s,land):
        s.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
        s.page_width=Cm(29.7 if land else 21);s.page_height=Cm(21 if land else 29.7)
        s.left_margin=Cm(3);s.right_margin=Cm(1.5);s.top_margin=Cm(2);s.bottom_margin=Cm(2);s.footer_distance=Cm(1)
    def page(self,title,land=False):
        if land!=self.land:
            s=self.doc.add_section(WD_SECTION_START.NEW_PAGE);self.setup(s,land);s.different_first_page_header_footer=False
            for el in list(s._sectPr.findall(qn('w:pgNumType'))):s._sectPr.remove(el)
            self.land=land
            # Section transition must not inherit large paragraph spacing.
            p=self.doc.paragraphs[-1];p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
            for r in p.runs:font(r,1)
            self.h(title)
        else:
            p=self.h(title);p.paragraph_format.page_break_before=True
        return self
    def h(self,s,level=1):return self.doc.add_paragraph(s,style='Heading '+str(level))
    def p(self,s):return self.doc.add_paragraph(s,style='Основной текст отчёта')
    def table(self,title,headers,rows,widths=None,compact=False):
        self.t+=1
        self.doc.add_paragraph(f'Таблица {self.t} – {title}',style='Подпись таблицы')
        t=self.doc.add_table(rows=0,cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
        widths=widths or [(25.2 if self.land else 16.5)/len(headers)]*len(headers)
        for c,w in zip(t.columns,widths):c.width=Cm(w)
        for i,vals in enumerate([headers]+rows):
            row=t.add_row();pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
            if i==0:pr.append(OxmlElement('w:tblHeader'))
            for j,(c,v,w) in enumerate(zip(row.cells,vals,widths)):
                c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp=c._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders');cp.append(bd)
                for side in ['top','left','bottom','right']:
                    x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'6');x.set(qn('w:color'),'B7B7B7');bd.append(x)
                mar=OxmlElement('w:tcMar');cp.append(mar)
                for side,val in [('top',25 if compact else 70),('bottom',25 if compact else 70),('left',85),('right',85)]:
                    x=OxmlElement('w:'+side);x.set(qn('w:w'),str(val));x.set(qn('w:type'),'dxa');mar.append(x)
                if i==0:
                    x=OxmlElement('w:shd');x.set(qn('w:fill'),'EDEDED');cp.append(x)
                p=c.paragraphs[0];p.style='Текст таблицы';p.paragraph_format.keep_with_next=False
                if len(str(v))<14:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                font(p.add_run(str(v)),12,i==0)
        return t
    def fig(self,path,title):
        self.f+=1
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.left_indent=Cm(0);pf.right_indent=Cm(0);pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=1;pf.keep_with_next=True
        p.add_run().add_picture(str(path),width=Cm(25.0 if self.land else 16.5))
        self.doc.add_paragraph(f'Рисунок {self.f} – {title}',style='Подпись рисунка')
    def sources(self,extra=[]):
        self.h('Источники',2)
        ranges={2:'9–15',3:'16–20',4:'21–27',5:'28–32'}
        self.p(f'1. Информационный менеджмент. Методические указания к практике. Практическая работа {self.n}. С. {ranges[self.n]}.')
        self.p('2. Албахтин И.В. Информационный менеджмент. Практическая работа 1. Экспресс-обследование ООО «Юкомс». Москва, 2026.')
        for i,s in enumerate(extra,3):self.p(f'{i}. {s}')
    def save(self):
        p=OUT/f'ИМ_Практическая_{self.n}_АлбахтинИВ.docx';self.doc.save(p);print(p)
        return p
