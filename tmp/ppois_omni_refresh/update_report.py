from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt
from docx.text.paragraph import Paragraph
from PIL import Image
from pypdf import PdfReader
import shutil,sys
BASE=Path('D:/GitHub/MIREA/tmp/ppois_omni_refresh')
sys.path.insert(0,str(BASE/'pydeps'))
import pymupdf
DEST=Path('D:/GitHub/MIREA/4/ППОИС')
NAMES=[('use','ПР3_01_Варианты_использования'),('classes','ПР3_02_Классы_анализа'),('bom_seq','ПР3_03_Последовательность_UC02'),('pub_seq','ПР3_04_Последовательность_UC05'),('comm','ПР3_05_Кооперация_UC05')]
for name,out in NAMES:
    pdf=BASE/(name+'.pdf')
    assert len(PdfReader(pdf).pages)==1,(name,len(PdfReader(pdf).pages))
    with pymupdf.open(pdf) as vector:
        (BASE/(name+'.svg')).write_text(vector[0].get_svg_image(text_as_path=True),encoding='utf8')
    for ext in ['png','svg']:shutil.copy2(BASE/(name+'.'+ext),DEST/'Схемы'/(out+'.'+ext))
shutil.copy2(BASE/'updated.omni',DEST/'Схемы'/'ПР3_Модель_анализа.omni')
d=Document(BASE/'before.docx')
captions=[p for p in d.paragraphs if p.text.startswith('Рисунок ')]
assert len(captions)==7
for cap,name in zip(captions,['use','classes','bom_seq_1','bom_seq_2','pub_seq_1','pub_seq_2','comm']):
    p=Paragraph(cap._p.getprevious(),cap._parent)
    assert p._p.xpath('.//w:drawing')
    for e in list(p._p):
        if e.tag.endswith('}r'):p._p.remove(e)
    w,h=Image.open(BASE/(name+'.png')).size
    maxh=12.9 if name=='comm' else 12.3 if name.endswith('_1') else 14.3
    width=min(24.4,maxh*w/h)
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
    p.add_run().add_picture(str(BASE/(name+'.png')),width=Cm(width))
d.save(DEST/'ППОИС_3_АлбахтинИВ.docx')
print('Replaced seven report panels; updated five PNG/SVG files and the native model.')
