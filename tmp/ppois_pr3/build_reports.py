from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import hashlib, json, re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from diagrams import BASE, ASSET, CLASSES, SEQ1, SEQ2, UC

ROOT=Path('D:/GitHub/MIREA');DEST=ROOT/'4/ППОИС';REF=ROOT/'4/ИСУРО/ИСУРО_1_АлбахтинИВ.docx'

def section(s,land=False,first=False):
 s.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT;s.page_width=Cm(29.7 if land else 21);s.page_height=Cm(21 if land else 29.7)
 s.left_margin=Cm(3);s.right_margin=Cm(1.5);s.top_margin=Cm(2);s.bottom_margin=Cm(2);s.footer_distance=Cm(1)
 s.different_first_page_header_footer=first
 for e in s._sectPr.findall(qn('w:pgNumType')):s._sectPr.remove(e)

def style_doc(d):
 for name in ['ReportBody','ReportTable','ReportTableCaption','ReportFigure','Heading 1','Heading 2']:
  if name not in d.styles:d.styles.add_style(name,1)
  st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(12 if name in ['ReportTable','ReportTableCaption'] else 14);st.font.color.rgb=RGBColor(0,0,0)
  st.font.bold=name.startswith('Heading');st.font.italic=name=='ReportTableCaption'
  pf=st.paragraph_format;pf.left_indent=0;pf.right_indent=0;pf.first_line_indent=Cm(1.25) if name=='ReportBody' else 0;pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=1.5 if name=='ReportBody' or name.startswith('Heading') else 1;pf.widow_control=True;pf.page_break_before=False
  pf.alignment=A.JUSTIFY if name=='ReportBody' else A.CENTER if name in ['ReportFigure','Heading 1'] else A.LEFT
  pf.keep_with_next=name.startswith('Heading') or name=='ReportTableCaption';pf.keep_together=name.startswith('Heading') or name=='ReportFigure'
  if name.startswith('Heading'):pf.space_before=Pt(12);pf.space_after=Pt(6)
  if name=='ReportTableCaption':pf.space_before=Pt(6);pf.space_after=Pt(3)
  if name=='ReportFigure':pf.space_before=Pt(3);pf.space_after=Pt(6)

def field(p,instr,cached='2'):
 el=OxmlElement('w:fldSimple');el.set(qn('w:instr'),instr);r=OxmlElement('w:r');rp=OxmlElement('w:rPr');fo=OxmlElement('w:rFonts');fo.set(qn('w:ascii'),'Times New Roman');fo.set(qn('w:hAnsi'),'Times New Roman');rp.append(fo);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'24');rp.append(sz);r.append(rp);t=OxmlElement('w:t');t.text=cached;r.append(t);el.append(r);p._p.append(el)

def footers(d):
 for i,s in enumerate(d.sections):
  s.header.is_linked_to_previous=False;s.footer.is_linked_to_previous=False;s.first_page_footer.is_linked_to_previous=False
  for hf in [s.header,s.footer,s.first_page_footer]:
   for p in hf.paragraphs:p.clear()
  p=s.footer.paragraphs[0];p.alignment=A.CENTER;p.paragraph_format.first_line_indent=0;p.paragraph_format.left_indent=0;p.paragraph_format.right_indent=0;p.paragraph_format.line_spacing=1;field(p,'PAGE')
  s.different_first_page_header_footer=i==0

def fix_pr2():
 d=Document(BASE/'input/original_pr2.docx');ps=d.paragraphs
 for pi,name in [(42,'ПР2_IDEF0_A-0'),(46,'ПР2_IDEF0_A0'),(50,'ПР2_IDEF0_A2')]:
  p=ps[pi];p.clear();p.alignment=A.CENTER;p.paragraph_format.first_line_indent=0;p.paragraph_format.keep_with_next=True;p.paragraph_format.line_spacing=1;p.add_run().add_picture(str(ASSET/(name+'.png')),width=Cm(23.0 if pi==42 else 24.8))
 for i in (43,47,51):
  p=ps[i];p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=True
 for i in (40,44,48):ps[i].paragraph_format.page_break_before=i!=40
 for i in (41,45,49):ps[i].paragraph_format.line_spacing=1;ps[i].paragraph_format.space_after=Pt(5)
 # End the portrait prose before the first diagram; all three diagrams use landscape pages.
 ep=OxmlElement('w:p');pr=OxmlElement('w:pPr');sp=deepcopy(d.sections[1]._sectPr);pr.append(sp);ep.append(pr);ps[40]._p.addprevious(ep)
 for i,s in enumerate(d.sections):section(s,land=i in (2,3),first=i==0)
 footers(d);d.save(DEST/'ППОИС_2_АлбахтинИВ.docx')

def base_doc():
 d=Document(REF);body=d.element.body;sect=deepcopy(body[15].xpath('.//w:sectPr')[0])
 for el in list(body)[16:]:body.remove(el)
 body.append(sect)
 teacher=d.tables[1].rows[1].cells[1]
 for pp in teacher.paragraphs:
  if pp.runs:
   pp.runs[0].text='Борзых Н.Ю.'
   for rr in pp.runs[1:]:rr.text=''
 for t in body.xpath('.//w:t'):
  if t.text:t.text=t.text.replace('ИНБО-01-17','').replace('Информационные системы управления ресурсами организации','Проектирование предметно-ориентированных информационных систем').replace('Шендяпин А.В.','Борзых Н.Ю.').replace('ОТЧЕТ ПО ПРАКТИЧЕСКИМ РАБОТАМ','ОТЧЕТ ПО ПРАКТИЧЕСКОЙ РАБОТЕ № 3')
 for rfont in body.xpath('.//w:rFonts'):
  for k in ['ascii','hAnsi','eastAsia','cs']:rfont.set(qn('w:'+k),'Times New Roman')
 if 'Title' not in d.styles:d.styles.add_style('Title',1)
 st=d.styles['Title'];st.font.name='Times New Roman';st.font.size=Pt(16);st.font.bold=True;st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.alignment=A.CENTER
 for p in d.paragraphs:
  if p.text.startswith('ОТЧЕТ'):p.style=d.styles['Title']
 style_doc(d)
 for s in d.sections:section(s)
 d.core_properties.author='Албахтин И.В.';d.core_properties.title='ППОИС Практическая работа 3 Модель анализа ООО Юкомс'
 return d

def create_pr3():
 d=base_doc();land=False;tn=0;fn=0;pending_break=False
 def page(wide=False,force=True):
  nonlocal land,pending_break
  if wide!=land:
   section(d.add_section(WD_SECTION_START.NEW_PAGE),wide);land=wide
   pp=d.paragraphs[-1];pp.paragraph_format.space_before=Pt(0);pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=Pt(1)
   for rr in pp.runs:rr.font.size=Pt(1)
  elif force:pending_break=True
 def p(t):return d.add_paragraph(t,'ReportBody')
 def h(t,l=1):
  nonlocal pending_break
  pp=d.add_paragraph(t,'Heading '+str(l));pp.paragraph_format.page_break_before=pending_break;pending_break=False;return pp
 def table(title,heads,rows,widths):
  nonlocal tn;tn+=1
  d.add_paragraph(f'Таблица {tn} – {title}','ReportTableCaption')
  tb=d.add_table(rows=1,cols=len(heads));tb.autofit=False;tb.alignment=WD_TABLE_ALIGNMENT.CENTER
  for col,w in zip(tb.columns,widths):col.width=Cm(w)
  for i,tx in enumerate(heads):tb.rows[0].cells[i].text=tx
  for rr in rows:
   for c,tx in zip(tb.add_row().cells,rr):c.text=str(tx)
  for ri,row in enumerate(tb.rows):
   trp=row._tr.get_or_add_trPr();trp.append(OxmlElement('w:cantSplit'))
   if ri==0:trp.append(OxmlElement('w:tblHeader'))
   for ci,c in enumerate(row.cells):
    c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tcp=c._tc.get_or_add_tcPr();ma=OxmlElement('w:tcMar')
    for side in ['top','bottom','left','right']:
     e=OxmlElement('w:'+side);e.set(qn('w:w'),'55' if side in ['top','bottom'] else '75');e.set(qn('w:type'),'dxa');ma.append(e)
    tcp.append(ma)
    borders=OxmlElement('w:tcBorders')
    for side in ['top','left','bottom','right']:
     e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'808080');borders.append(e)
    tcp.append(borders)
    if ri==0:
     sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEDED');tcp.append(sh)
    for pp in c.paragraphs:
     pp.style='ReportTable'
     for r in pp.runs:r.bold=ri==0
  return tb
 def figure(name,caption,width=24.8):
  nonlocal fn;fn+=1
  pp=d.add_paragraph();pp.alignment=A.CENTER;pp.paragraph_format.first_line_indent=0;pp.paragraph_format.left_indent=0;pp.paragraph_format.right_indent=0;pp.paragraph_format.line_spacing=1;pp.paragraph_format.space_before=Pt(0);pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.keep_with_next=True
  pp.add_run().add_picture(str(ASSET/(name+'.png')),width=Cm(width));d.add_paragraph(f'Рисунок {fn} – {caption}','ReportFigure')
 h('ПРАКТИЧЕСКАЯ РАБОТА № 3')
 pp=p('Формирование функциональных требований к ИС\nОбъектный анализ и модель анализа');pp.alignment=A.CENTER;pp.paragraph_format.first_line_indent=0
 h('1 Цель и постановка задачи')
 p('Цель работы — выполнить концептуальное и динамическое моделирование предметной области при разработке информационной системы формирования и актуализации карточек товаров компьютерных систем ООО «Юкомс».')
 p('По материалам практик № 1 и № 2 необходимо построить диаграммы вариантов использования, классов анализа, последовательности для двух сценариев и кооперации, а также таблицы взаимодействий [1–3].')
 h('2 Границы системы и участники')
 p('Система управляет составом компьютерных конфигураций BOM, предложениями поставщиков и данными карточек Ozon. Закупки, сборка компьютеров, бухгалтерский учёт и доставка находятся за границами модели. Участники и их цели приведены в таблице 1.')
 table('Актёры проектируемой системы',['Актёр','Цель взаимодействия'],[
 ('Менеджер маркетплейсов','Зарегистрировать запрос, обновить данные поставщиков и опубликовать карточку.'),('Технический специалист','Сформировать BOM, проверить совместимость и заменить компоненты.'),('Дизайнер','Подготовить контент и настроить шаблоны изображений.'),('Администратор','Настроить права доступа и параметры интеграций.'),('PC4Games и ITPartner','Передать ассортимент, характеристики, цены и наличие.'),('Ozon','Принять карточку и вернуть результат обработки.')],[5,11.5])
 p('BOM — спецификация состава компьютерной системы. На рисунке 1 приведены варианты использования и граница системы.')
 page(True);h('3 Диаграмма вариантов использования')
 figure('ПР3_01_Варианты_использования','Варианты использования информационной системы',24.7)
 page(False);h('3.1 Взаимодействия актёров и вариантов использования',2)
 p('Описание связей рисунка 1 приведено в таблице 2. Ассоциация обозначает участие актёра в сценарии. Связь «include» направлена от основного варианта к обязательной общей проверке; она не задаёт очередность выполнения самостоятельных сценариев.')
 table('Описание взаимодействий актёров и вариантов использования',['Актёр / ВИ','Тип связи','Вариант использования'],[
 ('Менеджер маркетплейсов','Ассоциация','UC01 — зарегистрировать запрос'),('Менеджер маркетплейсов','Ассоциация','UC05 — опубликовать карточку'),('Менеджер маркетплейсов','Ассоциация','UC06 — обновить предложения поставщиков'),('Технический специалист','Ассоциация','UC02 — сформировать BOM'),('Технический специалист','Ассоциация','UC03 — проверить совместимость'),('Технический специалист','Ассоциация','UC04 — массово заменить компонент'),('Дизайнер','Ассоциация','UC07 — рассчитать цену, вес и подготовить контент'),('Дизайнер','Ассоциация','UC09 — настроить шаблоны изображений'),('Администратор','Ассоциация','UC10 — настроить доступ и интеграции'),('PC4Games и ITPartner','Ассоциация','UC06 — обновить предложения поставщиков'),('Ozon','Ассоциация','UC05 — опубликовать карточку'),('UC02 — сформировать BOM','«include»','UC03 — проверить совместимость'),('UC04 — массово заменить компонент','«include»','UC03 — проверить совместимость'),('UC05 — опубликовать карточку','«include»','UC08 — проверить готовность карточки')],[5.2,3.1,8.2])
 p('UC02 и UC05 детализируются далее как два ключевых сценария. UC04 создаёт новые версии затронутых BOM после подтверждения области замены. UC07 использует согласованную BOM; публикация выполняется отдельным действием менеджера.')
 p('На рисунке 2 представлены классы анализа для формирования BOM и публикации карточки. Модель разделяет интерфейсы взаимодействия, управление сценариями и объекты предметной области.')
 page(True);h('4 Диаграмма классов анализа')
 figure('ПР3_02_Классы_анализа','Классы анализа двух ключевых сценариев',22.5)
 page(False);h('4.1 Ответственность классов и ограничения',2)
 p('Классы «boundary» представляют формы и внешнюю интеграцию, «control» координируют сценарии, «entity» хранят предметные данные. Назначение классов рисунка 2 приведено в таблице 3.')
 roles=[('Форма BOM','boundary','Принимает состав конфигурации и показывает результат проверки.'),('Форма карточки','boundary','Принимает команду публикации и показывает статус.'),('Адаптер Ozon','boundary','Передаёт данные и приводит ответ внешнего сервиса к результату сценария.'),('Контроллер BOM','control','Проверяет состав по правилам и сохраняет согласованную версию.'),('Контроллер публикации','control','Проверяет готовность, организует обмен и фиксирует результат.'),('Версия BOM','entity','Хранит номер версии, статус и состав конфигурации.'),('Слот BOM','entity','Описывает тип, количество и группу либо точную модель компонента.'),('Предложение поставщика','entity','Хранит модель, цену, наличие и момент актуальности.'),('Правило совместимости','entity','Определяет ограничения на сочетания компонентов.'),('Карточка товара','entity','Связывает атрибуты, цену и изображения с одной версией BOM.'),('Публикация','entity','Фиксирует запрос, ответ, время и результат отдельной попытки.')]
 table('Ответственность классов анализа',['Класс','Стереотип','Ответственность'],roles,[4.3,2.4,9.8])
 p('Версия BOM содержит один или несколько слотов (композиция). Каждый слот принадлежит одной версии и имеет не более одного выбранного предложения. Для согласованной BOM предложение обязательно; одно предложение может использоваться в нескольких BOM.')
 p('Карточка ссылается на одну версию BOM и имеет историю попыток публикации. Изменение BOM создаёт новую версию; связанные карточки требуют актуализации.')
 page(True);h('5 Динамическая модель');h('5.1 Последовательность формирования BOM',2)
 p('На рисунке 3 показано формирование BOM по зарегистрированному запросу и выбранному корпусу.')
 figure('ПР3_03_Последовательность_UC02','Диаграмма последовательности UC02',23.1)
 page(True);h('5.2 Сообщения сценария UC02',2)
 p('В таблице 4 приведены сообщения рисунка 3. Все вызовы синхронные. Шаги 8–11 и 12–13 — альтернативные ветви.')
 ns1=['Технический специалист','Форма BOM','Контроллер BOM','Правило совместимости','Предложение поставщика','Версия BOM']
 table('Взаимодействие элементов диаграммы последовательности UC02',['Отправитель','Тип сообщения','Наименование','Получатель'],[(ns1[a],'Ответ' if ret else 'Вызов',tx,ns1[b]) for a,b,tx,ret in SEQ1],[5.1,3.3,10.9,5.5])
 p('Контроллер проверяет заполнение обязательных слотов, совместимость и наличие допустимого предложения. Ошибки блокируют сохранение согласованной версии. Получение предложений использует данные, обновлённые в UC06; если они недоступны или устарели, специалисту возвращается замечание.')
 page(True);h('5.3 Последовательность публикации карточки',2)
 p('На рисунке 4 показана публикация подготовленной карточки с проверкой её готовности к отправке.')
 figure('ПР3_04_Последовательность_UC05','Диаграмма последовательности UC05',22.8)
 page(True);h('5.4 Сообщения сценария UC05',2)
 p('В таблице 5 приведены сообщения рисунка 4. Все вызовы синхронные. Ветвь 6–14 — готовая карточка, 15–16 — ошибки готовности.')
 ns2=['Менеджер маркетплейсов','Форма карточки','Контроллер публикации','Карточка товара','Адаптер Ozon','Ozon','Публикация']
 table('Взаимодействие элементов диаграммы последовательности UC05',['Отправитель','Тип сообщения','Наименование','Получатель'],[(ns2[a],'Ответ' if ret else 'Вызов',tx,ns2[b]) for a,b,tx,ret in SEQ2],[5.1,3.3,10.9,5.5])
 page(True);h('6 Диаграмма кооперации')
 p('На рисунке 5 показана ветвь готовой карточки. Сообщения 1–14 соответствуют таблице 5.')
 figure('ПР3_05_Кооперация_UC05','Кооперация объектов при публикации карточки',23.5)
 page(False);h('7 Согласованность модели и функциональные требования')
 p('В UC05 попытка регистрируется до отправки. Адаптер представляет обмен с Ozon как одну операцию анализа; при отложенной обработке итог уточняется отдельно. Подтверждённый приём даёт статус «опубликована», отказ — «ошибка публикации», отсутствие окончательного ответа — «ожидает подтверждения».')
 p('Статическая и динамическая модели используют одни и те же классы. Формы обращаются к контроллерам; внешнее взаимодействие проходит через адаптер, а результат сохраняется в предметных объектах. В таблице 6 показана связь модели анализа с функциями IDEF0 второй практической работы.')
 table('Трассировка функций IDEF0 и модели анализа',['Функция IDEF0','Варианты использования','Результат'],[
 ('A1 — инициировать карточку','UC01','Запрос и черновик карточки'),('A2 — сформировать BOM','UC02, UC04, UC06','Состав и версия BOM с предложениями'),('A3 — проверить совместимость','UC03','Согласованная BOM либо список конфликтов'),('A4 — рассчитать и подготовить контент','UC07, UC09','Цена, вес, изображения и атрибуты'),('A5 — проверить и опубликовать','UC05, UC08','Статус карточки и запись публикации')],[5.2,4.3,7])
 p('UC10 обеспечивает доступ и работоспособность интеграций для всех функций. Отдельного блока основного потока IDEF0 ему не соответствует: он поддерживает механизмы выполнения процесса.')
 h('7.1 Уточнённые требования',2)
 p('ФТ-01. При сохранении BOM система должна проверить обязательные слоты, совместимость и выбранные предложения. При наличии ошибок согласованная версия не создаётся, а пользователь получает замечания по конкретным слотам.')
 p('ФТ-02. Массовая замена должна показывать затронутые BOM до подтверждения, создавать новые версии и помечать связанные карточки как требующие актуализации.')
 p('ФТ-03. Публикация допускается только для готовой карточки. Перед отправкой создаётся запись попытки с идентификатором карточки и версии; после обмена сохраняются ответ и итоговый статус.')
 p('ФТ-04. Ошибка внешнего обмена не должна удалять подготовленные данные. До повторной отправки система должна проверить результат предыдущей попытки, если он не был подтверждён, и показать менеджеру доступное действие.')
 h('8 Вывод')
 p('Разработана модель анализа ИС ООО «Юкомс»: шесть типов актёров, десять вариантов использования, классы анализа, две диаграммы последовательности и диаграмма кооперации. Описаны условия сохранения BOM, проверки готовности карточки и фиксации результата публикации.')
 p('Таблицы взаимодействий согласованы с диаграммами и функциями IDEF0. Результаты служат основой для проектирования компонентов, базы данных и интерфейсов системы.')
 h('Список использованных источников')
 for tx in ['1. ППОИС. Практическая работа № 3. Объектный анализ. Модель анализа. Методические материалы.','2. Албахтин И.В. ППОИС. Отчёт по практической работе № 1. Москва, 2026.','3. Албахтин И.В. ППОИС. Отчёт по практической работе № 2. Москва, 2026.']:p(tx)
 footers(d)
 d.save(DEST/'ППОИС_3_АлбахтинИВ.docx')
 return tn,fn

def contract():
 with ZipFile(REF) as z:inventory={n:hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}
 (BASE/'reference_parts.json').write_text(json.dumps(inventory,indent=2),encoding='utf-8')
 (BASE/'artifact.md').write_text(f'''# Контракт оформления ППОИС 3
Источник титульника: {REF.as_posix()}
SHA256: {hashlib.sha256(REF.read_bytes()).hexdigest()}
Источник содержит 13 страниц и 2 раздела. По инструкции пользователя используется только титульная страница, body children 0..15, включая университетскую шапку с эмблемой и двойной линией, институт, кафедру, таблицу подписей, город и год. Остальная часть заменяется.
Титульный визуальный ориентир: input/cover-01.png. Образец сохранён без изменений.
Слоты: body child5 — название, одна практика №3, TNR16 bold; child6 — ППОИС, TNR14; table2 row2 cell2 — Борзых Н.Ю.; hidden ИНБО-01-17 удалить. Студент ИНБО-12-23 Албахтин И.В., Москва2026г. Размеры и интервалы титульника сохраняются. Liberation Serif заменить TNR.
Основная часть по Оформление_отчётов.md: А4, поля30/15/20/20мм, TNR14, по ширине,1.5,абзац1.25см; таблицы TNR12 одинарный. Подпись таблицы12 italic слева без отступа сверху, рисунка14 обычный центр снизу. Заголовки1/2 TNR14 bold центр/слева12/6пт,keep-next.
Схемы и таблицы сообщений — отдельные альбомные разделы, текст — книжные. Диаграммы встроены inline, editable SVG лежат в папке предмета Схемы. Сквозные таблицы1..6, рисунки1..5. PAGE внизу по центру,TNR12,титульный номер скрыт.
Новые стили ReportBody/ReportTable/ReportFigure не меняют стиль титульника. Допустимые изменения пакета:document/styles/settings/relationships/core/footer и добавляемые media. Исходные изображения титульника сохраняются byte-for-byte. Ссылки на исходные ненужные медиа могут остаться в пакете без отображения.
Содержание: цель,актёры,UseCase+связи,классы+ответственности,2sequence+таблицы,кооперация,трассировка,требования,вывод,источники. У UML нумерация сообщений совпадает с таблицами.
ПР2: исходник input/original_pr2.docx. Меняются только три схемы, их размер и альбомная ориентация разделов. Исправляются привязки подписей к выходам и входам. Текст и таблицы сохраняются.
QA: bundled render_docx.py через Word COM adapter, PDF и PNG всех страниц, визуальная проверка каждой страницы.
''',encoding='utf-8')

if __name__=='__main__':
 contract();fix_pr2();print('PR3 table/figure counts',create_pr3())
