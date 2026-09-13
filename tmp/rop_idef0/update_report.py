from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Cm

base=Path('D:/GitHub/MIREA/4/РОП')
report=base/'РОП_Единый_отчет_АлбахтинИВ.docx'
backup=Path('D:/GitHub/MIREA/tmp/rop_idef0/РОП_отчет_до_обновления.docx')
if not backup.exists(): copy2(report,backup)
doc=Document(backup)
for idx,name,width in [(33,'A-0.png',16),(38,'A0.png',24.6)]:
 p=doc.paragraphs[idx]
 p.clear()
 p.add_run().add_picture(str(base/'IDEF0'/name),width=Cm(width))
 p.paragraph_format.keep_with_next=True
 p.paragraph_format.space_before=0
 p.paragraph_format.space_after=0
 p.paragraph_format.first_line_indent=0
 p.paragraph_format.left_indent=0
 p.paragraph_format.right_indent=0

def text(p,value):
 if p.runs:
  p.runs[0].text=value
  for r in p.runs[1:]:r.text=''
 else:p.add_run(value)

text(doc.paragraphs[32],'Контекстная диаграмма показывает преобразование запроса и исходных данных в согласованную BOM, рассчитанные показатели, изображения и результат публикации карточки. Управления задают порядок работы, правила выбора, совместимости и расчётов, требования Ozon. Механизмы включают участников и используемые инструменты.')
text(doc.paragraphs[36],'На уровне A0 выделены пять функций. Их внешние интерфейсы согласованы с контекстом A-0 по кодам ICOM. Замечания по совместимости возвращаются из A3 в A2 для исправления конфигурации. Результат A5 — опубликованная или обновлённая карточка либо сведения об ошибке публикации. Модель AS IS выполнена в OmniNotation.')
text(doc.tables[2].rows[3].cells[2].paragraphs[0],'Согласованная BOM или замечания для исправления конфигурации')
doc.save(report)
print(report)
