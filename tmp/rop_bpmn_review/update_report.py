from pathlib import Path
from shutil import copy2
from copy import deepcopy
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
base=Path('D:/GitHub/MIREA/4/РОП');tmp=Path('D:/GitHub/MIREA/tmp/rop_bpmn_review')
report=base/'РОП_Единый_отчет_АлбахтинИВ.docx';backup=tmp/'original.docx'
if not backup.exists():copy2(report,backup)
d=Document(backup)
def settext(p,s):
 if p.runs:
  p.runs[0].text=s
  for r in p.runs[1:]:r.text=''
 else:p.add_run(s)
def picture(p,path):
 p.clear();p.add_run().add_picture(str(path),width=Cm(24.6))
 p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 f=p.paragraph_format;f.first_line_indent=f.left_indent=f.right_indent=0;f.line_spacing=1;f.space_before=f.space_after=0;f.keep_with_next=True
def caption(p,s):
 settext(p,s);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 f=p.paragraph_format;f.first_line_indent=f.left_indent=f.right_indent=0;f.line_spacing=1;f.space_before=Pt(3);f.space_after=Pt(6);f.keep_with_next=False;f.keep_together=True
 for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(14);r.bold=r.italic=False
settext(d.paragraphs[57],'На рисунке 3 показано взаимодействие ООО «Юкомс» с поставщиками PC4Games и ITPartner и площадкой Ozon. Внешние участники выделены в отдельные пулы; обмен данными показан потоками сообщений. Менеджер инициирует работу, отправляет карточку и фиксирует результат: публикацию либо ошибку с замечаниями. Повторная попытка после ошибки начинается с нового запроса на изменение.')
p=d.paragraphs[58].insert_paragraph_before('На рисунке 4 раскрыт подпроцесс подготовки данных. Дорожки обозначают технического специалиста, сотрудника склада и дизайнера. После исправления или замены компонента повторяются проверка совместимости и сбор предложений. Расчёт цены и веса и подготовка изображений выполняются после подтверждения доступности. Обычные задачи обозначены без специализации; модель AS IS не предполагает движок исполнения BPMN [8].')
p.style=d.paragraphs[57].style
for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(14)
f=p.paragraph_format;f.first_line_indent=Cm(1.25);f.line_spacing=1.5;f.space_after=0;p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
pics=[p for p in d.paragraphs if p._p.xpath('.//a:blip/@r:embed')==['rId28']]
pic=pics[0];picture(pic,base/'BPMN/BPMN_1_Публикация.png')
cap=Paragraph(pic._p.getnext(),pic._parent);caption(cap,'Рисунок 3 – BPMN AS IS: публикация и внешние участники')
e=OxmlElement('w:p');cap._p.addnext(e);pic2=Paragraph(e,cap._parent);picture(pic2,base/'BPMN/BPMN_2_Подготовка.png');pic2.paragraph_format.page_break_before=True
e2=OxmlElement('w:p');e.addnext(e2);caption(Paragraph(e2,cap._parent),'Рисунок 4 – BPMN AS IS: подпроцесс подготовки данных карточки')
for p in d.paragraphs:
 if p.text.startswith('Рисунок 4 – Средняя'):settext(p,p.text.replace('Рисунок 4','Рисунок 5',1))
settext(d.tables[2].rows[2].cells[1].paragraphs[0],'Карточка опубликована или обновлена в Ozon либо зарегистрирована ошибка публикации')
p=d.add_paragraph();p._p.getparent().remove(p._p)
last=d.paragraphs[-1];new=deepcopy(last._p);last._p.addnext(new);settext(Paragraph(new,last._parent),'Object Management Group. Business Process Model and Notation (BPMN), Version 2.0.2. URL: https://www.omg.org/spec/BPMN/2.0.2/PDF (дата обращения: 14.09.2026).')
d.save(tmp/'updated.docx');print(tmp/'updated.docx')
