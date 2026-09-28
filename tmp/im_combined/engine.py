from pathlib import Path
from copy import deepcopy
from io import BytesIO
import re, hashlib, json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH as A
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table

BASE=Path(__file__).resolve().parent
ROOT=BASE.parent.parent
OUT=ROOT/'4/ИМ'
REF=Path('C:/Users/VeX/Downloads/AyuGram Desktop/ИМ_ГубаревСА_Отчёт.docx')
DOCX=OUT/'ИМ_АлбахтинИВ_Отчёт.docx'

def font(run,size=14,bold=False,italic=False):
    run.font.name='Times New Roman';run.font.size=Pt(size)
    run.bold=bold;run.italic=italic;run.font.color.rgb=RGBColor(0,0,0)
    pr=run._element.get_or_add_rPr()
    for x in list(pr):
        if x.tag==qn('w:rFonts'):
            for a in list(x.attrib):
                if 'Theme' in a or 'theme' in a:del x.attrib[a]
    rf=pr.rFonts
    for a in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+a),'Times New Roman')

def txt(e):return ''.join(t.text or '' for t in e.iter(qn('w:t')))
def settext(p,s):
    nodes=list(p.iter(qn('w:t')))
    if nodes:
        nodes[0].text=s
        for t in nodes[1:]:t.text=''

class Report:
    def __init__(self):
        self.doc=Document(REF);self.land=False;self.n=0;self.t=0;self.f=0
        body=self.doc._element.body
        title=list(body)[:40]
        for e in list(body):body.remove(e)
        for e in title:body.append(e)
        # Preserve the source title, including its continuous column sections and shapes.
        for p in body.iter(qn('w:p')):
            s=txt(p)
            if s=='ОТЧЕТ ПО ПРАКТИЧЕСКОЙ РАБОТЕ':settext(p,'ОТЧЕТ ПО ПРАКТИЧЕСКИМ РАБОТАМ')
            if s.startswith('Тема «'):settext(p,'Тема «Проектирование информационной системы подготовки и актуализации карточек компьютерных систем на примере ООО «Юкомс»»')
            for t in p.iter(qn('w:t')):
                if t.text:t.text=t.text.replace('Губарев С.А.','Албахтин И.В.').replace('ИНБО-01-17','')
        lasttitle=title[-1].find('.//'+qn('w:sectPr'))
        sect=deepcopy(lasttitle);body.append(sect)
        typ=sect.find(qn('w:type'))
        if typ is None:typ=OxmlElement('w:type');sect.insert(0,typ)
        typ.set(qn('w:val'),'nextPage')
        # The body uses a single column; title columns remain intact.
        cols=sect.find(qn('w:cols'))
        if cols is not None:sect.remove(cols)
        self.geometry(self.doc.sections[-1],False)
        self.styles()
        for s in self.doc.sections:
            for footer in [s.footer,s.first_page_footer,s.even_page_footer]:
                for e in list(footer._element):footer._element.remove(e)
                footer.add_paragraph()
        self.p('Оглавление','IM TOC Title')
        p=self.p('','IM TOC')
        for typ,text in [('begin',None),(None,' TOC \\o "1-1" \\h \\z \\u '),('separate',None),(None,'Оглавление обновляется полем Word.'),('end',None)]:
            r=p.add_run()
            if typ:
                el=OxmlElement('w:fldChar');el.set(qn('w:fldCharType'),typ);r._r.append(el)
            else:
                el=OxmlElement('w:instrText' if text.startswith(' TOC') else 'w:t');el.text=text;r._r.append(el)
        settings=self.doc.settings._element
        u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)
        self.doc.core_properties.author='Албахтин И.В.'
        self.doc.core_properties.title='Информационный менеджмент. Практические работы 1–12'
        self.doc.core_properties.subject='Проектирование ИС подготовки карточек компьютерных систем ООО «Юкомс»'
        self.doc.core_properties.last_modified_by='Албахтин И.В.'
        (BASE/'artifact.md').write_text(f'''# Контракт единого отчёта ИМ
Образец: {REF.as_posix()}
SHA-256: {hashlib.sha256(REF.read_bytes()).hexdigest()}
Объём: практические работы 1–12, пользователь подтвердил 27.09.2026.
Сохранить титульный макет: body 0–39, эмблему, таблицу университетской шапки, двойную линию, блоки подписей и их непрерывные разделы. Заменить только ФИО, тему, множественное число названия отчёта. Чужое содержание не переносить.
Общее оглавление — поле TOC по стилю IM Practice; обновить в Word перед экспортом.
Основной текст: А4, поля слева 3, справа 2.5, сверху 2.72, снизу 0.49 см; TNR14, 1.5, первая строка1.25, по ширине. Альбомные листы по геометрии образца: слева0.49/справа2.72/сверху3/снизу2.5 см.
Заголовки практик:16 полужирный, по центру, новая страница. Подразделы:14 полужирный, слева1.27см, без цифровых префиксов; оглавление только верхнего уровня.
Таблицы:12пт, интервал1.05, чёрная сетка, без заливки; подписи12 курсив сверху,6пт до/после. Рисунки:по центру, подписи14 снизу (устойчивый вариант первой практики образца),6пт до/после. Нумерация:номер практики.порядковый номер.
Видимые номера страниц отсутствуют в образце; не добавлять. Автоматическое оглавление содержит номера страниц.
Избыточные разрывы и пустые страницы не воспроизводить. Таблицы и рисунки укладывать в поля, повторять шапки, сохранять текст и проверяемые расчёты.
Содержательные исходники1–5: существующие отчёты Албахтина. Дополнения6–12: методичка и документированные источники. Не утверждать фактических интервью, испытаний или внедрения.
''',encoding='utf-8')

    def geometry(self,s,land):
        s.orientation=WD_ORIENT.LANDSCAPE if land else WD_ORIENT.PORTRAIT
        s.page_width=Cm(29.7 if land else 21.0);s.page_height=Cm(21.0 if land else 29.7)
        for k,v in zip(['left_margin','right_margin','top_margin','bottom_margin'],[.49,2.72,3,2.5] if land else [3,2.5,2.72,.49]):setattr(s,k,Cm(v))
        s.header_distance=Cm(.3);s.footer_distance=Cm(.3)
    def styles(self):
        specs={'IM Body':(14,False,False,A.JUSTIFY,1.5,1.25,0,0,0),
        'IM Practice':(16,True,False,A.CENTER,1.5,0,0,0,6),
        'IM Head':(14,True,False,A.LEFT,1.5,0,1.27,6,0),
        'IM SmallHead':(14,True,False,A.LEFT,1.5,0,1.27,6,0),
        'IM Table':(12,False,False,A.LEFT,1.05,0,0,0,0),
        'IM Table Caption':(12,False,True,A.LEFT,1,0,0,6,6),
        'IM Figure Caption':(14,False,False,A.CENTER,1,0,0,6,6),
        'IM Figure':(14,False,False,A.CENTER,1,0,0,6,0),
        'IM TOC':(14,False,False,A.LEFT,1.5,0,0,0,0),
        'IM TOC Title':(16,True,False,A.CENTER,1.5,0,0,0,12)}
        for name,(size,bold,ital,align,line,first,left,before,after) in specs.items():
            s=self.doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH) if name not in self.doc.styles else self.doc.styles[name]
            s.base_style=self.doc.styles['Normal'];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=bold;s.font.italic=ital;s.font.color.rgb=RGBColor(0,0,0)
            p=s.paragraph_format;p.alignment=align;p.line_spacing=line;p.first_line_indent=Cm(first);p.left_indent=Cm(left);p.right_indent=Cm(0);p.space_before=Pt(before);p.space_after=Pt(after);p.widow_control=True
            p.keep_with_next=name in ['IM Practice','IM Head','IM SmallHead','IM Table Caption','IM Figure'];p.keep_together=True
        self.doc.styles['IM Practice'].paragraph_format.page_break_before=True
        for n in ['IM Practice','IM Head','IM SmallHead']:
            el=OxmlElement('w:outlineLvl');el.set(qn('w:val'),'0' if n=='IM Practice' else '1');self.doc.styles[n]._element.get_or_add_pPr().append(el)
    def orient(self,land):
        if self.land!=land:
            s=self.doc.add_section(WD_SECTION_START.NEW_PAGE);self.geometry(s,land);self.land=land
    def p(self,text,style='IM Body'):
        return self.doc.add_paragraph(text,style)
    def h(self,text):return self.p(text,'IM Head')
    def start(self,n,goal):
        self.orient(False);self.n=n;self.t=0;self.f=0
        self.p(f'Практическая работа № {n}','IM Practice')
        p=self.p('');font(p.add_run('Цель занятия: '),14,True);p.add_run(goal)
    def table(self,title,headers,rows,widths=None):
        self.t+=1;self.p(f'Таблица {self.n}.{self.t} – {title}','IM Table Caption')
        t=self.doc.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
        for vals in [headers]+list(rows):
            cells=t.rows[0].cells if vals is headers else t.add_row().cells
            for c,v in zip(cells,vals):c.text=str(v)
        self.style_table(t,widths)
        return t
    def style_table(self,t,widths=None):
        t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        available=26.49 if self.land else 15.5
        if widths is None:widths=[max(1,c.width.cm if c.width else 1) for c in t.rows[0].cells]
        widths=[w*available/sum(widths) for w in widths]
        for col,w in zip(t.columns,widths):col.width=Cm(w)
        pr=t._tbl.tblPr
        for e in list(pr):
            if e.tag in [qn('w:shd'),qn('w:tblStyle'),qn('w:tblInd'),qn('w:tblBorders')]:pr.remove(e)
        bd=OxmlElement('w:tblBorders')
        for name in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+name);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'000000');bd.append(e)
        pr.append(bd)
        for i,row in enumerate(t.rows):
            rp=row._tr.get_or_add_trPr()
            for e in list(rp):
                if e.tag in [qn('w:trHeight'),qn('w:cantSplit'),qn('w:tblHeader')]:rp.remove(e)
            rp.append(OxmlElement('w:cantSplit'))
            if i==0:rp.append(OxmlElement('w:tblHeader'))
            for j,c in enumerate(row.cells):
                c.width=Cm(widths[min(j,len(widths)-1)]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp=c._tc.get_or_add_tcPr()
                for e in list(cp):
                    if e.tag in [qn('w:shd'),qn('w:tcMar'),qn('w:tcBorders')]:cp.remove(e)
                mar=OxmlElement('w:tcMar')
                for name,w in [('top',35),('bottom',35),('left',70),('right',70)]:
                    e=OxmlElement('w:'+name);e.set(qn('w:w'),str(w));e.set(qn('w:type'),'dxa');mar.append(e)
                cp.append(mar)
                for p in c.paragraphs:
                    if p._p.pPr is not None:p._p.remove(p._p.pPr)
                    p.style='IM Table'
                    p.paragraph_format.keep_with_next=(i==0 or (i==len(t.rows)-2 and t.rows[-1].cells[0].text in ['Итого','Всего']))
                    if len(c.text)<13:p.alignment=A.CENTER
                    for run in p.runs:font(run,12,i==0)
    def fig(self,path,title,width=None):
        self.f+=1;p=self.p('','IM Figure');p.add_run().add_picture(str(path),width=Cm(width or (24.0 if self.land else 15.5)))
        self.p(f'Рисунок {self.n}.{self.f} – {title}','IM Figure Caption')
    def import_practice(self,n):
        src=Document(OUT/f'ИМ_Практическая_{n}_АлбахтинИВ.docx')
        nodes=list(src._element.body);nodes=nodes[16:]
        stop=next((i for i,e in enumerate(nodes) if txt(e).lower() in ['источники','список использованных источников']),len(nodes)-1)
        nodes=nodes[1:stop] # replace the original practice title
        # Assign each block the geometry of the section that contains it.
        orient=[];land=src.sections[-1].page_width>src.sections[-1].page_height
        for e in reversed(nodes):
            sect=e.find('.//'+qn('w:sectPr'))
            if sect is not None:
                ps=sect.find(qn('w:pgSz'));land=int(ps.get(qn('w:w')))>int(ps.get(qn('w:h')))
            orient.append(land)
        orient.reverse();self.orient(False);self.n=n;self.t=0;self.f=0
        self.p(f'Практическая работа № {n}','IM Practice')
        for raw,land in zip(nodes,orient):
            s=txt(raw)
            if not s and not list(raw.iter(qn('w:drawing'))):continue
            oldstyle=Paragraph(raw,src._body).style.name if raw.tag==qn('w:p') else ''
            self.orient(land);e=deepcopy(raw)
            for el in list(e.iter(qn('w:sectPr'))):el.getparent().remove(el)
            for el in list(e.iter(qn('w:br'))):
                if el.get(qn('w:type'))=='page':el.getparent().remove(el)
            for el in list(e.iter(qn('w:lastRenderedPageBreak'))):el.getparent().remove(el)
            # Import images, never reuse the source document's relationship IDs.
            for blip in e.iter(qn('a:blip')):
                rid=blip.get(qn('r:embed'))
                if rid:
                    blob=(BASE/'figures/pr2_flows.png').read_bytes() if n==2 and (BASE/'figures/pr2_flows.png').exists() else src.part.related_parts[rid].blob
                    newrid,_=self.doc.part.get_or_add_image(BytesIO(blob));blip.set(qn('r:embed'),newrid)
            for textnode in e.iter(qn('w:t')):
                if textnode.text:
                    textnode.text=re.sub(r'(?i)(таблиц\w*\s+|рисунк\w*\s+)(\d+)(?!\d|\.\d)',lambda m:m[1]+str(n)+'.'+m[2],textnode.text)
                    textnode.text=re.sub(r'(?i)(таблиц\w*\s+\d+\.\d+\s+и\s+)(\d+)',lambda m:m[1]+str(n)+'.'+m[2],textnode.text)
            self.doc._element.body.insert(-1,e)
            if e.tag==qn('w:tbl'):
                self.style_table(Table(e,self.doc._body));continue
            if e.tag!=qn('w:p'):continue
            p=Paragraph(e,self.doc._body)
            haspic=bool(list(e.iter(qn('w:drawing'))))
            if not haspic:
                # Source runs often split the word and its number, so update the whole paragraph.
                s=re.sub(r'(?i)(таблиц\w*\s+|рисун(?:ок|к)\w*\s+)(\d+)(?!\d|\.\d)',lambda m:m[1]+str(n)+'.'+m[2],s)
                s=re.sub(r'(?i)((?:таблиц\w*|рисун(?:ок|к)\w*)\s+\d+\.\d+\s*(?:и|–|—|-)\s*)(\d+)(?!\d|\.\d)',lambda m:m[1]+str(n)+'.'+m[2],s)
                if n>1:
                    s=re.sub(r'\[(\d+(?:,\s*\d+)*)\]',lambda m:'[1]' if m[1]=='1' else '(см. практику 1)' if n==2 else f'(см. практики 1–{n-1})',s)
                # Flatten obsolete REF/SEQ fields and bookmarks from separate reports.
                # Their cached field numbers otherwise reappear when the new TOC is refreshed.
                p.clear();p.add_run(s)
            if p._p.pPr is not None:p._p.remove(p._p.pPr)
            if haspic:
                p.style='IM Figure'
                width=Cm(26.49 if self.land else 15.5)
                for drawing in e.iter(qn('wp:inline')):
                    ex=drawing.find(qn('wp:extent'));cx,cy=int(ex.get('cx')),int(ex.get('cy'));scale=min(width/cx,Cm(11.0 if self.land else 17)/cy)
                    for dim in list(drawing.iter(qn('wp:extent')))+list(drawing.iter(qn('a:ext'))):
                        if dim.get('cx'):dim.set('cx',str(int(cx*scale)));dim.set('cy',str(int(cy*scale)))
            elif re.match(r'^Таблица \d+\.\d+\s*[–-]',s):p.style='IM Table Caption';self.t+=1
            elif s.startswith('Рисунок '):
                p.style='IM Figure Caption';self.f+=1
                settext(e,re.sub(r'^Рисунок \d+(?:\.\d+)?',f'Рисунок {n}.{self.f}',p.text))
            elif ('heading' in oldstyle.lower() or 'заголов' in oldstyle.lower()):
                p.style='IM Head';s=re.sub(r'^\d+(?:\.\d+)*\s+','',p.text);settext(e,s)
            else:p.style='IM Body'
            size=12 if p.style.name=='IM Table Caption' else 14
            for run in p.runs:font(run,size,p.style.name=='IM Head',p.style.name=='IM Table Caption')
        if n==3:self.p('Для последующего расчёта сметы условные часовые ставки понимаются как полные затраты на труд, включая отчисления. Расшифровка по экономическим элементам приведена в практике 12; повторное начисление отчислений поверх 512 000 руб. не выполняется.')
    def save(self):
        for i,e in enumerate(self.doc._element.iter(qn('wp:docPr')),1):e.set('id',str(i))
        self.doc.save(DOCX);return DOCX
